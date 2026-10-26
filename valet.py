# -*- coding: utf-8 -*-
"""
CLUB PARK VALET Y ESTACIONAMIENTOS S.A. DE C.V.
Sistema de control operativo con Streamlit.

IMPORTANTE:
- Para conservar archivos en Streamlit Cloud después de reinicios/despliegues,
  configura CLUBPARK_DATA_DIR con una ruta de almacenamiento persistente o usa
  un servicio externo (S3, Azure Blob, Supabase Storage, etc.).
- Cambia los tokens de ejemplo antes de usar la aplicación en producción.
- La impresión Bluetooth Web Bluetooth requiere navegador/dispositivo compatible,
  contexto HTTPS y una impresora que exponga un servicio GATT compatible.
"""

import io
import json
import math
import os
import re
import uuid
import zipfile
import urllib.parse
import html
import base64
from datetime import date, datetime, timedelta, timezone
from pathlib import Path

try:
    from zoneinfo import ZoneInfo
    _TZ_MX = ZoneInfo("America/Mexico_City")
except Exception:  # sin base de zonas horarias: México centro es UTC-6 fijo desde 2022
    _TZ_MX = timezone(timedelta(hours=-6))


def ahora_mx():
    """Fecha y hora actuales de Ciudad de México (el servidor suele estar en UTC)."""
    return datetime.now(_TZ_MX).replace(tzinfo=None)


def hoy_mx():
    return ahora_mx().date()


import pandas as pd
import streamlit as st
from PIL import Image, ImageOps

try:
    import qrcode
    QR_OK = True
except ImportError:
    qrcode = None
    QR_OK = False

try:
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    OPENPYXL_OK = True
except ImportError:
    OPENPYXL_OK = False

try:
    import xlsxwriter  # noqa: F401
    XLSXWRITER_OK = True
except ImportError:
    XLSXWRITER_OK = False


# ---------------------------------------------------------------------
# CONFIGURACIÓN
# ---------------------------------------------------------------------
pass  # set_page_config lo hace main.py

try:
    if st.query_params.get("ping"):
        st.write("ok")
        st.stop()
except AttributeError:
    pass

# ---------------------------------------------------------------------
# ESTILO CONSTITUCIONAL E INSTITUCIONAL - VERDE BOTELLA (#0B3B24)
# CLUB PARK VALET Y ESTACIONAMIENTOS S.A. DE C.V.
# ---------------------------------------------------------------------
CSS_INSTITUCIONAL = """
<style>
/* Paleta Institucional Club Park: Verde Botella Oficial */
:root {
  --cpv-verde-botella: #0B3B24;
  --cpv-verde-oscuro: #072617;
  --cpv-verde-medio: #145A32;
  --cpv-verde-claro: #EBF3EE;
  --cpv-dorado-oficial: #5B8F6E;
  --cpv-dorado-brillante: #7FA98C;
  --cpv-fondo-suave: #F4F7F5;
  --cpv-texto-principal: #1A2B20;
}

/* Encabezados y títulos */
h1, h2, h3 {
  color: #0B3B24 !important;
  font-family: -apple-system, BlinkMacSystemFont, "Segoe UI", Roboto, "Helvetica Neue", Arial, sans-serif;
  font-weight: 700;
}

/* Botones Primarios */
button[kind="primary"], .stButton > button[kind="primary"], div[data-testid="stFormSubmitButton"] > button {
  background-color: #0B3B24 !important;
  color: #FFFFFF !important;
  border: 1px solid #072617 !important;
  border-radius: 8px !important;
  font-weight: 600 !important;
  box-shadow: 0 2px 5px rgba(11, 59, 36, 0.25) !important;
  transition: all 0.2s ease !important;
}
button[kind="primary"]:hover, .stButton > button[kind="primary"]:hover {
  background-color: #145A32 !important;
  border-color: #0B3B24 !important;
  color: #FFFFFF !important;
}

/* Botones Secundarios */
button[kind="secondary"], .stButton > button[kind="secondary"] {
  background-color: #FFFFFF !important;
  color: #0B3B24 !important;
  border: 1px solid #C4D7CC !important;
  border-radius: 8px !important;
  font-weight: 600 !important;
}
button[kind="secondary"]:hover, .stButton > button[kind="secondary"]:hover {
  background-color: #EBF3EE !important;
  border-color: #0B3B24 !important;
  color: #0B3B24 !important;
}

/* Pestañas (Tabs) en Verde Botella */
button[data-baseweb="tab"] {
  color: #1A2B20 !important;
  font-weight: 600 !important;
  font-size: 14px !important;
}
button[data-baseweb="tab"][aria-selected="true"] {
  color: #0B3B24 !important;
  border-bottom-color: #0B3B24 !important;
  font-weight: 700 !important;
}

/* Alertas y métricas */
div[data-testid="stMetricValue"] {
  color: #0B3B24 !important;
}
div[data-testid="stAlert"] {
  border-radius: 8px !important;
}

/* Banners institucionales */
.clubpark-header-card {
  background: linear-gradient(135deg, #072617 0%, #0B3B24 65%, #145A32 100%);
  border: 2px solid #5B8F6E;
  border-radius: 12px;
  padding: 20px 24px;
  color: #FFFFFF;
  margin-bottom: 20px;
  box-shadow: 0 4px 14px rgba(11, 59, 36, 0.25);
}
.cpv-titulo-blanco {
  color: #FFFFFF !important;
  font-family: Georgia, "Palatino Linotype", "Times New Roman", serif;
  font-size: 1.35rem;
  font-weight: 800;
  line-height: 1.25;
  letter-spacing: 0.5px;
  text-shadow: 0 1px 3px rgba(0, 0, 0, 0.55);
}
.clubpark-header-card, .clubpark-header-card div, .clubpark-header-card span, .clubpark-header-card b,
.clubpark-header-card h1, .clubpark-header-card h2, .clubpark-header-card h3,
.clubpark-header-card h1 *, .clubpark-header-card h2 *, .clubpark-header-card h3 * {
  color: #FFFFFF !important;
}
.clubpark-badge {
  background: rgba(46, 107, 71, 0.30);
  border: 1px solid #5B8F6E;
  color: #E8F3EC;
  padding: 4px 10px;
  border-radius: 6px;
  font-size: 12px;
  font-weight: 700;
  display: inline-block;
  letter-spacing: 0.5px;
}

/* Presentación formal y elegante - Tipografía institucional */
h1, h2, h3 {
  font-family: Georgia, "Palatino Linotype", "Times New Roman", serif !important;
  letter-spacing: 0.4px;
}
h1 { font-size: 1.65rem !important; }
h2 { font-size: 1.35rem !important; }
h3 { font-size: 1.12rem !important; }
/* Barra lateral institucional en verde botella */
section[data-testid="stSidebar"] {
  background: linear-gradient(180deg, #072617 0%, #0B3B24 70%, #145A32 100%);
  border-right: 2px solid #5B8F6E;
}
section[data-testid="stSidebar"] * { color: #F4F7F5 !important; }
section[data-testid="stSidebar"] button[kind="primary"],
section[data-testid="stSidebar"] button[kind="secondary"],
section[data-testid="stSidebar"] .stButton > button {
  background-color: #F4F7F5 !important; color: #0B3B24 !important;
  border: 1px solid #5B8F6E !important; font-weight: 700 !important;
}
section[data-testid="stSidebar"] .stButton > button:hover {
  background-color: #EBF3EE !important;
}
/* Tarjetas de métricas formales */
div[data-testid="stMetric"] {
  background: #FFFFFF;
  border: 1px solid #C4D7CC;
  border-left: 3px solid #5B8F6E;
  border-radius: 10px;
  padding: 10px 14px;
  box-shadow: 0 1px 4px rgba(11, 59, 36, 0.10);
}
/* Expanders y contenedores */
div[data-testid="stExpander"] {
  border: 1px solid #C4D7CC !important;
  border-radius: 10px !important;
  background: #FFFFFF !important;
}
/* Pestaña activa con subrayado dorado */
button[data-baseweb="tab"][aria-selected="true"] {
  box-shadow: inset 0 -3px 0 #5B8F6E;
}
/* Entradas de texto con borde institucional */
input, textarea, div[data-baseweb="select"] > div {
  border: 1px solid #C4D7CC !important;
  border-radius: 8px !important;
}
input:focus, textarea:focus {
  border-color: #0B3B24 !important;
  box-shadow: 0 0 0 1px #0B3B24 !important;
}
/* Regla dorada institucional */
.clubpark-gold-rule { border: 0; border-top: 2px solid #5B8F6E; margin: 12px 0; }
</style>
"""
st.markdown(CSS_INSTITUCIONAL, unsafe_allow_html=True)


def activar_camara_trasera_por_defecto():
    """Hace que st.camera_input abra la cámara TRASERA del teléfono.
    Intercepta getUserMedia en la página: lo que Streamlit pide como cámara
    frontal se convierte en trasera y viceversa, así el botón de cambiar
    cámara sigue funcionando. En computadoras con una sola cámara no afecta."""
    st.components.v1.html("""
    <script>
    (function () {
      try {
        const w = window.parent;
        const md = w && w.navigator && w.navigator.mediaDevices;
        if (!md || !md.getUserMedia || w.__cpvCamaraTrasera) return;
        w.__cpvCamaraTrasera = true;
        const original = md.getUserMedia.bind(md);
        const leer = (fm) => (fm && typeof fm === "object") ? (fm.exact || fm.ideal || "") : (fm || "");
        md.getUserMedia = function (c) {
          try {
            if (c && c.video) {
              const v = (c.video === true) ? {} : Object.assign({}, c.video);
              if (!v.deviceId) {
                const pedido = String(leer(v.facingMode));
                v.facingMode = { ideal: pedido === "environment" ? "user" : "environment" };
              }
              c = Object.assign({}, c, { video: v });
            }
          } catch (e) {}
          return original(c);
        };
      } catch (e) {}
    })();
    </script>
    """, height=0)


activar_camara_trasera_por_defecto()


def encabezado_institucional(titulo, badge="", subtitulo=""):
    """Encabezado formal uniforme para todas las pestañas del sistema institucional."""
    insignia = f'<div class="clubpark-badge" style="margin-top:2px">{badge}</div>' if badge else ""
    sub = f'<div style="font-size:13px;color:#EBF3EE;margin-top:2px">{subtitulo}</div>' if subtitulo else ""
    st.markdown(
        f"""<div class="clubpark-header-card">
          {insignia}
          <div class="cpv-titulo-blanco" style="margin:6px 0 2px 0">{titulo}</div>
          {sub}
        </div>""",
        unsafe_allow_html=True,
    )

BASE_DIR = Path(__file__).resolve().parent.parent
DATA_DIR = Path(os.environ.get("CLUBPARK_DATA_DIR", str(BASE_DIR))) / "valet"
DATA_DIR.mkdir(parents=True, exist_ok=True)
PHOTO_DIR = DATA_DIR / "fotos_vehiculos"
PHOTO_DIR.mkdir(parents=True, exist_ok=True)
DB_FILE = DATA_DIR / "vehiculos_produccion_db.json"
CONFIG_FILE = DATA_DIR / "config_operativa.json"
AUDIT_FILE = DATA_DIR / "auditoria_operativa.json"
PAYMENTS_FILE = DATA_DIR / "pagos_registrados.json"

# Contador de llaves por ejecución: evita StreamlitDuplicateElementKey aunque una
# sección se renderice dos veces en la misma corrida (el contador se reinicia en cada ejecución del script).
_CONTADOR_LLAVES = 0

def llave_unica(base):
    global _CONTADOR_LLAVES
    _CONTADOR_LLAVES += 1
    return f"{base}__{_CONTADOR_LLAVES}"

PRECIOS_SUCURSALES = {
    "PAMPAS INSURGENTES": 50,
    "PAMPAS COAPA": 50,
    "RAMMA": 70,
    "SUSHITTO": 40,
    "CARAJILLO POLANCO": 100,
    "BENINGAN´S": 40,
    "SANATORIO DURANGO": 100,
    "BARRANCA": 25,  # solo base (1ª hora); el cobro real sale de TARIFAS (tiempo / promoción)
    "VETERINARIA": 0,
    "ROMANCE PALACE": 65,
    "JUNGLA TLALPAN": 100,
    "JUNGLA DIVISIÓN": 100,
    "JUNGLA PICACHO": 100,
    "JUNGLA EMPRESA": 100,
    "SIPIRILY COAPA": 80,
    "MAROMETA SATÉLITE": 70,
    "MAROMETA MARINA": 70,
    "SUMO STA. MONICA": 50,
    "SUMO ECHEGARAY": 50,
    "RYOSHI": 80,
    "ENCANTO": 60,
    "E-BNI": 50,
    "CHURRASQUERIA": 70,
}
SUCURSALES_FILE = DATA_DIR / "sucursales_config.json"


def cargar_sucursales_config():
    try:
        datos = json.loads(SUCURSALES_FILE.read_text(encoding="utf-8"))
        if isinstance(datos, dict):
            return {"agregadas": dict(datos.get("agregadas", {})), "quitadas": list(datos.get("quitadas", [])),
                    "supervisores": dict(datos.get("supervisores", {})),
                    "salones": datos.get("salones"),
                    "detalles": dict(datos.get("detalles", {})),
                    "tokens_extra": dict(datos.get("tokens_extra", {})),
                    "telefonos_sup": dict(datos.get("telefonos_sup", {}))}
    except (OSError, ValueError):
        pass
    return {"agregadas": {}, "quitadas": [], "supervisores": {}, "salones": None, "detalles": {}, "tokens_extra": {}, "telefonos_sup": {}}


def guardar_sucursales_config(cfg):
    temporal = SUCURSALES_FILE.with_suffix(".tmp")
    temporal.write_text(json.dumps(cfg, ensure_ascii=False, indent=2), encoding="utf-8")
    os.replace(temporal, SUCURSALES_FILE)


_SUC_CFG = cargar_sucursales_config()
for _n, _p in _SUC_CFG["agregadas"].items():
    PRECIOS_SUCURSALES[_n] = float(_p)
for _n in _SUC_CFG["quitadas"]:
    PRECIOS_SUCURSALES.pop(_n, None)
SUCURSALES = sorted(PRECIOS_SUCURSALES)

# ---------------------------------------------------------------------
# TARIFAS CONFIGURABLES POR SUCURSAL
# tipo "fija": cobra siempre el mismo precio (ej. $50).
# tipo "tiempo": 1ª hora + bloques de N minutos (ej. $25 + $7 c/15 min).
# promo_activa: ofrece además una promoción (ej. $80 por 10 horas).
# El Contador puede cambiar todo esto en la pestaña "💲 Tarifas".
# ---------------------------------------------------------------------
TARIFAS_FILE = DATA_DIR / "tarifas_config.json"
_TARIFA_BASE = {
    "tipo": "fija", "precio_fijo": 0.0,
    "primera_hora": 25.0, "bloque_min": 15, "precio_bloque": 7.0,
    "promo_activa": False, "promo_precio": 80.0, "promo_horas": 10,
    # Cobro al terminar la promoción: se cobra "promo_extra_precio" por cada
    # bloque de "promo_extra_min" minutos (fracción iniciada se cobra completa).
    "promo_extra_precio": 7.0, "promo_extra_min": 15,
}
ETIQUETAS_MODALIDAD = {"fija": "Tarifa fija", "tiempo": "Normal por tiempo", "promocion": "Promoción"}


def tarifas_por_defecto():
    tarifas = {suc: dict(_TARIFA_BASE, precio_fijo=float(precio)) for suc, precio in PRECIOS_SUCURSALES.items()}
    tarifas["BARRANCA"] = dict(_TARIFA_BASE, tipo="tiempo", precio_fijo=0.0, promo_activa=True)
    return tarifas


_CACHE_TARIFAS = {"firma": None, "datos": None}


def cargar_tarifas():
    import copy
    try:
        firma = (TARIFAS_FILE.stat().st_mtime_ns, len(PRECIOS_SUCURSALES)) if TARIFAS_FILE.exists() else (None, len(PRECIOS_SUCURSALES))
    except OSError:
        firma = None
    if firma is not None and _CACHE_TARIFAS["firma"] == firma:
        return copy.deepcopy(_CACHE_TARIFAS["datos"])
    datos = _cargar_tarifas_disco()
    _CACHE_TARIFAS.update(firma=firma, datos=copy.deepcopy(datos))
    return datos


def _cargar_tarifas_disco():
    tarifas = tarifas_por_defecto()
    if TARIFAS_FILE.exists():
        try:
            guardadas = json.loads(TARIFAS_FILE.read_text(encoding="utf-8"))
            if isinstance(guardadas, dict):
                for suc, cfg in guardadas.items():
                    if suc in tarifas and isinstance(cfg, dict):
                        tarifas[suc].update({k: v for k, v in cfg.items() if k in _TARIFA_BASE})
        except (OSError, ValueError):
            pass
    return tarifas


def guardar_tarifas(tarifas):
    temporal = TARIFAS_FILE.with_suffix(".tmp")
    temporal.write_text(json.dumps(tarifas, ensure_ascii=False, indent=2), encoding="utf-8")
    os.replace(temporal, TARIFAS_FILE)


def tarifa_de(sucursal):
    return cargar_tarifas().get(sucursal) or dict(_TARIFA_BASE, precio_fijo=float(PRECIOS_SUCURSALES.get(sucursal, 50)))


def modalidades_disponibles(cfg):
    if cfg.get("tipo") == "tiempo":
        return ["tiempo", "promocion"] if cfg.get("promo_activa") else ["tiempo"]
    return ["fija"]


def etiqueta_modalidad(cfg, modalidad):
    if modalidad == "promocion":
        return (f"Promoción ${float(cfg['promo_precio']):,.0f} por {int(cfg['promo_horas'])} h; después "
                f"${float(cfg.get('promo_extra_precio', cfg['precio_bloque'])):,.0f} cada "
                f"{int(cfg.get('promo_extra_min', cfg['bloque_min']))} min")
    if modalidad == "tiempo":
        return (f"Normal: ${float(cfg['primera_hora']):,.0f} la 1ª hora + "
                f"${float(cfg['precio_bloque']):,.0f} cada {int(cfg['bloque_min'])} min")
    return f"Tarifa fija ${float(cfg['precio_fijo']):,.0f}"


def precio_inicial(cfg, modalidad):
    if modalidad == "promocion":
        return float(cfg["promo_precio"])
    if modalidad == "tiempo":
        return float(cfg["primera_hora"])
    return float(cfg["precio_fijo"])


def _cobro_tiempo(cfg, minutos):
    primera = float(cfg["primera_hora"])
    if minutos <= 60:
        return primera
    bloque = max(1, int(cfg["bloque_min"]))
    return primera + math.ceil((minutos - 60) / bloque) * float(cfg["precio_bloque"])


def calcular_cobro(cfg, modalidad, minutos, precio_registrado=None):
    """Importe a cobrar. Cada fracción de bloque iniciada se cobra completa."""
    minutos = max(0, int(minutos or 0))
    if modalidad == "promocion":
        base = float(cfg["promo_precio"])
        limite = int(cfg["promo_horas"]) * 60
        if minutos > limite:  # terminada la promoción se cobra su costo excedente
            bloque = max(1, int(cfg.get("promo_extra_min", cfg["bloque_min"])))
            precio_extra = float(cfg.get("promo_extra_precio", cfg["precio_bloque"]))
            base += math.ceil((minutos - limite) / bloque) * precio_extra
        return base
    if modalidad == "tiempo":
        return _cobro_tiempo(cfg, minutos)
    return float(precio_registrado) if precio_registrado is not None else float(cfg["precio_fijo"])


def minutos_estancia(registro, fin=None):
    fin = fin or ahora_mx()
    entrada = None
    iso = registro.get("fecha_entrada_iso")
    if iso:
        try:
            entrada = datetime.fromisoformat(iso)
        except ValueError:
            entrada = None
    if entrada is None:
        for valor in (registro.get("fecha_entrada"), registro.get("fecha")):
            for formato in ("%d/%m/%Y %H:%M:%S", "%Y-%m-%d %H:%M:%S"):
                try:
                    entrada = datetime.strptime(str(valor or ""), formato)
                    break
                except ValueError:
                    continue
            if entrada:
                break
    if entrada is None:
        return 0
    return max(0, int((fin - entrada).total_seconds() // 60))


def texto_duracion(minutos):
    return f"{minutos // 60} h {minutos % 60} min"


def lineas_tarifa_ticket(sucursal, modalidad):
    """Texto de tarifas del ticket. Tarifa fija: solo bienvenida. Por tiempo o promoción: sus fracciones configuradas."""
    cfg = tarifa_de(sucursal)
    if cfg.get("tipo") == "tiempo" or modalidad == "promocion":
        if modalidad == "promocion":
            return [f"PROMOCION ${float(cfg['promo_precio']):,.2f} POR {int(cfg['promo_horas'])} HORAS.",
                    f"Al terminar: ${float(cfg.get('promo_extra_precio', cfg['precio_bloque'])):,.2f} cada "
                    f"{int(cfg.get('promo_extra_min', cfg['bloque_min']))} min."]
        return [f"BIENVENIDO. Primera hora ${float(cfg['primera_hora']):,.2f}.",
                f"Despues: ${float(cfg['precio_bloque']):,.2f} cada {int(cfg['bloque_min'])} min",
                "(fraccion iniciada se cobra)."]
    # Sucursales de tarifa fija (Pampas, Encanto, Carajillo, etc.): solo bienvenida, sin texto de fracciones.
    return ["BIENVENIDO."]


# Tokens de ejemplo: reemplázalos por secretos/variables de entorno en producción.
# Si defines CLUBPARK_TOKENS_JSON, debe contener un objeto JSON con el mismo formato.
TOKENS_POR_DEFECTO = {
    "CPV-GLOBAL-AUDITOR-2026": {
        "nombre": "Víctor (Contador)",
        "sucursal": "TODAS",
        "rol": "Contador Auditor (Máxima Jerarquía)",
    },
    "CPV-GLOBAL-JESUS-2026": {
        "nombre": "José de Jesús",
        "sucursal": "TODAS",
        "rol": "Supervisor Global",
    },
    "CPV-GLOBAL-VICTOR-2026": {
        "nombre": "Víctor Manuel",
        "sucursal": "TODAS",
        "rol": "Supervisor Global",
    },
    "CPV-QR-PAMPAS-INS": {"nombre": "Líder Pampas Ins.", "sucursal": "PAMPAS INSURGENTES", "rol": "Líder de Sucursal"},
    "CPV-QR-PAMPAS-COAPA": {"nombre": "Líder Pampas Coapa", "sucursal": "PAMPAS COAPA", "rol": "Líder de Sucursal"},
    "CPV-QR-RAMMA": {"nombre": "Líder Ramma", "sucursal": "RAMMA", "rol": "Líder de Sucursal"},
    "CPV-QR-SUSHITTO": {"nombre": "Líder Sushitto", "sucursal": "SUSHITTO", "rol": "Líder de Sucursal"},
    "CPV-QR-CARAJILLO": {"nombre": "Líder Carajillo", "sucursal": "CARAJILLO POLANCO", "rol": "Líder de Sucursal"},
    "CPV-QR-BENINGANS": {"nombre": "Líder Beningan's", "sucursal": "BENINGAN´S", "rol": "Líder de Sucursal"},
    "CPV-QR-DURANGO": {"nombre": "Líder Sanatorio Durango", "sucursal": "SANATORIO DURANGO", "rol": "Líder de Sucursal"},
    "CPV-QR-BARRANCA": {"nombre": "Líder Barranca", "sucursal": "BARRANCA", "rol": "Líder de Sucursal"},
    "CPV-QR-VETERINARIA": {"nombre": "Líder Veterinaria", "sucursal": "VETERINARIA", "rol": "Líder de Sucursal"},
    "CPV-QR-ROMANCE": {"nombre": "Líder Romance", "sucursal": "ROMANCE PALACE", "rol": "Líder de Sucursal"},
    "CPV-QR-TLALPAN": {"nombre": "Líder Jungla Tlalpan", "sucursal": "JUNGLA TLALPAN", "rol": "Líder de Sucursal"},
    "CPV-QR-DIVISION": {"nombre": "Líder Jungla División", "sucursal": "JUNGLA DIVISIÓN", "rol": "Líder de Sucursal"},
    "CPV-QR-PICACHO": {"nombre": "Líder Jungla Picacho", "sucursal": "JUNGLA PICACHO", "rol": "Líder de Sucursal"},
    "CPV-QR-EMPRESA": {"nombre": "Líder Jungla Empresa", "sucursal": "JUNGLA EMPRESA", "rol": "Líder de Sucursal"},
    "CPV-QR-SIPIRILY": {"nombre": "Líder Sipirily", "sucursal": "SIPIRILY COAPA", "rol": "Líder de Sucursal"},
    "CPV-QR-SATELITE": {"nombre": "Líder Marometa Satélite", "sucursal": "MAROMETA SATÉLITE", "rol": "Líder de Sucursal"},
    "CPV-QR-MARINA": {"nombre": "Líder Marometa Marina", "sucursal": "MAROMETA MARINA", "rol": "Líder de Sucursal"},
    "CPV-QR-SUMOSTA": {"nombre": "Líder Sumo Santa Mónica", "sucursal": "SUMO STA. MONICA", "rol": "Líder de Sucursal"},
    "CPV-QR-CHEGARAY": {"nombre": "Líder Sumo Echegaray", "sucursal": "SUMO ECHEGARAY", "rol": "Líder de Sucursal"},
    "CPV-QR-RYOSHI": {"nombre": "Líder Ryoshi", "sucursal": "RYOSHI", "rol": "Líder de Sucursal"},
    "CPV-QR-ENCANTO": {"nombre": "Líder Encanto", "sucursal": "ENCANTO", "rol": "Líder de Sucursal"},
    "CPV-QR-EBNI": {"nombre": "Líder E-BNI", "sucursal": "E-BNI", "rol": "Líder de Sucursal"},
    "CPV-QR-CHURRAS": {"nombre": "Líder Churrasquería", "sucursal": "CHURRASQUERIA", "rol": "Líder de Sucursal"},
}
# ---------------------------------------------------------------------
# CARGA PROTEGIDA DE CREDENCIALES QR
# Prioridad: 1) st.secrets  2) variable de entorno CLUBPARK_TOKENS_JSON
# 3) tokens por defecto del código (solo contingencia; la app avisa en
# pantalla para moverlos a st.secrets lo antes posible).
# ---------------------------------------------------------------------
def cargar_accesos_protegidos():
    fuente = "por_defecto"
    accesos = None
    try:
        if "CLUBPARK_TOKENS_JSON" in st.secrets:
            accesos = json.loads(str(st.secrets["CLUBPARK_TOKENS_JSON"]))
            fuente = "secrets"
        elif "CLUBPARK_TOKENS" in st.secrets:
            accesos = {clave: dict(valor) for clave, valor in st.secrets["CLUBPARK_TOKENS"].items()}
            fuente = "secrets"
    except Exception:
        accesos = None
    if not accesos:
        tokens_env = os.environ.get("CLUBPARK_TOKENS_JSON", "").strip()
        if tokens_env:
            try:
                accesos = json.loads(tokens_env)
                fuente = "entorno"
            except json.JSONDecodeError:
                st.warning("CLUBPARK_TOKENS_JSON no contiene JSON válido; se ignoró.")
                accesos = None
    if not (isinstance(accesos, dict) and accesos):
        accesos = dict(TOKENS_POR_DEFECTO)
        fuente = "por_defecto"
    return accesos, fuente


ACCESOS, ACCESOS_FUENTE = cargar_accesos_protegidos()
# Credenciales de líder generadas al dar de alta sucursales nuevas, y bloqueo de las sucursales quitadas
for _tok, _dat in _SUC_CFG.get("tokens_extra", {}).items():
    ACCESOS.setdefault(_tok, dict(_dat))
ACCESOS = {_t: _d for _t, _d in ACCESOS.items() if _d.get("sucursal") not in set(_SUC_CFG.get("quitadas", []))}

try:
    APP_URL = str(st.secrets.get("CLUBPARK_APP_URL", "") or os.environ.get("CLUBPARK_APP_URL", "")).strip()
except Exception:
    APP_URL = str(os.environ.get("CLUBPARK_APP_URL", "")).strip()


@st.cache_data(show_spinner=False, max_entries=500)
def qr_png_cache(contenido):
    """PNG verde institucional del QR; se genera una sola vez por contenido."""
    qr_c = qrcode.QRCode(version=1, box_size=8, border=2)
    qr_c.add_data(contenido)
    qr_c.make(fit=True)
    buf_c = io.BytesIO()
    qr_c.make_image(fill_color="#0B3B24", back_color="#FFFFFF").save(buf_c, format="PNG")
    return buf_c.getvalue()


def contenido_qr_para(token):
    """Contenido del QR de la credencial: si se configura CLUBPARK_APP_URL el QR
    abre la app con la sesión lista al escanearlo con la cámara del teléfono."""
    if APP_URL:
        separador = "&" if "?" in APP_URL else "?"
        return f"{APP_URL}{separador}qrscan={token}"
    return token

# Agrega aquí los tokens QR de encargados de sucursal o define todos en
# CLUBPARK_TOKENS_JSON. Formato:
# {"TOKEN":{"nombre":"Líder Ramma","sucursal":"RAMMA","rol":"Líder de Sucursal"}}


# ---------------------------------------------------------------------
# DATOS Y UTILIDADES
# ---------------------------------------------------------------------
def cargar_bd():
    if not DB_FILE.exists():
        return []
    try:
        with DB_FILE.open("r", encoding="utf-8") as f:
            data = json.load(f)
        return data if isinstance(data, list) else []
    except (OSError, json.JSONDecodeError):
        st.warning("No se pudo leer la base de datos JSON. Se iniciará una lista vacía.")
        return []


@st.cache_data(show_spinner=False, max_entries=4)
def _bd_cache(ruta_txt, mtime, tam):
    try:
        datos = json.loads(Path(ruta_txt).read_text(encoding="utf-8"))
        return datos if isinstance(datos, list) else []
    except (OSError, ValueError):
        return []


def cargar_bd_rapido():
    """Igual que cargar_bd, pero solo relee el archivo cuando cambió (para revisiones en vivo)."""
    if not DB_FILE.exists():
        return []
    est = DB_FILE.stat()
    return _bd_cache(str(DB_FILE), est.st_mtime_ns, est.st_size)


def guardar_bd(registros):
    temp_file = DB_FILE.with_suffix(".json.tmp")
    try:
        with temp_file.open("w", encoding="utf-8") as f:
            json.dump(registros, f, ensure_ascii=False, separators=(",", ":"))
        os.replace(temp_file, DB_FILE)
        return True
    except OSError as exc:
        st.error(f"No se pudo guardar la base de datos: {exc}")
        return False


def cargar_pagos_respaldo():
    """Carga el libro independiente de cobros, además del pago dentro del vehículo."""
    if not PAYMENTS_FILE.exists():
        return []
    try:
        with PAYMENTS_FILE.open("r", encoding="utf-8") as f:
            datos = json.load(f)
        return datos if isinstance(datos, list) else []
    except (OSError, json.JSONDecodeError):
        return []


def guardar_pago_respaldo(pago, registro):
    """Guarda cada cobro en un libro separado para facilitar respaldo y auditoría."""
    pagos = cargar_pagos_respaldo()
    pago_respaldo = dict(pago)
    pago_respaldo["registro_id"] = str(registro.get("id", ""))
    pago_respaldo["sucursal"] = registro.get("sucursal", "")
    pago_respaldo["marca"] = registro.get("marca", "")
    pago_respaldo["placas"] = registro.get("placas", "")
    # Evita duplicar el mismo pago si Streamlit repite el rerun.
    pago_id = pago_respaldo.get("id_pago")
    if pago_id and any(p.get("id_pago") == pago_id for p in pagos):
        return True
    pagos.insert(0, pago_respaldo)
    temporal = PAYMENTS_FILE.with_suffix(".tmp")
    try:
        with temporal.open("w", encoding="utf-8") as f:
            json.dump(pagos, f, ensure_ascii=False, indent=2)
        os.replace(temporal, PAYMENTS_FILE)
        return True
    except OSError as exc:
        st.error(f"No se pudo respaldar el cobro en el libro de pagos: {exc}")
        return False



def cargar_config():
    """Configuración persistente: siguiente folio por sucursal y secuencias."""
    config = {"siguiente_folio": {}, "historial_folios": []}
    if CONFIG_FILE.exists():
        try:
            with CONFIG_FILE.open("r", encoding="utf-8") as f:
                guardada = json.load(f)
            if isinstance(guardada, dict):
                config.update(guardada)
        except (OSError, json.JSONDecodeError):
            pass
    config.setdefault("siguiente_folio", {})
    config.setdefault("historial_folios", [])
    for suc in SUCURSALES:
        config["siguiente_folio"].setdefault(suc, 1)
    return config


def guardar_config(config):
    temporal = CONFIG_FILE.with_suffix(".tmp")
    with temporal.open("w", encoding="utf-8") as f:
        json.dump(config, f, ensure_ascii=False, indent=2)
    os.replace(temporal, CONFIG_FILE)


def cargar_auditoria():
    if not AUDIT_FILE.exists():
        return []
    try:
        with AUDIT_FILE.open("r", encoding="utf-8") as f:
            datos = json.load(f)
        return datos if isinstance(datos, list) else []
    except (OSError, json.JSONDecodeError):
        return []


def guardar_evento_auditoria(tipo, detalle):
    """Bitácora de auditoría institucional: usuario, rol, fecha, tipo y cambio realizado."""
    eventos = cargar_auditoria()
    eventos.insert(0, {
        "id_evento": ahora_mx().strftime("%Y%m%d%H%M%S") + f"-{len(eventos) + 1:04d}",
        "fecha_hora": ahora_mx().strftime("%Y-%m-%d %H:%M:%S"),
        "usuario": st.session_state.get("usuario_autenticado", "Sistema"),
        "rol": st.session_state.get("rol", ""),
        "tipo": tipo,
        "detalle": detalle,
    })
    temporal = AUDIT_FILE.with_suffix(".tmp")
    with temporal.open("w", encoding="utf-8") as f:
        json.dump(eventos, f, ensure_ascii=False, separators=(",", ":"))
    os.replace(temporal, AUDIT_FILE)


ARCHIVO_APRENDIZAJE = DATA_DIR / "aprendizaje_sistema.json"


def cargar_aprendizaje():
    try:
        datos = json.loads(ARCHIVO_APRENDIZAJE.read_text(encoding="utf-8"))
        return datos if isinstance(datos, list) else []
    except (OSError, json.JSONDecodeError):
        return []


def evolucion_aprendizaje(registros):
    """Aprendizaje continuo: el sistema observa su propia operación y documenta su evolución."""
    notas = cargar_aprendizaje()
    total = len(registros)
    if total == 0:
        return notas
    hito = (total // 25) * 25
    ultima_hito = notas[0].get("vehiculos_totales", 0) if notas else 0
    if hito <= ultima_hito:
        return notas
    por_sucursal = {}
    tarifas = []
    entregados = 0
    for registro in registros:
        clave = registro.get("sucursal", "SIN SUCURSAL") or "SIN SUCURSAL"
        por_sucursal[clave] = por_sucursal.get(clave, 0) + 1
        try:
            tarifas.append(float(registro.get("tarifa", 0) or 0))
        except (TypeError, ValueError):
            pass
        if registro.get("estado") == "Entregado":
            entregados += 1
    top = sorted(por_sucursal.items(), key=lambda par: -par[1])[:3]
    promedio = sum(tarifas) / len(tarifas) if tarifas else 0.0
    observaciones = (
        f"El sistema alcanzó {total} vehículos aprendidos. Sucursales con mayor actividad: "
        + ", ".join(f"{nombre} ({cantidad})" for nombre, cantidad in top)
        + f". Tarifa promedio ${promedio:,.2f} MXN. {entregados} de {total} vehículos ya entregados "
        + f"({entregados / total * 100:.0f}%)."
    )
    notas.insert(0, {
        "fecha": ahora_mx().strftime("%Y-%m-%d %H:%M:%S"),
        "vehiculos_totales": total,
        "observaciones": observaciones,
    })
    try:
        ARCHIVO_APRENDIZAJE.write_text(json.dumps(notas[:60], ensure_ascii=False, indent=2), encoding="utf-8")
    except OSError:
        pass
    return notas


CODIGOS_SUCURSAL = {
    "PAMPAS INSURGENTES": "PINS",
    "PAMPAS COAPA": "PCOA",
    "RAMMA": "RAM",
    "SUSHITTO": "SUS",
    "CARAJILLO POLANCO": "CARA",
    "BENINGAN´S": "BEN",
    "SANATORIO DURANGO": "SDUR",
    "BARRANCA": "BARR",
    "VETERINARIA": "VET",
    "ROMANCE PALACE": "ROM",
    "JUNGLA TLALPAN": "JTL",
    "JUNGLA DIVISIÓN": "JDV",
    "JUNGLA PICACHO": "JPI",
    "JUNGLA EMPRESA": "JEM",
    "SIPIRILY COAPA": "SIP",
    "MAROMETA SATÉLITE": "MSAT",
    "MAROMETA MARINA": "MMAR",
    "SUMO STA. MONICA": "SSTA",
    "SUMO ECHEGARAY": "SEC",
    "RYOSHI": "RYO",
    "ENCANTO": "ENC",
    "E-BNI": "EBN",
    "CHURRASQUERIA": "CHUR",
}


# ---------------------------------------------------------------------
# REGISTRO DE ERRORES DE IMPRESIÓN (RawBT / ESC-POS)
# ---------------------------------------------------------------------
ARCHIVO_REGISTRO_IMPRESION = DATA_DIR / "registro_impresion.json"
ETAPAS_IMPRESION = {
    "CONEXION": "Fallo de conexión: RawBT no se abrió o no hay impresora emparejada/configurada",
    "ENVIO_ESCPOS": "Fallo de envío ESC/POS: RawBT se abrió pero el binario no llegó a la impresora",
    "IMPRESORA": "Fallo de la impresora: el dato llegó pero no imprimió (papel, energía, cabezal o Bluetooth)",
    "IMPRESION_OK": "Impresión correcta",
    "INTENTO": "Intento de impresión en espera de confirmación del operador",
    "REINTENTO": "Reintento de impresión desde la bitácora",
}


def cargar_registro_impresion():
    """Lee la bitácora de impresión (lista de eventos, más recientes primero)."""
    try:
        datos = json.loads(ARCHIVO_REGISTRO_IMPRESION.read_text(encoding="utf-8"))
        return datos if isinstance(datos, list) else []
    except (OSError, json.JSONDecodeError):
        return []


def registrar_evento_impresion(etapa, detalle, resultado="FALLO", datos=None):
    """Guarda un evento de impresión clasificado (conexión, envío ESC/POS o impresora).
    Con «datos» conserva el binario ESC/POS y el folio para poder reintentar la
    impresión sin generar cobros, folios ni tickets duplicados."""
    eventos = cargar_registro_impresion()
    evento = {
        "id_impresion": uuid.uuid4().hex[:10],
        "fecha_hora": ahora_mx().strftime("%Y-%m-%d %H:%M:%S"),
        "usuario": st.session_state.get("usuario_autenticado") or "Sesión local",
        "etapa": etapa,
        "resultado": resultado,
        "diagnóstico": ETAPAS_IMPRESION.get(etapa, "Evento de impresión"),
        "detalle": str(detalle)[:300],
    }
    if datos:
        evento["folio"] = str(datos.get("folio", "") or "")
        evento["tipo_documento"] = str(datos.get("tipo_documento", "") or "BOLETO")
        if datos.get("payload_b64"):
            evento["payload_b64"] = datos["payload_b64"]
    eventos.insert(0, evento)
    try:
        ARCHIVO_REGISTRO_IMPRESION.write_text(
            json.dumps(eventos[:300], ensure_ascii=False, indent=2), encoding="utf-8"
        )
    except OSError:
        pass
    return eventos[:300]


def codigo_sucursal(sucursal):
    """Código estable y exclusivo para que los folios no choquen entre sucursales."""
    nombre = str(sucursal or "").strip().upper()
    if nombre in CODIGOS_SUCURSAL:
        return CODIGOS_SUCURSAL[nombre]
    palabras = re.findall(r"[A-ZÁÉÍÓÚÜÑ0-9]+", nombre)
    return ("".join(p[0] for p in palabras)[:4] or "CP")


def _completar_credenciales_lider():
    """Toda sucursal activa debe tener credencial de Líder. Las que falten (p. ej. sucursales
    agregadas al catálogo sin token) se crean una sola vez y se guardan en la configuración."""
    faltan = [s for s in PRECIOS_SUCURSALES if not any(d.get("sucursal") == s for d in ACCESOS.values())]
    if not faltan:
        return
    cfg_t = cargar_sucursales_config()
    for suc in faltan:
        existente = next(((t, d) for t, d in cfg_t.get("tokens_extra", {}).items() if d.get("sucursal") == suc), None)
        if existente:
            ACCESOS[existente[0]] = dict(existente[1])
            continue
        token = f"CPV-QR-{codigo_sucursal(suc)}-{uuid.uuid4().hex[:4].upper()}"
        datos = {"nombre": f"Líder {suc.title()}", "sucursal": suc, "rol": "Líder de Sucursal"}
        cfg_t.setdefault("tokens_extra", {})[token] = datos
        ACCESOS[token] = dict(datos)
    try:
        guardar_sucursales_config(cfg_t)
    except OSError:
        pass


_completar_credenciales_lider()


def siguiente_numero_usado(sucursal, registros):
    """Obtiene el mayor consecutivo de esa sucursal, incluidos folios de versiones anteriores."""
    mayor = 0
    for registro in registros:
        if registro.get("sucursal") != sucursal:
            continue
        folio = str(registro.get("id", ""))
        # La sucursal ya está filtrada; así conservamos secuencias anteriores aunque
        # haya cambiado el prefijo histórico del folio.
        coincidencia = re.search(r"-(\d+)$", folio)
        if coincidencia:
            try:
                mayor = max(mayor, int(coincidencia.group(1)))
            except ValueError:
                pass
        try:
            mayor = max(mayor, int(registro.get("folio_numero", 0) or 0))
        except (ValueError, TypeError):
            pass
    return mayor


def asignar_folio(sucursal, registros):
    config = cargar_config()
    configurado = int(config["siguiente_folio"].get(sucursal, 1))
    minimo_seguro = siguiente_numero_usado(sucursal, registros) + 1
    numero = max(configurado, minimo_seguro)
    folio = f"CP-{codigo_sucursal(sucursal)}-{numero:06d}"
    config["siguiente_folio"][sucursal] = numero + 1
    guardar_config(config)
    return folio, numero


def qr_payload(registro):
    return f"CLUBPARK|VEHICULO|{registro.get('id', '')}"


def generar_qr_bytes(registro):
    if not QR_OK:
        return None
    imagen = qrcode.make(qr_payload(registro))
    salida = io.BytesIO()
    imagen.save(salida, format="PNG")
    return salida.getvalue()


def extraer_folio_qr(valor):
    valor = (valor or "").strip()
    if valor.startswith("CLUBPARK|VEHICULO|"):
        return valor.split("|", 2)[-1].strip()
    return valor


def crear_partes_ticket(registro):
    """Devuelve el boleto del cliente y la segunda impresión de contraseñas."""
    ticket = registro.get("ticket_texto") or crear_ticket(registro)
    corte = "=" * 32 + "\n           [ CORTE ]\n" + "=" * 32 + "\n"
    partes = ticket.split(corte, 1)
    if len(partes) == 2:
        primera = partes[0] + "\n" + corte.rstrip() + "\n"
        segunda = partes[1]
    else:
        primera = ticket
        segunda = (
            "=" * 32 + "\n      BOLETO SIN VALOR LLAVERO\n"
            f"CLUB PARK VALET Y ESTACIONAMIENTOS\nFOLIO: {registro.get('id', '')}\n"
            f"MARCA: {registro.get('marca', '')}\n"
            f"PLACA: {registro.get('placas', '')} | COLOR: {registro.get('color', '')}\n"
            "UBICACIÓN: _____________________\n"
            + "=" * 32 + "\n          CONTRASEÑAS\n"
            "     BOLETO SIN VALOR PARABRISAS\n"
            f"CLUB PARK VALET Y ESTACIONAMIENTOS\nFOLIO: {registro.get('id', '')}\n"
            + "=" * 32 + "\n"
        )
    return primera, segunda



def vista_impresion_con_qr(registro, parte="cliente", key="print_qr"):
    """Vista imprimible del boleto con QR y soporte térmico directo (sin bloqueador de ventanas)."""
    primera, segunda = crear_partes_ticket(registro)
    texto = primera if parte == "cliente" else segunda
    qr_data = generar_qr_bytes(registro) if QR_OK else None
    qr_b64 = base64.b64encode(qr_data).decode("ascii") if qr_data else ""
    qr_img = (
        f'<img alt="QR del vehículo" style="width:130px;height:130px;margin:8px auto;display:block;image-rendering:pixelated" src="data:image/png;base64,{qr_b64}">'
        if qr_data and parte == "cliente" else ""
    )
    seguro = html.escape(texto)
    if parte == "cliente" and qr_img:
        marca_corte = "[ CORTE ]"
        if marca_corte in seguro:
            antes, despues = seguro.split(marca_corte, 1)
            seguro = antes + '<br><center>' + qr_img + '</center><center><b>ESCANEE PARA REGISTRAR LA SALIDA</b></center><br>' + marca_corte + despues
        else:
            seguro += '<br><center>' + qr_img + '</center><center><b>ESCANEE PARA REGISTRAR LA SALIDA</b></center>'

    # Base64 para RawBT (impresión Bluetooth instantánea en Android)
    texto_rawbt_b64 = base64.b64encode(texto.encode("utf-8")).decode("ascii")
    rawbt_url = f"rawbt:data:text/plain;base64,{texto_rawbt_b64}"
    safe_key = re.sub(r"[^a-zA-Z0-9_-]", "_", str(key))
    titulo_btn = 'Boleto del cliente + QR de salida' if parte == 'cliente' else 'Llavero y contraseña parabrisas'

    html_doc = f"""
    <div style="font-family:-apple-system,BlinkMacSystemFont,'Segoe UI',Roboto,sans-serif;color:#1A2B20;margin:6px 0">
      <div style="display:flex;gap:8px;flex-wrap:wrap;align-items:center">
        <!-- Botón 1: Impresión directa térmica (sin ventanas emergentes bloqueadas) -->
        <button id="btn_print_{safe_key}" type="button" style="padding:10px 16px;font-weight:700;border:1px solid #072617;border-radius:8px;background:#0B3B24;color:#FFFFFF;cursor:pointer;display:inline-flex;align-items:center;gap:6px">
          🖨️ Imprimir {titulo_btn}
        </button>
        <!-- Botón 2: Enlace directo RawBT para impresoras Bluetooth Android -->
        <a href="{rawbt_url}" target="_top" style="text-decoration:none;padding:10px 14px;font-weight:700;border:1px solid #5B8F6E;border-radius:8px;background:#FAF8F5;color:#0B3B24;display:inline-flex;align-items:center;gap:6px;font-size:13px">
          📱 Enviar a Mini Impresora Bluetooth (Android)
        </a>
      </div>
      <div id="status_{safe_key}" style="margin-top:6px;font-size:12px;color:#4A6B56">
        El botón verde usa el diálogo de impresión del navegador (impresoras de red o servicio Mopria). Para una impresora térmica <b>Bluetooth</b> como la MP210 (58 mm), usa los botones RawBT del panel de impresión Bluetooth de abajo.
      </div>
      <div id="ticket_html_{safe_key}" style="display:none">{seguro}</div>
    </div>
    <script>
      (function() {{
        const btn = document.getElementById("btn_print_{safe_key}");
        const status = document.getElementById("status_{safe_key}");
        btn.addEventListener("click", function() {{
          try {{
            const contenido = document.getElementById("ticket_html_{safe_key}").innerHTML;
            let iframe = document.getElementById("iframe_print_{safe_key}");
            if (!iframe) {{
              iframe = document.createElement("iframe");
              iframe.id = "iframe_print_{safe_key}";
              iframe.style.position = "fixed";
              iframe.style.right = "0";
              iframe.style.bottom = "0";
              iframe.style.width = "0";
              iframe.style.height = "0";
              iframe.style.border = "0";
              document.body.appendChild(iframe);
            }}
            const doc = iframe.contentWindow.document;
            doc.open();
            doc.write('<!DOCTYPE html><html><head><title>Ticket Club Park</title><style>@page{{size:58mm auto;margin:1.5mm}}body{{font-family:monospace;font-size:9.5pt;line-height:1.25;white-space:pre-wrap;width:54mm;margin:0;padding:2px;color:#000}}img{{display:block;margin:6px auto;max-width:130px}}center{{text-align:center}}</style></head><body>' + contenido + '</body></html>');
            doc.close();
            setTimeout(() => {{
              try {{
                iframe.contentWindow.focus();
                iframe.contentWindow.print();
                status.textContent = "Diálogo de impresión activado correctamente.";
              }} catch(err) {{
                // Fallback por si el navegador bloquea print() en iframe
                const w = window.open("", "_blank", "width=420,height=650");
                if (w) {{
                  w.document.write('<!DOCTYPE html><html><head><title>Ticket Club Park</title><style>@page{{size:58mm auto;margin:2mm}}body{{font-family:monospace;font-size:10pt;white-space:pre-wrap;width:54mm}}</style></head><body>' + contenido + '</body></html>');
                  w.document.close();
                  w.focus();
                  setTimeout(() => w.print(), 350);
                }} else {{
                  status.textContent = "El navegador bloqueó la impresión. Permite ventanas emergentes para este sitio.";
                }}
              }}
            }}, 400);
          }} catch(e) {{
            status.textContent = "Error al preparar impresión: " + (e.message || e);
          }}
        }});
      }})();
    </script>
    """
    st.components.v1.html(html_doc, height=105)


def texto_estancia_pago(pago):
    """'1 h 8 min' a partir de los minutos guardados en el cobro; vacío si el pago no los trae."""
    try:
        minutos = int(pago.get("minutos_estancia"))
    except (TypeError, ValueError):
        return ""
    return f"{minutos // 60} h {minutos % 60} min"


def vista_comprobante_cobro(pago, registro, key="imprimir_comprobante_cobro"):
    """Muestra un comprobante imprimible del cobro ya registrado con estilo institucional."""
    datos = {
        "Sucursal": registro.get("sucursal", ""),
        "Folio": registro.get("id", ""),
        "Vehículo": f"{registro.get('marca', '')} · {registro.get('placas', '')}",
        "Fecha y hora": pago.get("fecha_hora", ""),
        "Importe": f"${dinero(pago.get('importe')):,.2f} MXN",
        **({"Tiempo de estancia": texto_estancia_pago(pago)} if texto_estancia_pago(pago) else {}),
        "Método": pago.get("metodo", ""),
        "Estado": pago.get("estado", ""),
        "Referencia": pago.get("referencia", "") or "No aplica",
        "Atendió": pago.get("usuario", ""),
    }
    filas = "".join(
        f"<tr><td style='padding:3px 0;font-weight:bold;width:40%'>{html.escape(str(etiqueta))}</td><td style='padding:3px 0'>{html.escape(str(valor))}</td></tr>"
        for etiqueta, valor in datos.items()
    )
    safe_key = re.sub(r"[^a-zA-Z0-9_-]", "_", str(key))
    html_doc = f"""
    <div style="font-family:-apple-system,BlinkMacSystemFont,'Segoe UI',Roboto,sans-serif;color:#1A2B20;max-width:440px">
      <button id="{safe_key}" style="padding:10px 16px;font-weight:700;border:1px solid #072617;border-radius:8px;background:#0B3B24;color:#FFFFFF;cursor:pointer">
        🧾 Imprimir comprobante de cobro oficial
      </button>
      <div id="contenido_{safe_key}" style="display:none">
        <div style="width:72mm;padding:3mm;font-family:sans-serif;color:#000">
          <center>
            <b style="font-size:12pt">CLUB PARK</b><br>
            <span style="font-size:9pt">VALET Y ESTACIONAMIENTOS S.A. DE C.V.</span><br>
            <b style="font-size:10pt">COMPROBANTE DE COBRO OFICIAL</b>
          </center>
          <hr style="border:0;border-top:1px dashed #000;margin:6px 0">
          <table style="width:100%;font-size:9.5pt;border-collapse:collapse">{filas}</table>
          <hr style="border:0;border-top:1px dashed #000;margin:6px 0">
          <center style="font-size:8.5pt">Conserve este comprobante para cualquier aclaración.<br>Gracias por su preferencia.</center>
        </div>
      </div>
      <script>
      (function() {{
        const btn = document.getElementById("{safe_key}");
        btn.addEventListener("click", function() {{
          const contenido = document.getElementById("contenido_{safe_key}").innerHTML;
          let iframe = document.getElementById("iframe_{safe_key}");
          if (!iframe) {{
            iframe = document.createElement("iframe");
            iframe.id = "iframe_{safe_key}";
            iframe.style.position = "fixed";
            iframe.style.right = "0";
            iframe.style.bottom = "0";
            iframe.style.width = "0";
            iframe.style.height = "0";
            iframe.style.border = "0";
            document.body.appendChild(iframe);
          }}
          const doc = iframe.contentWindow.document;
          doc.open();
          doc.write('<!DOCTYPE html><html><head><title>Comprobante Club Park</title><style>@page{{size:80mm auto;margin:3mm}}body{{font-family:sans-serif;margin:0;padding:2px}}</style></head><body>' + contenido + '</body></html>');
          doc.close();
          setTimeout(() => {{
            try {{ iframe.contentWindow.focus(); iframe.contentWindow.print(); }}
            catch(e) {{ window.print(); }}
          }}, 350);
        }});
      }})();
      </script>
    </div>
    """
    st.components.v1.html(html_doc, height=75)


def imprimir_partes_ticket(registro, sufijo):
    primera, segunda = crear_partes_ticket(registro)
    folio = registro.get("id", "")
    st.markdown("**Impresión 1: Boleto del cliente**")
    st.code(primera, language="text")
    vista_impresion_con_qr(registro, parte="cliente", key=f"{sufijo}_cliente_directo")
    imprimir_bluetooth(primera, f"{sufijo}_cliente", qr_payload=qr_payload(registro),
                       contexto={"folio": folio, "tipo_documento": "BOLETO DEL CLIENTE"})

    st.markdown("**Impresión 2: Llavero y contraseña de parabrisas**")
    st.code(segunda, language="text")
    vista_impresion_con_qr(registro, parte="llavero", key=f"{sufijo}_contrasena_directo")
    imprimir_bluetooth(segunda, f"{sufijo}_contrasena",
                       contexto={"folio": folio, "tipo_documento": "LLAVERO Y CONTRASEÑA"})


def construir_payload_escpos(texto, qr_payload_txt=None):
    """Construye el binario ESC/POS completo: inicialización, texto, código QR y avance.
    El mismo binario se envía por RawBT (Bluetooth clásico SPP en Android) o por Web Bluetooth (BLE)."""
    partes = bytearray()
    partes += bytes((0x1B, 0x40))  # ESC @ : inicializar impresora
    try:
        partes += str(texto).encode("cp437", errors="replace")
    except Exception:
        partes += str(texto).encode("utf-8", errors="replace")
    partes += b"\n\n\n"
    if qr_payload_txt:
        datos_qr = str(qr_payload_txt).encode("utf-8")
        longitud = len(datos_qr) + 3
        pL, pH = longitud & 0xFF, (longitud >> 8) & 0xFF
        partes += bytes((0x1B, 0x61, 0x01))  # ESC a 1 : centrado
        partes += bytes((0x1D, 0x28, 0x6B, 0x04, 0x00, 0x31, 0x41, 0x32, 0x00))  # modelo QR
        partes += bytes((0x1D, 0x28, 0x6B, 0x03, 0x00, 0x31, 0x43, 0x06))  # tamaño módulo
        partes += bytes((0x1D, 0x28, 0x6B, 0x03, 0x00, 0x31, 0x45, 0x31))  # nivel corrección M
        partes += bytes((0x1D, 0x28, 0x6B, pL, pH, 0x31, 0x50, 0x30)) + datos_qr  # datos
        partes += bytes((0x1D, 0x28, 0x6B, 0x03, 0x00, 0x31, 0x51, 0x30))  # imprimir QR
        partes += bytes((0x1B, 0x61, 0x00))  # ESC a 0 : alineación izquierda
    partes += bytes((0x0A, 0x0A, 0x0A, 0x0A))  # avance de papel
    return bytes(partes)


def imprimir_bluetooth(texto, key_suffix, qr_payload=None, contexto=None):
    """Impresión térmica por Bluetooth: binario ESC/POS completo (texto + QR) para
    impresoras Bluetooth clásico SPP vía RawBT, botón de compartir, y Web Bluetooth BLE."""
    texto_json = json.dumps(str(texto), ensure_ascii=False)
    qr_json = json.dumps(str(qr_payload or ""), ensure_ascii=False)
    safe_key = re.sub(r"[^a-zA-Z0-9_-]", "_", str(key_suffix))

    payload_binario = construir_payload_escpos(texto, qr_payload)
    # Formato oficial RawBT: "rawbt:base64,<datos>" (sin MIME). El binario ESC/POS
    # llega intacto a la impresora Bluetooth clasica (SPP), incluido el QR.
    payload_b64 = base64.b64encode(payload_binario).decode("ascii")
    texto_b64 = base64.b64encode(str(texto).encode("utf-8")).decode("ascii")

    if contexto:
        try:
            registrar_evento_impresion(
                "INTENTO",
                f"{contexto.get('tipo_documento', '')} · folio {contexto.get('folio', '—')} · en espera de confirmación del operador",
                "PENDIENTE",
                datos={
                    "folio": str(contexto.get("folio", "")),
                    "tipo_documento": str(contexto.get("tipo_documento", "")),
                    "payload_b64": payload_b64,
                },
            )
        except Exception:
            pass

    html_doc = f"""
    <div style="font-family:-apple-system,BlinkMacSystemFont,'Segoe UI',Roboto,sans-serif;border:1px solid #C4D7CC;background:#F9FBF9;border-radius:10px;padding:12px;margin:8px 0">
      <div style="display:flex;gap:8px;flex-wrap:wrap">
        <a href="rawbt:base64,{payload_b64}" target="_top"
           style="flex:1;min-width:200px;text-decoration:none;background:#0B3B24;color:#FFFFFF;padding:12px;border:1px solid #072617;border-radius:8px;font-weight:700;display:inline-flex;align-items:center;justify-content:center">
          🖨️ Imprimir por Bluetooth (RawBT - Android)
        </a>
        <a href="rawbt:data:text/plain;base64,{texto_b64}" target="_top"
           style="text-decoration:none;background:#FAF8F5;color:#0B3B24;padding:12px;border:1px solid #5B8F6E;border-radius:8px;font-weight:700;display:inline-flex;align-items:center;justify-content:center">
          📄 RawBT solo texto
        </a>
        <button id="cp_share_{safe_key}" type="button"
           style="background:#FAF8F5;color:#0B3B24;padding:12px;border:1px solid #5B8F6E;border-radius:8px;font-weight:700;cursor:pointer">
          📤 Compartir / Copiar
        </button>
      </div>
      <div style="display:flex;gap:8px;flex-wrap:wrap;margin-top:8px">
        <button id="cp_bt_{safe_key}" type="button"
           style="background:#145A32;color:#FFFFFF;padding:11px;border:1px solid #072617;border-radius:8px;font-weight:700;cursor:pointer">
          📶 Conectar e imprimir por Bluetooth BLE (Chrome)
        </button>
        <button id="cp_dl_{safe_key}" type="button"
           style="background:#FAF8F5;color:#0B3B24;padding:11px;border:1px solid #5B8F6E;border-radius:8px;font-weight:700;cursor:pointer">
          ⬇️ Descargar boleto .bin (abrir con RawBT)
        </button>
      </div>
      <div id="cp_status_{safe_key}" role="status" aria-live="polite" style="font-size:12px;line-height:1.45;margin-top:10px;color:#2B4A37">
        <b>Impresora térmica Bluetooth (p. ej. MP210, 58 mm):</b> 1) Empareja la impresora en <b>Ajustes › Bluetooth</b> de Android. 2) Instala la app <b>RawBT</b>: <a href="https://play.google.com/store/apps/details?id=ru.a402d.rawbtprinter" target="_top" style="color:#0B3B24;font-weight:700">📲 Instalar RawBT</a>. 3) En RawBT, Ajustes ※ <b>Default printer</b>, elige la MP210. 4) Regresa aquí y pulsa <b>🖨🸏 Imprimir por Bluetooth (RawBT)</b>: se abre RawBT y envía el boleto completo con su QR. Si el botón no abre nada, RawBT no está instalada o no tiene impresora configurada: usa <b>⬇🸏 Descargar boleto .bin</b> y ábrelo o compártelo a RawBT. Impresoras solo BLE: botón BLE en Chrome.
      </div>
    </div>
    <script>
    (() => {{
      const ticket = {texto_json};
      const qr = {qr_json};
      const btnBle = document.getElementById('cp_bt_{safe_key}');
      const btnShare = document.getElementById('cp_share_{safe_key}');
      const status = document.getElementById('cp_status_{safe_key}');
      const btnDl = document.getElementById('cp_dl_{safe_key}');
      const say = t => status.innerHTML = t;
      const enc = new TextEncoder();
      const concat = (...parts) => {{ let n=parts.reduce((a,p)=>a+p.length,0), o=new Uint8Array(n), k=0; for(const p of parts){{o.set(p,k);k+=p.length;}} return o; }};
      const cmd = (...a) => new Uint8Array(a);
      function qrCommands(data) {{
        if (!data) return new Uint8Array();
        const b=enc.encode(data), len=b.length+3, pL=len&255, pH=(len>>8)&255;
        return concat(
          cmd(0x1B,0x61,0x01),
          cmd(0x1D,0x28,0x6B,0x04,0x00,0x31,0x41,0x32,0x00),
          cmd(0x1D,0x28,0x6B,0x03,0x00,0x31,0x43,0x06),
          cmd(0x1D,0x28,0x6B,0x03,0x00,0x31,0x45,0x31),
          cmd(0x1D,0x28,0x6B,pL,pH,0x31,0x50,0x30), b,
          cmd(0x1D,0x28,0x6B,0x03,0x00,0x31,0x51,0x30),
          cmd(0x1B,0x61,0x00)
        );
      }}
      async function sendChunks(ch, data) {{
        const max = 160;
        for(let i=0;i<data.length;i+=max) {{
          const chunk=data.slice(i,Math.min(i+max,data.length));
          if(ch.writeValueWithoutResponse) await ch.writeValueWithoutResponse(chunk);
          else if(ch.writeValueWithResponse) await ch.writeValueWithResponse(chunk);
          else await ch.writeValue(chunk);
          await new Promise(r=>setTimeout(r,15));
        }}
      }}
      btnShare.addEventListener('click', async () => {{
        try {{
          if (navigator.share) {{
            await navigator.share({{ title: 'Boleto Club Park', text: ticket }});
            say('<b>Boleto enviado al menú de compartir.</b> Elige RawBT (o tu app de impresión) para mandarlo a la impresora.');
          }} else {{
            await navigator.clipboard.writeText(ticket);
            say('<b>Boleto copiado al portapapeles.</b> Pégalo en RawBT o en la app de tu impresora.');
          }}
        }} catch(e) {{ say('No se pudo compartir: ' + (e.message || e)); }}
      }});
      btnDl.addEventListener('click', () => {{
        try {{
          const bin = Uint8Array.from(atob('{payload_b64}'), c => c.charCodeAt(0));
          const blob = new Blob([bin], {{type: 'application/octet-stream'}});
          const a = document.createElement('a');
          a.href = URL.createObjectURL(blob);
          a.download = 'boleto_clubpark_escpos.bin';
          document.body.appendChild(a); a.click(); a.remove();
          say('<b>Boleto descargado.</b> Ábrelo desde tus descargas con RawBT (o compártelo a RawBT) para imprimirlo por Bluetooth.');
        }} catch(e) {{ say('No se pudo descargar el binario: ' + (e.message || e)); }}
      }});
      btnBle.addEventListener('click', async () => {{
        if (!navigator.bluetooth) {{
          say('<span style="color:#B3261E">Este navegador no soporta Web Bluetooth (requiere Chrome en Android o desktop). Usa el botón verde <b>RawBT</b>, que funciona con TODAS las impresoras Bluetooth clásicas de Android.</span>');
          return;
        }}
        btnBle.disabled=true;
        try {{
          say('Selecciona tu impresora térmica en el cuadro de diálogo…');
          const device=await navigator.bluetooth.requestDevice({{
            acceptAllDevices: true,
            optionalServices: [
              '000018f0-0000-1000-8000-00805f9b34fb',
              '0000ffe0-0000-1000-8000-00805f9b34fb',
              '0000ff00-0000-1000-8000-00805f9b34fb',
              '0000ae00-0000-1000-8000-00805f9b34fb',
              '49535343-fe7d-4ae5-8fa9-9fafd205e455',
              'e7810a71-73ae-499d-8c15-faa9aef0c3f2',
              '0000af30-0000-1000-8000-00805f9b34fb'
            ]
          }});
          const nombre = device.name || 'impresora térmica';
          say('Conectando con <b>' + nombre + '</b>…');
          if (!device.gatt) throw new Error('El dispositivo no expone conexión GATT (es Bluetooth clásico). Usa el botón verde RawBT.');
          const server=await device.gatt.connect();
          const services=await server.getPrimaryServices();
          let characteristic=null;
          for(const service of services) {{
            try {{
              const chars=await service.getCharacteristics();
              characteristic=chars.find(c=>c.properties.writeWithoutResponse||c.properties.write||c.properties.writeWithResponse);
              if(characteristic) break;
            }} catch(err){{}}
          }}
          if(!characteristic) throw new Error('No se encontró canal de escritura ESC/POS en ' + nombre + '. Usa RawBT.');
          say('Enviando boleto y QR a <b>' + nombre + '</b>…');
          const init=cmd(0x1B,0x40);
          const body=enc.encode(ticket+'\n');
          const feed=cmd(0x0A,0x0A,0x0A);
          await sendChunks(characteristic,init);
          await sendChunks(characteristic,body);
          if(qr) await sendChunks(characteristic,qrCommands(qr));
          await sendChunks(characteristic,feed);
          say('<span style="color:#1E4620"><b>✅ Impresión enviada a ' + nombre + ' con éxito.</b></span>');
          try {{ device.gatt.disconnect(); }} catch(e){{}}
        }} catch(e) {{
          say('<span style="color:#B3261E">Aviso: ' + (e.message || e) + '. Si la impresora se empareja por los Ajustes de Android, es Bluetooth clásico y DEBES usar el botón verde <b>RawBT</b>.</span>');
        }} finally {{ btnBle.disabled=false; }}
      }});
    }})();
    </script>
    """
    st.components.v1.html(html_doc, height=310)


def generar_ticket_prueba_impresion():
    """Ticket corto de verificación: valida que la impresora reciba y procese ESC/POS + QR."""
    ahora = ahora_mx()
    lineas = [
        "       PRUEBA DE IMPRESIÓN",
        " CLUB PARK VALET Y ESTACIONAMIENTOS",
        "  SISTEMA INSTITUCIONAL DE CONTROL",
        "------------------------------------",
        f" FECHA  : {ahora.strftime('%d/%m/%Y %H:%M:%S')}",
        f" USUARIO: {st.session_state.get('usuario_autenticado') or 'Sesión local'}",
        " PRUEBA : BINARIO ESC/POS + QR",
        " Si lees este texto, la impresora",
        " recibe y procesa los datos.",
        "------------------------------------",
    ]
    return "\n".join(lineas), f"CLUBPARK|PRUEBA|{ahora:%Y%m%d%H%M%S}"


def panel_prueba_impresion(key_sufijo="diag", expandido=True):
    """Prueba de impresión con registro de errores RawBT: clasifica los fallos en
    conexión, envío ESC/POS o impresora y los guarda en la bitácora del sistema."""
    with st.expander("🧪 Prueba de impresión y registro de errores RawBT", expanded=expandido):
        texto_prueba, qr_prueba = generar_ticket_prueba_impresion()
        st.code(texto_prueba)
        imprimir_bluetooth(texto_prueba, key_suffix=f"prueba_{key_sufijo}", qr_payload=qr_prueba,
                           contexto={"folio": "PRUEBA", "tipo_documento": "TICKET DE PRUEBA DE IMPRESIÓN"})
        st.markdown("**¿Qué resultado obtuviste al pulsar el botón RawBT?**")
        resultado = st.radio(
            "Resultado de la prueba de impresión",
            [
                "✅ Imprimió correctamente",
                "🔌 No se abrió RawBT / no hay impresora (fallo de conexión)",
                "📡 RawBT se abrió pero el binario no llegó (fallo de envío ESC/POS)",
                "🖨️ El dato llegó pero no imprimió (fallo de la impresora)",
            ],
            index=None,
            key=f"res_prueba_{key_sufijo}",
            label_visibility="collapsed",
        )
        detalle_extra = st.text_input(
            "Detalle opcional (modelo de impresora, mensaje visto, paso donde se detuvo)",
            key=f"det_prueba_{key_sufijo}",
        )
        if resultado is None:
            st.caption("Selecciona el resultado y regístralo: el sistema distingue si los fallos son de conexión, envío o impresora.")
        else:
            if "Imprimió correctamente" in resultado:
                etapa = "IMPRESION_OK"
            elif "No se abrió RawBT" in resultado:
                etapa = "CONEXION"
            elif "no llegó" in resultado:
                etapa = "ENVIO_ESCPOS"
            else:
                etapa = "IMPRESORA"
            if st.button("📋 Registrar resultado en la bitácora de impresión",
                         key=f"reg_prueba_{key_sufijo}", use_container_width=True):
                eventos = registrar_evento_impresion(etapa, detalle_extra or resultado,
                                                     "OK" if etapa == "IMPRESION_OK" else "FALLO")
                st.success(f"Registrado: {ETAPAS_IMPRESION[etapa]}")
                st.dataframe(pd.DataFrame(eventos[:15]), use_container_width=True, hide_index=True)
        eventos_log = cargar_registro_impresion()
        if eventos_log:
            st.markdown("**Bitácora de impresión (últimos 15 eventos)**")
            st.dataframe(pd.DataFrame(eventos_log[:15]), use_container_width=True, hide_index=True)


def enviar_payload_guardado(payload_b64, folio, tipo_documento, indice):
    """Reenvía un binario ESC/POS guardado: mismo folio, sin cobros ni tickets nuevos."""
    safe_key = re.sub(r"[^a-zA-Z0-9_-]", "_", f"reintento_{indice}")
    html_doc = f"""
    <div style="font-family:-apple-system,BlinkMacSystemFont,'Segoe UI',Roboto,sans-serif;border:1px solid #5B8F6E;background:#F4F7F5;border-radius:10px;padding:12px;margin:8px 0">
      <div style="font-size:12px;color:#2B4A37;margin-bottom:8px"><b>Reintento · {tipo_documento} · folio {folio}</b> — el binario guardado se reenvía sin cambios: no se genera cobro, folio ni ticket nuevo.</div>
      <div style="display:flex;gap:8px;flex-wrap:wrap">
        <a href="rawbt:base64,{payload_b64}" target="_top" style="flex:1;min-width:220px;text-decoration:none;background:#0B3B24;color:#FFFFFF;padding:12px;border:1px solid #072617;border-radius:8px;font-weight:700;display:inline-flex;align-items:center;justify-content:center">🖨️ Reimprimir por Bluetooth (RawBT)</a>
        <button id="cp_dl_{safe_key}" type="button" style="background:#FAF8F5;color:#0B3B24;padding:12px;border:1px solid #5B8F6E;border-radius:8px;font-weight:700;cursor:pointer">⬇️ Descargar .bin (abrir con RawBT)</button>
      </div>
    </div>
    <script>
    (() => {{
      const btnDl = document.getElementById('cp_dl_{safe_key}');
      btnDl.addEventListener('click', () => {{
        try {{
          const bin = Uint8Array.from(atob('{payload_b64}'), c => c.charCodeAt(0));
          const blob = new Blob([bin], {{type: 'application/octet-stream'}});
          const a = document.createElement('a');
          a.href = URL.createObjectURL(blob);
          a.download = 'reintento_folio_{folio}_escpos.bin';
          document.body.appendChild(a); a.click(); a.remove();
        }} catch(e) {{}}
      }});
    }})();
    </script>
    """
    st.components.v1.html(html_doc, height=180)


def panel_reintentos_impresion(expandido=False, sufijo_panel="panel"):
    """Reintento de impresión desde la bitácora: mismo folio y mismo binario,
    sin crear cobros, folios ni tickets duplicados."""
    eventos = cargar_registro_impresion()
    reintentables = [e for e in eventos if e.get("payload_b64") and e.get("resultado") != "OK"]
    with st.expander("🔄 Reintento de impresiones fallidas (desde la bitácora)", expanded=expandido):
        if not reintentables:
            st.success("Sin impresiones pendientes de reintento.")
            return
        st.caption(
            "Cada reintento reenvía el binario ESC/POS guardado con el mismo folio. "
            "No se crean cobros, folios ni tickets nuevos. Cuando salga el papel, marca el evento como impreso."
        )
        for indice, evento in enumerate(reintentables[:8]):
            # Llave estable: depende solo del id del evento, nunca del orden de la lista ni de la corrida,
            # para que el clic no se pierda cuando la bitácora crece entre corridas.
            llave = f"reintento_{sufijo_panel}_{evento.get('id_impresion') or re.sub(r'[^0-9]', '', str(evento.get('fecha_hora', '')))}"
            with st.container(border=True):
                st.markdown(
                    f"**{evento.get('tipo_documento', 'BOLETO')} · folio {evento.get('folio', '—')} · "
                    f"{evento.get('fecha_hora', '')} · {evento.get('etapa', '')}**"
                )
                st.caption(f"Registrado por {evento.get('usuario', '—')} · {evento.get('diagnóstico', '')}")
                col_r1, col_r2 = st.columns(2)
                if col_r1.button("🔄 Reintentar impresión", key=f"reintento_enviar_{llave}", use_container_width=True, type="primary"):
                    enviar_payload_guardado(evento["payload_b64"], evento.get("folio", "—"), evento.get("tipo_documento", "BOLETO"), llave)
                    registrar_evento_impresion(
                        "REINTENTO",
                        f"Reenvío del binario guardado · folio {evento.get('folio', '—')}",
                        "PENDIENTE",
                        datos={"folio": evento.get("folio", ""), "tipo_documento": evento.get("tipo_documento", "BOLETO")},
                    )
                    st.success("Binario listo: pulsa «🖨️ Reimprimir por Bluetooth (RawBT)» del recuadro verde.")
                if col_r2.button("✅ Marcar como impreso", key=f"reintento_ok_{llave}", use_container_width=True):
                    registrar_evento_impresion(
                        "IMPRESION_OK",
                        f"Confirmado por el operador tras reintento · folio {evento.get('folio', '—')}",
                        "OK",
                        datos={"folio": evento.get("folio", ""), "tipo_documento": evento.get("tipo_documento", "BOLETO")},
                    )
                    st.rerun()


def enlace_whatsapp(telefono, mensaje):
    """Enlace wa.me con mensaje prellenado (la app original lo usaba sin definirlo)."""
    tel = "".join(c for c in str(telefono) if c.isdigit())
    if len(tel) == 10:
        tel = "52" + tel
    return f"https://wa.me/{tel}?text={urllib.parse.quote(mensaje)}"


def dinero(valor):
    try:
        return float(valor or 0)
    except (TypeError, ValueError):
        return 0.0


def normalizar_texto(valor):
    return str(valor or "").strip()


ROL_LIDER = "Líder de Sucursal"


def normalizar_rol(rol):
    """Credenciales antiguas con 'Supervisor de Sucursal' pasan al rol nuevo de Líder de Sucursal."""
    return ROL_LIDER if str(rol or "").strip() in ("Supervisor de Sucursal", "Encargado de Sucursal") else rol


def usuario_es_lider():
    return normalizar_rol(st.session_state.get("rol")) == ROL_LIDER


def usuario_es_contador():
    return "Contador" in str(st.session_state.get("rol", ""))


def usuario_es_global():
    return usuario_es_contador() or st.session_state.get("rol") == "Supervisor Global"


def vehiculos_visibles(registros):
    if usuario_es_global():
        return registros
    sucursal = st.session_state.get("sucursal_asignada")
    return [r for r in registros if r.get("sucursal") == sucursal]


def guardar_foto(uploaded_file, folio, vista):
    """Guarda una foto optimizada y devuelve su ruta relativa al directorio de datos."""
    if isinstance(uploaded_file, Image.Image):
        imagen = uploaded_file.copy()
    else:
        imagen = Image.open(uploaded_file)
    imagen = ImageOps.exif_transpose(imagen).convert("RGB")
    imagen.thumbnail((1920, 1920))
    nombre = f"{folio}_{vista}.jpg"
    ruta = PHOTO_DIR / nombre
    imagen.save(ruta, format="JPEG", quality=82, optimize=True)
    return str(ruta)


ETIQUETAS_VISTAS = [("frente", "Frente"), ("trasera", "Trasera"), ("izquierda", "Costado izq."),
                    ("derecha", "Costado der."), ("tablero", "Tablero"), ("interior", "Interior")]


def _ruta_absoluta(valor):
    try:
        p = Path(valor)
    except TypeError:
        return None
    if not p.is_absolute():
        p = DATA_DIR / p
    return p if p.exists() and p.is_file() else None


@st.cache_data(show_spinner=False, max_entries=800)
def _miniatura_bytes(ruta_txt, mtime, lado=420):
    img = Image.open(ruta_txt)
    img = ImageOps.exif_transpose(img).convert("RGB")
    img.thumbnail((lado, lado))
    buf = io.BytesIO()
    img.save(buf, format="JPEG", quality=72)
    return buf.getvalue()


@st.cache_data(show_spinner=False, max_entries=60)
def _foto_base64(ruta_txt, mtime):
    return base64.b64encode(Path(ruta_txt).read_bytes()).decode("ascii")


def fotos_evidencia_registro(registro):
    """Lista [(etiqueta, ruta)] con todas las fotos del vehículo y de pertenencias."""
    fotos = []
    for vista, etiqueta in ETIQUETAS_VISTAS:
        p = _ruta_absoluta((registro.get("fotos") or {}).get(vista, "")) if (registro.get("fotos") or {}).get(vista) else None
        if p:
            fotos.append((etiqueta, p))
    for i, extra in enumerate(registro.get("fotos_custodia") or [], start=1):
        p = _ruta_absoluta(extra)
        if p:
            fotos.append((f"Pertenencia {i}", p))
    return fotos


_HTML_VISOR = """
<div id="vz" style="position:relative;width:100%;height:__ALTO__px;background:#0d1410;border-radius:10px;overflow:hidden;touch-action:none;user-select:none">
  <img id="vzi" src="data:image/jpeg;base64,__B64__" draggable="false"
       style="position:absolute;left:0;top:0;transform-origin:0 0;will-change:transform;max-width:none">
  <div style="position:absolute;top:8px;left:8px;background:rgba(11,59,36,.85);color:#fff;font:600 13px sans-serif;padding:4px 10px;border-radius:6px">__TITULO__</div>
  <div style="position:absolute;bottom:8px;right:8px;display:flex;gap:6px">
    <button id="vzm" class="b">−</button><button id="vzp" class="b">+</button>
    <button id="vzr" class="b">⟲</button><button id="vzf" class="b">⛶</button>
  </div>
  <div id="vzz" style="position:absolute;bottom:12px;left:10px;color:#fff;font:600 12px sans-serif;background:rgba(0,0,0,.5);padding:3px 8px;border-radius:5px">100%</div>
</div>
<style>.b{width:42px;height:42px;border:0;border-radius:8px;background:rgba(255,255,255,.92);color:#0B3B24;font:700 20px sans-serif;cursor:pointer}</style>
<script>
(function(){
  const box=document.getElementById('vz'), img=document.getElementById('vzi'), lab=document.getElementById('vzz');
  let s=1, x=0, y=0, base=1, ptrs=new Map(), d0=0, s0=1, mid0=null, last=null, tap=0;
  function lim(){ const W=box.clientWidth,H=box.clientHeight,w=img.naturalWidth*base*s,h=img.naturalHeight*base*s;
    x = w<=W ? (W-w)/2 : Math.min(0, Math.max(W-w, x));
    y = h<=H ? (H-h)/2 : Math.min(0, Math.max(H-h, y)); }
  function draw(){ lim(); img.style.transform=`translate(${x}px,${y}px) scale(${base*s})`; lab.textContent=Math.round(s*100)+'%'; }
  function fit(){ base=Math.min(box.clientWidth/img.naturalWidth, box.clientHeight/img.naturalHeight); s=1; x=0; y=0; draw(); }
  function zoomAt(ns,cx,cy){ ns=Math.max(1,Math.min(8,ns)); const k=ns/s; x=cx-(cx-x)*k; y=cy-(cy-y)*k; s=ns; draw(); }
  function rel(e){ const r=box.getBoundingClientRect(); return [e.clientX-r.left, e.clientY-r.top]; }
  img.onload=fit; if(img.complete) fit(); window.addEventListener('resize', fit);
  box.addEventListener('wheel', e=>{ e.preventDefault(); const [cx,cy]=rel(e); zoomAt(s*(e.deltaY<0?1.2:1/1.2),cx,cy); }, {passive:false});
  box.addEventListener('pointerdown', e=>{ if(e.target.tagName==='BUTTON') return; box.setPointerCapture(e.pointerId); ptrs.set(e.pointerId, rel(e));
    if(ptrs.size===2){ const [a,b]=[...ptrs.values()]; d0=Math.hypot(a[0]-b[0],a[1]-b[1]); s0=s; mid0=[(a[0]+b[0])/2,(a[1]+b[1])/2]; }
    else { last=rel(e); const t=Date.now(); if(t-tap<300){ const [cx,cy]=rel(e); zoomAt(s>1.05?1:3,cx,cy); } tap=t; } });
  box.addEventListener('pointermove', e=>{ if(!ptrs.has(e.pointerId)) return; ptrs.set(e.pointerId, rel(e));
    if(ptrs.size===2){ const [a,b]=[...ptrs.values()]; const d=Math.hypot(a[0]-b[0],a[1]-b[1]); const m=[(a[0]+b[0])/2,(a[1]+b[1])/2];
      x+=m[0]-mid0[0]; y+=m[1]-mid0[1]; mid0=m; zoomAt(s0*d/d0, m[0], m[1]); }
    else if(last){ const p=rel(e); x+=p[0]-last[0]; y+=p[1]-last[1]; last=p; draw(); } });
  const fin=e=>{ ptrs.delete(e.pointerId); if(ptrs.size<2){ d0=0; } if(ptrs.size===1){ last=[...ptrs.values()][0]; } if(ptrs.size===0) last=null; };
  box.addEventListener('pointerup', fin); box.addEventListener('pointercancel', fin);
  const c=()=>[box.clientWidth/2, box.clientHeight/2];
  document.getElementById('vzp').onclick=()=>zoomAt(s*1.5,...c());
  document.getElementById('vzm').onclick=()=>zoomAt(s/1.5,...c());
  document.getElementById('vzr').onclick=fit;
  document.getElementById('vzf').onclick=()=>{ const f=box.requestFullscreen||box.webkitRequestFullscreen; if(document.fullscreenElement){document.exitFullscreen();} else if(f){ f.call(box); } };
  document.addEventListener('fullscreenchange', ()=>setTimeout(fit,150));
})();
</script>
"""


def visor_zoom_fotos(fotos, key, alto=440, titulo=""):
    """Miniaturas de todas las fotos + visor con zoom (pellizcar, doble toque, rueda, + / −, pantalla completa)."""
    fotos = [(e, p) for e, p in fotos if p]
    if not fotos:
        st.info("No hay fotografías disponibles.")
        return
    if titulo:
        st.markdown(titulo)
    cols_m = st.columns(min(len(fotos), 6))
    for i, (etq, p) in enumerate(fotos):
        try:
            cols_m[i % len(cols_m)].image(_miniatura_bytes(str(p), p.stat().st_mtime), caption=etq, use_container_width=True)
        except Exception:  # archivo dañado o incompleto: no tumba la pantalla
            cols_m[i % len(cols_m)].caption(f"⚠️ {etq}: foto dañada")
    idx = st.radio("🔍 Foto para ampliar", list(range(len(fotos))), format_func=lambda i: fotos[i][0],
                   horizontal=True, key=f"zoom_sel_{key}")
    etq, p = fotos[idx]
    html_v = (_HTML_VISOR.replace("__ALTO__", str(int(alto)))
              .replace("__B64__", _foto_base64(str(p), p.stat().st_mtime))
              .replace("__TITULO__", str(etq).replace("<", "")))
    st.components.v1.html(html_v, height=int(alto) + 6)
    st.caption("Pellizca o usa + / − para acercar, arrastra para moverte, doble toque para zoom rápido y ⛶ para pantalla completa.")


def ruta_foto(registro, vista):
    ruta = (registro.get("fotos") or {}).get(vista, "")
    if not ruta:
        return None
    p = Path(ruta)
    if not p.is_absolute():
        p = DATA_DIR / p
    return p if p.exists() else None


DIAS_TICKET = ["LUNES", "MARTES", "MIÉRCOLES", "JUEVES", "VIERNES", "SÁBADO", "DOMINGO"]


def formato_dia_hora(valor):
    """'23/10/2026 12:00:00' -> 'JUEVES 23/10/2026 12:00:00 hrs'. Si no se puede leer, devuelve el texto original."""
    texto = str(valor or "").strip()
    if not texto or texto == "PENDIENTE":
        return texto or "PENDIENTE"
    for formato in ("%d/%m/%Y %H:%M:%S", "%Y-%m-%d %H:%M:%S"):
        try:
            dt = datetime.strptime(texto, formato)
            return f"{DIAS_TICKET[dt.weekday()]} {dt:%d/%m/%Y %H:%M:%S} hrs"
        except ValueError:
            continue
    return texto


def crear_ticket(registro):
    """Crea ticket térmico con marcas temporales explícitas de entrada y salida."""
    folio = registro.get("id", "")
    sucursal = registro.get("sucursal", "")
    marca = registro.get("marca", "")
    placas = registro.get("placas", "")
    color = registro.get("color", "")
    tarifa = float(registro.get("tarifa", 0) or 0)
    entrada_dt = registro.get("fecha_entrada") or registro.get("fecha", "")
    salida_dt = registro.get("salida_fecha") or registro.get("salida_fecha_hora") or "PENDIENTE"
    custodia = registro.get("custodia", "") or "SIN OBJETOS DECLARADOS"
    linea = "=" * 32
    guion = "-" * 32
    try:
        texto_tarifa_ticket = "\n".join(lineas_tarifa_ticket(sucursal, registro.get("modalidad_tarifa")))
    except Exception:
        texto_tarifa_ticket = ""
    return f"""{linea}
 CLUB PARK VALET Y ESTACIONAMIENTOS
         S.A. DE C.V.
{linea}
SUCURSAL: {sucursal}
FOLIO TICKET: {folio}
TARIFA: ${tarifa:,.2f} MXN
{guion}
MARCA: {marca}
PLACA: {placas}  COLOR: {color}
ENTRADA:
{formato_dia_hora(entrada_dt)}
SALIDA:
{formato_dia_hora(salida_dt)}
{guion}
{texto_tarifa_ticket}
{guion}
[X] OBSERVACIONES Y CUSTODIA:
{custodia}
{guion}
       BOLETO DEL CLIENTE
{guion}
NO NOS HACEMOS RESPONSABLES POR
OBJETOS COMO SON: EFECTIVO, CHEQUES,
ARMAS, BOLSAS, TARJETAS, CELULARES,
APARATOS ELECTRONICOS.

SIN LA PRESENTACION DE ESTE TALON,
NO SE ENTREGARA EL VEHICULO.
LEER CLAUSULAS AL REVERSO.
{linea}
           [ CORTE ]
{linea}
      BOLETO SIN VALOR LLAVERO
CLUB PARK VALET Y ESTACIONAMIENTOS
FOLIO: {folio}
MARCA: {marca}
PLACA: {placas} | COLOR: {color}
UBICACIÓN: _____________________
{linea}
           [ CORTE ]
{linea}
          CONTRASEÑAS
     BOLETO SIN VALOR PARABRISAS
CLUB PARK VALET Y ESTACIONAMIENTOS
FOLIO: {folio}
{linea}
"""


def generar_excel_oficial(df):
    """Genera un archivo XLSX usando el motor disponible."""
    output = io.BytesIO()
    columnas = ["id", "fecha", "sucursal", "marca", "placas", "tarifa", "operador", "estado"]
    encabezados = ["FOLIO ID", "FECHA Y HORA", "SUCURSAL", "MARCA/MODELO", "PLACAS", "TARIFA ($)", "OPERADOR", "ESTADO"]
    datos = df.reindex(columns=columnas).copy()
    datos.columns = encabezados

    if OPENPYXL_OK:
        from openpyxl.styles import Alignment, Font, PatternFill

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Reporte Operativo"
        ws.merge_cells("A1:H1")
        ws["A1"] = "CLUB PARK VALET Y ESTACIONAMIENTOS S.A. DE C.V."
        ws["A1"].font = Font(name="Arial", size=14, bold=True, color="FFFFFF")
        ws["A1"].fill = PatternFill("solid", fgColor="1F4E78")
        ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 28

        ws.append([])
        ws.append(encabezados)
        header_fill = PatternFill("solid", fgColor="D9EAF7")
        for cell in ws[3]:
            cell.font = Font(bold=True)
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")

        for row in datos.itertuples(index=False, name=None):
            ws.append(["" if pd.isna(v) else v for v in row])

        for row_num in range(4, ws.max_row + 1):
            ws.cell(row_num, 6).number_format = '"$"#,##0.00'
        for idx, width in enumerate([17, 22, 26, 25, 16, 16, 25, 18], start=1):
            ws.column_dimensions[openpyxl.utils.get_column_letter(idx)].width = width
        ws.freeze_panes = "A4"
        ws.auto_filter.ref = f"A3:H{ws.max_row}"
        wb.save(output)
        output.seek(0)
        return output

    if XLSXWRITER_OK:
        with pd.ExcelWriter(output, engine="xlsxwriter") as writer:
            datos.to_excel(writer, sheet_name="Reporte Operativo", index=False, startrow=2)
            workbook = writer.book
            worksheet = writer.sheets["Reporte Operativo"]
            titulo = workbook.add_format({
                "bold": True, "font_size": 14, "font_color": "white",
                "bg_color": "#1F4E78", "align": "center", "valign": "vcenter"
            })
            header = workbook.add_format({"bold": True, "bg_color": "#D9EAF7", "border": 1})
            money = workbook.add_format({"num_format": '$#,##0.00'})
            worksheet.merge_range("A1:H1", "CLUB PARK VALET Y ESTACIONAMIENTOS S.A. DE C.V.", titulo)
            for col, name in enumerate(encabezados):
                worksheet.write(2, col, name, header)
            worksheet.set_column("A:A", 17)
            worksheet.set_column("B:B", 22)
            worksheet.set_column("C:D", 26)
            worksheet.set_column("E:E", 16)
            worksheet.set_column("F:F", 16, money)
            worksheet.set_column("G:G", 25)
            worksheet.set_column("H:H", 18)
            worksheet.freeze_panes(3, 0)
            if len(datos):
                worksheet.autofilter(2, 0, len(datos) + 2, len(encabezados) - 1)
        output.seek(0)
        return output

    raise RuntimeError(
        "No hay motor Excel instalado. Agrega openpyxl y xlsxwriter a requirements.txt."
    )


def leer_codigo_qr(imagen_pil):
    """Lector de QR robusto: prepara la imagen (contraste, ampliacion, binarizacion,
    inversiones y rotaciones) y prueba tres decodificadores antes de rendirse."""
    if imagen_pil is None:
        return None
    if not isinstance(imagen_pil, Image.Image):
        try:
            imagen_pil = Image.open(imagen_pil)
        except Exception:
            return None
    try:
        base = ImageOps.exif_transpose(imagen_pil).convert("RGB")
    except Exception:
        base = imagen_pil.convert("RGB")

    candidatos = [base]
    try:
        gris = ImageOps.autocontrast(base.convert("L"))
        ancho, alto = gris.size
        if max(ancho, alto) < 1100:
            factor = max(2, int(1100 / max(ancho, alto)))
            candidatos.append(gris.resize((ancho * factor, alto * factor), Image.LANCZOS))
        candidatos.append(gris)
        candidatos.append(ImageOps.invert(gris))
        candidatos.append(ImageOps.autocontrast(base))
        try:
            import cv2
            import numpy as np
            arr = np.array(gris)
            _, otsu = cv2.threshold(arr, 0, 255, cv2.THRESH_BINARY + cv2.THRESH_OTSU)
            candidatos.append(Image.fromarray(otsu))
            candidatos.append(Image.fromarray(cv2.adaptiveThreshold(
                arr, 255, cv2.ADAPTIVE_THRESH_MEAN_C, cv2.THRESH_BINARY, 31, 2)))
        except Exception:
            pass
    except Exception:
        pass

    def _texto_de(resultados):
        for r in resultados or []:
            texto = getattr(r, "text", None)
            if texto is None:
                datos = getattr(r, "data", None)
                if isinstance(datos, bytes):
                    texto = datos.decode("utf-8", "replace")
                else:
                    texto = str(datos or "")
            if str(texto).strip():
                return str(texto).strip()
        return None

    # 1) zxing-cpp con TRY_HARDER y rotaciones
    try:
        import zxingcpp
        for imagen in candidatos:
            try:
                hallado = _texto_de(zxingcpp.read_barcodes(imagen, try_harder=True, try_rotate=True))
                if hallado:
                    return hallado
            except Exception:
                pass
    except Exception:
        pass

    # 2) pyzbar con rotaciones manuales
    try:
        from pyzbar.pyzbar import decode
        for imagen in candidatos:
            for angulo in (0, 90, 180, 270):
                prueba = imagen if angulo == 0 else imagen.rotate(angulo, expand=True)
                try:
                    hallado = _texto_de(decode(prueba))
                    if hallado:
                        return hallado
                except Exception:
                    pass
    except Exception:
        pass

    # 3) detector de OpenCV
    try:
        import cv2
        import numpy as np
        detector = cv2.QRCodeDetector()
        for imagen in candidatos:
            arr = cv2.cvtColor(np.array(imagen), cv2.COLOR_RGB2BGR)
            for angulo in (0, 90, 180, 270):
                prueba = arr if angulo == 0 else np.rot90(arr, angulo // 90)
                try:
                    valor, _, _ = detector.detectAndDecode(prueba)
                    if valor and valor.strip():
                        return valor.strip()
                except Exception:
                    pass
    except Exception:
        pass
    return None

_HTML_LECTOR_QR = r'''<!doctype html>
<html><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<style>
  html,body{margin:0;padding:0;background:transparent;font-family:-apple-system,Segoe UI,Roboto,sans-serif}
  #caja{position:relative;width:100%;height:300px;background:#0d1410;border-radius:12px;overflow:hidden}
  video{width:100%;height:100%;object-fit:cover;display:block}
  #marco{position:absolute;left:50%;top:50%;width:62%;max-width:240px;aspect-ratio:1;transform:translate(-50%,-50%);
         border:3px solid rgba(255,255,255,.9);border-radius:14px;box-shadow:0 0 0 2000px rgba(0,0,0,.35)}
  #linea{position:absolute;left:8%;right:8%;height:2px;background:#4ade80;top:10%;animation:barre 1.6s ease-in-out infinite}
  @keyframes barre{0%{top:10%}50%{top:88%}100%{top:10%}}
  #estado{position:absolute;left:0;right:0;bottom:0;padding:8px 10px;color:#fff;font-size:14px;font-weight:600;
          background:linear-gradient(transparent,rgba(0,0,0,.75));text-align:center}
  .barra{display:flex;gap:8px;margin-top:8px}
  button{flex:1;height:44px;border:0;border-radius:10px;background:#0B3B24;color:#fff;font-size:15px;font-weight:700}
  button.sec{background:#EBF3EE;color:#0B3B24}
  #ok{display:none;position:absolute;inset:0;background:rgba(11,59,36,.92);color:#fff;align-items:center;justify-content:center;
      flex-direction:column;font-size:18px;font-weight:800;text-align:center;padding:10px}
</style></head>
<body>
<div id="caja">
  <video id="v" playsinline muted autoplay></video>
  <div id="marco"><div id="linea"></div></div>
  <div id="estado">Abriendo cámara…</div>
  <div id="ok"><div style="font-size:46px">✅</div><div>QR leído</div><div id="okt" style="font-size:12px;font-weight:500;margin-top:6px;word-break:break-all"></div></div>
</div>
<div class="barra">
  <button id="otro" class="sec" style="display:none">🔄 Escanear otro</button>
  <button id="luz" class="sec" style="display:none">🔦 Linterna</button>
</div>
<canvas id="c" style="display:none"></canvas>
<script>
(function(){
  const enviar = (tipo, extra) => window.parent.postMessage(Object.assign({isStreamlitMessage:true, type:tipo}, extra||{}), "*");
  enviar("streamlit:componentReady", {apiVersion:1});
  const alto = () => enviar("streamlit:setFrameHeight", {height: document.body.scrollHeight + 4});
  window.addEventListener("message", e => { if (e.data && e.data.type === "streamlit:render") alto(); });

  const v = document.getElementById("v"), c = document.getElementById("c"), ctx = c.getContext("2d", {willReadFrequently:true});
  const est = document.getElementById("estado"), ok = document.getElementById("ok");
  let flujo = null, pista = null, detector = null, corriendo = false, luzOn = false, ultimo = 0;

  function cargarJsQR(){
    return new Promise(res => {
      if (window.jsQR) return res(true);
      const fuentes = ["https://cdn.jsdelivr.net/npm/jsqr@1.4.0/dist/jsQR.js", "https://unpkg.com/jsqr@1.4.0/dist/jsQR.js"];
      const intenta = i => {
        if (i >= fuentes.length) return res(false);
        const s = document.createElement("script"); s.src = fuentes[i];
        s.onload = () => res(true); s.onerror = () => intenta(i + 1); document.head.appendChild(s);
      };
      intenta(0);
    });
  }

  async function preparar(){
    try {
      if ("BarcodeDetector" in window) {
        const f = await window.BarcodeDetector.getSupportedFormats();
        if (f.includes("qr_code")) detector = new window.BarcodeDetector({formats:["qr_code"]});
      }
    } catch (e) { detector = null; }
    if (!detector) await cargarJsQR();
  }

  function aviso(){
    try { navigator.vibrate && navigator.vibrate(120); } catch (e) {}
    try {
      const a = new (window.AudioContext || window.webkitAudioContext)(), o = a.createOscillator(), g = a.createGain();
      o.frequency.value = 1400; o.connect(g); g.connect(a.destination); g.gain.value = .15; o.start(); o.stop(a.currentTime + .12);
    } catch (e) {}
  }

  function detener(){
    corriendo = false;
    if (flujo) flujo.getTracks().forEach(t => t.stop());
    flujo = null; pista = null;
  }

  function encontrado(texto){
    detener(); aviso();
    document.getElementById("okt").textContent = texto.length > 80 ? texto.slice(0, 80) + "…" : texto;
    ok.style.display = "flex";
    document.getElementById("otro").style.display = "block";
    document.getElementById("luz").style.display = "none";
    enviar("streamlit:setComponentValue", {value: {texto: texto, t: Date.now()}, dataType: "json"});
    alto();
  }

  async function cuadro(){
    if (!corriendo) return;
    const ahora = performance.now();
    if (ahora - ultimo > 110 && v.readyState >= 2 && v.videoWidth) {
      ultimo = ahora;
      try {
        if (detector) {
          const r = await detector.detect(v);
          if (r && r.length && r[0].rawValue) return encontrado(r[0].rawValue);
        } else if (window.jsQR) {
          // zona central ampliada: más rápido y lee QR pequeños
          const lado = Math.min(v.videoWidth, v.videoHeight) * 0.8, sx = (v.videoWidth - lado) / 2, sy = (v.videoHeight - lado) / 2;
          const t = Math.min(720, Math.round(lado));
          c.width = t; c.height = t;
          ctx.drawImage(v, sx, sy, lado, lado, 0, 0, t, t);
          const d = ctx.getImageData(0, 0, t, t);
          const r = window.jsQR(d.data, t, t, {inversionAttempts: "attemptBoth"});
          if (r && r.data) return encontrado(r.data);
        }
      } catch (e) {}
    }
    requestAnimationFrame(cuadro);
  }

  async function iniciar(){
    ok.style.display = "none"; document.getElementById("otro").style.display = "none";
    est.textContent = "Abriendo cámara…";
    try {
      flujo = await navigator.mediaDevices.getUserMedia({audio:false, video:{facingMode:{ideal:"environment"},
                                                          width:{ideal:1280}, height:{ideal:720}}});
    } catch (e) {
      est.textContent = "No se pudo abrir la cámara: permite el acceso o usa «Tomar foto» abajo.";
      enviar("streamlit:setComponentValue", {value: {error: String(e && e.name || e), t: Date.now()}, dataType: "json"});
      return;
    }
    v.srcObject = flujo;
    try { await v.play(); } catch (e) {}
    pista = flujo.getVideoTracks()[0];
    try {
      const cap = pista.getCapabilities ? pista.getCapabilities() : {};
      const ajustes = [];
      if (cap.focusMode && cap.focusMode.includes("continuous")) ajustes.push({focusMode:"continuous"});
      if (cap.zoom && cap.zoom.max >= 1.5) ajustes.push({zoom: Math.min(1.5, cap.zoom.max)});
      if (ajustes.length) await pista.applyConstraints({advanced: ajustes});
      if (cap.torch) document.getElementById("luz").style.display = "block";
    } catch (e) {}
    est.textContent = detector || window.jsQR ? "Apunta al código QR dentro del cuadro" : "Lector no disponible: usa «Tomar foto» abajo";
    corriendo = true; requestAnimationFrame(cuadro); alto();
  }

  document.getElementById("otro").onclick = iniciar;
  document.getElementById("luz").onclick = async () => {
    if (!pista) return;
    luzOn = !luzOn;
    try { await pista.applyConstraints({advanced:[{torch: luzOn}]}); } catch (e) {}
  };
  window.addEventListener("pagehide", detener);
  preparar().then(iniciar);
  alto();
})();
</script>
</body></html>
'''


@st.cache_resource(show_spinner=False)
def _componente_lector_qr():
    """Lector de QR en vivo: lee directo del video de la cámara trasera (sin tomar foto).
    Usa el lector nativo del teléfono (BarcodeDetector) y, si no existe, jsQR."""
    import tempfile
    carpeta = Path(tempfile.gettempdir()) / "clubpark_lector_qr_v1"
    carpeta.mkdir(parents=True, exist_ok=True)
    (carpeta / "index.html").write_text(_HTML_LECTOR_QR, encoding="utf-8")
    return st.components.v1.declare_component("lector_qr_vivo", path=str(carpeta))


def lector_qr_camara_trasera(key="lector_qr_trasero", height=390):
    """Lectura de QR: primero el lector en vivo (instantáneo); si no funciona, foto o imagen del QR."""
    valor = None
    try:
        valor = _componente_lector_qr()(key=f"{key}_vivo", default=None)
    except Exception:
        valor = None
    if isinstance(valor, dict) and str(valor.get("texto") or "").strip():
        return str(valor["texto"]).strip()
    if isinstance(valor, dict) and valor.get("error"):
        st.warning("La cámara en vivo no abrió (permiso o navegador). Usa la foto del QR de abajo.")
    usar_foto = st.checkbox("📸 ¿No lo lee? Tomar foto del QR", key=f"{key}_usar_foto")
    if usar_foto:
        foto_qr = st.camera_input("Capturar el código QR", key=key,
                                  help="Enfoca el QR a 10-15 cm con buena luz.")
        if foto_qr is not None:
            try:
                token = leer_codigo_qr(Image.open(foto_qr))
                if token:
                    return str(token).strip()
                st.warning("No se pudo leer el QR. Acércate, mejora la luz y vuelve a tomar la foto.")
            except Exception:
                st.warning("No se pudo leer el QR. Acércate, mejora la luz y vuelve a tomar la foto.")
    subida = st.file_uploader("O sube una imagen del código QR", type=["png", "jpg", "jpeg"], key=f"{key}_subida")
    if subida is not None:
        try:
            token = leer_codigo_qr(Image.open(subida))
            return str(token).strip() if token else None
        except Exception:
            return None
    return None


def camara_trasera_foto_pertenencias(key="foto_pertenencias_trasera", height=430):
    """Captura fotográfica con la cámara del dispositivo, sin componentes externos."""
    captura = st.camera_input("Tomar fotografía", key=key)
    if captura is None:
        st.caption("Se usa la cámara trasera del teléfono. Si la cámara no abre, cambia el modo a «📁 Subir fotos».")
        return None
    try:
        return Image.open(captura).copy()
    except Exception:
        st.warning("No se pudo procesar la foto. Vuelve a tomarla.")
        return None


# ---------------------------------------------------------------------
# PRUEBAS OPERATIVAS (autodiagnóstico institucional)
# ---------------------------------------------------------------------
def ejecutar_pruebas_operativas():
    """Ejecuta una batería de pruebas sobre los componentes críticos de la app."""
    resultados = []

    def registrar(nombre, funcion):
        try:
            detalle = funcion()
            resultados.append({"prueba": nombre, "estado": "✅", "detalle": detalle or "OK"})
        except Exception as exc:
            resultados.append({"prueba": nombre, "estado": "❌", "detalle": str(exc)[:220]})

    def prueba_base_datos():
        archivo = DATA_DIR / "_prueba_operativa.json"
        datos = [{"id": "CP-TEST-000001", "sucursal": "RAMMA"}]
        with archivo.open("w", encoding="utf-8") as f:
            json.dump(datos, f, ensure_ascii=False)
        with archivo.open("r", encoding="utf-8") as f:
            leido = json.load(f)
        archivo.unlink()
        assert leido == datos
        return "Escritura y lectura JSON correctas"

    def prueba_folios():
        registros = [
            {"id": "CP-RAM-000005", "sucursal": "RAMMA", "folio_numero": 5},
            {"id": "CP-SUS-000002", "sucursal": "SUSHITTO", "folio_numero": 2},
        ]
        folio1, num1 = asignar_folio("RAMMA", registros)
        assert folio1.startswith("CP-RAM-") and num1 >= 6, f"folio inesperado {folio1}"
        folio2, num2 = asignar_folio("SUSHITTO", registros)
        assert folio2.startswith("CP-SUS-") and num2 >= 3
        config = cargar_config()
        config["siguiente_folio"]["RAMMA"] = 6
        config["siguiente_folio"]["SUSHITTO"] = 3
        guardar_config(config)
        return f"Secuencias independientes por sucursal ({folio1} / {folio2})"

    def prueba_registro():
        """Prueba de registro: simula el alta completa de un vehículo (folio, ticket y foto)."""
        config_previa = cargar_config()
        folio_previo = int(config_previa["siguiente_folio"].get("RAMMA", 1))
        registros_sim = [{"id": "CP-RAM-000005", "sucursal": "RAMMA", "folio_numero": 5}]
        folio, numero = asignar_folio("RAMMA", registros_sim)
        # Restaura el contador real para no consumir folios con la prueba
        config_post = cargar_config()
        config_post["siguiente_folio"]["RAMMA"] = folio_previo
        guardar_config(config_post)
        assert re.match(r"^CP-[A-Z0-9]+-\d{6}$", folio), f"Formato de folio inválido: {folio}"
        ahora = ahora_mx()
        registro = {
            "id": folio,
            "folio_numero": numero,
            "qr_payload": f"CLUBPARK|VEHICULO|{folio}",
            "fecha": ahora.strftime("%d/%m/%Y %H:%M:%S"),
            "sucursal": "RAMMA",
            "marca": "Vehículo de prueba", "color": "Verde Botella", "placas": "PRB-000-0",
            "tarifa": 70, "operador": "Prueba operativa",
            "custodia": "SIN OBJETOS DE VALOR DECLARADOS",
            "fotos": {},
        }
        registro["ticket_texto"] = crear_ticket(registro)
        assert folio in registro["ticket_texto"], "El ticket no contiene el folio asignado"
        primera, segunda = crear_partes_ticket(registro)
        assert folio in primera and folio in segunda, "El folio falta en el boleto o en el llavero"
        imagen_prueba = Image.new("RGB", (320, 240), (11, 59, 36))
        ruta_foto_prueba = guardar_foto(imagen_prueba, folio, "prueba_registro")
        assert Path(ruta_foto_prueba).exists(), "La foto del registro no se guardó"
        Path(ruta_foto_prueba).unlink()
        return f"Registro simulado válido ({folio}): folio, boleto, llavero y foto correctos"

    def prueba_qr():
        if not QR_OK:
            raise RuntimeError("Falta el paquete qrcode[pil] en el servidor: actualiza requirements.txt en el repositorio y reinicia la app (Manage app → Reboot)")
        registro = {"id": "CP-RAM-000099"}
        png = generar_qr_bytes(registro)
        assert png and png[:8] == b"\x89PNG\r\n\x1a\n", "El QR no es un PNG válido"
        texto = None
        try:
            import zxingcpp
            decodificados = zxingcpp.read_barcodes(Image.open(io.BytesIO(png)))
            texto = decodificados[0].text if decodificados else None
        except Exception:
            try:
                from pyzbar.pyzbar import decode as decode_zbar
                detectados = decode_zbar(Image.open(io.BytesIO(png)))
                texto = detectados[0].data.decode() if detectados else None
            except Exception:
                raise RuntimeError("Falta el paquete zxing-cpp en el servidor: actualiza requirements.txt en el repositorio y reinicia la app")
        assert texto == qr_payload(registro), f"QR decodificado {texto} != {qr_payload(registro)}"
        return f"QR generado y releído correctamente ({texto})"

    def prueba_partes_ticket():
        registro = {
            "id": "CP-RAM-000099", "sucursal": "RAMMA", "marca": "Aveo", "color": "Azúl",
            "placas": "ABC-123-D", "tarifa": 70, "operador": "Prueba",
            "fecha": "01/01/2026 10:00:00", "custodia": "Prueba",
        }
        registro["ticket_texto"] = crear_ticket(registro)
        primera, segunda = crear_partes_ticket(registro)
        assert registro["id"] in primera, "El folio no aparece en el boleto del cliente"
        assert "LLAVERO" in segunda.upper() or "CONTRASE" in segunda.upper(), "Falta el llavero"
        return "Boleto del cliente y llavero separados con folio"

    def prueba_escpos():
        payload = construir_payload_escpos("PRUEBA BOLETO", "CLUBPARK|VEHICULO|CP-RAM-000099")
        assert payload[:2] == b"\x1b@", "Falta inicialización ESC @"
        assert b"\x1d(\x6b" in payload, "Faltan comandos del código QR"
        assert b"PRUEBA BOLETO" in payload
        return f"Binario ESC/POS válido de {len(payload)} bytes con QR"

    def prueba_impresion_rawbt():
        """Valida el payload de prueba y la bitácora de errores de impresión (ida y vuelta)."""
        texto, qr = generar_ticket_prueba_impresion()
        payload = construir_payload_escpos(texto, qr)
        assert payload[:2] == b"\x1b@", "Falta inicialización ESC @"
        assert b"\x1d(\x6b" in payload, "Faltan comandos del QR en el payload de prueba"
        marca_prueba = f"Prueba operativa automática de impresión {uuid.uuid4().hex[:6]}"
        registrar_evento_impresion("IMPRESION_OK", marca_prueba)
        eventos = cargar_registro_impresion()  # la bitácora guarda los 300 más recientes
        assert eventos and eventos[0].get("detalle") == marca_prueba, "La bitácora de impresión no guardó el evento"
        ultimo = eventos[0]
        assert ultimo.get("etapa") == "IMPRESION_OK" and "diagnóstico" in ultimo, "Evento incompleto"
        eventos.pop(0)  # limpia el evento de prueba para no ensuciar la bitácora real
        ARCHIVO_REGISTRO_IMPRESION.write_text(
            json.dumps(eventos, ensure_ascii=False, indent=2), encoding="utf-8"
        )
        return f"Payload ESC/POS de prueba ({len(payload)} bytes) y bitácora RawBT operativos"

    def prueba_credenciales():
        assert isinstance(ACCESOS, dict) and ACCESOS, "No hay credenciales cargadas"
        for token, datos in ACCESOS.items():
            assert datos.get("nombre") and datos.get("sucursal") and datos.get("rol"), f"Credencial incompleta: {token}"
        faltantes = [suc for suc in PRECIOS_SUCURSALES if not any(d.get("sucursal") == suc for d in ACCESOS.values())]
        assert not faltantes, f"Sucursales sin credencial: {faltantes}"
        return f"{len(ACCESOS)} credenciales completas (fuente: {ACCESOS_FUENTE})"

    def prueba_auditoria():
        guardar_evento_auditoria("PRUEBA_OPERATIVA", {"mensaje": "Prueba automática de bitácora"})
        eventos = cargar_auditoria()
        assert any(e.get("tipo") == "PRUEBA_OPERATIVA" for e in eventos), "El evento no se registró"
        return f"Bitácora con {len(eventos)} eventos"

    def prueba_fotos():
        assert PHOTO_DIR.exists() and os.access(PHOTO_DIR, os.W_OK), "La carpeta de fotos no es escribible"
        return f"Carpeta de fotografías escribible ({PHOTO_DIR.name})"

    def prueba_config():
        config = cargar_config()
        assert isinstance(config.get("siguiente_folio"), dict)
        return f"Configuración con {len(config['siguiente_folio'])} sucursales"

    registrar("Base de datos (JSON) escritura y lectura", prueba_base_datos)
    registrar("Folios independientes por sucursal", prueba_folios)
    registrar("Registro de ingreso de vehículo (simulado)", prueba_registro)
    registrar("Generación y lectura de código QR", prueba_qr)
    registrar("Partes del ticket (boleto + llavero)", prueba_partes_ticket)
    registrar("Binario de impresión térmica ESC/POS", prueba_escpos)
    registrar("Prueba de impresión y bitácora RawBT", prueba_impresion_rawbt)
    registrar("Integridad de credenciales QR", prueba_credenciales)
    registrar("Bitácora de auditoría", prueba_auditoria)
    registrar("Carpeta de fotografías del vehículo", prueba_fotos)
    registrar("Configuración de folios", prueba_config)
    return resultados


# ---------------------------------------------------------------------
# ESTADO DE SESIÓN
# ---------------------------------------------------------------------
if "usuario_autenticado" not in st.session_state:
    st.session_state.usuario_autenticado = None
if "sucursal_asignada" not in st.session_state:
    st.session_state.sucursal_asignada = None
if "rol" not in st.session_state:
    st.session_state.rol = None
if "vehiculos" not in st.session_state:
    st.session_state.vehiculos = cargar_bd()


# ---------------------------------------------------------------------
# LOGIN
# ---------------------------------------------------------------------
if st.session_state.usuario_autenticado is None:
    left, center, right = st.columns([1, 2.2, 1])
    with center:
        st.markdown(f"""
        <div class="clubpark-header-card" style="text-align:center;padding:18px 20px;margin-bottom:16px">
          <div class="clubpark-badge" style="margin-bottom:6px">Acceso Operativo Oficial</div>
          <div class="cpv-titulo-blanco" style="margin:4px 0 2px 0;font-size:22px">CLUB PARK VALET</div>
          <div style="font-size:13px;color:#FFFFFF;letter-spacing:1px;font-weight:700">Y ESTACIONAMIENTOS S.A. DE C.V.</div>
        </div>
        """, unsafe_allow_html=True)

        st.info("Identifícate mediante credencial QR o ingresa tu token operativo.")

        opciones_login = [
            "📷 Foto de QR con la cámara del dispositivo (Recomendado)",
            "🖼️ Subir imagen QR",
            "🔑 Token manual",
        ]
        metodo = st.radio(
            "Método de autenticación:",
            opciones_login,
            horizontal=False,
            key="metodo_login_operativo_v6"
        )

        token = str(st.query_params.get("qrscan", "") or "").strip() or None

        if metodo.startswith("📷"):
            if not token:
                st.caption("Pulsa el botón: se abrirá la cámara de tu teléfono. Se abre la cámara trasera. Enfoca la credencial QR a 10-15 cm, con buena luz.")
                token_leido = lector_qr_camara_trasera("lector_login_trasero", height=390)
                if token_leido:
                    token = str(token_leido).strip()
                    st.success("Código QR leído correctamente.")

        elif metodo.startswith("🖼️"):
            archivo_qr = st.file_uploader("Selecciona el archivo de tu credencial QR", type=["png", "jpg", "jpeg"])
            if archivo_qr:
                token = leer_codigo_qr(Image.open(archivo_qr))

        else:
            token_input = st.text_input("Ingresa tu token de acceso asignado", type="password")
            if st.button("Validar acceso", type="primary", use_container_width=True):
                token = token_input.strip()

        if token:
            info = ACCESOS.get(token)
            if info:
                st.session_state.usuario_autenticado = info.get("nombre", "Usuario")
                st.session_state.sucursal_asignada = info.get("sucursal", "")
                st.session_state.rol = normalizar_rol(info.get("rol", ROL_LIDER))
                if "qrscan" in st.query_params:
                    del st.query_params["qrscan"]
                st.success(f"Acceso concedido: **{info.get('nombre')}** ({info.get('rol')})")
                st.rerun()
            else:
                st.error(f"Token o código QR no reconocido: `{token}`. Verifica con el Contador Auditor.")

    st.stop()


# ---------------------------------------------------------------------
# ENCABEZADO Y NAVEGACIÓN
# ---------------------------------------------------------------------
st.session_state.rol = normalizar_rol(st.session_state.rol)
es_contador = usuario_es_contador()
es_global = usuario_es_global()
es_lider = usuario_es_lider()


# ---------------------------------------------------------------------
# CONSERVACIÓN DE FOTOS: cada foto de vehículo dura 7 días desde que se tomó.
# No se borran las de vehículos que siguen en custodia ni las de vehículos con
# «Incidencia» (se conservan hasta que el Contador cierre el caso).
# La limpieza corre sola, como máximo una vez cada 6 horas.
# ---------------------------------------------------------------------
DIAS_CONSERVAR_FOTOS = 7
ESTADOS_FOTOS_PROTEGIDAS = {"En Resguardo", "Incidencia"}
MARCA_LIMPIEZA_FOTOS = DATA_DIR / "ultima_limpieza_fotos.txt"


def _momento_foto(registro):
    try:
        return datetime.fromisoformat(str(registro.get("fecha_entrada_iso") or ""))
    except ValueError:
        try:
            return datetime.strptime(str(registro.get("fecha_entrada") or registro.get("fecha") or ""), "%d/%m/%Y %H:%M:%S")
        except ValueError:
            return None


def limpiar_fotos_vencidas(forzar=False):
    """Borra las fotos de vehículos tomadas hace más de DIAS_CONSERVAR_FOTOS días. Devuelve (vehículos, archivos)."""
    ahora_l = ahora_mx()
    if not forzar:
        try:
            ultima = datetime.fromisoformat(MARCA_LIMPIEZA_FOTOS.read_text(encoding="utf-8").strip())
            if (ahora_l - ultima).total_seconds() < 6 * 3600:
                return 0, 0
        except (OSError, ValueError):
            pass
    try:
        MARCA_LIMPIEZA_FOTOS.write_text(ahora_l.isoformat(timespec="seconds"), encoding="utf-8")
    except OSError:
        return 0, 0
    limite = ahora_l - timedelta(days=DIAS_CONSERVAR_FOTOS)
    base_l = cargar_bd()
    protegidas, vehiculos_l, archivos_l = set(), 0, 0
    for r in base_l:
        rutas_r = [rp for rp in list((r.get("fotos") or {}).values()) + list(r.get("fotos_custodia") or []) if rp]
        momento = _momento_foto(r)
        vencida = momento is not None and momento < limite
        if r.get("estado", "En Resguardo") in ESTADOS_FOTOS_PROTEGIDAS or not vencida:
            for rp in rutas_r:
                p_abs = _ruta_absoluta(rp)
                if p_abs:
                    protegidas.add(str(p_abs.resolve()))
            continue
        if rutas_r:
            archivos_l += _borrar_archivos(rutas_r)
            r["fotos"], r["fotos_custodia"] = {}, []
            r["fotos_borradas_el"] = ahora_l.strftime("%Y-%m-%d %H:%M:%S")
            vehiculos_l += 1
    if vehiculos_l:
        if guardar_bd(base_l):
            st.session_state.vehiculos = base_l
    # Archivos sueltos (de registros borrados o intentos fallidos) con más de 7 días
    try:
        for p_s in PHOTO_DIR.glob("*.jpg"):
            if str(p_s.resolve()) in protegidas:
                continue
            if datetime.fromtimestamp(p_s.stat().st_mtime) < limite - timedelta(hours=6):
                try:
                    p_s.unlink()
                    archivos_l += 1
                except OSError:
                    pass
    except OSError:
        pass
    if vehiculos_l or archivos_l:
        guardar_evento_auditoria("LIMPIEZA AUTOMÁTICA DE FOTOS",
                                 f"{vehiculos_l} vehículo(s) · {archivos_l} archivo(s) con más de {DIAS_CONSERVAR_FOTOS} días")
    return vehiculos_l, archivos_l


sucursal_texto = "Todas (Global)" if es_global else st.session_state.sucursal_asignada

st.markdown(f"""
<div class="clubpark-header-card">
  <div style="display:flex;justify-content:space-between;align-items:center;flex-wrap:wrap;gap:10px">
    <div>
      <div class="clubpark-badge">SISTEMA INSTITUCIONAL DE CONTROL OPERATIVO</div>
      <div class="cpv-titulo-blanco" style="margin:6px 0 2px 0;font-size:24px">CLUB PARK VALET Y ESTACIONAMIENTOS S.A. DE C.V.</div>
      <div style="font-size:13px;color:#EBF3EE">
        <b>Usuario:</b> {st.session_state.usuario_autenticado} &nbsp;|&nbsp;
        <b>Rol:</b> {st.session_state.rol} &nbsp;|&nbsp;
        <b>Sucursal:</b> <span style="color:#FFFFFF;font-weight:700">{sucursal_texto}</span>
      </div>
    </div>
    <div style="text-align:right">
      <div style="font-size:11px;color:#D9EBDF;text-transform:uppercase;font-weight:700">Identidad Corporativa</div>
      <div style="font-size:13px;font-weight:800;color:#FFFFFF;letter-spacing:0.5px">SERVIR CON HONESTIDAD, ACTUAR CON INTEGRIDAD Y GENERAR CONFIANZA</div>
    </div>
  </div>
</div>
""", unsafe_allow_html=True)

with st.sidebar:
    st.markdown('<div style="font-weight:800;color:#FFFFFF;letter-spacing:2px;font-size:16px;text-align:center;background:#0B3B24;border-radius:10px;padding:10px 8px;border:1px solid #2F6B4C">CLUB PARK VALET</div>', unsafe_allow_html=True)
    st.subheader("Panel de sesión")
    st.write(f"**Usuario:** {st.session_state.usuario_autenticado}")
    st.write(f"**Rol:** {st.session_state.rol}")
    st.write(f"**Vehículos registrados:** {len(st.session_state.vehiculos)}")
    if st.button("🚪 Cerrar sesión", use_container_width=True):
        st.session_state.usuario_autenticado = None
        st.session_state.sucursal_asignada = None
        st.session_state.rol = None
        st.rerun()

def analisis_inteligente(registros):
    """Análisis automático de la operación: KPIs, alertas tempranas y resumen
    ejecutivo. Se recalcula con la información real cada vez que se consulta."""
    ahora = ahora_mx()
    activos = [r for r in registros if r.get("estado") != "Entregado"]
    entregados = [r for r in registros if r.get("estado") == "Entregado"]

    def horas_estancia(r):
        try:
            iso = r.get("fecha_entrada_iso") or ""
            if not iso:
                return None
            return (ahora - datetime.fromisoformat(iso)).total_seconds() / 3600
        except Exception:
            return None

    estancias = [(r, h) for r, h in ((r, horas_estancia(r)) for r in activos) if h is not None]

    alertas = []
    for r, h in sorted(estancias, key=lambda par: -par[1]):
        ident = r.get("id", "")
        placas = r.get("placas", "")
        suc = r.get("sucursal", "")
        if h >= 12:
            alertas.append({"tipo": "🚨 Estancia crítica", "sucursal": suc, "detalle": f"{ident} ({placas}) en {suc}: {h:.1f} h sin salida", "accion": "Verificar con el encargado y registrar salida o incidencia hoy."})
        elif h >= 6:
            alertas.append({"tipo": "⚠️ Estancia prolongada", "sucursal": suc, "detalle": f"{ident} ({placas}) en {suc}: {h:.1f} h", "accion": "Dar seguimiento en la próxima revisión."})
    placas_activas = {}
    for r in activos:
        clave = normalizar_texto(r.get("placas")).upper()
        if clave:
            placas_activas.setdefault(clave, []).append(r)
    for clave, lista in placas_activas.items():
        if len(lista) > 1:
            alertas.append({"tipo": "⚠️ Placa repetida en resguardo", "sucursal": ", ".join(sorted({str(r.get("sucursal", "")) for r in lista})), "detalle": f"Placa {clave} en {len(lista)} vehículos activos: " + ", ".join(str(r.get('id', '')) for r in lista), "accion": "Revisar duplicados: posible error de captura."})
    for r in entregados:
        if normalizar_texto(r.get("pago_estado", "")).lower() not in ("pagado", "cortado", ""):
            alertas.append({"tipo": "⚠️ Entrega sin cobro confirmado", "sucursal": str(r.get("sucursal", "")), "detalle": f"{r.get('id', '')} ({r.get('placas', '')}) en {r.get('sucursal', '')}", "accion": "Conciliar el cobro en el corte de caja."})

    entregados_hoy = [r for r in entregados if str(r.get("fecha_salida", "")).startswith(ahora.strftime("%d/%m/%Y"))]
    ingresos_hoy = sum(dinero(r.get("tarifa")) for r in entregados_hoy)
    estancia_prom = (sum(h for _, h in estancias) / len(estancias)) if estancias else 0.0
    sucursales_activas = sorted({r.get("sucursal", "Sin sucursal") for r in activos})
    por_sucursal = max(sucursales_activas, key=lambda s: sum(1 for r in activos if r.get("sucursal") == s)) if sucursales_activas else "—"
    conteo_top = sum(1 for r in activos if r.get("sucursal") == por_sucursal) if por_sucursal != "—" else 0
    ticket_prom = (sum(dinero(r.get("tarifa")) for r in entregados_hoy) / len(entregados_hoy)) if entregados_hoy else 0.0

    kpis = [
        ("Vehículos en resguardo", f"{len(activos)}"),
        ("Entregados hoy", f"{len(entregados_hoy)}"),
        ("Ingresos de hoy", f"${ingresos_hoy:,.0f}"),
        ("Estancia promedio", f"{estancia_prom:.1f} h" if activos else "—"),
    ]

    resumen = (
        f"Hay {len(activos)} vehículos en resguardo en {len(sucursales_activas)} sucursal(es). "
        f"Hoy se han entregado {len(entregados_hoy)} vehículos con ingresos de ${ingresos_hoy:,.0f} "
        f"(ticket promedio ${ticket_prom:,.0f}). "
        + (f"La sucursal con más vehículos en sitio es {por_sucursal} ({conteo_top}). " if por_sucursal != "—" else "")
        + (f"La estancia promedio en sitio es de {estancia_prom:.1f} horas. " if activos else "")
        + f"Análisis generado automáticamente el {ahora.strftime('%d/%m/%Y %H:%M')}."
    )

    recomendaciones = []
    if len(activos) >= 20:
        recomendaciones.append("Volumen alto en sitio: refuerza personal en la zona de resguardo durante horas pico.")
    if estancia_prom >= 5 and activos:
        recomendaciones.append("La estancia promedio supera 5 horas: revisa tarifas por tiempo o promueve pagos anticipados.")
    sucursales_alertadas = sorted({a.get("sucursal", "") for a in alertas if a.get("sucursal")})
    if sucursales_alertadas:
        recomendaciones.append("Concentra la revisión en: " + ", ".join(sucursales_alertadas[:5]) + ".")
    if not entregados_hoy and activos:
        recomendaciones.append("Aún sin salidas registradas hoy: verifica que los encargados estén usando el módulo de Salida y Cobro.")
    if not recomendaciones:
        recomendaciones.append("Operación dentro de parámetros normales. Continúa con el protocolo fotográfico establecido.")

    return kpis, alertas, resumen, recomendaciones



# =====================================================================
# REPORTES FÍSICOS (ingresos / gastos) + EXCEL OFICIAL DE FLUJO
# (Integrado desde la app de reporte de flujo diario por sitio.)
# =====================================================================
CATALOGO_OFICIAL_23 = [
    {"num": 1, "sitio": "PAMPAS INSURGENTES", "dias": "LUNES A DOMINGO", "horarios": "11:30 AM A 12:00 PM", "empleados": 5, "supervisor": "JESÚS", "tarifa": 50},
    {"num": 2, "sitio": "PAMPAS COAPA", "dias": "LUNES A DOMINGO", "horarios": "-", "empleados": 0, "supervisor": "JESÚS", "tarifa": 50},
    {"num": 3, "sitio": "RAMMA", "dias": "LUNES A DOMINGO", "horarios": "12:00 PM A 02:00 AM", "empleados": 2, "supervisor": "JESÚS", "tarifa": 70},
    {"num": 4, "sitio": "SUSHITTO", "dias": "LUNES A DOMINGO", "horarios": "01:00 PM A 08:00 PM", "empleados": 1, "supervisor": "JESÚS", "tarifa": 40},
    {"num": 5, "sitio": "CARAJILLO POLANCO", "dias": "LUNES A DOMINGO", "horarios": "01:00 PM A 01:00 AM", "empleados": 2, "supervisor": "JESÚS", "tarifa": 100},
    {"num": 6, "sitio": "BENINGAN´S", "dias": "LUNES A DOMINGO", "horarios": "12:00 PM A 09:00 PM", "empleados": 2, "supervisor": "JESÚS", "tarifa": 40},
    {"num": 7, "sitio": "SANATORIO DURANGO", "dias": "LUNES A DOMINGO", "horarios": "12:00 PM A 09:00 PM", "empleados": 4, "supervisor": "JESÚS", "tarifa": 100},
    {"num": 8, "sitio": "BARRANCA", "dias": "LUNES A VIERNES", "horarios": "08:00 AM A 08:00 PM", "empleados": 2, "supervisor": "JESÚS", "tarifa": 0},
    {"num": 9, "sitio": "VETERINARIA", "dias": "LUNES A DOMINGO", "horarios": "09:00 AM A 19:00 PM", "empleados": 1, "supervisor": "JESÚS", "tarifa": 0},
    {"num": 10, "sitio": "ROMANCE PALACE", "dias": "VIERNES A DOMINGO", "horarios": "-", "empleados": 0, "supervisor": "JESÚS", "tarifa": 65},
    {"num": 11, "sitio": "JUNGLA TLALPAN", "dias": "SÁBADO Y DOMINGO", "horarios": "-", "empleados": 0, "supervisor": "JESÚS", "tarifa": 100},
    {"num": 12, "sitio": "JUNGLA DIVISIÓN", "dias": "SÁBADO Y DOMINGO", "horarios": "-", "empleados": 0, "supervisor": "JESÚS", "tarifa": 100},
    {"num": 13, "sitio": "JUNGLA PICACHO", "dias": "SÁBADO Y DOMINGO", "horarios": "-", "empleados": 0, "supervisor": "JESÚS", "tarifa": 100},
    {"num": 14, "sitio": "JUNGLA EMPRESA", "dias": "SÁBADO Y DOMINGO", "horarios": "-", "empleados": 0, "supervisor": "JESÚS", "tarifa": 100},
    {"num": 15, "sitio": "SIPIRILY COAPA", "dias": "SÁBADO Y DOMINGO", "horarios": "-", "empleados": 0, "supervisor": "JESÚS", "tarifa": 80},
    {"num": 16, "sitio": "MAROMETA SATÉLITE", "dias": "SÁBADO Y DOMINGO", "horarios": "-", "empleados": 0, "supervisor": "JESÚS", "tarifa": 70},
    {"num": 17, "sitio": "MAROMETA MARINA", "dias": "SÁBADO Y DOMINGO", "horarios": "-", "empleados": 0, "supervisor": "JESÚS", "tarifa": 70},
    {"num": 18, "sitio": "SUMO STA. MONICA", "dias": "LUNES A DOMINGO", "horarios": "12:30 PM A 11:00 PM", "empleados": 1, "supervisor": "VÍCTOR", "tarifa": 50},
    {"num": 19, "sitio": "SUMO ECHEGARAY", "dias": "LUNES A DOMINGO", "horarios": "12:30 PM A 11:00 PM", "empleados": 1, "supervisor": "VÍCTOR", "tarifa": 50},
    {"num": 20, "sitio": "RYOSHI", "dias": "LUNES A DOMINGO", "horarios": "12:00 PM A 1:00 AM", "empleados": 4, "supervisor": "VÍCTOR", "tarifa": 80},
    {"num": 21, "sitio": "ENCANTO", "dias": "LUNES A DOMINGO", "horarios": "12:30 PM A 01 AM", "empleados": 1, "supervisor": "VÍCTOR", "tarifa": 60},
    {"num": 22, "sitio": "E-BNI", "dias": "LUNES A DOMINGO", "horarios": "-", "empleados": 0, "supervisor": "VÍCTOR", "tarifa": 50},
    {"num": 23, "sitio": "CHURRASQUERIA", "dias": "LUNES A DOMINGO", "horarios": "12:30 PM A 11:00 PM", "empleados": 1, "supervisor": "VÍCTOR", "tarifa": 70},
]

_nombres_cat = {i["sitio"] for i in CATALOGO_OFICIAL_23}
CATALOGO_OFICIAL_23[:] = [i for i in CATALOGO_OFICIAL_23 if i["sitio"] not in _SUC_CFG["quitadas"]]
for _n, _pr in _SUC_CFG["agregadas"].items():
    if _n not in _nombres_cat:
        CATALOGO_OFICIAL_23.append({"num": len(CATALOGO_OFICIAL_23) + 1, "sitio": _n, "dias": "LUNES A DOMINGO", "horarios": "-",
                                    "empleados": 0, "supervisor": "-", "tarifa": float(_pr)})
for _i, _it in enumerate(CATALOGO_OFICIAL_23, start=1):
    _it["num"] = _i
    if _it["sitio"] in _SUC_CFG.get("supervisores", {}):
        _it["supervisor"] = _SUC_CFG["supervisores"][_it["sitio"]]
    for _campo in ("dias", "horarios"):
        if _SUC_CFG.get("detalles", {}).get(_it["sitio"], {}).get(_campo):
            _it[_campo] = _SUC_CFG["detalles"][_it["sitio"]][_campo]

# ---------------------------------------------------------------------
# SUPERVISORES: sucursales a cargo de cada supervisor (José de Jesús / Víctor)
# ---------------------------------------------------------------------
SUPERVISORES_NOMBRE = {"JESÚS": "José de Jesús", "VÍCTOR": "Víctor Manuel"}


def _sin_acentos(texto):
    import unicodedata
    return unicodedata.normalize("NFKD", str(texto or "")).encode("ASCII", "ignore").decode("ASCII").upper()


def sucursales_de_supervisor(clave):
    return [i["sitio"] for i in CATALOGO_OFICIAL_23
            if _sin_acentos(i.get("supervisor")) == _sin_acentos(clave) and i["sitio"] in PRECIOS_SUCURSALES]


def supervisor_del_usuario():
    """Clave del supervisor (JESÚS / VÍCTOR) de la sesión actual; None si no es supervisor global."""
    if st.session_state.get("rol") != "Supervisor Global":
        return None
    nombre = _sin_acentos(st.session_state.get("usuario_autenticado"))
    for clave in SUPERVISORES_NOMBRE:
        if _sin_acentos(clave) in nombre:
            return clave
    return None

MESES_ESP = {1: "ENE", 2: "FEB", 3: "MAR", 4: "ABR", 5: "MAY", 6: "JUN", 7: "JUL", 8: "AGO", 9: "SEP", 10: "OCT", 11: "NOV", 12: "DIC"}
DIAS_ESP = {0: "LUNES", 1: "MARTES", 2: "MIERCOLES", 3: "JUEVES", 4: "VIERNES", 5: "SÁBADO", 6: "DOMINGO"}

def normalizar_a_sucursal_oficial(nombre):
    if not nombre:
        return "PAMPAS INSURGENTES"
    n = str(nombre).strip().upper()
    mapa = {
        "PAMPAS INS": "PAMPAS INSURGENTES", "PAMPAS COA": "PAMPAS COAPA",
        "PAMPAS": "PAMPAS INSURGENTES", "RAMMA": "RAMMA", "RAMA": "RAMMA",
        "SUSHITTO": "SUSHITTO", "SUSH": "SUSHITTO",
        "CARAJILLO": "CARAJILLO POLANCO", "CARAJ": "CARAJILLO POLANCO",
        "BENINGANS": "BENINGAN´S", "BENINGAN´S": "BENINGAN´S", "BENI": "BENINGAN´S",
        "SANATORIO": "SANATORIO DURANGO", "DURANGO": "SANATORIO DURANGO", "CLÍNICA": "SANATORIO DURANGO",
        "BARRANCA": "BARRANCA", "VETERINARIA": "VETERINARIA",
        "ROMAS": "ROMANCE PALACE", "ROMANCE": "ROMANCE PALACE",
        "J.TLALPAN": "JUNGLA TLALPAN", "TLALPAN": "JUNGLA TLALPAN",
        "J.DIVISION": "JUNGLA DIVISIÓN", "DIVISIÓN": "JUNGLA DIVISIÓN",
        "J.PICACHO": "JUNGLA PICACHO", "PICACHO": "JUNGLA PICACHO",
        "J.EMPRESA": "JUNGLA EMPRESA", "EMPRESA": "JUNGLA EMPRESA",
        "SIPIRILY": "SIPIRILY COAPA", "SIPILY": "SIPIRILY COAPA",
        "M SATÉLITE": "MAROMETA SATÉLITE", "SATÉLITE": "MAROMETA SATÉLITE",
        "M MARINA": "MAROMETA MARINA", "MARINA": "MAROMETA MARINA",
        "SUMO STA": "SUMO STA. MONICA", "MONICA": "SUMO STA. MONICA",
        "CHEGARAY": "SUMO ECHEGARAY", "ECHEGARAY": "SUMO ECHEGARAY",
        "RYOSHI": "RYOSHI", "RYO": "RYOSHI", "ENCANTO": "ENCANTO", "ENCANT": "ENCANTO",
        "E-BNI": "E-BNI", "CHURRASCA": "CHURRASQUERIA", "CHURRASQUERIA": "CHURRASQUERIA"
    }
    for k, v in mapa.items():
        if k in n:
            return v
    return n

def generar_excel_formato_oficial(registros, etiqueta_periodo="SEMANAL"):
    """Genera el libro oficial idéntico a CONTROL_FLUJO_DIARIO: hojas DETALLE, CONCENTRADO X DÍA y Rentabilidad."""
    wb = openpyxl.Workbook()
    ws_det = wb.active
    ws_det.title = "DETALLE "
    ws_conc = wb.create_sheet(title="CONCENTRADO X DÍA")
    ws_aud = wb.create_sheet(title="Rentabilidad y Auditoría")

    fill_green = PatternFill(start_color="003300", end_color="003300", fill_type="solid")
    font_title = Font(name="Calibri", size=15, bold=True, color="003300")
    font_subtitle = Font(name="Calibri", size=11, bold=True, color="000000")
    font_head = Font(name="Calibri", size=8, bold=True, color="FFFFFF")
    font_bold = Font(name="Calibri", size=8, bold=True, color="000000")
    font_data = Font(name="Calibri", size=8, color="000000")
    align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    align_left = Alignment(horizontal="left", vertical="center")
    align_right = Alignment(horizontal="right", vertical="center")
    thin_border = Border(
        left=Side(style="thin", color="B2B2B2"),
        right=Side(style="thin", color="B2B2B2"),
        top=Side(style="thin", color="B2B2B2"),
        bottom=Side(style="thin", color="B2B2B2")
    )

    fechas_dt = []
    for r in registros:
        f_str = r.get("Fecha")
        if f_str:
            for fmt in ("%d/%m/%Y", "%Y-%m-%d"):
                try:
                    fechas_dt.append(datetime.strptime(f_str, fmt).date())
                    break
                except ValueError:
                    pass

    if fechas_dt:
        fechas_unicas = sorted(list(set(fechas_dt)))
        if len(fechas_unicas) <= 7:
            min_d = fechas_unicas[0]
            lunes = min_d - timedelta(days=min_d.weekday())
            dias_periodo = [lunes + timedelta(days=i) for i in range(7)]
        else:
            dias_periodo = fechas_unicas
    else:
        hoy = hoy_mx()
        lunes = hoy - timedelta(days=hoy.weekday())
        dias_periodo = [lunes + timedelta(days=i) for i in range(7)]

    txt_periodo = f"CORTE {dias_periodo[0].strftime('%d/%m/%Y')} AL {dias_periodo[-1].strftime('%d/%m/%Y')} ({etiqueta_periodo.upper()})"

    matriz_reg = {}
    for r in registros:
        f_str = r.get("Fecha")
        d_obj = None
        if f_str:
            for fmt in ("%d/%m/%Y", "%Y-%m-%d"):
                try:
                    d_obj = datetime.strptime(f_str, fmt).date()
                    break
                except ValueError:
                    pass
        if d_obj:
            suc_norm = normalizar_a_sucursal_oficial(r.get("Sitio / Unidad"))
            key = (suc_norm, d_obj)
            matriz_reg[key] = r

    # --- HOJA 1: DETALLE ---
    ws_det["A1"] = "CLUB PARK VALET Y ESTACIONAMIENTOS, S.A. DE C.V."
    ws_det["A1"].font = font_title
    ws_det["A3"] = txt_periodo
    ws_det["A3"].font = font_subtitle
    ws_det["A5"] = "INGRESOS POR PERÍODO"
    ws_det["A5"].font = font_subtitle

    meta_headers = ["No.", "SITIO", "DÍAS QUE SE OPERA", "HORARIOS", "No. EMPLEADOS", "SUPERVISOR"]
    for i, h in enumerate(meta_headers, start=1):
        cell = ws_det.cell(8, i, h)
        cell.fill = fill_green; cell.font = font_head; cell.alignment = align_center; cell.border = thin_border
        ws_det.merge_cells(start_row=8, start_column=i, end_row=10, end_column=i)

    col_idx = 7
    dia_cols_map = {}
    for d in dias_periodo:
        d_str = f"{DIAS_ESP[d.weekday()]} {d.day:02d}/{MESES_ESP[d.month]}/{d.year}"
        dia_cols_map[d] = col_idx
        c8 = ws_det.cell(8, col_idx, d_str)
        c8.fill = fill_green; c8.font = font_head; c8.alignment = align_center
        ws_det.merge_cells(start_row=8, start_column=col_idx, end_row=8, end_column=col_idx+4)
        for c in range(col_idx, col_idx+5):
            ws_det.cell(8, c).fill = fill_green; ws_det.cell(8, c).border = thin_border

        ws_det.cell(9, col_idx, "CONSECUTIVO BOLETAJE").fill = fill_green
        ws_det.cell(9, col_idx).font = font_head; ws_det.cell(9, col_idx).alignment = align_center
        ws_det.merge_cells(start_row=9, start_column=col_idx, end_row=9, end_column=col_idx+1)
        ws_det.cell(9, col_idx+1).fill = fill_green

        c_tb = ws_det.cell(9, col_idx+2, "TOTAL DE BOLETOS")
        c_tb.fill = fill_green; c_tb.font = font_head; c_tb.alignment = align_center
        ws_det.merge_cells(start_row=9, start_column=col_idx+2, end_row=10, end_column=col_idx+2)

        c_cb = ws_det.cell(9, col_idx+3, "COSTO X BOLETO")
        c_cb.fill = fill_green; c_cb.font = font_head; c_cb.alignment = align_center
        ws_det.merge_cells(start_row=9, start_column=col_idx+3, end_row=10, end_column=col_idx+3)

        c_ing = ws_det.cell(9, col_idx+4, "INGRESO POR DÍA")
        c_ing.fill = fill_green; c_ing.font = font_head; c_ing.alignment = align_center
        ws_det.merge_cells(start_row=9, start_column=col_idx+4, end_row=10, end_column=col_idx+4)

        c_fi = ws_det.cell(10, col_idx, "FOLIO INICIAL")
        c_fi.fill = fill_green; c_fi.font = font_head; c_fi.alignment = align_center
        c_ff = ws_det.cell(10, col_idx+1, "FOLIO FINAL")
        c_ff.fill = fill_green; c_ff.font = font_head; c_ff.alignment = align_center

        for r_h in [9, 10]:
            for c_h in range(col_idx, col_idx+5):
                ws_det.cell(r_h, c_h).border = thin_border
        col_idx += 5

    col_tot_bol = col_idx
    col_tot_ing = col_idx + 1
    c_tbl = ws_det.cell(8, col_tot_bol, "TOTAL DE BOLETOS UTILIZADOS X PERÍODO")
    c_tbl.fill = fill_green; c_tbl.font = font_head; c_tbl.alignment = align_center; c_tbl.border = thin_border
    ws_det.merge_cells(start_row=8, start_column=col_tot_bol, end_row=10, end_column=col_tot_bol)

    c_tin = ws_det.cell(8, col_tot_ing, "TOTAL DE INGRESOS X PERÍODO")
    c_tin.fill = fill_green; c_tin.font = font_head; c_tin.alignment = align_center; c_tin.border = thin_border
    ws_det.merge_cells(start_row=8, start_column=col_tot_ing, end_row=10, end_column=col_tot_ing)

    r_actual = 11
    fila_inicio_ing = r_actual
    for item in CATALOGO_OFICIAL_23:
        ws_det.cell(r_actual, 1, item["num"]).alignment = align_center
        ws_det.cell(r_actual, 2, item["sitio"]).alignment = align_left
        ws_det.cell(r_actual, 3, item["dias"]).alignment = align_left
        ws_det.cell(r_actual, 4, item["horarios"]).alignment = align_center
        ws_det.cell(r_actual, 5, item["empleados"]).alignment = align_center
        ws_det.cell(r_actual, 6, item["supervisor"]).alignment = align_center

        for ci in range(1, 7):
            cell = ws_det.cell(r_actual, ci)
            cell.font = font_data; cell.border = thin_border

        tot_bol_suc = 0
        tot_ing_suc = 0.0
        for d in dias_periodo:
            c_d = dia_cols_map[d]
            rep = matriz_reg.get((item["sitio"], d))
            if rep:
                f_ini = rep.get("Folio Inicial", 0)
                f_fin = rep.get("Folio Final", 0)
                bol = rep.get("Boletos Cobrados", rep.get("Total Boletos", 0))
                costo = rep.get("Costo Boleto", item["tarifa"])
                ing = rep.get("Ingreso Bruto", float(bol) * float(costo))
            else:
                f_ini = 0; f_fin = 0; bol = 0; costo = item["tarifa"]; ing = 0.0

            ws_det.cell(r_actual, c_d, f_ini).alignment = align_center
            ws_det.cell(r_actual, c_d+1, f_fin).alignment = align_center
            ws_det.cell(r_actual, c_d+2, bol).alignment = align_center
            ws_det.cell(r_actual, c_d+3, costo).alignment = align_right
            ws_det.cell(r_actual, c_d+3).number_format = "$#,##0.00"
            ws_det.cell(r_actual, c_d+4, ing).alignment = align_right
            ws_det.cell(r_actual, c_d+4).number_format = "$#,##0.00"

            tot_bol_suc += bol
            tot_ing_suc += ing

            for c_sub in range(c_d, c_d+5):
                cell = ws_det.cell(r_actual, c_sub)
                cell.font = font_data; cell.border = thin_border

        c_tb_cell = ws_det.cell(r_actual, col_tot_bol, tot_bol_suc)
        c_tb_cell.alignment = align_center; c_tb_cell.font = font_bold; c_tb_cell.border = thin_border
        c_ti_cell = ws_det.cell(r_actual, col_tot_ing, tot_ing_suc)
        c_ti_cell.alignment = align_right; c_ti_cell.font = font_bold; c_ti_cell.border = thin_border
        c_ti_cell.number_format = "$#,##0.00"
        r_actual += 1

    fila_fin_ing = r_actual - 1

    r_tot_ing = r_actual
    ws_det.cell(r_tot_ing, 2, "TOTAL").font = font_bold
    ws_det.cell(r_tot_ing, 2).alignment = align_left
    for ci in range(1, col_tot_ing + 1):
        ws_det.cell(r_tot_ing, ci).border = thin_border

    for d in dias_periodo:
        c_d = dia_cols_map[d]
        col_let_bol = get_column_letter(c_d+2)
        col_let_ing = get_column_letter(c_d+4)
        c_sb = ws_det.cell(r_tot_ing, c_d+2, f"=SUM({col_let_bol}{fila_inicio_ing}:{col_let_bol}{fila_fin_ing})")
        c_sb.font = font_bold; c_sb.alignment = align_center
        c_si = ws_det.cell(r_tot_ing, c_d+4, f"=SUM({col_let_ing}{fila_inicio_ing}:{col_let_ing}{fila_fin_ing})")
        c_si.font = font_bold; c_si.alignment = align_right; c_si.number_format = "$#,##0.00"

    let_tbl = get_column_letter(col_tot_bol)
    let_tin = get_column_letter(col_tot_ing)
    ws_det.cell(r_tot_ing, col_tot_bol, f"=SUM({let_tbl}{fila_inicio_ing}:{let_tbl}{fila_fin_ing})").font = font_bold
    ws_det.cell(r_tot_ing, col_tot_bol).alignment = align_center
    c_tot_ing_fin = ws_det.cell(r_tot_ing, col_tot_ing, f"=SUM({let_tin}{fila_inicio_ing}:{let_tin}{fila_fin_ing})")
    c_tot_ing_fin.font = font_bold; c_tot_ing_fin.alignment = align_right; c_tot_ing_fin.number_format = "$#,##0.00"

    # Sección Gastos
    r_gastos_head = r_tot_ing + 3
    ws_det.cell(r_gastos_head - 2, 1, "GASTOS POR PERÍODO").font = font_subtitle

    for i, h in enumerate(meta_headers, start=1):
        cell = ws_det.cell(r_gastos_head, i, h)
        cell.fill = fill_green; cell.font = font_head; cell.alignment = align_center; cell.border = thin_border
        ws_det.merge_cells(start_row=r_gastos_head, start_column=i, end_row=r_gastos_head+1, end_column=i)

    for d in dias_periodo:
        c_d = dia_cols_map[d]
        d_str = f"{DIAS_ESP[d.weekday()]} {d.day:02d}/{MESES_ESP[d.month]}/{d.year}"
        c_gh = ws_det.cell(r_gastos_head, c_d, d_str)
        c_gh.fill = fill_green; c_gh.font = font_head; c_gh.alignment = align_center
        ws_det.merge_cells(start_row=r_gastos_head, start_column=c_d, end_row=r_gastos_head, end_column=c_d+4)
        for c in range(c_d, c_d+5):
            ws_det.cell(r_gastos_head, c).fill = fill_green; ws_det.cell(r_gastos_head, c).border = thin_border

        sub_g = ["NÓMINA LIDER DE UNIDAD", "NÓMINA OPERADOR", "GASTOS SUPERVISIÓN", "VARIOS", "TOTAL GASTOS X DÍA"]
        for sg_i, sg_txt in enumerate(sub_g):
            cell = ws_det.cell(r_gastos_head+1, c_d+sg_i, sg_txt)
            cell.fill = fill_green; cell.font = font_head; cell.alignment = align_center; cell.border = thin_border

    c_tg_per = ws_det.cell(r_gastos_head, col_tot_bol, "TOTAL DE GASTOS X PERÍODO")
    c_tg_per.fill = fill_green; c_tg_per.font = font_head; c_tg_per.alignment = align_center; c_tg_per.border = thin_border
    ws_det.merge_cells(start_row=r_gastos_head, start_column=col_tot_bol, end_row=r_gastos_head+1, end_column=col_tot_ing)

    r_actual_g = r_gastos_head + 2
    fila_inicio_gas = r_actual_g
    for item in CATALOGO_OFICIAL_23:
        ws_det.cell(r_actual_g, 1, item["num"]).alignment = align_center
        ws_det.cell(r_actual_g, 2, item["sitio"]).alignment = align_left
        ws_det.cell(r_actual_g, 3, item["dias"]).alignment = align_left
        ws_det.cell(r_actual_g, 4, item["horarios"]).alignment = align_center
        ws_det.cell(r_actual_g, 5, item["empleados"]).alignment = align_center
        ws_det.cell(r_actual_g, 6, item["supervisor"]).alignment = align_center

        for ci in range(1, 7):
            cell = ws_det.cell(r_actual_g, ci)
            cell.font = font_data; cell.border = thin_border

        tot_gas_suc = 0.0
        for d in dias_periodo:
            c_d = dia_cols_map[d]
            rep = matriz_reg.get((item["sitio"], d))
            if rep:
                nom = rep.get("Gastos Generales / Nómina", 0.0)
                esp = rep.get("Gastos Especiales", 0.0)
                tot_g = rep.get("Total Gastos", nom + esp)
            else:
                nom = 0.0; esp = 0.0; tot_g = 0.0

            ws_det.cell(r_actual_g, c_d, nom).alignment = align_right
            ws_det.cell(r_actual_g, c_d).number_format = "$#,##0.00"
            ws_det.cell(r_actual_g, c_d+1, 0).alignment = align_right
            ws_det.cell(r_actual_g, c_d+1).number_format = "$#,##0.00"
            ws_det.cell(r_actual_g, c_d+2, 0).alignment = align_right
            ws_det.cell(r_actual_g, c_d+2).number_format = "$#,##0.00"
            ws_det.cell(r_actual_g, c_d+3, esp).alignment = align_right
            ws_det.cell(r_actual_g, c_d+3).number_format = "$#,##0.00"
            ws_det.cell(r_actual_g, c_d+4, tot_g).alignment = align_right
            ws_det.cell(r_actual_g, c_d+4).number_format = "$#,##0.00"

            tot_gas_suc += tot_g
            for c_sub in range(c_d, c_d+5):
                cell = ws_det.cell(r_actual_g, c_sub)
                cell.font = font_data; cell.border = thin_border

        c_tg_cell = ws_det.cell(r_actual_g, col_tot_bol, tot_gas_suc)
        c_tg_cell.alignment = align_right; c_tg_cell.font = font_bold; c_tg_cell.border = thin_border
        c_tg_cell.number_format = "$#,##0.00"
        ws_det.merge_cells(start_row=r_actual_g, start_column=col_tot_bol, end_row=r_actual_g, end_column=col_tot_ing)
        r_actual_g += 1

    fila_fin_gas = r_actual_g - 1

    r_tot_gas = r_actual_g
    ws_det.cell(r_tot_gas, 2, "TOTAL").font = font_bold
    for ci in range(1, col_tot_ing + 1):
        ws_det.cell(r_tot_gas, ci).border = thin_border

    for d in dias_periodo:
        c_d = dia_cols_map[d]
        col_let_gas = get_column_letter(c_d+4)
        c_sg = ws_det.cell(r_tot_gas, c_d+4, f"=SUM({col_let_gas}{fila_inicio_gas}:{col_let_gas}{fila_fin_gas})")
        c_sg.font = font_bold; c_sg.alignment = align_right; c_sg.number_format = "$#,##0.00"

    ws_det.cell(r_tot_gas, col_tot_bol, f"=SUM({get_column_letter(col_tot_bol)}{fila_inicio_gas}:{get_column_letter(col_tot_bol)}{fila_fin_gas})").font = font_bold
    ws_det.cell(r_tot_gas, col_tot_bol).alignment = align_right
    ws_det.cell(r_tot_gas, col_tot_bol).number_format = "$#,##0.00"
    ws_det.merge_cells(start_row=r_tot_gas, start_column=col_tot_bol, end_row=r_tot_gas, end_column=col_tot_ing)

    r_util = r_tot_gas + 2
    ws_det.cell(r_util, 2, "UTILIDAD O PÉRDIDA X DÍA").font = font_bold
    for ci in range(1, col_tot_ing + 1):
        ws_det.cell(r_util, ci).border = thin_border

    for d in dias_periodo:
        c_d = dia_cols_map[d]
        let_ing = get_column_letter(c_d+4)
        c_ut = ws_det.cell(r_util, c_d+4, f"=+{let_ing}{r_tot_ing}-{let_ing}{r_tot_gas}")
        c_ut.font = font_bold; c_ut.alignment = align_right; c_ut.number_format = "$#,##0.00"

    c_ut_tot = ws_det.cell(r_util, col_tot_bol, f"=+{get_column_letter(col_tot_ing)}{r_tot_ing}-{get_column_letter(col_tot_bol)}{r_tot_gas}")
    c_ut_tot.font = font_bold; c_ut_tot.alignment = align_right; c_ut_tot.number_format = "$#,##0.00"
    ws_det.merge_cells(start_row=r_util, start_column=col_tot_bol, end_row=r_util, end_column=col_tot_ing)

    # --- HOJA 2: CONCENTRADO X DÍA ---
    ws_conc["A1"] = "CLUB PARK VALET Y ESTACIONAMIENTOS, S.A. DE C.V."
    ws_conc["A1"].font = font_title
    ws_conc["A2"] = txt_periodo
    ws_conc["A2"].font = font_subtitle
    ws_conc["A3"] = "CONCENTRADO DE INGRESOS Y GASTOS"
    ws_conc["A3"].font = font_subtitle

    for i, h in enumerate(["No.", "SITIO", "SUPERVISOR"], start=1):
        cell = ws_conc.cell(5, i, h)
        cell.fill = fill_green; cell.font = font_head; cell.alignment = align_center; cell.border = thin_border
        ws_conc.merge_cells(start_row=5, start_column=i, end_row=6, end_column=i)

    c_col_conc = 4
    conc_dia_cols = {}
    for d in dias_periodo:
        conc_dia_cols[d] = c_col_conc
        d_str = f"{DIAS_ESP[d.weekday()]} {d.day:02d}/{MESES_ESP[d.month]}/{d.year}"
        cell = ws_conc.cell(5, c_col_conc, d_str)
        cell.fill = fill_green; cell.font = font_head; cell.alignment = align_center
        ws_conc.merge_cells(start_row=5, start_column=c_col_conc, end_row=5, end_column=c_col_conc+2)
        for c in range(c_col_conc, c_col_conc+3):
            ws_conc.cell(5, c).fill = fill_green; ws_conc.cell(5, c).border = thin_border

        for s_i, s_txt in enumerate(["TOTAL INGRESOS X DÍA", "TOTAL GASTOS X DÍA", "UTILIDAD O PÉRDIDA X DÍA"]):
            sc = ws_conc.cell(6, c_col_conc+s_i, s_txt)
            sc.fill = fill_green; sc.font = font_head; sc.alignment = align_center; sc.border = thin_border
        c_col_conc += 3

    col_conc_tot_ing = c_col_conc
    col_conc_tot_gas = c_col_conc + 1
    col_conc_tot_ut = c_col_conc + 2

    c_ct = ws_conc.cell(5, col_conc_tot_ing, "CONCENTRADO DEL PERÍODO")
    c_ct.fill = fill_green; c_ct.font = font_head; c_ct.alignment = align_center
    ws_conc.merge_cells(start_row=5, start_column=col_conc_tot_ing, end_row=5, end_column=col_conc_tot_ut)
    for c in range(col_conc_tot_ing, col_conc_tot_ut+1):
        ws_conc.cell(5, c).fill = fill_green; ws_conc.cell(5, c).border = thin_border

    for s_i, s_txt in enumerate(["TOTAL INGRESOS", "TOTAL GASTOS", "UTILIDAD O PÉRDIDA"]):
        sc = ws_conc.cell(6, col_conc_tot_ing+s_i, s_txt)
        sc.fill = fill_green; sc.font = font_head; sc.alignment = align_center; sc.border = thin_border

    r_c_act = 8
    fila_inicio_c = r_c_act
    for i_suc, item in enumerate(CATALOGO_OFICIAL_23):
        ws_conc.cell(r_c_act, 1, item["num"]).alignment = align_center
        ws_conc.cell(r_c_act, 2, item["sitio"]).alignment = align_left
        ws_conc.cell(r_c_act, 3, item["supervisor"]).alignment = align_center
        for ci in range(1, 4):
            ws_conc.cell(r_c_act, ci).font = font_data; ws_conc.cell(r_c_act, ci).border = thin_border

        fila_det_ing = fila_inicio_ing + i_suc
        fila_det_gas = fila_inicio_gas + i_suc

        for d in dias_periodo:
            cd_conc = conc_dia_cols[d]
            cd_det = dia_cols_map[d]
            let_ing = get_column_letter(cd_det+4)
            let_gas = get_column_letter(cd_det+4)

            ws_conc.cell(r_c_act, cd_conc, f"='DETALLE '!{let_ing}{fila_det_ing}").number_format = "$#,##0.00"
            ws_conc.cell(r_c_act, cd_conc+1, f"='DETALLE '!{let_gas}{fila_det_gas}").number_format = "$#,##0.00"
            let_ci = get_column_letter(cd_conc)
            let_cg = get_column_letter(cd_conc+1)
            ws_conc.cell(r_c_act, cd_conc+2, f"=+{let_ci}{r_c_act}-{let_cg}{r_c_act}").number_format = "$#,##0.00"

            for sub_c in range(cd_conc, cd_conc+3):
                cell = ws_conc.cell(r_c_act, sub_c)
                cell.font = font_data; cell.alignment = align_right; cell.border = thin_border

        let_tin_det = get_column_letter(col_tot_ing)
        let_tga_det = get_column_letter(col_tot_bol)
        ws_conc.cell(r_c_act, col_conc_tot_ing, f"='DETALLE '!{let_tin_det}{fila_det_ing}").number_format = "$#,##0.00"
        ws_conc.cell(r_c_act, col_conc_tot_gas, f"='DETALLE '!{let_tga_det}{fila_det_gas}").number_format = "$#,##0.00"
        let_ti_c = get_column_letter(col_conc_tot_ing)
        let_tg_c = get_column_letter(col_conc_tot_gas)
        ws_conc.cell(r_c_act, col_conc_tot_ut, f"=+{let_ti_c}{r_c_act}-{let_tg_c}{r_c_act}").number_format = "$#,##0.00"

        for sub_c in range(col_conc_tot_ing, col_conc_tot_ut+1):
            cell = ws_conc.cell(r_c_act, sub_c)
            cell.font = font_bold; cell.alignment = align_right; cell.border = thin_border

        r_c_act += 1

    fila_fin_c = r_c_act - 1

    r_c_tot = r_c_act
    ws_conc.cell(r_c_tot, 1, "TOTAL").font = font_bold
    for ci in range(1, col_conc_tot_ut + 1):
        ws_conc.cell(r_c_tot, ci).border = thin_border

    for d in dias_periodo:
        cd_conc = conc_dia_cols[d]
        for sub_i in range(3):
            col_l = get_column_letter(cd_conc + sub_i)
            c = ws_conc.cell(r_c_tot, cd_conc + sub_i, f"=SUM({col_l}{fila_inicio_c}:{col_l}{fila_fin_c})")
            c.font = font_bold; c.alignment = align_right; c.number_format = "$#,##0.00"

    for sub_i in range(3):
        col_l = get_column_letter(col_conc_tot_ing + sub_i)
        c = ws_conc.cell(r_c_tot, col_conc_tot_ing + sub_i, f"=SUM({col_l}{fila_inicio_c}:{col_l}{fila_fin_c})")
        c.font = font_bold; c.alignment = align_right; c.number_format = "$#,##0.00"

    # --- HOJA 3: Rentabilidad y Auditoría ---
    ws_aud["A1"] = "ANÁLISIS DE RENTABILIDAD Y AUDITORÍA INTELIGENTE"
    ws_aud["A1"].font = font_title
    headers_aud = ["ID", "Fecha", "Sitio / Unidad", "Supervisor", "Boletos Cobrados", "Tarifa", "Ingreso Bruto", "Total Gastos", "Ingreso Neto", "Margen (%)", "Estado Auditoría"]
    for c_i, h in enumerate(headers_aud, start=1):
        cell = ws_aud.cell(3, c_i, h)
        cell.fill = fill_green; cell.font = font_head; cell.alignment = align_center; cell.border = thin_border

    r_aud = 4
    for r in registros:
        bruto = float(r.get("Ingreso Bruto", 0.0))
        neto = float(r.get("Ingreso Neto", 0.0))
        margen = (neto / bruto * 100) if bruto > 0 else 0.0
        ws_aud.cell(r_aud, 1, r.get("ID", "")).alignment = align_center
        ws_aud.cell(r_aud, 2, r.get("Fecha", "")).alignment = align_center
        ws_aud.cell(r_aud, 3, r.get("Sitio / Unidad", "")).alignment = align_left
        ws_aud.cell(r_aud, 4, r.get("Elaboró", r.get("Usuario", ""))).alignment = align_center
        ws_aud.cell(r_aud, 5, r.get("Boletos Cobrados", 0)).alignment = align_center
        ws_aud.cell(r_aud, 6, r.get("Costo Boleto", 0)).number_format = "$#,##0.00"
        ws_aud.cell(r_aud, 7, bruto).number_format = "$#,##0.00"
        ws_aud.cell(r_aud, 8, float(r.get("Total Gastos", 0.0))).number_format = "$#,##0.00"
        ws_aud.cell(r_aud, 9, neto).number_format = "$#,##0.00"
        ws_aud.cell(r_aud, 10, f"{margen:.1f}%").alignment = align_center
        ws_aud.cell(r_aud, 11, r.get("Nivel Auditoría", "Conforme")).alignment = align_center
        for ci in range(1, 12):
            ws_aud.cell(r_aud, ci).font = font_data; ws_aud.cell(r_aud, ci).border = thin_border
        r_aud += 1

    # Presentación corporativa: impresión horizontal, títulos legibles y encabezados congelados.
    for ws in (ws_det, ws_conc, ws_aud):
        ws.sheet_view.showGridLines = False
        ws.freeze_panes = "G11" if ws is ws_det else ("D8" if ws is ws_conc else "A4")
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.sheet_properties.pageSetUpPr.autoPageBreaks = False
        ws.sheet_format.defaultRowHeight = 16
    for col, width in {"A": 7, "B": 26, "C": 20, "D": 21, "E": 13, "F": 16}.items():
        ws_det.column_dimensions[col].width = width
    for c in range(7, col_tot_ing + 1):
        ws_det.column_dimensions[get_column_letter(c)].width = 14
    for col, width in {"A": 7, "B": 28, "C": 16}.items():
        ws_conc.column_dimensions[col].width = width
    for c in range(4, col_conc_tot_ut + 1):
        ws_conc.column_dimensions[get_column_letter(c)].width = 18
    for col, width in {"A": 14, "B": 13, "C": 25, "D": 21, "E": 16, "F": 13, "G": 16, "H": 16, "I": 16, "J": 12, "K": 23}.items():
        ws_aud.column_dimensions[col].width = width

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()

def generar_csv_oficial(registros, etiqueta_periodo="SEMANAL"):
    """Genera CSV con el esquema oficial corporativo de corte diario por sucursal y supervisor."""
    if not registros:
        return "No hay registros disponibles para el período seleccionado.\n".encode("utf-8-sig")
    
    filas = []
    for r in registros:
        suc = normalizar_a_sucursal_oficial(r.get("Sitio / Unidad", ""))
        info_cat = next((item for item in CATALOGO_OFICIAL_23 if item["sitio"] == suc), None)
        dias_op = info_cat["dias"] if info_cat else "-"
        horarios = info_cat["horarios"] if info_cat else "-"
        empleados = info_cat["empleados"] if info_cat else 0
        sup_cat = info_cat["supervisor"] if info_cat else (r.get("Elaboró") or r.get("Usuario") or "JESÚS")
        
        bruto = float(r.get("Ingreso Bruto", 0.0))
        nom = float(r.get("Gastos Generales / Nómina", 0.0))
        esp = float(r.get("Gastos Especiales", 0.0))
        tot_g = float(r.get("Total Gastos", nom + esp))
        neto = float(r.get("Ingreso Neto", bruto - tot_g))
        
        filas.append({
            "No.": info_cat["num"] if info_cat else "",
            "SITIO": suc,
            "DÍAS QUE SE OPERA": dias_op,
            "HORARIOS": horarios,
            "No. EMPLEADOS": empleados,
            "SUPERVISOR": sup_cat,
            "FECHA": r.get("Fecha", ""),
            "FOLIO INICIAL": r.get("Folio Inicial", 0),
            "FOLIO FINAL": r.get("Folio Final", 0),
            "TOTAL BOLETOS": r.get("Boletos Cobrados", r.get("Total Boletos", 0)),
            "COSTO X BOLETO": r.get("Costo Boleto", info_cat["tarifa"] if info_cat else 0),
            "INGRESO X DÍA": bruto,
            "NÓMINA LIDER": nom,
            "NÓMINA OPERADOR": 0.0,
            "GASTOS SUPERVISIÓN": 0.0,
            "VARIOS": esp,
            "TOTAL GASTOS X DÍA": tot_g,
            "UTILIDAD O PÉRDIDA X DÍA": neto,
            "ESTADO AUDITORÍA": r.get("Nivel Auditoría", "Conforme"),
            "USUARIO REGISTRO": r.get("Usuario") or r.get("Elaboró") or sup_cat
        })
    df_csv = pd.DataFrame(filas)
    return df_csv.to_csv(index=False).encode("utf-8-sig")


# ---------------------------------------------------------------------
# LECTURA CON IA (Claude con visión): placas de vehículo y reportes físicos
# Requiere el secreto ANTHROPIC_API_KEY (st.secrets o variable de entorno) y `anthropic` en requirements.txt.
# ---------------------------------------------------------------------
def _clave_ia():
    try:
        valor = st.secrets.get("ANTHROPIC_API_KEY", "")
    except Exception:
        valor = ""
    return str(valor or os.environ.get("ANTHROPIC_API_KEY", "")).strip()


def _modelo_ia():
    try:
        valor = st.secrets.get("ANTHROPIC_MODEL", "")
    except Exception:
        valor = ""
    return str(valor or os.environ.get("ANTHROPIC_MODEL", "") or "claude-sonnet-5-5").strip()


def ia_disponible():
    return bool(_clave_ia())


def _imagen_a_base64(archivo):
    imagen = archivo.copy() if isinstance(archivo, Image.Image) else Image.open(archivo)
    imagen = ImageOps.exif_transpose(imagen).convert("RGB")
    imagen.thumbnail((1600, 1600))
    buffer = io.BytesIO()
    imagen.save(buffer, format="JPEG", quality=85)
    return base64.b64encode(buffer.getvalue()).decode("ascii")


def ia_extraer(archivos, instrucciones, esquema, nombre_herramienta):
    """Envía una o varias fotos a Claude y devuelve un dict con los campos del esquema.
    Lanza RuntimeError con un mensaje legible si algo falla."""
    clave = _clave_ia()
    if not clave:
        raise RuntimeError("Falta configurar ANTHROPIC_API_KEY en los secretos de la app.")
    try:
        import anthropic
    except ImportError:
        raise RuntimeError("Falta instalar la librería `anthropic` (agrégala a requirements.txt).")
    contenido = []
    for archivo in archivos:
        if hasattr(archivo, "seek"):
            archivo.seek(0)
        contenido.append({"type": "image", "source": {"type": "base64", "media_type": "image/jpeg", "data": _imagen_a_base64(archivo)}})
    contenido.append({"type": "text", "text": instrucciones})
    try:
        cliente = anthropic.Anthropic(api_key=clave, timeout=60.0)
        respuesta = cliente.messages.create(
            model=_modelo_ia(), max_tokens=2000,
            tools=[{"name": nombre_herramienta, "description": "Registra los datos leídos de la imagen.", "input_schema": esquema}],
            tool_choice={"type": "tool", "name": nombre_herramienta},
            messages=[{"role": "user", "content": contenido}],
        )
    except Exception as exc:
        raise RuntimeError(f"No se pudo consultar la IA: {exc}")
    for bloque in respuesta.content:
        if getattr(bloque, "type", "") == "tool_use":
            return dict(bloque.input)
    raise RuntimeError("La IA no devolvió datos legibles. Intenta con una foto más clara.")


ESQUEMA_VEHICULO = {
    "type": "object",
    "properties": {
        "placas": {"type": "string", "description": "Placas visibles, en mayúsculas, sin espacios extra. Vacío si no se leen."},
        "marca_modelo": {"type": "string", "description": "Marca y modelo si se identifican (ej. NISSAN VERSA). Vacío si no."},
        "color": {"type": "string", "description": "Color principal en español (ej. Gris, Blanco). Vacío si no."},
        "legible": {"type": "boolean", "description": "False si las placas no se alcanzan a leer con certeza."},
    },
    "required": ["placas", "legible"],
}


def ia_leer_vehiculo(archivos):
    return ia_extraer(
        archivos,
        "Estas son fotos de un vehículo que entra a un estacionamiento valet en México. Lee las placas con cuidado "
        "(no inventes caracteres: si dudas, marca legible=false) e identifica marca/modelo y color.",
        ESQUEMA_VEHICULO, "registrar_vehiculo",
    )


ESQUEMA_REPORTE_INGRESOS = {
    "type": "object",
    "properties": {
        "sitio": {"type": "string", "description": "Nombre de la sucursal/sitio tal como aparece en la hoja. Vacío si no aparece."},
        "fecha": {"type": "string", "description": "Fecha de operación en formato DD/MM/AAAA. Vacío si no aparece."},
        "folio_inicial": {"type": "integer", "description": "Folio inicial del boletaje. 0 si no aparece."},
        "folio_final": {"type": "integer", "description": "Folio final del boletaje. 0 si no aparece."},
        "boletos_sin_cobro": {"type": "integer", "description": "Boletos de cortesía / sin cobro. 0 si no aparece."},
        "boletos_cancelados": {"type": "integer", "description": "Boletos cancelados. 0 si no aparece."},
        "boletos_perdidos": {"type": "integer", "description": "Boletos perdidos. 0 si no aparece."},
        "ingreso_bruto": {"type": "number", "description": "Ingreso total del día en pesos si la hoja lo muestra. 0 si no."},
        "nomina": {"type": "number", "description": "Gasto de nómina del día en pesos. 0 si no aparece."},
        "otros_gastos": {"type": "number", "description": "Otros gastos del día en pesos (sin nómina). 0 si no aparece."},
        "dudas": {"type": "string", "description": "Campos ilegibles o dudosos que el usuario debe revisar. Vacío si todo es claro."},
    },
    "required": ["folio_inicial", "folio_final"],
}

ESQUEMA_REPORTE_GASTOS = {
    "type": "object",
    "properties": {
        "sitio": {"type": "string", "description": "Nombre de la sucursal/sitio. Vacío si no aparece."},
        "fecha": {"type": "string", "description": "Fecha en formato DD/MM/AAAA. Vacío si no aparece."},
        "gastos": {
            "type": "array",
            "items": {
                "type": "object",
                "properties": {
                    "categoria": {"type": "string", "description": "Una de: Nómina, Renta, Insumos, Mantenimiento, Comisiones, Servicios, Otros."},
                    "concepto": {"type": "string"},
                    "importe": {"type": "number", "description": "Importe en pesos."},
                },
                "required": ["categoria", "concepto", "importe"],
            },
        },
        "dudas": {"type": "string", "description": "Renglones ilegibles o dudosos. Vacío si todo es claro."},
    },
    "required": ["gastos"],
}


def ia_leer_reporte(archivos, es_ingresos):
    if es_ingresos:
        return ia_extraer(
            archivos,
            "Es la hoja física de control de un día de un sitio de valet parking en México. Extrae folio inicial y final del boletaje, "
            "boletos sin cobro, cancelados y perdidos, y los gastos del día (nómina y otros) si aparecen. No inventes números: "
            "si algo no se lee con certeza déjalo en 0 y descríbelo en 'dudas'.",
            ESQUEMA_REPORTE_INGRESOS, "registrar_reporte_ingresos",
        )
    return ia_extraer(
        archivos,
        "Es una hoja o comprobante de gastos de un sitio de valet parking en México. Extrae cada gasto con su categoría, concepto e importe. "
        "No inventes renglones ni importes: lo ilegible va en 'dudas'.",
        ESQUEMA_REPORTE_GASTOS, "registrar_reporte_gastos",
    )


def _entero_seguro(valor):
    try:
        return max(0, int(float(valor or 0)))
    except (TypeError, ValueError):
        return 0


def _flotante_seguro(valor):
    try:
        return max(0.0, float(valor or 0))
    except (TypeError, ValueError):
        return 0.0


FISICOS_FILE = DATA_DIR / "reportes_fisicos.json"
GASTOS_FILE_RF = DATA_DIR / "gastos_operativos.json"
FOTOS_REPORTES_DIR = DATA_DIR / "fotos_reportes"
FOTOS_REPORTES_DIR.mkdir(parents=True, exist_ok=True)
CATEGORIAS_GASTO_RF = ["Nómina", "Renta", "Insumos", "Mantenimiento", "Comisiones", "Servicios", "Otros"]


def _leer_lista_json(ruta):
    if not ruta.exists():
        return []
    try:
        datos = json.loads(ruta.read_text(encoding="utf-8"))
        return datos if isinstance(datos, list) else []
    except (OSError, json.JSONDecodeError):
        return []


def _guardar_lista_json(ruta, datos):
    temporal = ruta.with_suffix(".tmp")
    temporal.write_text(json.dumps(datos, ensure_ascii=False, indent=2), encoding="utf-8")
    os.replace(temporal, ruta)


# ---------------------------------------------------------------------
# EVENTOS EXTRAORDINARIOS EN SALONES
# Ingreso del evento = costo del evento + vehículos × costo por boleto.
# Utilidad = ingreso − (nómina + otros gastos). Los eventos cancelados no suman.
# ---------------------------------------------------------------------
EVENTOS_FILE = DATA_DIR / "eventos_extraordinarios.json"
SALONES_EVENTOS_BASE = ["JUNGLA DIVISIÓN", "JUNGLA EMPRESA", "JUNGLA PICACHO", "JUNGLA TLALPAN",
                        "MAROMETA SATÉLITE", "MAROMETA MARINA", "ROMANCE PALACE", "SIPIRILY COAPA"]
_salones_cfg = _SUC_CFG.get("salones")
SALONES_EVENTOS = [s for s in (SALONES_EVENTOS_BASE if _salones_cfg is None else _salones_cfg) if s in PRECIOS_SUCURSALES]
ESTADOS_EVENTO = ["Programado", "Realizado", "Cancelado"]


def calcular_evento(ev):
    """Agrega ingreso de boletaje, ingreso total, gastos y utilidad al evento."""
    vehiculos = _entero_seguro(ev.get("vehiculos"))
    costo_boleto = _flotante_seguro(ev.get("costo_boleto"))
    costo_evento = _flotante_seguro(ev.get("costo_evento"))
    gastos = _flotante_seguro(ev.get("nomina")) + _flotante_seguro(ev.get("otros_gastos"))
    boletaje = vehiculos * costo_boleto
    ingreso = costo_evento + boletaje
    ev.update({"vehiculos": vehiculos, "costo_boleto": costo_boleto, "costo_evento": costo_evento,
               "ingreso_boletaje": round(boletaje, 2), "ingreso_total": round(ingreso, 2),
               "gastos_total": round(gastos, 2), "utilidad": round(ingreso - gastos, 2)})
    return ev


def cargar_eventos():
    return [calcular_evento(e) for e in _leer_lista_json(EVENTOS_FILE)]


def complementar_eventos_con_reportes(eventos):
    """Agrega a cada evento lo REAL de los reportes físicos del salón y día.
    - Reporte por evento (Evento k de n): se asigna al evento k (en orden de hora de inicio).
    - Reporte global del día: se reparte entre los eventos del día según su número de eventos.
    - Gastos: nómina y otros gastos del reporte de ingresos + reportes de gastos del salón ese día.
    Sin reporte, el evento se queda con lo programado (fuente «Programado»)."""
    reportes = _leer_lista_json(FISICOS_FILE)
    ingresos_dia, gastos_dia = {}, {}
    for r in reportes:
        dia = fecha_flexible(r.get("Fecha"))
        clave = (r.get("Sitio / Unidad"), dia)
        if r.get("tipo_reporte") == "ingresos":
            ingresos_dia.setdefault(clave, []).append(r)
            gastos_dia[clave] = gastos_dia.get(clave, 0.0) + _flotante_seguro(r.get("Nómina")) + _flotante_seguro(r.get("Otros gastos"))
        elif r.get("tipo_reporte") == "gastos":
            gastos_dia[clave] = gastos_dia.get(clave, 0.0) + _flotante_seguro(r.get("Total Gastos"))
    activos = [e for e in eventos if e.get("estado") != "Cancelado"]
    por_dia = {}
    for e in activos:
        por_dia.setdefault((e.get("sucursal"), fecha_flexible(e.get("fecha"))), []).append(e)
    salida = []
    for e in eventos:
        e = dict(e)
        e.update({"vehiculos_est": e["vehiculos"], "boletaje_est": e["ingreso_boletaje"], "ingreso_est": e["ingreso_total"],
                  "gastos_est": e["gastos_total"], "utilidad_est": e["utilidad"], "fuente": "Programado", "reportes_ids": ""})
        clave = (e.get("sucursal"), fecha_flexible(e.get("fecha")))
        reps = ingresos_dia.get(clave, [])
        if e.get("estado") == "Cancelado" or not reps:
            salida.append(e)
            continue
        hermanos = sorted(por_dia.get(clave, [e]), key=lambda x: (x.get("hora_inicio", ""), x.get("id", "")))
        # Lugares (números de evento) que le tocan a este registro: uno por cada evento que agrupa
        lugares, siguiente = {}, 1
        for h in hermanos:
            n_h = max(1, _entero_seguro(h.get("num_eventos")))
            lugares[h.get("id")] = list(range(siguiente, siguiente + n_h))
            siguiente += n_h
        total_eventos = siguiente - 1
        mis_lugares = set(lugares.get(e.get("id"), []))
        por_evento = [r for r in reps if _entero_seguro(r.get("evento_num"))]
        globales = [r for r in reps if not _entero_seguro(r.get("evento_num"))]
        if por_evento:
            mios = [r for r in por_evento if _entero_seguro(r.get("evento_num")) in mis_lugares]
            parte = len(mios) / max(1, len(por_evento))
            bol = sum(_flotante_seguro(r.get("Ingreso Bruto")) for r in mios)
            veh = sum(_entero_seguro(r.get("Boletos Cobrados")) for r in mios)
            gas = sum(_flotante_seguro(r.get("Nómina")) + _flotante_seguro(r.get("Otros gastos")) for r in mios)
            # gastos sueltos del día (reportes de gastos) se reparten por número de eventos
            gas_sueltos = gastos_dia.get(clave, 0.0) - sum(_flotante_seguro(r.get("Nómina")) + _flotante_seguro(r.get("Otros gastos")) for r in reps)
            gas += max(0.0, gas_sueltos) * len(mis_lugares) / max(1, total_eventos)
            ids = [r.get("ID") for r in mios]
            if not mios:
                salida.append(e)
                continue
        else:
            parte = len(mis_lugares) / max(1, total_eventos)
            bol = sum(_flotante_seguro(r.get("Ingreso Bruto")) for r in globales) * parte
            veh = round(sum(_entero_seguro(r.get("Boletos Cobrados")) for r in globales) * parte)
            gas = gastos_dia.get(clave, 0.0) * parte
            ids = [r.get("ID") for r in globales]
        if veh == 0 and bol > 0 and e.get("costo_boleto"):
            veh = round(bol / float(e["costo_boleto"]))  # salones con ingreso directo: vehículos estimados por el boletaje
        ingreso_real = e["costo_evento"] + bol
        e.update({"vehiculos": int(veh), "ingreso_boletaje": round(bol, 2), "ingreso_total": round(ingreso_real, 2),
                  "gastos_total": round(gas, 2), "utilidad": round(ingreso_real - gas, 2),
                  "fuente": "Reporte físico" + (" (global repartido)" if globales and not por_evento else ""),
                  "reportes_ids": ", ".join(str(i) for i in ids)})
        salida.append(e)
    return salida


def guardar_eventos(eventos):
    _guardar_lista_json(EVENTOS_FILE, sorted(eventos, key=lambda e: (e.get("fecha", ""), e.get("hora_inicio", ""))))


# ---------------------------------------------------------------------
# CONTROL DE DUPLICADOS EN EL INGRESO DE VEHÍCULOS
# ---------------------------------------------------------------------
PLACAS_GENERICAS = {"SINPLACA", "SINPLACAS", "SP", "NA", "NOAPLICA", "SINPLACAS0", "PERMISO", "0"}
VENTANA_DUPLICADO_MIN = 3  # minutos: mismo vehículo registrado dos veces casi al mismo tiempo


def normalizar_placa(valor):
    return re.sub(r"[^A-Z0-9]", "", str(valor or "").upper())


def buscar_duplicado_ingreso(registros, sucursal, placas, marca="", color=""):
    """Devuelve (registro, motivo) si el vehículo ya está registrado; (None, "") si no.
    1) Placa real que ya está En Resguardo (en cualquier sucursal).
    2) Mismo vehículo en la misma sucursal registrado hace menos de VENTANA_DUPLICADO_MIN minutos
       (doble toque al botón o doble captura), incluido el caso de autos sin placa."""
    placa_n = normalizar_placa(placas)
    marca_n, color_n = normalizar_placa(marca), normalizar_placa(color)
    generica = placa_n in PLACAS_GENERICAS or len(placa_n) < 3
    ahora_d = ahora_mx()
    for r in registros:
        if not generica and r.get("estado", "En Resguardo") == "En Resguardo" and normalizar_placa(r.get("placas")) == placa_n:
            return r, "ya está en resguardo"
        if r.get("sucursal") != sucursal:
            continue
        try:
            entrada_r = datetime.fromisoformat(str(r.get("fecha_entrada_iso") or ""))
        except ValueError:
            continue
        if abs((ahora_d - entrada_r).total_seconds()) > VENTANA_DUPLICADO_MIN * 60:
            continue
        mismo = (normalizar_placa(r.get("placas")) == placa_n) if not generica else (
            normalizar_placa(r.get("marca")) == marca_n and normalizar_placa(r.get("color")) == color_n)
        if mismo:
            return r, f"se registró hace menos de {VENTANA_DUPLICADO_MIN} minutos"
    return None, ""


# ---------------------------------------------------------------------
# MERCADO PAGO: link y QR de cobro (Checkout Pro) + verificación del pago
# Requiere el secreto MP_ACCESS_TOKEN (Access Token de producción de la cuenta
# de Club Park). El folio del vehículo viaja como referencia externa.
# ---------------------------------------------------------------------
METODO_MP = "Mercado Pago"


def _mp_token():
    try:
        valor = st.secrets.get("MP_ACCESS_TOKEN", "")
    except Exception:
        valor = ""
    return str(valor or os.environ.get("MP_ACCESS_TOKEN", "")).strip()


def mp_disponible():
    return bool(_mp_token())


def _mp_api(metodo, ruta_api, datos=None):
    import urllib.error
    import urllib.request
    token = _mp_token()
    if not token:
        raise RuntimeError("Falta configurar MP_ACCESS_TOKEN en los secretos de la app.")
    peticion = urllib.request.Request(
        "https://api.mercadopago.com" + ruta_api,
        data=json.dumps(datos).encode("utf-8") if datos is not None else None,
        method=metodo,
        headers={"Authorization": f"Bearer {token}", "Content-Type": "application/json",
                 "X-Idempotency-Key": uuid.uuid4().hex},
    )
    try:
        with urllib.request.urlopen(peticion, timeout=25) as resp:
            return json.loads(resp.read().decode("utf-8") or "{}")
    except urllib.error.HTTPError as exc:
        detalle = exc.read().decode("utf-8", "replace")[:250]
        raise RuntimeError(f"Mercado Pago respondió {exc.code}: {detalle}")
    except Exception as exc:
        raise RuntimeError(f"No se pudo conectar con Mercado Pago: {exc}")


def mp_crear_cobro(folio, sucursal, importe, minutos_vigencia=30):
    """Crea un link de pago por el importe exacto. Devuelve {'pref_id', 'link', 'importe', 'creado'}."""
    if float(importe) <= 0:
        raise RuntimeError("El importe debe ser mayor a $0 para cobrar con Mercado Pago.")
    ahora_p = datetime.now(_TZ_MX)
    respuesta = _mp_api("POST", "/checkout/preferences", {
        "items": [{"id": str(folio), "title": f"Estacionamiento Club Park · {sucursal} · folio {folio}",
                   "quantity": 1, "unit_price": round(float(importe), 2), "currency_id": "MXN"}],
        "external_reference": str(folio),
        "statement_descriptor": "CLUB PARK",
        "expires": True,
        "expiration_date_from": ahora_p.isoformat(timespec="milliseconds"),
        "expiration_date_to": (ahora_p + timedelta(minutes=minutos_vigencia)).isoformat(timespec="milliseconds"),
    })
    link = respuesta.get("init_point")
    if not link:
        raise RuntimeError("Mercado Pago no devolvió el link de pago.")
    return {"pref_id": respuesta.get("id"), "link": link, "importe": round(float(importe), 2),
            "creado": ahora_mx().strftime("%Y-%m-%d %H:%M:%S")}


def mp_buscar_pago(folio):
    """Busca el pago más reciente con la referencia del folio. Devuelve el aprobado si existe; si no, el último."""
    respuesta = _mp_api("GET", "/v1/payments/search?" + urllib.parse.urlencode(
        {"external_reference": str(folio), "sort": "date_created", "criteria": "desc", "limit": 20}))
    pagos_mp = respuesta.get("results") or []
    aprobado = next((p for p in pagos_mp if p.get("status") == "approved"), None)
    return aprobado or (pagos_mp[0] if pagos_mp else None)


def _borrar_archivos(rutas):
    borradas = 0
    for ruta_b in rutas or []:
        try:
            p_b = Path(ruta_b)
            if not p_b.is_absolute():
                p_b = DATA_DIR / p_b
            if p_b.exists() and p_b.is_file():
                p_b.unlink()
                borradas += 1
        except (OSError, TypeError, ValueError):
            pass
    return borradas


def borrar_reporte_fisico(id_reporte):
    """Borra un reporte físico, los gastos que generó y sus fotos. Solo el Contador lo usa."""
    reportes_b = _leer_lista_json(FISICOS_FILE)
    reporte_b = next((r for r in reportes_b if r.get("ID") == id_reporte), None)
    if not reporte_b:
        return None
    _guardar_lista_json(FISICOS_FILE, [r for r in reportes_b if r.get("ID") != id_reporte])
    gastos_b = _leer_lista_json(GASTOS_FILE_RF)
    _guardar_lista_json(GASTOS_FILE_RF, [g for g in gastos_b if g.get("reporte_fisico_id") != id_reporte])
    if reporte_b.get("origen") != "Cierre del sistema":  # las fotos del cierre siguen siendo del cierre
        _borrar_archivos(reporte_b.get("fotos"))
    return reporte_b


def calcular_boletos_block100(f_ini, f_fin):
    """Conteo inclusivo y cíclico de folios cuyo consecutivo se reinicia cada 100 (595→303 = 9 boletos)."""
    inicio, fin = int(f_ini), int(f_fin)
    if inicio < 0 or fin < 0 or (inicio == 0 and fin == 0):
        return 0
    return ((fin % 100 - inicio % 100) % 100) + 1


def evaluar_rentabilidad(margen):
    if margen >= 40.0:
        return "🟢 Alta"
    if margen >= 15.0:
        return "🟡 Moderada"
    if margen > 0.0:
        return "🟠 Baja"
    return "🔴 Pérdida"


def es_sucursal_directa(sucursal):
    """Sucursales que reportan ingreso bruto directo, sin control por folio."""
    cfg = tarifa_de(sucursal)
    return cfg.get("tipo") == "tiempo" or float(cfg.get("precio_fijo", 0) or 0) == 0 or sucursal == "E-BNI"


def fecha_flexible(valor):
    for formato in ("%d/%m/%Y %H:%M:%S", "%Y-%m-%d %H:%M:%S", "%d/%m/%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(str(valor or ""), formato).date()
        except ValueError:
            continue
    return None


def _limpiar_formulario_rf():
    """Borra de la sesión todos los valores del formulario de reportes físicos y abre uno nuevo."""
    for k in [k for k in list(st.session_state.keys()) if str(k).startswith("rf_") and k not in ("rf_nonce", "rf_apartado",
              "rf_excel_bytes", "rf_excel_nombre", "rf_aviso_guardado", "rf_filtro_suc", "rf_ver_detalle")]:
        st.session_state.pop(k, None)
    st.session_state.rf_nonce = int(st.session_state.get("rf_nonce", 0)) + 1


def guardar_fotos_reporte(archivos, id_reporte):
    rutas = []
    for i, archivo in enumerate(archivos, start=1):
        imagen = archivo.copy() if isinstance(archivo, Image.Image) else Image.open(archivo)
        imagen = ImageOps.exif_transpose(imagen).convert("RGB")
        imagen.thumbnail((1600, 1600))
        ruta = FOTOS_REPORTES_DIR / f"{id_reporte}_{i}.jpg"
        imagen.save(ruta, format="JPEG", quality=80, optimize=True)
        rutas.append(str(ruta))
    return rutas


def _fila_flujo(suc, dia, bruto=0.0, cobrados=0, costo=0.0, f_ini=0, f_fin=0, total=None, elaboro="", usuario="", rid=""):
    return {
        "ID": rid, "Fecha": dia.strftime("%d/%m/%Y"), "Sitio / Unidad": suc,
        "Costo Boleto": float(costo), "Folio Inicial": int(f_ini), "Folio Final": int(f_fin),
        "Total Boletos": int(cobrados if total is None else total), "Sin Cobro": 0, "Cancelados": 0, "Perdidos": 0,
        "Boletos Cobrados": int(cobrados), "Ingreso Bruto": float(bruto),
        "Gastos Especiales": 0.0, "Gastos Generales / Nómina": 0.0, "Total Gastos": 0.0, "Ingreso Neto": float(bruto),
        "Elaboró": elaboro, "Usuario": usuario, "Nivel Auditoría": "Conforme", "Tipo": "Operación Regular",
    }


def registros_flujo_periodo(desde, hasta, sucursales):
    """Une reportes físicos de ingresos, cobros digitales y gastos en el esquema del Excel oficial.
    Si una sucursal/fecha tiene reporte físico de ingresos, ese manda; si no, se usan los cobros digitales."""
    permitidas = set(sucursales)
    filas = {}
    for r in _leer_lista_json(FISICOS_FILE):
        if r.get("tipo_reporte") != "ingresos":
            continue
        dia = fecha_flexible(r.get("Fecha"))
        suc = r.get("Sitio / Unidad")
        if dia and suc in permitidas and desde <= dia <= hasta:
            fila = _fila_flujo(suc, dia, r.get("Ingreso Bruto", 0), r.get("Boletos Cobrados", 0), r.get("Costo Boleto", 0),
                               r.get("Folio Inicial", 0), r.get("Folio Final", 0), r.get("Total Boletos"),
                               r.get("Elaboró", ""), r.get("Usuario", ""), r.get("ID", ""))
            for k in ("Sin Cobro", "Cancelados", "Perdidos"):
                fila[k] = int(r.get(k, 0) or 0)
            previa = filas.get((suc, dia))
            if previa:  # salón con un reporte por evento: se suman en un solo renglón del día
                for k in ("Ingreso Bruto", "Ingreso Neto", "Boletos Cobrados", "Total Boletos", "Sin Cobro", "Cancelados", "Perdidos"):
                    previa[k] = (previa.get(k) or 0) + (fila.get(k) or 0)
                ini_v = [x for x in (previa.get("Folio Inicial"), fila.get("Folio Inicial")) if x]
                previa["Folio Inicial"] = min(ini_v) if ini_v else 0
                previa["Folio Final"] = max(previa.get("Folio Final") or 0, fila.get("Folio Final") or 0)
                previa["ID"] = f"{previa.get('ID', '')}+{fila.get('ID', '')}"
                previa["Tipo"] = "Eventos (suma de reportes)"
            else:
                filas[(suc, dia)] = fila
    digitales = {}
    for reg in st.session_state.get("vehiculos", []):
        for pago in reg.get("pagos", []) or []:
            if pago.get("estado") != "Pagado":
                continue
            dia = fecha_flexible(pago.get("fecha_hora"))
            suc = pago.get("sucursal") or reg.get("sucursal")
            if dia and suc in permitidas and desde <= dia <= hasta:
                acum = digitales.setdefault((suc, dia), [0.0, 0])
                acum[0] += dinero(pago.get("importe"))
                acum[1] += 1
    for clave, (importe, n) in digitales.items():
        if clave not in filas:
            suc, dia = clave
            filas[clave] = _fila_flujo(suc, dia, importe, n, (importe / n) if n else 0, usuario="Cobro digital")
    for g in _leer_lista_json(GASTOS_FILE_RF):
        dia = fecha_flexible(g.get("fecha"))
        suc = g.get("sucursal")
        if not (dia and suc in permitidas and desde <= dia <= hasta):
            continue
        fila = filas.setdefault((suc, dia), _fila_flujo(suc, dia, usuario="Gastos"))
        monto = dinero(g.get("importe"))
        if g.get("categoria") == "Nómina":
            fila["Gastos Generales / Nómina"] += monto
        else:
            fila["Gastos Especiales"] += monto
    resultado = []
    for fila in filas.values():
        fila["Total Gastos"] = fila["Gastos Generales / Nómina"] + fila["Gastos Especiales"]
        fila["Ingreso Neto"] = fila["Ingreso Bruto"] - fila["Total Gastos"]
        resultado.append(fila)
    return sorted(resultado, key=lambda f: (fecha_flexible(f["Fecha"]), f["Sitio / Unidad"]))



# ---------------------------------------------------------------------
# REPORTES POR SERVICIO (SUCURSAL) Y GLOBAL: semana / quincena / mes
# ---------------------------------------------------------------------
NOMBRES_DIA = ["LUNES", "MARTES", "MIÉRCOLES", "JUEVES", "VIERNES", "SÁBADO", "DOMINGO"]
COLUMNAS_NO_MONEDA = {"Fecha", "Día", "Sucursal", "Categoría", "Días con venta", "Días del periodo", "Margen %", "Sin información",
                      "N° de día", "Día actual", "Día anterior", "Fecha actual", "Fecha anterior"}


def rango_periodo(tipo, referencia, desde_pers=None, hasta_pers=None):
    """Devuelve (desde, hasta) para Semana (lun-dom), Quincena (1-15 / 16-fin), Mes o Personalizado."""
    if tipo == "Semana":
        lunes = referencia - timedelta(days=referencia.weekday())
        return lunes, lunes + timedelta(days=6)
    primer_dia_mes = referencia.replace(day=1)
    primer_dia_sig = (primer_dia_mes + timedelta(days=32)).replace(day=1)
    ultimo_dia_mes = primer_dia_sig - timedelta(days=1)
    if tipo == "Quincena":
        if referencia.day <= 15:
            return primer_dia_mes, referencia.replace(day=15)
        return referencia.replace(day=16), ultimo_dia_mes
    if tipo == "Mes":
        return primer_dia_mes, ultimo_dia_mes
    return desde_pers, hasta_pers


def construir_reporte(desde, hasta, sucursales):
    """Estado de resultados por sucursal: vendió, gastos, utilidad; venta diaria; promedio por día de la semana."""
    sucursales = list(sucursales)
    registros = registros_flujo_periodo(desde, hasta, sucursales)
    dias = [desde + timedelta(days=i) for i in range((hasta - desde).days + 1)]
    ventas, gastos = {}, {}
    for r in registros:
        clave = (r["Sitio / Unidad"], fecha_flexible(r["Fecha"]))
        ventas[clave] = ventas.get(clave, 0.0) + float(r["Ingreso Bruto"])
        gastos[clave] = gastos.get(clave, 0.0) + float(r["Total Gastos"])

    filas = []
    for suc in sucursales:
        v = sum(ventas.get((suc, d), 0.0) for d in dias)
        g = sum(gastos.get((suc, d), 0.0) for d in dias)
        dias_venta = sum(1 for d in dias if ventas.get((suc, d), 0.0) > 0)
        filas.append({
            "Sucursal": suc, "Ventas": v, "Gastos": g, "Utilidad": v - g,
            "Margen %": round((v - g) / v * 100, 1) if v else 0.0,
            "Días con venta": dias_venta,
            "Promedio por día": (v / dias_venta) if dias_venta else 0.0,
        })
    resumen = pd.DataFrame(filas, columns=["Sucursal", "Ventas", "Gastos", "Utilidad", "Margen %", "Días con venta", "Promedio por día"])
    v_tot, g_tot = float(resumen["Ventas"].sum()), float(resumen["Gastos"].sum())
    dias_con_venta_global = sum(1 for d in dias if any(ventas.get((suc, d), 0.0) > 0 for suc in sucursales))
    fila_total = {
        "Sucursal": "TOTAL", "Ventas": v_tot, "Gastos": g_tot, "Utilidad": v_tot - g_tot,
        "Margen %": round((v_tot - g_tot) / v_tot * 100, 1) if v_tot else 0.0,
        "Días con venta": dias_con_venta_global,
        "Promedio por día": (v_tot / dias_con_venta_global) if dias_con_venta_global else 0.0,
    }
    resumen_total = pd.concat([resumen, pd.DataFrame([fila_total])], ignore_index=True)

    sucursales_con_venta = [suc for suc in sucursales if any(ventas.get((suc, d), 0.0) > 0 for d in dias)] or sucursales
    filas_dia, filas_util = [], []
    acumulado = 0.0
    for d in dias:
        fila = {"Fecha": d.strftime("%d/%m/%Y"), "Día": NOMBRES_DIA[d.weekday()]}
        for suc in sucursales_con_venta:
            fila[suc] = ventas.get((suc, d), 0.0)
        v_dia = sum(ventas.get((suc, d), 0.0) for suc in sucursales)
        g_dia = sum(gastos.get((suc, d), 0.0) for suc in sucursales)
        fila["TOTAL"] = v_dia
        filas_dia.append(fila)
        acumulado += v_dia - g_dia
        filas_util.append({"Fecha": d.strftime("%d/%m/%Y"), "Día": NOMBRES_DIA[d.weekday()], "Ventas": v_dia,
                           "Gastos": g_dia, "Utilidad": v_dia - g_dia, "Utilidad acumulada": acumulado})
    diario_sucursal = pd.DataFrame(filas_dia)
    utilidad_diaria = pd.DataFrame(filas_util)

    filas_sem = []
    for idx, nombre in enumerate(NOMBRES_DIA):
        fechas_idx = [d for d in dias if d.weekday() == idx]
        total_dia = sum(fila["Ventas"] for fila, d in zip(filas_util, dias) if d.weekday() == idx)
        filas_sem.append({"Día": nombre, "Días del periodo": len(fechas_idx), "Ventas": total_dia,
                          "Promedio por día": (total_dia / len(fechas_idx)) if fechas_idx else 0.0})
    por_dia_semana = pd.DataFrame(filas_sem)

    acum_cat = {}
    for g in _leer_lista_json(GASTOS_FILE_RF):
        dia = fecha_flexible(g.get("fecha"))
        if dia and desde <= dia <= hasta and g.get("sucursal") in sucursales:
            clave = (g.get("sucursal"), g.get("categoria") or "Otros")
            acum_cat[clave] = acum_cat.get(clave, 0.0) + dinero(g.get("importe"))
    gastos_categoria = pd.DataFrame(
        [{"Sucursal": k[0], "Categoría": k[1], "Importe": v} for k, v in sorted(acum_cat.items())],
        columns=["Sucursal", "Categoría", "Importe"],
    )
    return {"resumen": resumen, "resumen_total": resumen_total, "diario_sucursal": diario_sucursal,
            "utilidad_diaria": utilidad_diaria, "por_dia_semana": por_dia_semana, "gastos_categoria": gastos_categoria,
            "ventas_total": v_tot, "gastos_total": g_tot, "sucursales_con_venta": sucursales_con_venta}


def rango_anterior(tipo, desde, hasta):
    """Periodo inmediato anterior: semana-7 días, quincena anterior, mes anterior o igual duración."""
    if tipo == "Semana":
        return desde - timedelta(days=7), hasta - timedelta(days=7)
    if tipo == "Quincena":
        if desde.day == 1:
            fin_prev = desde - timedelta(days=1)
            return fin_prev.replace(day=16), fin_prev
        return desde.replace(day=1), desde.replace(day=15)
    if tipo == "Mes":
        fin_prev = desde - timedelta(days=1)
        return fin_prev.replace(day=1), fin_prev
    dias = (hasta - desde).days + 1
    fin_prev = desde - timedelta(days=1)
    return fin_prev - timedelta(days=dias - 1), fin_prev


def _variacion_pct(actual, anterior):
    return round((actual - anterior) / anterior * 100, 1) if anterior else None


def construir_comparativo(rep_act, rep_ant):
    """Compara dos reportes (actual vs anterior): por sucursal y día a día."""
    a = rep_act["resumen_total"].set_index("Sucursal")
    b = rep_ant["resumen_total"].set_index("Sucursal")
    filas = []
    for suc in a.index:
        va, ga, ua = float(a.loc[suc, "Ventas"]), float(a.loc[suc, "Gastos"]), float(a.loc[suc, "Utilidad"])
        vb = float(b.loc[suc, "Ventas"]) if suc in b.index else 0.0
        gb = float(b.loc[suc, "Gastos"]) if suc in b.index else 0.0
        ub = float(b.loc[suc, "Utilidad"]) if suc in b.index else 0.0
        filas.append({
            "Sucursal": suc, "Ventas actual": va, "Ventas anterior": vb, "Var. ventas $": va - vb,
            "Var. ventas %": _variacion_pct(va, vb), "Gastos actual": ga, "Gastos anterior": gb,
            "Utilidad actual": ua, "Utilidad anterior": ub, "Var. utilidad $": ua - ub,
            "Var. utilidad %": _variacion_pct(ua, ub),
        })
    resumen = pd.DataFrame(filas)
    da, db = rep_act["utilidad_diaria"], rep_ant["utilidad_diaria"]
    filas_d = []
    for i in range(max(len(da), len(db))):
        fa = da.iloc[i] if i < len(da) else None
        fb = db.iloc[i] if i < len(db) else None
        va = float(fa["Ventas"]) if fa is not None else 0.0
        vb = float(fb["Ventas"]) if fb is not None else 0.0
        filas_d.append({
            "N° de día": i + 1,
            "Día actual": fa["Día"] if fa is not None else "", "Fecha actual": fa["Fecha"] if fa is not None else "",
            "Ventas actual": va,
            "Día anterior": fb["Día"] if fb is not None else "", "Fecha anterior": fb["Fecha"] if fb is not None else "",
            "Ventas anterior": vb, "Variación $": va - vb,
        })
    return {"resumen": resumen, "diario": pd.DataFrame(filas_d)}


def generar_excel_reporte(rep, titulo, periodo_txt, comparativo=None):
    """Excel con Resumen, Venta diaria, Utilidad diaria, Por día de la semana y Gastos por categoría."""
    hojas = [("Resumen", rep["resumen_total"])]
    if len(rep["sucursales_con_venta"]) > 1:
        hojas.append(("Venta diaria por sucursal", rep["diario_sucursal"]))
    hojas += [("Utilidad diaria", rep["utilidad_diaria"]), ("Por día de la semana", rep["por_dia_semana"]),
              ("Gastos por categoría", rep["gastos_categoria"])]
    if comparativo:
        hojas += [("Comparativo", comparativo["resumen"]), ("Comparativo diario", comparativo["diario"])]
    salida = io.BytesIO()
    verde = PatternFill(start_color="0B3B24", end_color="0B3B24", fill_type="solid")
    with pd.ExcelWriter(salida, engine="openpyxl") as writer:
        for nombre, df in hojas:
            if df.empty:
                df = pd.DataFrame({"Sin información": ["No hay datos en el periodo"]})
            df.to_excel(writer, sheet_name=nombre, index=False, startrow=3)
            ws = writer.sheets[nombre]
            ws["A1"] = titulo
            ws["A1"].font = Font(size=14, bold=True, color="0B3B24")
            ws["A2"] = periodo_txt
            ws["A2"].font = Font(size=10, italic=True)
            for celda in ws[4]:
                celda.fill = verde
                celda.font = Font(bold=True, color="FFFFFF")
                celda.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            for i, col in enumerate(df.columns, start=1):
                letra = get_column_letter(i)
                ancho = max([len(str(col))] + [len(str(v)) for v in df[col].head(300)]) + 3
                ws.column_dimensions[letra].width = min(max(ancho, 12), 40)
                for celda in list(ws[letra])[4:]:
                    if isinstance(celda.value, (int, float)) and str(col).endswith("%"):
                        celda.number_format = '0.0"%"'
                    elif isinstance(celda.value, (int, float)) and col not in COLUMNAS_NO_MONEDA:
                        celda.number_format = '"$"#,##0.00'
            for fila in ws.iter_rows(min_row=5):
                if str(fila[0].value) == "TOTAL":
                    for celda in fila:
                        celda.font = Font(bold=True)
            ws.freeze_panes = "A5"
    return salida.getvalue()


def formato_columnas_df(df):
    config = {}
    for col in df.columns:
        if str(col).endswith("%"):
            config[col] = st.column_config.NumberColumn(col, format="%.1f%%")
        elif col not in COLUMNAS_NO_MONEDA:
            config[col] = st.column_config.NumberColumn(col, format="$%.2f")
    return config


try:
    limpiar_fotos_vencidas()  # fotos con más de 7 días (corre como máximo cada 6 horas)
except Exception:
    pass  # la limpieza nunca debe impedir que la app abra

PESTANAS = [
    ("ing", "📥 Ingreso y Fotografías"),
    ("list", "📋 Vehículos Registrados"),
    ("salida", "🚘 Salida y Cobro"),
    ("balance", "📈 Balances y Rentabilidad"),
    ("imp", "🖨️ Re-Impresión Térmica"),
    ("rep", "📊 Reportes y Excel"),
    ("corte", "💵 Cortes de Caja"),
    ("cierre_diario", "📒 Cierre diario por sucursal"),
    ("aud", "🕵️ Auditoría y WhatsApp"),
    ("qrs", "🏷️ Credenciales QR Institucionales"),
    ("pruebas", "🧪 Pruebas Operativas"),
    ("intel", "🧠 Inteligencia Operativa"),
    ("reportes", "📄 Reportes físicos y Excel"),
    ("eventos", "🎉 Eventos extraordinarios"),
    ("tarifas", "💲 Tarifas"),
]
# El Líder de Sucursal solo opera su sucursal: ingreso, vehículos, salida y cobro,
# reimpresión, reportes y Excel, cortes de caja y cierre diario. Los reportes físicos son solo
# para supervisores y el Contador.
PESTANAS_LIDER = {"ing", "list", "salida", "imp", "rep", "corte", "cierre_diario"}
_pestanas_visibles = [(k, e) for k, e in PESTANAS if not es_lider or k in PESTANAS_LIDER]
_claves_sec = [k for k, _ in _pestanas_visibles]
_etiquetas_sec = dict(_pestanas_visibles)
if "nav_seccion" not in st.session_state or st.session_state.nav_seccion not in _claves_sec:
    _sec_url = str(st.query_params.get("sec", "") or "")
    if st.query_params.get("qrscan") and "salida" in _claves_sec:
        _sec_url = "salida"  # al volver del lector QR, abrir directo Salida y Cobro
    st.session_state.nav_seccion = _sec_url if _sec_url in _claves_sec else _claves_sec[0]
_SECCION = st.radio("Sección", _claves_sec, format_func=lambda k: _etiquetas_sec[k], horizontal=True,
                    key="nav_seccion", label_visibility="collapsed")
if st.query_params.get("sec") != _SECCION:
    st.query_params["sec"] = _SECCION  # al recargar la página se conserva la sección
st.markdown("<hr style='margin:4px 0 10px 0;border-color:#0B3B24'>", unsafe_allow_html=True)
import contextlib
_contenedor_sec = contextlib.nullcontext()  # las secciones se dibujan aquí mismo, debajo del menú
_fragmento = getattr(st, "fragment", None) or (lambda f: f)  # versiones viejas de Streamlit: sin fragmentos
_TABS = {k: _contenedor_sec for k in _claves_sec}
tab_ing, tab_list, tab_salida, tab_imp, tab_rep, tab_corte, tab_cierre_diario, tab_reportes = (
    _TABS["ing"], _TABS["list"], _TABS["salida"], _TABS["imp"], _TABS["rep"], _TABS["corte"], _TABS["cierre_diario"], _TABS.get("reportes"))
tab_balance, tab_aud, tab_qrs, tab_pruebas, tab_intel, tab_tarifas = (
    _TABS.get("balance"), _TABS.get("aud"), _TABS.get("qrs"), _TABS.get("pruebas"), _TABS.get("intel"), _TABS.get("tarifas"))


# ---------------------------------------------------------------------
# TAB 1: INGRESO
# ---------------------------------------------------------------------
# Sin fragmento: las cámaras y la subida de fotos fallan dentro de fragmentos de Streamlit
def _seccion_ing():
    global archivo, color, custodia, exc, faltantes, folio, indice, marca, placas, qr_data, registro, rutas, sucursal, tarifa, ultimo, vista
    globals()["_CONTADOR_LLAVES"] = 1000  # llaves estables entre recargas completas y de sección
    with tab_ing:
        encabezado_institucional("REGISTRO DE INGRESO DE VEHÍCULO", "OPERACIÓN 1 · RECEPCIÓN", "Protocolo institucional de recepción con evidencia fotográfica obligatoria")
        if st.session_state.pop("limpiar_ingreso", False):
            st.session_state.ing_nonce = int(st.session_state.get("ing_nonce", 0)) + 1
            for _clave_l in ["ing_placas", "ing_marca", "ing_color", "ing_custodia_texto", "ing_sin_valores",
                             "foto_pertenencias_trasera", "fotos_pertenencias_archivos", "ia_veh_cam", "ia_veh_gal",
                             "foto_veh_frente", "foto_veh_trasera", "foto_veh_izquierda", "foto_veh_derecha",
                             "foto_veh_tablero", "foto_veh_interior"]:
                st.session_state.pop(_clave_l, None)
            for _clave_l in [k for k in list(st.session_state.keys()) if str(k).startswith((
                    "ing_fotos_", "ing_retomas_", "ing_marca_", "ing_color_", "ing_placas_", "ing_sin_valores_",
                    "ing_custodia_texto_", "foto_veh_", "foto_pertenencias_", "fotos_pertenencias_", "ia_veh_",
                    "selector_camara_", "modo_fotos_vehiculo_", "retomar_"))]:
                st.session_state.pop(_clave_l, None)
        _ing_n = int(st.session_state.get("ing_nonce", 0))  # al cambiar, cámaras y archivos quedan vacíos
        with st.expander("🤖 Leer placas y datos del vehículo con IA (opcional)", expanded=False):
            st.caption("Toma una foto donde se vean las placas (y de preferencia el auto completo). La IA llena placas, marca y color; tú los revisas antes de registrar.")
            if not ia_disponible():
                st.info("La lectura con IA aún no está activada: falta el secreto ANTHROPIC_API_KEY.")
            else:
                modo_ia_v = st.radio("Origen de la foto", ["Elegir…", "📸 Cámara", "📁 Galería"], horizontal=True, key=f"ia_veh_modo_{_ing_n}")
                fotos_ia_v = []
                if modo_ia_v.startswith("Elegir"):
                    pass
                elif modo_ia_v.startswith("📸"):
                    foto_ia_v = st.camera_input("Foto del vehículo", key=f"ia_veh_cam_{_ing_n}")
                    fotos_ia_v = [foto_ia_v] if foto_ia_v else []
                else:
                    fotos_ia_v = list(st.file_uploader("Foto(s) del vehículo", type=["jpg", "jpeg", "png", "webp"], accept_multiple_files=True, key=f"ia_veh_gal_{_ing_n}") or [])
                if fotos_ia_v and st.button("🤖 Leer con IA y prellenar", key="ia_veh_btn", type="primary"):
                    try:
                        with st.spinner("Leyendo la foto…"):
                            lectura = ia_leer_vehiculo(fotos_ia_v)
                        placas_ia = re.sub(r"[^A-Z0-9-]", "", str(lectura.get("placas", "")).upper())
                        if placas_ia:
                            st.session_state[f"ing_placas_{_ing_n}"] = placas_ia
                        if str(lectura.get("marca_modelo", "")).strip():
                            st.session_state[f"ing_marca_{_ing_n}"] = str(lectura["marca_modelo"]).strip().upper()
                        if str(lectura.get("color", "")).strip():
                            st.session_state[f"ing_color_{_ing_n}"] = str(lectura["color"]).strip().capitalize()
                        guardar_evento_auditoria("LECTURA IA VEHICULO", f"placas leídas: {placas_ia or 'ninguna'}")
                        if lectura.get("legible") and placas_ia:
                            st.success(f"Placas leídas: {placas_ia}. Verifícalas contra el vehículo.")
                        else:
                            st.warning("La IA no pudo leer las placas con certeza. Captúralas a mano o toma otra foto más cerca.")
                    except RuntimeError as exc:
                        st.error(str(exc))
        c1, c2 = st.columns(2)
        with c1:
            if es_global:
                sucursal = st.selectbox("Sucursal operativa", ["-- Selecciona la Sucursal --"] + SUCURSALES)
            else:
                sucursal = st.session_state.sucursal_asignada
                st.info(f"Sucursal asignada: **{sucursal}**")
            marca = st.text_input("Marca y modelo", key=f"ing_marca_{_ing_n}")
            color = st.text_input("Color", key=f"ing_color_{_ing_n}")
        with c2:
            placas = st.text_input("Placas / folio de placa", key=f"ing_placas_{_ing_n}").strip().upper()
            cfg_tarifa = tarifa_de(sucursal)
            mods_ingreso = modalidades_disponibles(cfg_tarifa)
            modalidad_tarifa = mods_ingreso[0]
            if len(mods_ingreso) > 1:
                modalidad_tarifa = st.radio(
                    "Tarifa a aplicar", mods_ingreso,
                    format_func=lambda m: etiqueta_modalidad(cfg_tarifa, m),
                    key=f"ing_modalidad_{sucursal}",
                )
            tarifa = precio_inicial(cfg_tarifa, modalidad_tarifa)
            if modalidad_tarifa == "fija":
                st.metric("Tarifa fija de la sucursal", f"${tarifa:,.2f} MXN")
            else:
                st.metric(ETIQUETAS_MODALIDAD[modalidad_tarifa], f"${tarifa:,.2f} MXN")
                st.caption(etiqueta_modalidad(cfg_tarifa, modalidad_tarifa) + ". El total se calcula al dar salida.")
            operador = st.text_input("Valet / operador", value=st.session_state.usuario_autenticado)
        st.subheader("Observaciones y custodia")
        sin_valores = st.checkbox("⚡ Sin objetos de valor declarados", key=f"ing_sin_valores_{_ing_n}")
        if sin_valores:
            custodia = "SIN OBJETOS DE VALOR DECLARADOS"
            st.success("Se registrará: SIN OBJETOS DE VALOR DECLARADOS")
        else:
            custodia = st.text_area(
                "Objetos declarados y observaciones de custodia",
                placeholder="Ej. computadora portátil, mochila negra, cargador...",
                key=f"ing_custodia_texto_{_ing_n}",
            )
        st.caption("Marca la opción rápida o describe los objetos que el cliente deja bajo resguardo.")
        foto_pertenencias_camara = None
        fotos_pertenencias_archivos = []
        if sin_valores:
            st.caption("📷 Sin objetos de valor: no se requieren fotos de pertenencias.")
        else:
            st.markdown("**📷 Fotografías de pertenencias bajo resguardo (opcional)**")
            st.caption("Puedes capturar con la cámara del celular o adjuntar fotos desde la galería.")

            opc_cam_p = st.radio(
                "Herramienta de captura:",
                ["Sin fotos", "📸 Cámara del dispositivo", "📁 Subir fotos desde galería"],
                horizontal=True,
                key=f"selector_camara_pertenencias_v7_{_ing_n}"
            )
            foto_pertenencias_camara = None
            fotos_pertenencias_archivos = []

            if opc_cam_p == "Sin fotos":
                pass
            elif opc_cam_p.startswith("📸"):
                foto_pertenencias_camara = camara_trasera_foto_pertenencias(f"foto_pertenencias_trasera_{_ing_n}", height=430)
                if foto_pertenencias_camara:
                    st.image(foto_pertenencias_camara, caption="Fotografía de pertenencias capturada", width=260)
            else:
                fotos_pertenencias_archivos = st.file_uploader(
                    "Selecciona fotos de pertenencias", type=["jpg", "jpeg", "png"],
                    accept_multiple_files=True, key=f"fotos_pertenencias_archivos_{_ing_n}"
                )
                if fotos_pertenencias_archivos:
                    st.caption(f"Fotos seleccionadas: {len(fotos_pertenencias_archivos)}")

        st.subheader("2. Fotografías del vehículo (protocolo de 4 o 6 fotos)")
        st.caption("Las vistas fotográficas son OBLIGATORIAS antes de registrar el ingreso: protegen al operador y a la empresa frente a reclamaciones de daños o pérdida de pertenencias. Elige el protocolo completo (6) o el rápido (4) según la operación.")
        opc_cam_v = st.radio(
            "Herramienta de captura del vehículo:",
            ["📸 Cámara del dispositivo (Recomendado)", "📁 Subir fotos desde galería"],
            horizontal=True,
            key=f"selector_camara_vehiculo_v6_{_ing_n}",
        )
        st.caption("Se abre una cámara a la vez y pasa sola a la siguiente vista. Para máxima nitidez (zoom a golpes), usa «📁 Subir fotos» y elige «Cámara» en el teléfono: toma la foto en resolución completa.")

        VISTAS_COMPLETAS = [
            ("frente", "1. Vista frontal"),
            ("trasera", "2. Vista trasera"),
            ("izquierda", "3. Costado izquierdo"),
            ("derecha", "4. Costado derecho"),
            ("tablero", "5. Tablero / kilometraje"),
            ("interior", "6. Interior / asientos"),
        ]
        VISTAS_RAPIDAS = VISTAS_COMPLETAS[:4]

        modo_fotos = st.radio(
            "Protocolo de evidencia fotográfica:",
            [
                "📋 Protocolo completo (6 fotos): frente, trasera, ambos costados, tablero/kilometraje e interior",
                "⚡ Protocolo rápido (4 fotos): frente, trasera y ambos costados",
            ],
            index=0,
            key=f"modo_fotos_vehiculo_{_ing_n}",
            help="El protocolo completo es el oficial: documenta el kilometraje y el interior. El rápido agiliza el ingreso en horas de alta demanda.",
        )
        VISTAS_VEHICULO = VISTAS_COMPLETAS if modo_fotos.startswith("📋") else VISTAS_RAPIDAS
        total_fotos = len(VISTAS_VEHICULO)

        fotos_vehiculo = {}
        if opc_cam_v.startswith("📸"):
            almacen_f = st.session_state.setdefault(f"ing_fotos_{_ing_n}", {})      # vista -> bytes ya capturados
            retomas_f = st.session_state.setdefault(f"ing_retomas_{_ing_n}", {})    # vista -> número de retoma
            claves_validas = dict(VISTAS_VEHICULO)
            tomadas = [(c, t) for c, t in VISTAS_VEHICULO if c in almacen_f]
            if tomadas:
                cols_t = st.columns(min(len(tomadas), 3))
                for i_t, (c_t, t_t) in enumerate(tomadas):
                    with cols_t[i_t % len(cols_t)]:
                        st.image(almacen_f[c_t], caption=f"✅ {t_t}", use_container_width=True)
                        if st.button("↺ Retomar", key=f"retomar_{c_t}_{_ing_n}_{retomas_f.get(c_t, 0)}", use_container_width=True):
                            almacen_f.pop(c_t, None)
                            retomas_f[c_t] = retomas_f.get(c_t, 0) + 1
                            st.rerun()
            pendientes_f = [(c, t) for c, t in VISTAS_VEHICULO if c not in almacen_f]
            if pendientes_f:
                c_act, t_act = pendientes_f[0]
                st.markdown(f"#### 📸 {t_act} · foto {len(VISTAS_VEHICULO) - len(pendientes_f) + 1} de {len(VISTAS_VEHICULO)}")
                captura_f = st.camera_input(f"{t_act} (obligatoria)",
                                            key=f"foto_veh_{c_act}_{_ing_n}_{retomas_f.get(c_act, 0)}")
                if captura_f is not None:
                    almacen_f[c_act] = captura_f.getvalue()
                    st.rerun()  # pasa sola a la siguiente vista
            else:
                st.success("✅ Todas las fotos del vehículo están tomadas. Puedes retomar cualquiera.")
            fotos_vehiculo = {c: io.BytesIO(b) for c, b in almacen_f.items() if c in claves_validas}
        else:
            for clave_vista, titulo_vista in VISTAS_VEHICULO:
                fotos_vehiculo[clave_vista] = st.file_uploader(
                    f"{titulo_vista} (obligatoria)", type=["jpg", "jpeg", "png"], key=f"foto_veh_{clave_vista}_{_ing_n}"
                )
        faltan_fotos = [titulo_vista for clave_vista, titulo_vista in VISTAS_VEHICULO if not fotos_vehiculo.get(clave_vista)]
        capturadas = total_fotos - len(faltan_fotos)
        st.progress(capturadas / total_fotos, text=f"Fotografías del vehículo capturadas: {capturadas} de {total_fotos}")
        if faltan_fotos:
            st.info("Faltan las siguientes vistas: " + ", ".join(faltan_fotos))

        duplicado_vivo, motivo_vivo = (None, "")
        if placas and not (es_global and sucursal == "-- Selecciona la Sucursal --"):
            duplicado_vivo, motivo_vivo = buscar_duplicado_ingreso(cargar_bd_rapido(), sucursal, placas, marca, color)
        if duplicado_vivo:
            st.error(f"🚫 Este vehículo {motivo_vivo}: folio **{duplicado_vivo.get('id')}** · {duplicado_vivo.get('sucursal')} · "
                     f"placas {duplicado_vivo.get('placas')} · entrada {duplicado_vivo.get('fecha_entrada') or duplicado_vivo.get('fecha')}. "
                     "No se puede registrar dos veces. Si necesitas el ticket, reimprímelo en 🖨️ Re-Impresión Térmica.")
        if st.button("💾 Registrar ingreso y generar ticket", type="primary", use_container_width=True, disabled=bool(duplicado_vivo)):
            faltantes = []
            if not es_global and not sucursal:
                faltantes.append("sucursal asignada")
            if es_global and sucursal == "-- Selecciona la Sucursal --":
                faltantes.append("sucursal")
            if not marca.strip():
                faltantes.append("marca y modelo")
            if not placas:
                faltantes.append("placas")
            if not operador.strip():
                faltantes.append("operador")
            faltantes.extend(faltan_fotos)

            if faltantes:
                st.error("Completa los siguientes campos: " + ", ".join(faltantes) + ".")
            else:
                # Se relee la base del disco: incluye lo que registraron otros teléfonos y un posible doble toque
                st.session_state.vehiculos = cargar_bd()
                duplicado, motivo_dup = buscar_duplicado_ingreso(st.session_state.vehiculos, sucursal, placas, marca, color)
            if not faltantes and duplicado:
                st.error(f"🚫 No se registró: este vehículo {motivo_dup} con el folio {duplicado.get('id')}. "
                         "No se generó un ticket nuevo.")
                guardar_evento_auditoria("INGRESO DUPLICADO BLOQUEADO",
                                         f"{sucursal} · placas {placas} · folio existente {duplicado.get('id')} · {motivo_dup}")
            elif not faltantes:
                folio, folio_numero = asignar_folio(sucursal, st.session_state.vehiculos)
                ahora_entrada = ahora_mx()
                fecha = ahora_entrada.strftime("%d/%m/%Y %H:%M:%S")
                fecha_entrada_iso = ahora_entrada.isoformat(timespec="seconds")
                rutas = {}
                archivos = [
                    (clave_vista, fotos_vehiculo[clave_vista])
                    for clave_vista, _ in VISTAS_VEHICULO
                    if fotos_vehiculo.get(clave_vista)
                ]
                try:
                    for vista, archivo in archivos:
                        rutas[vista] = guardar_foto(archivo, folio, vista)
                    fotos_custodia = []
                    if foto_pertenencias_camara:
                        fotos_custodia.append(guardar_foto(foto_pertenencias_camara, folio, "pertenencia_camara"))
                    for indice, archivo_extra in enumerate(fotos_pertenencias_archivos or [], start=1):
                        fotos_custodia.append(guardar_foto(archivo_extra, folio, f"pertenencia_{indice}"))
                    registro = {
                        "id": folio,
                        "folio_numero": folio_numero,
                        "qr_payload": f"CLUBPARK|VEHICULO|{folio}",
                        "fotos_requeridas": total_fotos,
                        "fecha": fecha,
                        "fecha_entrada": fecha,
                        "fecha_entrada_iso": fecha_entrada_iso,
                        "hora_entrada": ahora_entrada.strftime("%H:%M:%S"),
                        "fecha_salida": "",
                        "hora_salida": "",
                        "salida_fecha": "",
                        "sucursal": sucursal,
                        "marca": marca.strip(),
                        "color": color.strip(),
                        "placas": placas,
                        "tarifa": float(tarifa),
                        "modalidad_tarifa": modalidad_tarifa,
                        "operador": operador.strip(),
                        "custodia": custodia.strip() or "Sin objetos declarados",
                        "fotos": rutas,
                        "fotos_custodia": fotos_custodia,
                        "estado": "En Resguardo",
                        "registrado_por": st.session_state.usuario_autenticado,
                    }
                    registro["ticket_texto"] = crear_ticket(registro)
                    st.session_state.vehiculos.insert(0, registro)
                    if guardar_bd(st.session_state.vehiculos):
                        st.session_state.ultimo_ticket = registro["ticket_texto"]
                        st.session_state.ultimo_folio = folio
                        st.session_state.limpiar_ingreso = True
                        guardar_evento_auditoria(
                            "ALTA DE VEHICULO",
                            f"Folio {folio} · {sucursal} · {marca} placas {placas} · ${tarifa:,.2f} · operador {operador or 'no indicado'}",
                        )
                        st.success(f"Ingreso registrado. Folio: {folio}")
                        st.rerun()
                except Exception as exc:
                    st.error(f"No se pudieron guardar las fotografías o el registro: {exc}")

        if st.session_state.get("ultimo_ticket"):
            st.divider()
            st.subheader(f"Ticket generado: {st.session_state.get('ultimo_folio', '')}")
            ultimo = next((r for r in st.session_state.vehiculos if r.get("id") == st.session_state.get("ultimo_folio")), None)
            if ultimo:
                if QR_OK:
                    qr_data = generar_qr_bytes(ultimo)
                    st.image(qr_data, caption="QR único del vehículo (identifica el folio; no es QR de pago)", width=180)
                    st.download_button(
                        "⬇️ Descargar QR del vehículo",
                        data=qr_data,
                        file_name=f"QR_{ultimo.get('id')}.png",
                        mime="image/png",
                        key="descargar_qr_ingreso",
                    )
                else:
                    st.warning("Instala qrcode[pil] para generar imágenes QR.")
                vista_impresion_con_qr(ultimo, "cliente", "imprimir_cliente_ingreso")
                vista_impresion_con_qr(ultimo, "llavero", "imprimir_contrasena_ingreso")
                st.warning("IMPORTANTE: para imprimir el QR gráfico usa el botón 'Imprimir boleto del cliente + QR de salida' de arriba. El botón Bluetooth de texto solo imprime texto y no puede dibujar el QR. El QR se coloca antes del primer corte.")
                imprimir_partes_ticket(ultimo, "ingreso")
            else:
                st.code(st.session_state.ultimo_ticket, language="text")

if _SECCION == "ing":
    _seccion_ing()


# ---------------------------------------------------------------------
# TAB 2: VEHÍCULOS REGISTRADOS
# ---------------------------------------------------------------------
@_fragmento
def _seccion_list():
    global detalle, exc, p, r, s
    globals()["_CONTADOR_LLAVES"] = 1000  # llaves estables entre recargas completas y de sección
    with tab_list:
        encabezado_institucional("VEHÍCULOS REGISTRADOS", "CONTROL DE RESGUARDO", "Consulta, seguimiento y estado de cada vehículo bajo resguardo")
        if es_contador:
            with st.expander("🧹 Reiniciar sucursal desde cero (solo Contador Auditor)"):
                st.caption("Borra los registros vehiculares de las sucursales elegidas, con sus fotografías y cobros. "
                           "Descarga el respaldo antes de borrar: esta acción no se puede deshacer.")
                todas_rv = st.checkbox("Todas las sucursales", key="rv_todas")
                sucs_rv = list(SUCURSALES) if todas_rv else st.multiselect("Sucursales a reiniciar", SUCURSALES, key="rv_sucs")
                alcance_rv = st.radio("¿Qué se borra?", ["Todos los registros vehiculares (en resguardo y entregados)",
                                                         "Solo los vehículos que siguen en resguardo"], key="rv_alcance")
                solo_resguardo_rv = alcance_rv.startswith("Solo")
                folios_rv = st.checkbox("También reiniciar el folio de ticket a 000001", value=False, key="rv_folios",
                                        disabled=solo_resguardo_rv)
                base_rv = cargar_bd_rapido()
                borrar_rv = [r for r in base_rv if r.get("sucursal") in set(sucs_rv)
                             and (not solo_resguardo_rv or r.get("estado", "En Resguardo") == "En Resguardo")]
                if sucs_rv:
                    resumen_rv = pd.DataFrame([{
                        "Sucursal": s, "Registros a borrar": sum(1 for r in borrar_rv if r.get("sucursal") == s),
                        "En resguardo": sum(1 for r in borrar_rv if r.get("sucursal") == s and r.get("estado", "En Resguardo") == "En Resguardo"),
                    } for s in sucs_rv])
                    st.dataframe(resumen_rv[resumen_rv["Registros a borrar"] > 0] if len(sucs_rv) > 5 else resumen_rv,
                                 use_container_width=True, hide_index=True)
                    st.download_button("⬇️ Descargar respaldo de lo que se va a borrar (JSON)",
                                       data=json.dumps(borrar_rv, ensure_ascii=False, indent=2).encode("utf-8"),
                                       file_name=f"RESPALDO_VEHICULOS_{ahora_mx():%Y%m%d_%H%M}.json", mime="application/json",
                                       key="rv_respaldo", disabled=not borrar_rv)
                palabra_rv = st.text_input("Escribe BORRAR para confirmar", key="rv_confirmar")
                if st.button(f"🧹 Borrar {len(borrar_rv)} registro(s)", type="primary", key="rv_btn",
                             disabled=not (sucs_rv and borrar_rv and palabra_rv.strip().upper() == "BORRAR")):
                    base_rv = cargar_bd()  # al borrar, lectura fresca del disco
                    borrar_rv = [r for r in base_rv if r.get("sucursal") in set(sucs_rv)
                                 and (not solo_resguardo_rv or r.get("estado", "En Resguardo") == "En Resguardo")]
                    ids_rv = {r.get("id") for r in borrar_rv}
                    fotos_rv = 0
                    for r in borrar_rv:
                        fotos_rv += _borrar_archivos(list((r.get("fotos") or {}).values()) + list(r.get("fotos_custodia") or []))
                    conservar_rv = [r for r in base_rv if r.get("id") not in ids_rv]
                    if guardar_bd(conservar_rv):
                        st.session_state.vehiculos = conservar_rv
                        try:
                            pagos_rv = cargar_pagos_respaldo()
                            PAYMENTS_FILE.write_text(json.dumps([p for p in pagos_rv if p.get("folio") not in ids_rv],
                                                                ensure_ascii=False, indent=2), encoding="utf-8")
                        except (OSError, TypeError) as exc:
                            st.warning(f"Los vehículos se borraron, pero no se pudo limpiar el libro de cobros: {exc}")
                        if folios_rv and not solo_resguardo_rv:
                            cfg_fol = cargar_config()
                            for s in sucs_rv:
                                anterior = int(cfg_fol["siguiente_folio"].get(s, 1))
                                cfg_fol["siguiente_folio"][s] = 1
                                cfg_fol.setdefault("historial_folios", []).insert(0, {
                                    "fecha_hora": ahora_mx().strftime("%Y-%m-%d %H:%M:%S"),
                                    "usuario": st.session_state.usuario_autenticado, "sucursal": s,
                                    "anterior": anterior, "nuevo": 1, "motivo": "Reinicio de sucursal desde Vehículos registrados"})
                            guardar_config(cfg_fol)
                        if st.session_state.get("ultimo_folio") in ids_rv:
                            st.session_state.pop("ultimo_ticket", None)
                            st.session_state.pop("ultimo_folio", None)
                        guardar_evento_auditoria("REINICIO DE VEHÍCULOS POR SUCURSAL",
                                                 f"{'TODAS' if todas_rv else ', '.join(sucs_rv)} · {len(ids_rv)} registros · {fotos_rv} fotos · "
                                                 f"{'solo en resguardo' if solo_resguardo_rv else 'todos'}" + (" · folios reiniciados" if folios_rv and not solo_resguardo_rv else ""))
                        st.success(f"Listo: se borraron {len(ids_rv)} registros y {fotos_rv} fotografías.")
                        st.rerun()
        visibles = vehiculos_visibles(st.session_state.vehiculos)

        st.markdown("#### 📊 Control de resguardo")
        hoy_cr = ahora_mx().date()
        cr1, cr2, cr3 = st.columns([1, 1, 2])
        if es_global:
            # Contador y supervisores: cualquier día, semana, quincena o mes
            tipo_cr = cr1.selectbox("Periodo", ["Día", "Semana", "Quincena", "Mes", "Personalizado"], key="cr_tipo")
            ref_cr = cr2.date_input("Día de referencia", value=hoy_cr, key="cr_ref")
            if tipo_cr == "Día":
                desde_cr = hasta_cr = ref_cr
            elif tipo_cr == "Personalizado":
                pcr1, pcr2 = st.columns(2)
                desde_cr = pcr1.date_input("Desde", value=hoy_cr - timedelta(days=6), key="cr_desde")
                hasta_cr = pcr2.date_input("Hasta", value=hoy_cr, key="cr_hasta")
            else:
                desde_cr, hasta_cr = rango_periodo(tipo_cr, ref_cr, None, None)
            st.caption(f"Periodo: {desde_cr:%d/%m/%Y} al {hasta_cr:%d/%m/%Y}")
        else:
            # Líder de sucursal: solo la operación del día, hasta su cierre diario
            desde_cr = hasta_cr = hoy_cr
            cierre_hoy_cr = next((c for c in _leer_lista_json(DATA_DIR / "cierres_diarios_sucursal.json") if c.get("fecha") == str(hoy_cr)
                                  and c.get("sucursal") == st.session_state.sucursal_asignada), None)
            cr1.text_input("Día de operación", value=f"{hoy_cr:%d/%m/%Y}", disabled=True, key="cr_dia_lider")
            cr2.text_input("Cierre diario", value="Guardado" if cierre_hoy_cr else "Pendiente", disabled=True, key="cr_cierre_lider")
        sucs_cr = sorted({r.get("sucursal", "") for r in visibles if r.get("sucursal")})
        if es_global and len(sucs_cr) > 1:
            suc_cr = cr3.selectbox("Sucursal", ["Todas"] + sucs_cr, key="cr_suc")
        else:
            suc_cr = sucs_cr[0] if sucs_cr else "Todas"
            cr3.text_input("Sucursal", value=suc_cr, disabled=True, key="cr_suc_fija")
        en_rango = lambda d: d is not None and desde_cr <= d <= hasta_cr
        base_cr = [r for r in visibles if suc_cr == "Todas" or r.get("sucursal") == suc_cr]
        recibidos_cr = [r for r in base_cr if en_rango(fecha_flexible(r.get("fecha_entrada") or r.get("fecha")))]
        custodia_cr = [r for r in base_cr if r.get("estado", "En Resguardo") == "En Resguardo"]
        entregados_cr = [r for r in base_cr if r.get("estado") == "Entregado" and en_rango(fecha_flexible(r.get("salida_fecha")))]
        cobros_cr = [(r, p) for r in base_cr for p in (r.get("pagos") or [])
                     if p.get("estado") == "Pagado" and en_rango(fecha_flexible(p.get("fecha_hora")))]
        por_metodo = {}
        for _r, p in cobros_cr:
            por_metodo[p.get("metodo", "Efectivo")] = por_metodo.get(p.get("metodo", "Efectivo"), 0.0) + dinero(p.get("importe"))
        efectivo_cr = por_metodo.get("Efectivo", 0.0)
        total_cr = sum(por_metodo.values())
        k1, k2, k3, k4 = st.columns(4)
        k1.metric("🚗 Recibidos", len(recibidos_cr))
        k2.metric("🅿️ En custodia ahora", len(custodia_cr))
        k3.metric("✅ Entregados", len(entregados_cr))
        k4.metric("💵 Efectivo cobrado", f"${efectivo_cr:,.2f}")
        otros_txt = " · ".join(f"{m}: ${v:,.2f}" for m, v in por_metodo.items() if m != "Efectivo" and v)
        st.caption(f"Total cobrado en el periodo: ${total_cr:,.2f}" + (f" ({otros_txt})" if otros_txt else "") +
                   ". «En custodia» cuenta todos los vehículos que siguen dentro, sin importar el día en que entraron.")

        with st.expander("🧮 Cuadre de caja: efectivo contra vehículos entregados", expanded=True):
            folios_cobrados = {r.get("id") for r, _p in cobros_cr}
            sin_cobro_cr = [r for r in entregados_cr if not any(p.get("estado") == "Pagado" for p in (r.get("pagos") or []))]
            entregados_efectivo = [r for r in entregados_cr if any(p.get("estado") == "Pagado" and p.get("metodo") == "Efectivo"
                                                                   for p in (r.get("pagos") or []))]
            q1, q2, q3 = st.columns(3)
            q1.metric("Entregados en efectivo", len(entregados_efectivo))
            q2.metric("Efectivo que debe haber", f"${efectivo_cr:,.2f}")
            contado_cr = q3.number_input("Efectivo contado en caja ($)", min_value=0.0, step=50.0, key="cr_contado")
            if contado_cr > 0:
                dif_cr = round(contado_cr - efectivo_cr, 2)
                if abs(dif_cr) < 0.01:
                    st.success("✅ La caja cuadra: el efectivo contado coincide con los vehículos entregados en efectivo.")
                elif dif_cr > 0:
                    st.warning(f"⚠️ Sobrante de ${dif_cr:,.2f}: hay más efectivo del registrado en cobros.")
                else:
                    st.error(f"🚨 Faltante de ${-dif_cr:,.2f}: hay menos efectivo que los cobros registrados.")
            if sin_cobro_cr:
                st.error(f"🚨 {len(sin_cobro_cr)} vehículo(s) entregados sin cobro registrado: " +
                         ", ".join(str(r.get("id")) for r in sin_cobro_cr[:15]) + ("…" if len(sin_cobro_cr) > 15 else ""))
            elif entregados_cr:
                st.caption(f"✔️ Los {len(entregados_cr)} vehículos entregados tienen su cobro registrado.")
            if entregados_efectivo:
                st.dataframe(pd.DataFrame([{
                    "Folio": r.get("id"), "Sucursal": r.get("sucursal"), "Placas": r.get("placas"),
                    "Salida": r.get("salida_fecha"),
                    "Efectivo": sum(dinero(p.get("importe")) for p in (r.get("pagos") or [])
                                    if p.get("estado") == "Pagado" and p.get("metodo") == "Efectivo"),
                    "Entregó": r.get("entregado_por", ""),
                } for r in entregados_efectivo]), use_container_width=True, hide_index=True,
                    column_config={"Efectivo": st.column_config.NumberColumn("Efectivo", format="$%.2f")})
        st.divider()

        if not es_global:
            visibles = [r for r in visibles if r.get("estado", "En Resguardo") == "En Resguardo"
                        or fecha_flexible(r.get("fecha_entrada") or r.get("fecha")) == hoy_cr
                        or fecha_flexible(r.get("salida_fecha")) == hoy_cr]
            st.caption("Ves los vehículos de hoy y los que siguen en custodia. El historial de días anteriores lo consultan supervisores y Contador.")
        if not visibles:
            st.info("No hay vehículos registrados para tu alcance.")
        else:
            buscar_v = st.text_input("🔍 Buscar por folio, placas o marca", key="lista_buscar").strip().upper()
            if buscar_v:
                visibles = [v for v in visibles if buscar_v in " ".join(str(v.get(c, "")) for c in ("id", "placas", "marca", "color")).upper()]
            visibles = sorted(visibles, key=lambda v: str(v.get("fecha_entrada_iso") or ""), reverse=True)
            if len(visibles) > 300:
                st.caption(f"Se muestran los 300 más recientes de {len(visibles)}. Usa el buscador para encontrar otros.")
            columnas_mostrar = ["id", "fecha", "sucursal", "marca", "placas", "tarifa", "operador", "estado", "pago_estado", "metodo_pago"]
            df_v = pd.DataFrame([{c: v.get(c, "") for c in columnas_mostrar} for v in visibles[:300]])
            if not df_v.empty:
                st.dataframe(df_v, use_container_width=True, hide_index=True)
            if not visibles:
                st.info("No se encontró ningún vehículo con ese texto.")

            folios = [v.get("id", "") for v in visibles[:300]]
            folio_sel = st.selectbox("Selecciona un folio para ver sus datos y fotografías", folios, key="lista_folio")
            detalle = next((v for v in visibles if v.get("id") == folio_sel), None)
            if detalle:
                st.write(
                    f"**Sucursal:** {detalle.get('sucursal', '')} · "
                    f"**Placas:** {detalle.get('placas', '')} · "
                    f"**Color:** {detalle.get('color', '')} · "
                    f"**Custodia:** {detalle.get('custodia', '')}"
                )
                faltan_lista = [etq for v, etq in ETIQUETAS_VISTAS[:4] if not ruta_foto(detalle, v)]
                if detalle.get("fotos_borradas_el"):
                    st.info(f"🗓️ Las fotos de este vehículo se borraron automáticamente el {detalle['fotos_borradas_el'][:10]} "
                            f"(se conservan {DIAS_CONSERVAR_FOTOS} días desde que se toman).")
                    faltan_lista = []
                if faltan_lista:
                    st.warning("Fotos no disponibles: " + ", ".join(faltan_lista))
                visor_zoom_fotos(fotos_evidencia_registro(detalle), key=f"lista_{folio_sel}",
                                 titulo="**📷 Evidencia fotográfica del vehículo**")

                if es_contador:
                    with st.expander("🗑️ Borrar información de este vehículo (solo Contador Auditor)"):
                        fotos_borrables = fotos_evidencia_registro(detalle)
                        if fotos_borrables:
                            st.markdown("**Borrar fotos**")
                            sel_fotos_b = st.multiselect("Fotos a borrar", [e for e, _ in fotos_borrables], key=f"del_fotos_sel_{folio_sel}")
                            if st.button("🗑️ Borrar fotos seleccionadas", disabled=not sel_fotos_b, key=f"del_fotos_btn_{folio_sel}"):
                                rutas_sel = {str(p) for e, p in fotos_borrables if e in sel_fotos_b}
                                base_b = cargar_bd()
                                for r_b in base_b:
                                    if r_b.get("id") == folio_sel:
                                        r_b["fotos"] = {v: rp for v, rp in (r_b.get("fotos") or {}).items()
                                                        if str(_ruta_absoluta(rp) or "") not in rutas_sel}
                                        r_b["fotos_custodia"] = [rp for rp in (r_b.get("fotos_custodia") or [])
                                                                 if str(_ruta_absoluta(rp) or "") not in rutas_sel]
                                n_b = _borrar_archivos(list(rutas_sel))
                                if guardar_bd(base_b):
                                    st.session_state.vehiculos = base_b
                                    guardar_evento_auditoria("FOTOS DE VEHÍCULO BORRADAS", f"Folio {folio_sel} · {', '.join(sel_fotos_b)} · {n_b} archivo(s)")
                                    st.success(f"Se borraron {n_b} foto(s).")
                                    st.rerun()
                        st.markdown("**Borrar el registro completo**")
                        st.caption("Se borra el vehículo, todas sus fotos y sus cobros del libro de pagos.")
                        ok_reg_b = st.checkbox(f"Confirmo borrar el folio {folio_sel}", key=f"del_reg_ok_{folio_sel}")
                        if st.button("🗑️ Borrar registro", disabled=not ok_reg_b, key=f"del_reg_btn_{folio_sel}"):
                            base_b = cargar_bd()
                            reg_b = next((r for r in base_b if r.get("id") == folio_sel), None)
                            if reg_b:
                                n_b = _borrar_archivos(list((reg_b.get("fotos") or {}).values()) + list(reg_b.get("fotos_custodia") or []))
                                base_b = [r for r in base_b if r.get("id") != folio_sel]
                                if guardar_bd(base_b):
                                    st.session_state.vehiculos = base_b
                                    try:
                                        PAYMENTS_FILE.write_text(json.dumps([p for p in cargar_pagos_respaldo() if p.get("folio") != folio_sel],
                                                                            ensure_ascii=False, indent=2), encoding="utf-8")
                                    except OSError:
                                        pass
                                    guardar_evento_auditoria("REGISTRO DE VEHÍCULO BORRADO",
                                                             f"Folio {folio_sel} · {reg_b.get('sucursal')} · placas {reg_b.get('placas')} · {n_b} foto(s)")
                                    st.success(f"Folio {folio_sel} borrado.")
                                    st.rerun()

                if es_global:
                    estados = ["En Resguardo", "Entregado", "Incidencia", "Cancelado"]
                    estado_actual = detalle.get("estado", "En Resguardo")
                    estado_nuevo = st.selectbox(
                        "Cambiar estado",
                        estados,
                        index=estados.index(estado_actual) if estado_actual in estados else 0,
                        key=f"estado_{folio_sel}",
                    )
                    if st.button("Actualizar estado", key=f"actualizar_{folio_sel}"):
                        for r in st.session_state.vehiculos:
                            if r.get("id") == folio_sel:
                                r["estado"] = estado_nuevo
                                r["actualizado_por"] = st.session_state.usuario_autenticado
                                r["actualizado_en"] = ahora_mx().strftime("%d/%m/%Y %H:%M:%S")
                                break
                        if guardar_bd(st.session_state.vehiculos):
                            st.success("Estado actualizado.")
                            st.rerun()

if _SECCION == "list":
    _seccion_list()



# ---------------------------------------------------------------------
# TAB: SALIDA Y COBRO
# ---------------------------------------------------------------------
# Sin fragmento: las cámaras y la subida de fotos fallan dentro de fragmentos de Streamlit
def _seccion_salida():
    global a, ahora, b, exc, pago
    globals()["_CONTADOR_LLAVES"] = 1000  # llaves estables entre recargas completas y de sección
    with tab_salida:
        encabezado_institucional("SALIDA DE VEHÍCULO Y REGISTRO DE COBRO", "OPERACIÓN 2 · ENTREGA", "Cobro verificado, comprobante oficial y entrega del vehículo")
        st.caption("Escanea el QR único del boleto o captura el folio. La salida solo se registra una vez.")
        if st.session_state.pop("limpiar_salida", False):
            st.session_state.pop("folio_buscar_salida", None)
            st.session_state.qr_scan_nonce_salida = int(st.session_state.get("qr_scan_nonce_salida", 0)) + 1
            st.session_state.archivo_qr_nonce = int(st.session_state.get("archivo_qr_nonce", 0)) + 1
            if "qrscan" in st.query_params:
                del st.query_params["qrscan"]
        metodo_busqueda = st.radio(
            "Cómo localizar el vehículo",
            ["Capturar folio", "Escanear QR con cámara", "Subir imagen QR"],
            horizontal=True,
            key="metodo_busqueda_salida",
        )
        valor_busqueda = ""
        if metodo_busqueda == "Capturar folio":
            valor_busqueda = st.text_input("Folio del boleto", key="folio_buscar_salida").strip()
        elif metodo_busqueda == "Escanear QR con cámara":
            valor_query = str(st.query_params.get("qrscan", "") or "").strip()
            st.info("Apunta la cámara al QR del boleto: se lee solo, sin tomar foto. Si no lo detecta, usa «Tomar foto del QR» o sube una imagen.")
            nonce_lector = int(st.session_state.get("qr_scan_nonce_salida", 0))
            valor_leido = lector_qr_camara_trasera(f"lector_salida_trasero_{nonce_lector}", height=430)
            valor_busqueda = str(valor_leido or valor_query or "").strip()
            if valor_busqueda:
                st.success(f"QR leído: {valor_busqueda}")
                if st.button("🧹 Limpiar QR leído", key="limpiar_qr_salida"):
                    st.session_state.qr_scan_nonce_salida = int(st.session_state.get("qr_scan_nonce_salida", 0)) + 1
                    if "qrscan" in st.query_params:
                        del st.query_params["qrscan"]
                    st.rerun()
        else:
            archivo_qr_salida = st.file_uploader("Sube la imagen QR del ticket", type=["png", "jpg", "jpeg"], key=f"archivo_qr_salida_{st.session_state.get('archivo_qr_nonce', 0)}")
            if archivo_qr_salida:
                valor_busqueda = leer_codigo_qr(Image.open(archivo_qr_salida)) or ""

        folio_busqueda = extraer_folio_qr(valor_busqueda)
        registro_salida = next(
            (r for r in st.session_state.vehiculos if str(r.get("id", "")).upper() == str(folio_busqueda).upper()),
            None,
        ) if folio_busqueda else None

        if folio_busqueda and not registro_salida:
            st.error("No se encontró un vehículo con ese folio/QR.")
        if registro_salida and not usuario_es_global() and registro_salida.get("sucursal") != st.session_state.sucursal_asignada:
            st.error("Este vehículo pertenece a otra sucursal. No tienes permiso para cobrarlo ni entregarlo.")
            registro_salida = None

        if registro_salida:
            st.divider()
            a, b = st.columns([2, 1])
            with a:
                st.markdown(f"### Folio {registro_salida.get('id')}")
                st.write(f"**Sucursal:** {registro_salida.get('sucursal')}")
                st.write(f"**Marca/modelo:** {registro_salida.get('marca', '')}")
                st.write(f"**Placas:** {registro_salida.get('placas', '')} · **Color:** {registro_salida.get('color', '')}")
                st.write(f"**Fecha/hora de entrada:** {registro_salida.get('fecha_entrada', registro_salida.get('fecha', ''))}")
                st.write(f"**Fecha/hora de salida:** {registro_salida.get('salida_fecha', 'Pendiente de salida') or 'Pendiente de salida'}")
                st.write(f"**Estado:** {registro_salida.get('estado', 'En Resguardo')}")
            with b:
                if QR_OK:
                    st.image(generar_qr_bytes(registro_salida), caption="QR del vehículo", width=150)
            with st.expander("📷 Evidencia del ingreso: revisa golpes antes de entregar", expanded=False):
                visor_zoom_fotos(fotos_evidencia_registro(registro_salida), key=f"salida_{registro_salida.get('id')}", alto=400)
            if registro_salida.get("estado") == "Entregado":
                st.success(f"Este vehículo ya fue entregado el {registro_salida.get('salida_fecha', 'fecha no disponible')}.")
                st.info("Para evitar cobros duplicados, no se permite registrar otra salida.")
            elif registro_salida.get("estado") != "En Resguardo":
                st.warning("El vehículo no está en estado 'En Resguardo'. Revisa el registro antes de continuar.")
            else:
                cfg_sal = tarifa_de(registro_salida.get("sucursal"))
                modalidad_reg = registro_salida.get("modalidad_tarifa") or ("fija" if cfg_sal["tipo"] == "fija" else "tiempo")
                modalidad_sal = modalidad_reg
                if es_contador:
                    opciones_mod = list(dict.fromkeys([modalidad_reg] + modalidades_disponibles(cfg_sal)))
                    modalidad_sal = st.selectbox(
                        "Modalidad de cobro (solo Contador puede cambiarla)", opciones_mod,
                        format_func=lambda m: etiqueta_modalidad(cfg_sal, m),
                        key=f"modalidad_sal_{registro_salida['id']}",
                    )
                minutos_est = minutos_estancia(registro_salida)
                precio_reg = dinero(registro_salida.get("tarifa", cfg_sal["precio_fijo"])) if modalidad_sal == "fija" and modalidad_reg == "fija" else None
                tarifa_calc = calcular_cobro(cfg_sal, modalidad_sal, minutos_est, precio_reg)
                tarifa_salida = tarifa_calc
                ajuste_manual = False
                motivo_ajuste = ""
                if es_contador:
                    ajuste_manual = st.checkbox("Cambiar el importe manualmente", key=f"ajuste_manual_{registro_salida['id']}")
                    if ajuste_manual:
                        tarifa_salida = st.number_input(
                            "Importe a cobrar ($)", min_value=0.0, value=float(tarifa_calc), step=5.0,
                            key=f"ajuste_importe_{registro_salida['id']}",
                        )
                        motivo_ajuste = st.text_input("Motivo del cambio (obligatorio)", key=f"ajuste_motivo_{registro_salida['id']}")
                st.metric(ETIQUETAS_MODALIDAD.get(modalidad_sal, "Tarifa") + " · a cobrar", f"${tarifa_salida:,.2f} MXN")
                if modalidad_sal != "fija":
                    st.caption(f"Tiempo en sitio: {texto_duracion(minutos_est)} · {etiqueta_modalidad(cfg_sal, modalidad_sal)}. El importe final se recalcula al confirmar.")
                else:
                    st.caption("La tarifa se toma del registro de ingreso. Verifica cualquier cambio autorizado antes de cobrar.")
                opciones_pago = ["Efectivo", "Tarjeta", "QR"] + ([METODO_MP] if mp_disponible() else [])
                forma_pago = st.radio("Método de pago", opciones_pago, horizontal=True, key=f"pago_metodo_{registro_salida['id']}")
                es_mp = forma_pago == METODO_MP
                mp_ok = False
                mp_info = st.session_state.setdefault("mp_cobros", {}).get(registro_salida["id"])
                if es_mp:
                    st.markdown("**💳 Cobro con Mercado Pago (tarjeta de crédito, débito o saldo)**")
                    if mp_info and abs(float(mp_info.get("importe", 0)) - float(tarifa_salida)) >= 0.01 and not mp_info.get("pago"):
                        st.warning(f"El link generado es por ${mp_info['importe']:,.2f} y ahora el importe es ${tarifa_salida:,.2f}. "
                                   "Genera un link nuevo si quieres cobrar el importe actualizado.")
                    if st.button(f"🔗 Generar link y QR de cobro por ${tarifa_salida:,.2f}", key=f"mp_generar_{registro_salida['id']}",
                                 disabled=bool(mp_info and mp_info.get("pago"))):
                        try:
                            with st.spinner("Creando el cobro en Mercado Pago…"):
                                mp_info = mp_crear_cobro(registro_salida["id"], registro_salida.get("sucursal", ""), tarifa_salida)
                            st.session_state["mp_cobros"][registro_salida["id"]] = mp_info
                            guardar_evento_auditoria("LINK MERCADO PAGO", f"Folio {registro_salida['id']} · ${mp_info['importe']:,.2f} · {mp_info.get('pref_id')}")
                        except RuntimeError as exc:
                            st.error(str(exc))
                    if mp_info:
                        if QR_OK:
                            qr_mp = qrcode.make(mp_info["link"])
                            buf_mp = io.BytesIO()
                            qr_mp.save(buf_mp, format="PNG")
                            st.image(buf_mp.getvalue(), width=230,
                                     caption=f"El cliente escanea con la cámara de su celular y paga ${mp_info['importe']:,.2f} (vigencia 30 min)")
                        m_mp1, m_mp2 = st.columns(2)
                        m_mp1.link_button("🌐 Abrir link de pago", mp_info["link"], use_container_width=True)
                        m_mp2.link_button("💬 Enviar link por WhatsApp", "https://wa.me/?text=" + urllib.parse.quote(
                            f"Club Park · Pago de estacionamiento folio {registro_salida['id']}: ${mp_info['importe']:,.2f}. Paga aquí: {mp_info['link']}"),
                            use_container_width=True)
                    # Se puede verificar aunque el link se haya generado en otro teléfono (se busca por folio)
                    if True:
                        if st.button("🔄 Verificar pago", type="primary", key=f"mp_verificar_{registro_salida['id']}"):
                            try:
                                with st.spinner("Consultando Mercado Pago…"):
                                    pago_mp = mp_buscar_pago(registro_salida["id"])
                                if pago_mp and pago_mp.get("status") == "approved":
                                    mp_info = dict(mp_info or {})
                                    mp_info["pago"] = {"id": str(pago_mp.get("id")), "importe": float(pago_mp.get("transaction_amount") or 0),
                                                       "tipo": pago_mp.get("payment_type_id", ""), "fecha": pago_mp.get("date_approved", "")}
                                    st.session_state["mp_cobros"][registro_salida["id"]] = mp_info
                                elif pago_mp:
                                    st.warning(f"El pago está en estado «{pago_mp.get('status')}» ({pago_mp.get('status_detail', '')}). Aún no está aprobado.")
                                else:
                                    st.info("Todavía no hay ningún pago de este folio. Pide al cliente que termine el pago y vuelve a verificar.")
                            except RuntimeError as exc:
                                st.error(str(exc))
                        if mp_info and mp_info.get("pago"):
                            mp_ok = True
                            st.success(f"✅ Pago aprobado por Mercado Pago · operación {mp_info['pago']['id']} · ${mp_info['pago']['importe']:,.2f}")
                referencia_pago = st.text_input("Referencia / autorización (si aplica)", key=f"pago_ref_{registro_salida['id']}")
                nota_pago = st.text_input("Observaciones del cobro", key=f"pago_nota_{registro_salida['id']}")
                if forma_pago == "Efectivo":
                    st.info("El efectivo se registrará como recibido al confirmar la entrega.")
                elif forma_pago == "Tarjeta":
                    st.warning("Registra el pago con la terminal bancaria. Esta versión guarda el comprobante y la referencia capturada, pero no consulta automáticamente al banco.")
                elif es_mp:
                    st.info("La salida se registra solo cuando Mercado Pago confirma el pago: el número de operación se guarda solo.")
                else:
                    st.warning("Cobra con la app/proveedor QR. Esta versión guarda la referencia y la confirmación del operador; no valida automáticamente la aprobación del proveedor.")
                referencia_obligatoria = forma_pago in ("Tarjeta", "QR")
                pago_aprobado_externamente = False
                if referencia_obligatoria:
                    pago_aprobado_externamente = st.checkbox(
                        f"Confirmo que revisé la aprobación en la terminal/app de {forma_pago} y que el importe es correcto.",
                        key=f"pago_externo_aprobado_{registro_salida['id']}",
                    )
                confirmar_pago = st.checkbox(
                    "Confirmo que verifiqué el importe y que el vehículo puede ser entregado.",
                    key=f"confirmar_salida_{registro_salida['id']}",
                )
                if st.button("💳 Registrar cobro y salida", type="primary", use_container_width=True, disabled=not confirmar_pago, key=f"btn_salida_{registro_salida['id']}"):
                    if ajuste_manual and not motivo_ajuste.strip():
                        st.error("Captura el motivo del cambio de importe.")
                    elif referencia_obligatoria and not referencia_pago.strip():
                        st.error("Captura la referencia de autorización del proveedor.")
                    elif referencia_obligatoria and not pago_aprobado_externamente:
                        st.error("No se registró la salida: confirma primero que el proveedor aprobó el pago.")
                    elif es_mp and not mp_ok:
                        st.error("No se registró la salida: Mercado Pago todavía no confirma el pago. Pulsa «Verificar pago».")
                    else:
                        momento_operacion = ahora_mx()
                        ahora = momento_operacion.strftime("%Y-%m-%d %H:%M:%S")
                        minutos_final = minutos_estancia(registro_salida, momento_operacion)
                        if es_mp and mp_ok:
                            tarifa_salida = float(mp_info["pago"]["importe"])  # lo que realmente cobró Mercado Pago
                            referencia_pago = f"MP {mp_info['pago']['id']}"
                        elif not ajuste_manual:
                            tarifa_salida = calcular_cobro(cfg_sal, modalidad_sal, minutos_final, precio_reg)
                        pago = {
                            "id_pago": str(uuid.uuid4()),
                            "fecha_hora": ahora,
                            "fecha_hora_iso": momento_operacion.isoformat(timespec="seconds"),
                            "sucursal": registro_salida.get("sucursal"),
                            "folio": registro_salida.get("id"),
                            "importe": tarifa_salida,
                            "metodo": forma_pago,
                            "estado": "Pagado",
                            "referencia": referencia_pago.strip(),
                            "observaciones": nota_pago.strip(),
                            "usuario": st.session_state.usuario_autenticado,
                            "confirmacion_manual": forma_pago in ("Tarjeta", "QR"),
                            **({"mp_payment_id": mp_info["pago"]["id"], "mp_preferencia": mp_info.get("pref_id"),
                                "mp_tipo_pago": mp_info["pago"].get("tipo", "")} if es_mp and mp_ok else {}),
                            "modalidad_tarifa": modalidad_sal,
                            "minutos_estancia": minutos_final,
                            "importe_calculado": tarifa_calc if ajuste_manual else tarifa_salida,
                            "ajuste_manual": bool(ajuste_manual),
                            "motivo_ajuste": motivo_ajuste.strip(),
                        }
                        registro_salida.setdefault("pagos", []).append(pago)
                        registro_salida["modalidad_tarifa"] = modalidad_sal
                        registro_salida["pago_estado"] = "Pagado"
                        registro_salida["metodo_pago"] = forma_pago
                        registro_salida["importe_cobrado"] = tarifa_salida
                        registro_salida["salida_fecha"] = ahora
                        registro_salida["fecha_salida"] = momento_operacion.strftime("%d/%m/%Y")
                        registro_salida["hora_salida"] = momento_operacion.strftime("%H:%M:%S")
                        registro_salida["salida_fecha_hora_iso"] = momento_operacion.isoformat(timespec="seconds")
                        registro_salida["entregado_por"] = st.session_state.usuario_autenticado
                        registro_salida["estado"] = "Entregado"
                        # Actualiza el ticket para que las reimpresiones posteriores reflejen la salida real.
                        registro_salida["ticket_texto"] = crear_ticket(registro_salida)
                        respaldo_ok = guardar_pago_respaldo(pago, registro_salida)
                        if respaldo_ok and guardar_bd(st.session_state.vehiculos):
                            st.session_state.ultimo_comprobante_cobro = {
                                "pago": pago.copy(),
                                "registro": {
                                    "id": registro_salida.get("id", ""),
                                    "sucursal": registro_salida.get("sucursal", ""),
                                    "marca": registro_salida.get("marca", ""),
                                    "placas": registro_salida.get("placas", ""),
                                },
                            }
                            guardar_evento_auditoria("SALIDA Y COBRO", f"Folio {registro_salida.get('id')} · ${tarifa_salida:.2f} · {forma_pago} · {registro_salida.get('sucursal')} · {modalidad_sal}" + (f" · AJUSTE MANUAL (calculado ${tarifa_calc:.2f}): {motivo_ajuste.strip()}" if ajuste_manual else ""))
                            st.success(f"Salida registrada. Folio {registro_salida.get('id')}; cobro de ${tarifa_salida:,.2f} MXN por {forma_pago}.")
                            st.session_state.limpiar_salida = True
                            st.rerun()


        comprobante_reciente = st.session_state.get("ultimo_comprobante_cobro")
        if comprobante_reciente:
            st.divider()
            st.subheader("🧾 Último comprobante de cobro")
            st.caption("El comprobante se genera después del cobro y la salida. Los cobros también se guardan en un libro independiente de pagos para respaldo. En tarjeta y QR, la aprobación se confirma manualmente en la terminal/app.")
            vista_comprobante_cobro(
                comprobante_reciente.get("pago", {}),
                comprobante_reciente.get("registro", {}),
                "imprimir_comprobante_cobro_reciente",
            )
            pago_reciente = comprobante_reciente.get("pago", {})
            registro_reciente = comprobante_reciente.get("registro", {})
            registro_completo_reciente = next(
                (r for r in vehiculos_visibles(st.session_state.vehiculos) if r.get("id") == registro_reciente.get("id")),
                {},
            ) or {}
            texto_comprobante = "\n".join([
                "       TICKET DE SALIDA",
                f"N° {registro_reciente.get('id', '')}",
                "RUT: CPV-ESTACIONAMIENTOS",
                "CLUB PARK VALET Y",
                "ESTACIONAMIENTOS",
                "--------------------------------",
                f"Placa: {registro_reciente.get('placas', '')}",
                f"Ubicacion: {registro_reciente.get('sucursal', '')}",
                f"Ingreso: {registro_completo_reciente.get('fecha_entrada') or registro_completo_reciente.get('fecha', '') or pago_reciente.get('fecha_hora', '')}",
                f"Pago: {pago_reciente.get('fecha_hora', '')}",
                f"Operador: {pago_reciente.get('usuario', '')}",
                "--------------------------------",
                f"Sucursal: {registro_reciente.get('sucursal', '')}",
                f"Folio: {registro_reciente.get('id', '')}",
                f"Vehiculo: {registro_reciente.get('marca', '')} {registro_reciente.get('placas', '')}",
                f"Fecha: {pago_reciente.get('fecha_hora', '')}",
                f"Total: ${dinero(pago_reciente.get('importe')):,.2f}",
                *([f"Tiempo de estancia: {texto_estancia_pago(pago_reciente)}"] if texto_estancia_pago(pago_reciente) else []),
                "SALDO PENDIENTE: $0.00",
                "--------------------------------",
                f"Metodo: {pago_reciente.get('metodo', '')}",
                f"Estado: {pago_reciente.get('estado', '')}",
                f"Referencia: {pago_reciente.get('referencia', '') or 'No aplica'}",
                f"Atendio: {pago_reciente.get('usuario', '')}",
                "--------------------------------",
                *lineas_tarifa_ticket(registro_reciente.get('sucursal', ''), pago_reciente.get('modalidad_tarifa')),
                "Factura: facturacionclubpark@gmail.com",
                "Conserve este ticket.",
                "Gracias por su preferencia.",
            ])
            st.markdown("**Impresión térmica Bluetooth del comprobante**")
            imprimir_bluetooth(texto_comprobante, f"comprobante_{registro_reciente.get('id', 'reciente')}",
                               contexto={"folio": registro_reciente.get('id', ''), "tipo_documento": "TICKET DE SALIDA"})
            st.download_button(
                "⬇️ Descargar comprobante en texto",
                data="CLUB PARK\nCOMPROBANTE DE COBRO\n" + "\n".join(
                    f"{k}: {v}" for k, v in {
                        "Sucursal": comprobante_reciente.get("registro", {}).get("sucursal", ""),
                        "Folio": comprobante_reciente.get("registro", {}).get("id", ""),
                        "Fecha y hora": comprobante_reciente.get("pago", {}).get("fecha_hora", ""),
                        "Importe": f"${dinero(comprobante_reciente.get('pago', {}).get('importe')):,.2f} MXN",
                        "Tiempo de estancia": texto_estancia_pago(comprobante_reciente.get("pago", {})) or "No disponible",
                        "Método": comprobante_reciente.get("pago", {}).get("metodo", ""),
                        "Estado": comprobante_reciente.get("pago", {}).get("estado", ""),
                        "Referencia": comprobante_reciente.get("pago", {}).get("referencia", "") or "No aplica",
                    }.items()
                ),
                file_name=f"Comprobante_{comprobante_reciente.get('registro', {}).get('id', 'CLUBPARK')}.txt",
                mime="text/plain",
                key="descargar_comprobante_cobro_reciente",
            )

if _SECCION == "salida":
    _seccion_salida()


# ---------------------------------------------------------------------
# TAB: BALANCES Y RENTABILIDAD POR SUCURSAL
# ---------------------------------------------------------------------
@_fragmento
def _seccion_balance():
    global exc, filas, gastos, margen, pago, pagos, reg, registro, salida_dt, suc
    globals()["_CONTADOR_LLAVES"] = 1000  # llaves estables entre recargas completas y de sección
    with tab_balance:
        encabezado_institucional("BALANCES OPERATIVOS Y RENTABILIDAD", "FINANZAS", "Rentabilidad por sucursal, gastos operativos y tendencias")
        st.caption(
            "El conteo de vehículos y los ingresos se calculan desde los registros. "
            "La rentabilidad usa los gastos operativos capturados por el supervisor o auditor."
        )

        COSTOS_FILE = DATA_DIR / "gastos_operativos.json"

        def cargar_gastos_operativos():
            if not COSTOS_FILE.exists():
                return []
            try:
                datos = json.loads(COSTOS_FILE.read_text(encoding="utf-8"))
                return datos if isinstance(datos, list) else []
            except (OSError, json.JSONDecodeError):
                return []

        def guardar_gastos_operativos(datos):
            temporal = COSTOS_FILE.with_suffix(".tmp")
            temporal.write_text(json.dumps(datos, ensure_ascii=False, indent=2), encoding="utf-8")
            os.replace(temporal, COSTOS_FILE)

        if not (es_contador or "Supervisor" in str(st.session_state.get("rol", ""))):
            st.error("No tienes permisos para consultar balances de rentabilidad.")
        else:
            sucursales_balance = SUCURSALES if es_global else [st.session_state.sucursal_asignada]
            colf1, colf2 = st.columns(2)
            with colf1:
                fecha_desde = st.date_input("Desde", value=ahora_mx().date(), key="balance_desde")
            with colf2:
                fecha_hasta = st.date_input("Hasta", value=ahora_mx().date(), key="balance_hasta")

            if fecha_desde > fecha_hasta:
                st.error("La fecha inicial no puede ser posterior a la fecha final.")
            else:
                registros_balance = vehiculos_visibles(st.session_state.vehiculos)

                def fecha_registro(registro):
                    for campo in ("salida_fecha", "fecha"):
                        valor = registro.get(campo, "")
                        for formato in ("%Y-%m-%d %H:%M:%S", "%d/%m/%Y %H:%M:%S", "%d/%m/%Y"):
                            try:
                                return datetime.strptime(valor, formato).date()
                            except (ValueError, TypeError):
                                continue
                    return None

                registros_periodo = []
                for registro in registros_balance:
                    f_ingreso = fecha_registro(registro)
                    try:
                        f_ingreso = datetime.strptime(registro.get("fecha", ""), "%d/%m/%Y %H:%M:%S").date()
                    except (ValueError, TypeError):
                        try:
                            f_ingreso = datetime.strptime(registro.get("fecha", ""), "%Y-%m-%d %H:%M:%S").date()
                        except (ValueError, TypeError):
                            f_ingreso = None
                    if f_ingreso and fecha_desde <= f_ingreso <= fecha_hasta:
                        registros_periodo.append(registro)

                gastos = cargar_gastos_operativos()
                gastos_periodo = []
                for gasto in gastos:
                    if not es_global and gasto.get("sucursal") != st.session_state.sucursal_asignada:
                        continue
                    try:
                        f_gasto = datetime.strptime(gasto.get("fecha", ""), "%Y-%m-%d").date()
                    except (ValueError, TypeError):
                        continue
                    if fecha_desde <= f_gasto <= fecha_hasta:
                        gastos_periodo.append(gasto)

                filas = []
                for suc in sucursales_balance:
                    todos_suc = [r for r in registros_balance if r.get("sucursal") == suc]
                    entradas = 0
                    salidas_periodo = []
                    ingresos = 0.0
                    folios_cobrados = set()

                    for reg in todos_suc:
                        try:
                            ingreso_dt = datetime.strptime(reg.get("fecha", ""), "%d/%m/%Y %H:%M:%S").date()
                        except (ValueError, TypeError):
                            try:
                                ingreso_dt = datetime.strptime(reg.get("fecha", ""), "%Y-%m-%d %H:%M:%S").date()
                            except (ValueError, TypeError):
                                ingreso_dt = None
                        if ingreso_dt and fecha_desde <= ingreso_dt <= fecha_hasta:
                            entradas += 1

                        salida_dt = None
                        try:
                            salida_dt = datetime.strptime(reg.get("salida_fecha", ""), "%Y-%m-%d %H:%M:%S").date()
                        except (ValueError, TypeError):
                            pass
                        if reg.get("estado") == "Entregado" and salida_dt and fecha_desde <= salida_dt <= fecha_hasta:
                            salidas_periodo.append(reg)
                            pagos = [p for p in reg.get("pagos", []) if p.get("estado") == "Pagado"]
                            if pagos:
                                for pago in pagos:
                                    pago_dt = None
                                    try:
                                        pago_dt = datetime.strptime(pago.get("fecha_hora", ""), "%Y-%m-%d %H:%M:%S").date()
                                    except (ValueError, TypeError):
                                        pass
                                    if pago_dt is None or fecha_desde <= pago_dt <= fecha_hasta:
                                        ingresos += dinero(pago.get("importe", 0))
                                        folios_cobrados.add(reg.get("id"))
                            elif reg.get("pago_estado") == "Pagado":
                                ingresos += dinero(reg.get("importe_cobrado", reg.get("tarifa", 0)))
                                folios_cobrados.add(reg.get("id"))

                    gastos_suc = sum(dinero(g.get("importe", 0)) for g in gastos_periodo if g.get("sucursal") == suc)
                    utilidad = ingresos - gastos_suc
                    margen = (utilidad / ingresos * 100) if ingresos > 0 else 0.0
                    filas.append({
                        "Sucursal": suc,
                        "Vehículos ingresados": entradas,
                        "Vehículos que salieron": len(salidas_periodo),
                        "Vehículos cobrados": len(folios_cobrados),
                        "Pendientes de cobro": max(0, len(salidas_periodo) - len(folios_cobrados)),
                        "Ingresos ($)": ingresos,
                        "Gastos operativos ($)": gastos_suc,
                        "Utilidad estimada ($)": utilidad,
                        "Margen (%)": margen,
                    })

                df_balance = pd.DataFrame(filas)
                if not es_global:
                    suc_sel = st.session_state.sucursal_asignada
                    df_balance = df_balance[df_balance["Sucursal"] == suc_sel]

                total_ingresos = float(df_balance["Ingresos ($)"].sum()) if not df_balance.empty else 0
                total_gastos = float(df_balance["Gastos operativos ($)"].sum()) if not df_balance.empty else 0
                total_utilidad = total_ingresos - total_gastos
                vehiculos_ingresados = int(df_balance["Vehículos ingresados"].sum()) if not df_balance.empty else 0
                vehiculos_salieron = int(df_balance["Vehículos que salieron"].sum()) if not df_balance.empty else 0
                vehiculos_cobrados = int(df_balance["Vehículos cobrados"].sum()) if not df_balance.empty else 0
                margen_total = total_utilidad / total_ingresos * 100 if total_ingresos else 0

                k1, k2, k3, k4 = st.columns(4)
                k1.metric("Vehículos ingresados", vehiculos_ingresados)
                k2.metric("Vehículos que salieron", vehiculos_salieron)
                k3.metric("Vehículos cobrados", vehiculos_cobrados)
                k4.metric("Pendientes de cobro", max(0, vehiculos_salieron - vehiculos_cobrados))
                k5, k6, k7 = st.columns(3)
                k5.metric("Ingresos confirmados", f"${total_ingresos:,.2f} MXN")
                k6.metric("Gastos operativos", f"${total_gastos:,.2f} MXN")
                k7.metric("Utilidad estimada", f"${total_utilidad:,.2f} MXN", delta=f"{margen_total:.1f}% margen")

                st.markdown("#### Balance por sucursal")
                st.dataframe(
                    df_balance.style.format({
                        "Ingresos ($)": "${:,.2f}",
                        "Gastos operativos ($)": "${:,.2f}",
                        "Utilidad estimada ($)": "${:,.2f}",
                        "Margen (%)": "{:.1f}%",
                    }),
                    use_container_width=True,
                    hide_index=True,
                )

                if es_contador or "Supervisor" in str(st.session_state.get("rol", "")):
                    st.markdown("#### Registrar gasto operativo")
                    st.caption("El Contador Auditor puede registrar gastos de cualquier sucursal; cada supervisor solo puede registrar gastos de su propia sucursal.")
                    with st.form("form_gasto_operativo", clear_on_submit=True):
                        opciones_gasto = SUCURSALES if es_contador else [st.session_state.sucursal_asignada]
                        suc_gasto = st.selectbox("Sucursal del gasto", opciones_gasto, key="gasto_sucursal")
                        fecha_gasto = st.date_input("Fecha del gasto", value=ahora_mx().date(), key="gasto_fecha")
                        categoria_gasto = st.selectbox(
                            "Categoría",
                            ["Nómina", "Renta", "Insumos", "Mantenimiento", "Comisiones", "Servicios", "Otros"],
                            key="gasto_categoria",
                        )
                        concepto_gasto = st.text_input("Concepto / descripción", key="gasto_concepto")
                        importe_gasto = st.number_input("Importe ($)", min_value=0.01, value=1.0, step=50.0, key="gasto_importe")
                        if st.form_submit_button("Guardar gasto"):
                            if not concepto_gasto.strip():
                                st.error("Escribe una descripción para el gasto.")
                            else:
                                nuevos_gastos = cargar_gastos_operativos()
                                nuevo_gasto = {
                                    "id": uuid.uuid4().hex[:12].upper(),
                                    "fecha": str(fecha_gasto),
                                    "sucursal": suc_gasto,
                                    "categoria": categoria_gasto,
                                    "concepto": concepto_gasto.strip(),
                                    "importe": float(importe_gasto),
                                    "registrado_por": st.session_state.usuario_autenticado,
                                    "registrado_en": ahora_mx().strftime("%Y-%m-%d %H:%M:%S"),
                                }
                                nuevos_gastos.insert(0, nuevo_gasto)
                                try:
                                    guardar_gastos_operativos(nuevos_gastos)
                                    guardar_evento_auditoria("GASTO OPERATIVO", f"{suc_gasto} · {categoria_gasto} · ${importe_gasto:.2f} · {concepto_gasto.strip()}")
                                    st.success("Gasto registrado y añadido al balance.")
                                    st.rerun()
                                except OSError as exc:
                                    st.error(f"No se pudo guardar el gasto: {exc}")

                st.markdown("#### Rentabilidad por sucursal")
                st.caption("Margen = (ingresos confirmados − gastos registrados) ÷ ingresos confirmados. No incluye costos que todavía no se hayan capturado.")
                if not df_balance.empty:
                    st.bar_chart(df_balance.set_index("Sucursal")[["Ingresos ($)", "Gastos operativos ($)", "Utilidad estimada ($)"]])
                    st.download_button(
                        "⬇️ Descargar balance CSV",
                        data=df_balance.to_csv(index=False).encode("utf-8-sig"),
                        file_name=f"balance_clubpark_{fecha_desde}_{fecha_hasta}.csv",
                        mime="text/csv",
                        key="descargar_balance_csv",
                    )

if _SECCION == "balance" and "balance" in _TABS:
    _seccion_balance()


# ---------------------------------------------------------------------
# TAB: CORTES DE CAJA
# ---------------------------------------------------------------------
@_fragmento
def _seccion_corte():
    global corte, exc, pago, reg, salida_dt
    globals()["_CONTADOR_LLAVES"] = 1000  # llaves estables entre recargas completas y de sección
    with tab_corte:
        encabezado_institucional("CORTE DE CAJA", "CONTROL DE EFECTIVO", "Conciliación de cobros en efectivo por sucursal y fecha")
        st.caption("Los totales se calculan con salidas registradas y pagos confirmados. Pagos QR/tarjeta no confirmados no se suman a lo cobrado.")
        registros_entregados = [
            r for r in vehiculos_visibles(st.session_state.vehiculos)
            if r.get("estado") == "Entregado" and r.get("salida_fecha")
        ]
        sucursales_corte = SUCURSALES if es_global else [st.session_state.sucursal_asignada]
        suc_corte = st.selectbox("Sucursal", sucursales_corte, key="corte_sucursal")
        fecha_corte = st.date_input("Fecha del corte", value=ahora_mx().date(), key="fecha_corte")
        inicio_corte = st.text_input("Hora inicial (HH:MM)", value="00:00", key="hora_inicio_corte")
        fin_corte = st.text_input("Hora final (HH:MM)", value="23:59", key="hora_fin_corte")
        try:
            desde_dt = datetime.strptime(f"{fecha_corte:%Y-%m-%d} {inicio_corte}", "%Y-%m-%d %H:%M")
            hasta_dt = datetime.strptime(f"{fecha_corte:%Y-%m-%d} {fin_corte}", "%Y-%m-%d %H:%M")
        except ValueError:
            desde_dt = hasta_dt = None
            st.error("Usa formato de hora HH:MM, por ejemplo 08:00.")
        movimientos_corte = []
        if desde_dt and hasta_dt and desde_dt <= hasta_dt:
            for reg in registros_entregados:
                if reg.get("sucursal") != suc_corte:
                    continue
                try:
                    salida_dt = datetime.strptime(reg.get("salida_fecha", ""), "%Y-%m-%d %H:%M:%S")
                except ValueError:
                    continue
                if not (desde_dt <= salida_dt <= hasta_dt):
                    continue
                for pago in reg.get("pagos", []):
                    if pago.get("estado") == "Pagado":
                        movimientos_corte.append(pago)
        elif desde_dt and hasta_dt:
            st.error("La hora inicial no puede ser posterior a la hora final.")

        efectivo_total = sum(dinero(p.get("importe")) for p in movimientos_corte if p.get("metodo") == "Efectivo")
        qr_total = sum(dinero(p.get("importe")) for p in movimientos_corte if p.get("metodo") == "QR")
        tarjeta_total = sum(dinero(p.get("importe")) for p in movimientos_corte if p.get("metodo") == "Tarjeta")
        mp_total = sum(dinero(p.get("importe")) for p in movimientos_corte if p.get("metodo") == METODO_MP)
        total_cobrado = efectivo_total + qr_total + tarjeta_total + mp_total
        m1, m2, m3 = st.columns(3)
        m1.metric("Efectivo", f"${efectivo_total:,.2f}")
        m2.metric("QR confirmado", f"${qr_total:,.2f}")
        m3.metric("Tarjeta confirmada", f"${tarjeta_total:,.2f}")
        m4, m5 = st.columns(2)
        m4.metric("Mercado Pago", f"${mp_total:,.2f}")
        m5.metric("Total cobrado", f"${total_cobrado:,.2f}")
        st.write(f"**Movimientos confirmados:** {len(movimientos_corte)}")
        if movimientos_corte:
            st.dataframe(pd.DataFrame(movimientos_corte), use_container_width=True, hide_index=True)
        else:
            st.info("No hay pagos confirmados en ese periodo.")

        st.markdown("#### Arqueo de efectivo")
        apertura = st.number_input("Fondo inicial de caja ($)", min_value=0.0, value=0.0, step=50.0, key="fondo_apertura")
        efectivo_contado = st.number_input("Efectivo contado físicamente al cierre ($)", min_value=0.0, value=0.0, step=50.0, key="efectivo_contado")
        efectivo_esperado = apertura + efectivo_total
        diferencia = efectivo_contado - efectivo_esperado
        c1, c2, c3 = st.columns(3)
        c1.metric("Efectivo esperado", f"${efectivo_esperado:,.2f}")
        c2.metric("Efectivo contado", f"${efectivo_contado:,.2f}")
        c3.metric("Diferencia", f"${diferencia:,.2f}")
        if st.button("🔒 Guardar corte de caja", type="primary", key="guardar_corte"):
            corte = {
                "fecha_hora": ahora_mx().strftime("%Y-%m-%d %H:%M:%S"),
                "fecha_operativa": str(fecha_corte),
                "desde": inicio_corte,
                "hasta": fin_corte,
                "sucursal": suc_corte,
                "usuario": st.session_state.usuario_autenticado,
                "efectivo": efectivo_total,
                "qr_confirmado": qr_total,
                "tarjeta_confirmada": tarjeta_total,
                "mercado_pago": mp_total,
                "total_cobrado": total_cobrado,
                "fondo_inicial": apertura,
                "efectivo_contado": efectivo_contado,
                "diferencia": diferencia,
                "folios": [p.get("folio") for p in movimientos_corte],
            }
            cortes_file = DATA_DIR / "cortes_caja.json"
            try:
                cortes = []
                if cortes_file.exists():
                    cortes = json.loads(cortes_file.read_text(encoding="utf-8"))
                cortes.insert(0, corte)
                cortes_file.write_text(json.dumps(cortes, ensure_ascii=False, indent=2), encoding="utf-8")
                guardar_evento_auditoria("CORTE DE CAJA", f"{suc_corte} · {fecha_corte} · total ${total_cobrado:.2f} · diferencia ${diferencia:.2f}")
                st.success("Corte guardado con usuario, sucursal, periodo, métodos de pago y diferencia.")
            except (OSError, json.JSONDecodeError) as exc:
                st.error(f"No se pudo guardar el corte: {exc}")

if _SECCION == "corte":
    _seccion_corte()


# ---------------------------------------------------------------------
# CIERRE DIARIO POR SUCURSAL
# ---------------------------------------------------------------------
# Sin fragmento: las cámaras y la subida de fotos fallan dentro de fragmentos de Streamlit
def _seccion_cierre_diario():
    global col, exc, pago, reg, writer, ws
    globals()["_CONTADOR_LLAVES"] = 1000  # llaves estables entre recargas completas y de sección
    with tab_cierre_diario:
        encabezado_institucional("REPORTE DIARIO DE CIERRE POR SUCURSAL", "CIERRE OPERATIVO", "Balance formal del día: ingresos, gastos y resultado por sucursal")
        st.caption(
            "Captura y entrega el balance diario de boletos, cobros y gastos. "
            "Cada sucursal puede guardar un cierre por fecha; el contador puede consultar todas."
        )

        CIERRES_DIARIOS_FILE = DATA_DIR / "cierres_diarios_sucursal.json"
        GASTOS_CIERRE_FILE = DATA_DIR / "gastos_operativos.json"

        def cargar_cierres_diarios():
            if not CIERRES_DIARIOS_FILE.exists():
                return []
            try:
                contenido = json.loads(CIERRES_DIARIOS_FILE.read_text(encoding="utf-8"))
                return contenido if isinstance(contenido, list) else []
            except (OSError, json.JSONDecodeError):
                return []

        def guardar_cierres_diarios(contenido):
            temporal = CIERRES_DIARIOS_FILE.with_suffix(".tmp")
            temporal.write_text(json.dumps(contenido, ensure_ascii=False, indent=2), encoding="utf-8")
            os.replace(temporal, CIERRES_DIARIOS_FILE)

        def fecha_a_date(valor):
            for formato in ("%d/%m/%Y %H:%M:%S", "%Y-%m-%d %H:%M:%S", "%d/%m/%Y", "%Y-%m-%d"):
                try:
                    return datetime.strptime(str(valor or ""), formato).date()
                except (ValueError, TypeError):
                    continue
            return None

        def numero_folio(valor):
            try:
                return int(str(valor or "").strip().rsplit("-", 1)[-1])
            except (ValueError, TypeError):
                return None

        sucursales_cierre = SUCURSALES if es_global else [st.session_state.sucursal_asignada]
        fecha_cierre = st.date_input("Fecha operativa del cierre", value=ahora_mx().date(), key="fecha_reporte_cierre_diario")
        sucursal_cierre = st.selectbox("Sucursal del reporte", sucursales_cierre, key="sucursal_reporte_cierre_diario")

        registros_cierre = [
            r for r in st.session_state.vehiculos
            if r.get("sucursal") == sucursal_cierre and fecha_a_date(r.get("fecha")) == fecha_cierre
        ]
        gastos_cierre = []
        if GASTOS_CIERRE_FILE.exists():
            try:
                datos_gastos = json.loads(GASTOS_CIERRE_FILE.read_text(encoding="utf-8"))
                if isinstance(datos_gastos, list):
                    gastos_cierre = [
                        g for g in datos_gastos
                        if g.get("sucursal") == sucursal_cierre and fecha_a_date(g.get("fecha")) == fecha_cierre
                        and g.get("origen") != "Cierre diario"
                    ]
            except (OSError, json.JSONDecodeError):
                gastos_cierre = []

        total_gastos_registrados = sum(dinero(g.get("importe")) for g in gastos_cierre)
        pagos_confirmados_cierre = []
        folios_pagados = set()
        ingresos_confirmados_cierre = 0.0
        for reg in registros_cierre:
            pagos_reg = [pago for pago in reg.get("pagos", []) if pago.get("estado") == "Pagado"]
            if pagos_reg:
                for pago in pagos_reg:
                    pagos_confirmados_cierre.append(pago)
                    ingresos_confirmados_cierre += dinero(pago.get("importe"))
                folios_pagados.add(reg.get("id"))
            elif reg.get("pago_estado") == "Pagado":
                folios_pagados.add(reg.get("id"))
                ingresos_confirmados_cierre += dinero(reg.get("importe_cobrado", reg.get("tarifa", 0)))

        cancelados_auto = sum(1 for r in registros_cierre if str(r.get("estado", "")).lower() in ("cancelado", "cancelada"))
        perdidos_auto = sum(1 for r in registros_cierre if str(r.get("estado", "")).lower() in ("perdido", "perdida", "extraviado", "extraviada"))
        pares_folios = [(numero_folio(r.get("id")), str(r.get("id", ""))) for r in registros_cierre]
        pares_folios = [(n, folio) for n, folio in pares_folios if n is not None]
        folio_inicial_default = min(pares_folios, key=lambda x: x[0])[1] if pares_folios else ""
        folio_final_default = max(pares_folios, key=lambda x: x[0])[1] if pares_folios else ""
        clave_widget_cierre = f"{sucursal_cierre}_{fecha_cierre:%Y%m%d}".replace(" ", "_")

        st.markdown("#### Datos del turno y control de boletos")
        st.caption("Los conteos se precargan desde los registros del sistema. Puedes corregirlos si hubo boletos físicos no capturados digitalmente.")
        c1, c2, c3 = st.columns(3)
        with c1:
            fecha_precio = st.number_input(
                "Costo unitario del boleto ($)", min_value=0.0,
                value=float(PRECIOS_SUCURSALES.get(sucursal_cierre, 0)), step=5.0,
                key=llave_unica(f"cierre_precio_{clave_widget_cierre}")
            )
            folio_inicial = st.text_input(
                "Folio inicial de la sesión", value=folio_inicial_default,
                key=llave_unica(f"cierre_folio_ini_{clave_widget_cierre}")
            )
        with c2:
            folio_final = st.text_input(
                "Folio final de la sesión", value=folio_final_default,
                key=llave_unica(f"cierre_folio_fin_{clave_widget_cierre}")
            )
            total_boletos_auto = len(registros_cierre)
            ni = numero_folio(folio_inicial)
            nf = numero_folio(folio_final)
            if ni is not None and nf is not None and nf >= ni:
                total_boletos_auto = nf - ni + 1
            total_boletos = st.number_input(
                "Total de boletos utilizados", min_value=0, value=int(total_boletos_auto), step=1,
                key=llave_unica(f"cierre_total_{clave_widget_cierre}")
            )
        with c3:
            boletos_pagados = st.number_input(
                "Boletos pagados", min_value=0, value=int(len(folios_pagados)), step=1,
                key=llave_unica(f"cierre_pagados_{clave_widget_cierre}")
            )
            boletos_cancelados = st.number_input(
                "Boletos cancelados", min_value=0, value=int(cancelados_auto), step=1,
                key=llave_unica(f"cierre_cancelados_{clave_widget_cierre}")
            )
            boletos_perdidos = st.number_input(
                "Boletos perdidos", min_value=0, value=int(perdidos_auto), step=1,
                key=llave_unica(f"cierre_perdidos_{clave_widget_cierre}")
            )

        st.markdown("#### Gastos generales y balance")
        st.write(f"**Gastos operativos ya registrados en el sistema:** ${total_gastos_registrados:,.2f} MXN")
        st.caption("Si agregas un gasto adicional aquí, quedará en este cierre diario y no se duplicará en el catálogo de gastos operativos.")
        colg1, colg2 = st.columns(2)
        with colg1:
            gasto_adicional = st.number_input(
                "Gastos generales adicionales de este cierre ($)", min_value=0.0, value=0.0, step=50.0,
                key=llave_unica(f"cierre_gasto_extra_{clave_widget_cierre}")
            )
        with colg2:
            detalle_gasto_adicional = st.text_input(
                "Concepto / observaciones de gastos", key=llave_unica(f"cierre_gasto_concepto_{clave_widget_cierre}")
            )

        total_gastos_cierre = total_gastos_registrados + float(gasto_adicional)
        balance_cierre = ingresos_confirmados_cierre - total_gastos_cierre
        m1, m2, m3, m4 = st.columns(4)
        m1.metric("Vehículos registrados ese día", len(registros_cierre))
        m2.metric("Ingresos confirmados", f"${ingresos_confirmados_cierre:,.2f}")
        m3.metric("Gastos totales", f"${total_gastos_cierre:,.2f}")
        m4.metric("Balance del día", f"${balance_cierre:,.2f}")
        st.caption("El balance se calcula con cobros confirmados menos gastos registrados y gastos adicionales del cierre; los boletos cancelados/perdidos no se cuentan como cobros salvo que exista un pago confirmado.")

        st.markdown("#### 📷 Evidencia fotográfica del cierre")
        st.caption("Sube fotos de la hoja de cierre, boletos y efectivo. Quedan guardadas con este cierre y las puede consultar el Contador.")
        fotos_cierre_archivos = st.file_uploader(
            "Subir fotos desde la galería", type=["png", "jpg", "jpeg"], accept_multiple_files=True,
            key=llave_unica(f"cierre_fotos_{clave_widget_cierre}"),
        )
        foto_cierre_camara = None
        llave_cam_cierre = llave_unica(f"cierre_camara_{clave_widget_cierre}")  # misma cantidad de llaves en cada ejecución
        if st.checkbox("📸 Tomar foto con la cámara", key=f"cierre_cam_on_{clave_widget_cierre}"):
            foto_cierre_camara = st.camera_input("Cámara", key=llave_cam_cierre)
        nuevas_fotos_cierre = list(fotos_cierre_archivos or []) + ([foto_cierre_camara] if foto_cierre_camara else [])
        if nuevas_fotos_cierre:
            cols_prev = st.columns(3)
            for i_prev, f_prev in enumerate(nuevas_fotos_cierre):
                cols_prev[i_prev % 3].image(f_prev, use_container_width=True)
        else:
            st.warning("Aún no has adjuntado evidencia fotográfica de este cierre.")

        fecha_cierre_txt = fecha_cierre.strftime("%d/%m/%Y")
        fisico_cierre = next((r for r in _leer_lista_json(FISICOS_FILE) if r.get("tipo_reporte") == "ingresos"
                              and r.get("Sitio / Unidad") == sucursal_cierre and r.get("Fecha") == fecha_cierre_txt
                              and r.get("origen") != "Cierre del sistema"), None)
        if fisico_cierre:
            st.info(f"Esta sucursal ya tiene reporte físico de ingresos del {fecha_cierre_txt} ({fisico_cierre.get('ID')}). El cierre se guarda como respaldo; en auditoría y Excel manda el reporte físico.")
        else:
            st.info("No hay reporte físico de ingresos de esta sucursal y fecha: al guardar, el cierre se sube automáticamente a Reportes físicos para auditarlo.")

        cierres_guardados = cargar_cierres_diarios()
        cierre_existente = next((c for c in cierres_guardados if c.get("fecha") == str(fecha_cierre) and c.get("sucursal") == sucursal_cierre), None)
        if cierre_existente and es_contador:
            st.warning(f"Ya existe un cierre para {sucursal_cierre} el {fecha_cierre}. Marca la casilla para reemplazarlo con los datos actuales.")
        elif cierre_existente:
            st.error(f"Ya existe un cierre para {sucursal_cierre} el {fecha_cierre}. No se puede subir dos veces: solo el Contador Auditor puede reemplazarlo o borrarlo.")
        reemplazar_cierre = st.checkbox(
            "Confirmo que deseo reemplazar el cierre existente de esta sucursal y fecha",
            value=False, key=llave_unica(f"cierre_reemplazar_{clave_widget_cierre}"), disabled=not (cierre_existente and es_contador)
        )
        if st.button("💾 Guardar cierre diario", type="primary", key=llave_unica(f"guardar_cierre_diario_{clave_widget_cierre}")):
            if not folio_inicial.strip() or not folio_final.strip():
                st.error("Captura el folio inicial y el folio final de la sesión.")
            elif int(boletos_pagados) + int(boletos_cancelados) + int(boletos_perdidos) > int(total_boletos):
                st.error("La suma de boletos pagados, cancelados y perdidos no puede superar el total de boletos utilizados.")
            elif cierre_existente and not es_contador:
                st.error("El cierre ya existe. Solo el Contador Auditor puede reemplazarlo.")
            elif cierre_existente and not reemplazar_cierre:
                st.error("El cierre ya existe. Marca la casilla de reemplazo si realmente necesitas corregirlo.")
            else:
                cierre = {
                    "id": cierre_existente.get("id") if cierre_existente else uuid.uuid4().hex[:12].upper(),
                    "fecha": str(fecha_cierre),
                    "sucursal": sucursal_cierre,
                    "costo_unitario_boleto": float(fecha_precio),
                    "folio_inicial": folio_inicial.strip(),
                    "folio_final": folio_final.strip(),
                    "total_boletos": int(total_boletos),
                    "boletos_pagados": int(boletos_pagados),
                    "boletos_cancelados": int(boletos_cancelados),
                    "boletos_perdidos": int(boletos_perdidos),
                    "ingresos_confirmados": round(ingresos_confirmados_cierre, 2),
                    "gastos_operativos_registrados": round(total_gastos_registrados, 2),
                    "gastos_generales_adicionales": round(float(gasto_adicional), 2),
                    "concepto_gastos_adicionales": detalle_gasto_adicional.strip(),
                    "gastos_totales": round(total_gastos_cierre, 2),
                    "balance_neto": round(balance_cierre, 2),
                    "usuario": st.session_state.usuario_autenticado,
                    "rol": st.session_state.rol,
                    "guardado_en": ahora_mx().strftime("%Y-%m-%d %H:%M:%S"),
                }
                if cierre_existente:
                    cierres_guardados = [cierre if c.get("id") == cierre_existente.get("id") else c for c in cierres_guardados]
                else:
                    cierres_guardados.insert(0, cierre)
                try:
                    rutas_evidencia = list((cierre_existente or {}).get("fotos_evidencia", []))
                    if nuevas_fotos_cierre:
                        rutas_evidencia = []
                        base_foto = re.sub(r"[^A-Za-z0-9_-]+", "_", f"CIERRE_{sucursal_cierre}_{fecha_cierre:%Y%m%d}")
                        for i_ev, f_ev in enumerate(nuevas_fotos_cierre, start=1):
                            rutas_evidencia.append(guardar_foto(f_ev, base_foto, f"evidencia_{i_ev}"))
                    cierre["fotos_evidencia"] = rutas_evidencia
                    guardar_cierres_diarios(cierres_guardados)
                    # Gasto adicional del cierre: pasa al libro de gastos (sin duplicar al reemplazar)
                    gastos_libro = [g for g in _leer_lista_json(GASTOS_FILE_RF) if g.get("cierre_id") != cierre["id"]]
                    if float(gasto_adicional) > 0:
                        gastos_libro.insert(0, {
                            "id": uuid.uuid4().hex[:12].upper(), "fecha": str(fecha_cierre), "sucursal": sucursal_cierre,
                            "categoria": "Otros", "concepto": detalle_gasto_adicional.strip() or "Gasto adicional del cierre",
                            "importe": float(gasto_adicional), "registrado_por": st.session_state.usuario_autenticado,
                            "registrado_en": ahora_mx().strftime("%Y-%m-%d %H:%M:%S"),
                            "origen": "Cierre diario", "cierre_id": cierre["id"],
                        })
                    _guardar_lista_json(GASTOS_FILE_RF, gastos_libro)
                    # Si no hay reporte físico de ingresos, el cierre se sube como reporte para auditar
                    reportes_libro = _leer_lista_json(FISICOS_FILE)
                    hay_fisico = any(r.get("tipo_reporte") == "ingresos" and r.get("Sitio / Unidad") == sucursal_cierre
                                     and r.get("Fecha") == fecha_cierre_txt and r.get("origen") != "Cierre del sistema" for r in reportes_libro)
                    if not hay_fisico:
                        reportes_libro = [r for r in reportes_libro if not (r.get("origen") == "Cierre del sistema"
                                          and r.get("Sitio / Unidad") == sucursal_cierre and r.get("Fecha") == fecha_cierre_txt)]
                        rid_auto = f"CIE_{cierre['id'][:8]}"
                        reporte_auto = _fila_flujo(
                            sucursal_cierre, fecha_cierre, ingresos_confirmados_cierre, int(boletos_pagados), float(fecha_precio),
                            numero_folio(folio_inicial) or 0, numero_folio(folio_final) or 0, int(total_boletos),
                            st.session_state.usuario_autenticado, st.session_state.usuario_autenticado, rid_auto)
                        reporte_auto.update({
                            "Cancelados": int(boletos_cancelados), "Perdidos": int(boletos_perdidos),
                            "tipo_reporte": "ingresos", "origen": "Cierre del sistema", "cierre_id": cierre["id"],
                            "fotos": rutas_evidencia, "Nivel Auditoría": "Sin reporte físico (cierre del sistema)",
                            "Timestamp": ahora_mx().strftime("%Y-%m-%d %H:%M:%S"),
                        })
                        reportes_libro.insert(0, reporte_auto)
                        _guardar_lista_json(FISICOS_FILE, reportes_libro)
                        guardar_evento_auditoria("CIERRE SUBIDO PARA AUDITORÍA", f"{sucursal_cierre} · {fecha_cierre_txt} · {rid_auto} · sin reporte físico")
                    guardar_evento_auditoria("CIERRE DIARIO POR SUCURSAL", f"{sucursal_cierre} · {fecha_cierre} · ingresos ${ingresos_confirmados_cierre:.2f} · gastos ${total_gastos_cierre:.2f} · balance ${balance_cierre:.2f}")
                    st.success("Cierre diario guardado. Ya puedes descargar la hoja de cálculo para entregar el balance.")
                    st.rerun()
                except Exception as exc:
                    st.error(f"No se pudo guardar el cierre diario: {exc}")

        st.divider()
        st.markdown("#### Historial y entrega de hoja de cálculo")
        cierres_guardados = cargar_cierres_diarios()
        cierres_visibles = cierres_guardados if es_global else [c for c in cierres_guardados if c.get("sucursal") == st.session_state.sucursal_asignada]
        if cierres_visibles:
            df_cierres = pd.DataFrame(cierres_visibles)
            columnas_mostrar_cierres = [
                "fecha", "sucursal", "costo_unitario_boleto", "folio_inicial", "folio_final",
                "total_boletos", "boletos_pagados", "boletos_cancelados", "boletos_perdidos",
                "ingresos_confirmados", "gastos_totales", "balance_neto", "usuario", "guardado_en"
            ]
            st.dataframe(df_cierres.reindex(columns=columnas_mostrar_cierres), use_container_width=True, hide_index=True)
            opciones_cierre = {
                f"{c.get('fecha')} · {c.get('sucursal')} · Balance ${dinero(c.get('balance_neto')):,.2f}": c
                for c in cierres_visibles
            }
            etiqueta_cierre = st.selectbox("Selecciona el cierre que vas a entregar", list(opciones_cierre.keys()), key="seleccionar_cierre_descarga")
            cierre_descarga = opciones_cierre[etiqueta_cierre]
            if es_contador:
                with st.expander("🗑️ Borrar este cierre (solo Contador Auditor)"):
                    st.caption("Se borra el cierre, su gasto adicional, el reporte automático que generó para auditoría y sus fotos de evidencia.")
                    ok_borrar_cierre = st.checkbox("Confirmo que quiero borrar este cierre", key=f"borrar_cierre_ok_{cierre_descarga.get('id')}")
                    if st.button("🗑️ Borrar cierre", disabled=not ok_borrar_cierre, key=f"borrar_cierre_btn_{cierre_descarga.get('id')}"):
                        id_c = cierre_descarga.get("id")
                        guardar_cierres_diarios([c for c in cargar_cierres_diarios() if c.get("id") != id_c])
                        _guardar_lista_json(GASTOS_FILE_RF, [g for g in _leer_lista_json(GASTOS_FILE_RF) if g.get("cierre_id") != id_c])
                        _guardar_lista_json(FISICOS_FILE, [r for r in _leer_lista_json(FISICOS_FILE)
                                                           if not (r.get("origen") == "Cierre del sistema" and r.get("cierre_id") == id_c)])
                        _borrar_archivos(cierre_descarga.get("fotos_evidencia"))
                        guardar_evento_auditoria("CIERRE DIARIO BORRADO", f"{cierre_descarga.get('sucursal')} · {cierre_descarga.get('fecha')} · {id_c}")
                        st.success("Cierre borrado.")
                        st.rerun()
            fotos_ev = cierre_descarga.get("fotos_evidencia") or []
            if fotos_ev:
                st.markdown("**📷 Evidencia fotográfica del cierre**")
                cols_ev = st.columns(3)
                for i_ev, ruta_ev in enumerate(fotos_ev):
                    p_ev = Path(ruta_ev)
                    if not p_ev.is_absolute():
                        p_ev = DATA_DIR / p_ev
                    if p_ev.exists():
                        cols_ev[i_ev % 3].image(str(p_ev), use_container_width=True)
            else:
                st.caption("Este cierre no tiene evidencia fotográfica.")
            fecha_exportacion = fecha_a_date(cierre_descarga.get("fecha"))
            registros_excel_cierre = [
                r for r in st.session_state.vehiculos
                if r.get("sucursal") == cierre_descarga.get("sucursal")
                and fecha_a_date(r.get("fecha")) == fecha_exportacion
            ]
            gastos_excel_cierre = []
            if GASTOS_CIERRE_FILE.exists():
                try:
                    gastos_archivo_excel = json.loads(GASTOS_CIERRE_FILE.read_text(encoding="utf-8"))
                    if isinstance(gastos_archivo_excel, list):
                        gastos_excel_cierre = [
                            g for g in gastos_archivo_excel
                            if g.get("sucursal") == cierre_descarga.get("sucursal")
                            and fecha_a_date(g.get("fecha")) == fecha_exportacion
                            and g.get("origen") != "Cierre diario"
                        ]
                except (OSError, json.JSONDecodeError):
                    gastos_excel_cierre = []
            salida_excel = io.BytesIO()
            df_resumen_cierre = pd.DataFrame([{
                "Fecha": cierre_descarga.get("fecha"),
                "Sucursal": cierre_descarga.get("sucursal"),
                "Costo unitario boleto ($)": cierre_descarga.get("costo_unitario_boleto"),
                "Folio inicial": cierre_descarga.get("folio_inicial"),
                "Folio final": cierre_descarga.get("folio_final"),
                "Total boletos": cierre_descarga.get("total_boletos"),
                "Boletos pagados": cierre_descarga.get("boletos_pagados"),
                "Boletos cancelados": cierre_descarga.get("boletos_cancelados"),
                "Boletos perdidos": cierre_descarga.get("boletos_perdidos"),
                "Ingresos confirmados ($)": cierre_descarga.get("ingresos_confirmados"),
                "Gastos operativos registrados ($)": cierre_descarga.get("gastos_operativos_registrados"),
                "Gastos generales adicionales ($)": cierre_descarga.get("gastos_generales_adicionales"),
                "Concepto gastos adicionales": cierre_descarga.get("concepto_gastos_adicionales"),
                "Gastos totales ($)": cierre_descarga.get("gastos_totales"),
                "Balance neto ($)": cierre_descarga.get("balance_neto"),
                "Elaboró": cierre_descarga.get("usuario"),
                "Rol": cierre_descarga.get("rol"),
                "Fecha/hora de elaboración": cierre_descarga.get("guardado_en"),
            }])
            try:
                with pd.ExcelWriter(salida_excel, engine="openpyxl") as writer:
                    df_resumen_cierre.to_excel(writer, sheet_name="Cierre diario", index=False, startrow=2)
                    ws = writer.sheets["Cierre diario"]
                    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(df_resumen_cierre.columns))
                    ws.cell(row=1, column=1, value="CLUB PARK VALET Y ESTACIONAMIENTOS S.A. DE C.V. — CIERRE DIARIO")
                    ws.cell(row=1, column=1).font = openpyxl.styles.Font(bold=True, size=14, color="FFFFFF")
                    ws.cell(row=1, column=1).fill = openpyxl.styles.PatternFill("solid", fgColor="1F4E78")
                    ws.freeze_panes = "A4"
                    ws.auto_filter.ref = f"A3:{openpyxl.utils.get_column_letter(len(df_resumen_cierre.columns))}4"
                    for col_cells in ws.columns:
                        max_len = max(len(str(cell.value or "")) for cell in col_cells)
                        ws.column_dimensions[openpyxl.utils.get_column_letter(col_cells[0].column)].width = min(max(max_len + 2, 14), 36)
                    ws.append([])
                    ws.append(["Detalle de gastos operativos registrados en la sucursal y fecha"])
                    ws.append(["Categoría", "Concepto", "Importe ($)", "Registrado por"])
                    for gasto in gastos_excel_cierre:
                        ws.append([gasto.get("categoria", ""), gasto.get("concepto", ""), dinero(gasto.get("importe")), gasto.get("registrado_por", "")])
                    ws2 = writer.book.create_sheet("Boletos del día")
                    ws2.append(["Folio", "Fecha de entrada", "Sucursal", "Tarifa ($)", "Estado", "Pagado ($)"])
                    for reg in registros_excel_cierre:
                        pagado_reg = sum(dinero(p.get("importe")) for p in reg.get("pagos", []) if p.get("estado") == "Pagado")
                        if not reg.get("pagos") and reg.get("pago_estado") == "Pagado":
                            pagado_reg = dinero(reg.get("importe_cobrado", reg.get("tarifa", 0)))
                        ws2.append([reg.get("id", ""), reg.get("fecha", ""), reg.get("sucursal", ""), dinero(reg.get("tarifa")), reg.get("estado", ""), pagado_reg])
                    ws2.freeze_panes = "A2"
                    ws2.auto_filter.ref = f"A1:F{max(ws2.max_row, 1)}"
                    for col in ws2.columns:
                        max_len = max(len(str(cell.value or "")) for cell in col)
                        ws2.column_dimensions[openpyxl.utils.get_column_letter(col[0].column)].width = min(max(max_len + 2, 14), 34)
                st.download_button(
                    "📥 Descargar cierre diario en Excel (.xlsx)",
                    data=salida_excel.getvalue(),
                    file_name=f"ClubPark_Cierre_{cierre_descarga.get('sucursal','Sucursal').replace(' ','_')}_{cierre_descarga.get('fecha','')}.xlsx",
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    use_container_width=True,
                    key="descargar_excel_cierre_diario",
                )
                st.download_button(
                    "📄 Descargar cierre diario en CSV",
                    data=df_resumen_cierre.to_csv(index=False).encode("utf-8-sig"),
                    file_name=f"ClubPark_Cierre_{cierre_descarga.get('fecha','')}.csv",
                    mime="text/csv",
                    key="descargar_csv_cierre_diario",
                )
            except Exception as exc:
                st.error(f"No se pudo preparar el Excel del cierre diario: {exc}")
                st.info("Confirma que openpyxl esté instalado en requirements.txt.")
        else:
            st.info("Todavía no hay cierres diarios guardados para esta sucursal.")

if _SECCION == "cierre_diario":
    _seccion_cierre_diario()


# ---------------------------------------------------------------------
# TAB 3: REIMPRESIÓN
# ---------------------------------------------------------------------
@_fragmento
def _seccion_imp():
    global qr_data, ticket
    globals()["_CONTADOR_LLAVES"] = 1000  # llaves estables entre recargas completas y de sección
    with tab_imp:
        encabezado_institucional("RE-IMPRESIÓN TÉRMICA DE TICKETS", "IMPRESIÓN ESC/POS", "Reemisión de boletos y llaveros con el folio original")
        panel_prueba_impresion(key_sufijo="reimpresion", expandido=False)
        visibles = vehiculos_visibles(st.session_state.vehiculos)
        if not visibles:
            st.info("No hay tickets disponibles.")
        else:
            folio_imp = st.selectbox(
                "Selecciona folio a imprimir",
                [r.get("id", "") for r in visibles],
                key="folio_reimpresion",
            )
            registro_imp = next((r for r in visibles if r.get("id") == folio_imp), None)
            if registro_imp:
                ticket = registro_imp.get("ticket_texto") or crear_ticket(registro_imp)
                st.markdown("**Las dos impresiones son independientes.**")
                vista_impresion_con_qr(registro_imp, "cliente", "imprimir_cliente_reimp")
                vista_impresion_con_qr(registro_imp, "llavero", "imprimir_contrasena_reimp")
                imprimir_partes_ticket(registro_imp, "reimpresion")
                st.download_button(
                    "📄 Descargar ticket completo TXT",
                    data=ticket.encode("utf-8"),
                    file_name=f"ticket_{folio_imp}.txt",
                    mime="text/plain",
                    use_container_width=True,
                )
                if QR_OK:
                    qr_data = generar_qr_bytes(registro_imp)
                    st.image(qr_data, caption="QR único para localizar el vehículo", width=160)
                    st.download_button(
                        "⬇️ Descargar QR del vehículo",
                        data=qr_data,
                        file_name=f"QR_{folio_imp}.png",
                        mime="image/png",
                        key=f"qr_reimp_{folio_imp}",
                    )

if _SECCION == "imp":
    _seccion_imp()


# ---------------------------------------------------------------------
# TAB 4: REPORTES Y EXCEL
# ---------------------------------------------------------------------
@_fragmento
def _seccion_rep():
    global exc
    globals()["_CONTADOR_LLAVES"] = 1000  # llaves estables entre recargas completas y de sección
    with tab_rep:
        encabezado_institucional("REPORTE OPERATIVO Y EXCEL OFICIAL", "DOCUMENTACIÓN", "Exportación institucional para dirección y contaduría")
        visibles = vehiculos_visibles(st.session_state.vehiculos)
        if not visibles:
            st.info("No hay registros para generar un reporte.")
        else:
            df_rep = pd.DataFrame(visibles)
            if "fecha" in df_rep.columns:
                fechas = df_rep["fecha"].fillna("").astype(str)
                opciones_sucursal = ["Todas"] + sorted(df_rep["sucursal"].dropna().astype(str).unique().tolist())
                filtro_sucursal = st.selectbox("Filtrar por sucursal", opciones_sucursal, key="rep_sucursal")
                if filtro_sucursal != "Todas":
                    df_rep = df_rep[df_rep["sucursal"] == filtro_sucursal]
                st.metric("Vehículos en el reporte", len(df_rep))
                ingresos = pd.to_numeric(df_rep.get("tarifa", pd.Series(dtype=float)), errors="coerce").fillna(0).sum()
                st.metric("Suma de tarifas registradas", f"${ingresos:,.2f} MXN")
                st.dataframe(
                    df_rep.reindex(columns=["id", "fecha", "sucursal", "marca", "placas", "tarifa", "operador", "estado"]),
                    use_container_width=True,
                    hide_index=True,
                )
                try:
                    excel_data = generar_excel_oficial(df_rep)
                    st.download_button(
                        "📥 Descargar Excel Oficial",
                        data=excel_data.getvalue(),
                        file_name=f"Club_Park_Reporte_{ahora_mx():%Y%m%d_%H%M}.xlsx",
                        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                        use_container_width=True,
                    )
                except Exception as exc:
                    st.error(f"No se pudo generar el Excel: {exc}")
                    st.info("Verifica requirements.txt y vuelve a desplegar la aplicación.")
                csv_data = df_rep.to_csv(index=False).encode("utf-8-sig")
                st.download_button(
                    "📥 Descargar respaldo CSV",
                    data=csv_data,
                    file_name=f"Club_Park_Reporte_{ahora_mx():%Y%m%d_%H%M}.csv",
                    mime="text/csv",
                    use_container_width=True,
                )

if _SECCION == "rep":
    _seccion_rep()


# ---------------------------------------------------------------------
# TAB 5: AUDITORÍA Y WHATSAPP
# ---------------------------------------------------------------------
@_fragmento
def _seccion_aud():
    global entregados, exc, p, r, registro, ruta, rutas, suc, total, vista
    globals()["_CONTADOR_LLAVES"] = 1000  # llaves estables entre recargas completas y de sección
    with tab_aud:
        encabezado_institucional("AUDITORÍA Y COMUNICACIÓN", "GOBIERNO DEL SISTEMA", "Bitácora de altas, cobros, salidas y cierres con usuario y fecha")
        visibles = vehiculos_visibles(st.session_state.vehiculos)
        total = len(visibles)
        entregados = sum(1 for r in visibles if r.get("estado") == "Entregado")
        resguardo = sum(1 for r in visibles if r.get("estado", "En Resguardo") == "En Resguardo")
        c1, c2, c3 = st.columns(3)
        c1.metric("Registros visibles", total)
        c2.metric("En resguardo", resguardo)
        c3.metric("Entregados", entregados)

        if es_contador or "Supervisor" in str(st.session_state.get("rol", "")):
            st.markdown("#### 🏢 Resumen por sucursal")
            st.caption("En custodia = vehículos que están ahora en resguardo. Entregados y recaudado = del periodo elegido (cobros confirmados).")
            hoy_a = ahora_mx().date()
            ra1, ra2 = st.columns(2)
            desde_a = ra1.date_input("Desde", value=hoy_a, key="aud_res_desde")
            hasta_a = ra2.date_input("Hasta", value=hoy_a, key="aud_res_hasta")
            clave_sup_a = supervisor_del_usuario()
            if es_contador:
                sucursales_a = SUCURSALES
            elif clave_sup_a:
                ver_a = st.radio("Sucursales", ["Mis sucursales a cargo", "Todas"], horizontal=True, key="aud_res_alcance")
                sucursales_a = sucursales_de_supervisor(clave_sup_a) if ver_a.startswith("Mis") else SUCURSALES
            elif es_global:
                sucursales_a = SUCURSALES
            else:
                sucursales_a = [st.session_state.sucursal_asignada]
            if hasta_a < desde_a:
                st.error("La fecha final no puede ser anterior a la inicial.")
            else:
                filas_a = []
                for suc_a in sucursales_a:
                    regs_a = [r for r in st.session_state.vehiculos if r.get("sucursal") == suc_a]
                    custodia_a = sum(1 for r in regs_a if r.get("estado", "En Resguardo") == "En Resguardo")
                    ingresados_a = sum(1 for r in regs_a if (lambda d: d and desde_a <= d <= hasta_a)(fecha_flexible(r.get("fecha"))))
                    entregados_a = sum(1 for r in regs_a if r.get("estado") == "Entregado"
                                       and (lambda d: d and desde_a <= d <= hasta_a)(fecha_flexible(r.get("salida_fecha"))))
                    metodos_a = {"Efectivo": 0.0, "Tarjeta": 0.0, "QR": 0.0, METODO_MP: 0.0}
                    for r in regs_a:
                        for p in r.get("pagos", []) or []:
                            d_p = fecha_flexible(p.get("fecha_hora"))
                            if p.get("estado") == "Pagado" and d_p and desde_a <= d_p <= hasta_a:
                                metodo_p = p.get("metodo") if p.get("metodo") in metodos_a else "Efectivo"
                                metodos_a[metodo_p] += dinero(p.get("importe"))
                    filas_a.append({"Sucursal": suc_a, "En custodia": custodia_a, "Ingresados": ingresados_a,
                                    "Entregados": entregados_a, "Recaudado": sum(metodos_a.values()), **metodos_a})
                df_a = pd.DataFrame(filas_a)
                solo_actividad = st.checkbox("Mostrar solo sucursales con actividad", value=True, key="aud_res_activas")
                df_a_ver = df_a[(df_a["En custodia"] > 0) | (df_a["Ingresados"] > 0) | (df_a["Entregados"] > 0) | (df_a["Recaudado"] > 0)] if solo_actividad else df_a
                t1a, t2a, t3a, t4a = st.columns(4)
                t1a.metric("En custodia", int(df_a["En custodia"].sum()))
                t2a.metric("Ingresados", int(df_a["Ingresados"].sum()))
                t3a.metric("Entregados", int(df_a["Entregados"].sum()))
                t4a.metric("Recaudado", f"${df_a['Recaudado'].sum():,.2f}")
                if df_a_ver.empty:
                    st.info("Sin actividad en el periodo para estas sucursales.")
                else:
                    total_a = {"Sucursal": "TOTAL", **{c: df_a_ver[c].sum() for c in df_a_ver.columns if c != "Sucursal"}}
                    df_a_tot = pd.concat([df_a_ver, pd.DataFrame([total_a])], ignore_index=True)
                    st.dataframe(df_a_tot, use_container_width=True, hide_index=True, column_config={
                        c: st.column_config.NumberColumn(c, format="$%.2f") for c in ("Recaudado", "Efectivo", "Tarjeta", "QR", METODO_MP)})
                    st.download_button("📄 Descargar resumen por sucursal (CSV)", data=df_a_tot.to_csv(index=False).encode("utf-8-sig"),
                                       file_name=f"Resumen_sucursales_{desde_a:%Y%m%d}_{hasta_a:%Y%m%d}.csv", mime="text/csv", key="aud_res_csv")
                suc_det_a = st.selectbox("Ver vehículos en custodia de una sucursal", ["—"] + list(sucursales_a), key="aud_res_det")
                if suc_det_a != "—":
                    en_custodia_a = [r for r in st.session_state.vehiculos
                                     if r.get("sucursal") == suc_det_a and r.get("estado", "En Resguardo") == "En Resguardo"]
                    if en_custodia_a:
                        st.dataframe(pd.DataFrame([{
                            "Folio": r.get("id"), "Entrada": r.get("fecha_entrada") or r.get("fecha"), "Placas": r.get("placas"),
                            "Marca": r.get("marca"), "Color": r.get("color"), "Tiempo": texto_duracion(minutos_estancia(r)),
                            "Custodia": r.get("custodia"),
                        } for r in en_custodia_a]), use_container_width=True, hide_index=True)
                    else:
                        st.success(f"{suc_det_a} no tiene vehículos en custodia.")
            st.divider()

        if visibles:
            st.markdown("#### Revisión de registros")
            df_aud = pd.DataFrame(visibles)
            sucursales_aud = ["Todas"] + sorted(df_aud["sucursal"].dropna().astype(str).unique().tolist())
            suc_aud = st.selectbox("Sucursal para auditar", sucursales_aud, key="aud_sucursal")
            if suc_aud != "Todas":
                df_aud = df_aud[df_aud["sucursal"] == suc_aud]
            st.dataframe(
                df_aud.reindex(columns=["id", "fecha", "sucursal", "marca", "placas", "tarifa", "operador", "estado"]),
                use_container_width=True,
                hide_index=True,
            )

            st.markdown("#### Validación de las cuatro fotografías")
            for registro in visibles[:100]:
                faltan = []
                for vista in ("frente", "izquierda", "derecha", "trasera"):
                    if not ruta_foto(registro, vista):
                        faltan.append(vista)
                if faltan:
                    st.warning(f"{registro.get('id')} — faltan/no están disponibles: {', '.join(faltan)}")

        st.divider()
        st.markdown("#### Enviar recordatorio por WhatsApp")
        st.caption("El enlace abre WhatsApp con el texto preparado; debes revisar y enviarlo manualmente.")
        telefono = st.text_input("Teléfono con clave de país o celular mexicano de 10 dígitos", key="wa_telefono")
        mensaje = st.text_area(
            "Mensaje",
            value="Buen día. Favor de revisar y completar los registros operativos pendientes de CLUB PARK.",
            key="wa_mensaje",
        )
        if telefono.strip():
            st.link_button(
                "💬 Abrir WhatsApp con mensaje",
                enlace_whatsapp(telefono, mensaje),
                use_container_width=True,
            )

        if es_contador:
            st.divider()
            st.markdown("#### Respaldo independiente de cobros")
            st.caption("Cada cobro guardado queda también en pagos_registrados.json, separado de la ficha del vehículo. Descarga una copia periódicamente y guárdala fuera del servidor.")
            pagos_respaldo = cargar_pagos_respaldo()
            st.metric("Cobros en libro de respaldo", len(pagos_respaldo))
            if pagos_respaldo:
                df_pagos_respaldo = pd.DataFrame(pagos_respaldo)
                st.dataframe(df_pagos_respaldo, use_container_width=True, hide_index=True)
                st.download_button(
                    "⬇️ Descargar respaldo de cobros (JSON)",
                    data=json.dumps(pagos_respaldo, ensure_ascii=False, indent=2).encode("utf-8"),
                    file_name=f"ClubPark_Respaldo_Cobros_{ahora_mx().strftime('%Y%m%d_%H%M%S')}.json",
                    mime="application/json",
                    key="descargar_respaldo_cobros_json",
                )
                st.download_button(
                    "📊 Descargar libro de cobros (CSV)",
                    data=df_pagos_respaldo.to_csv(index=False).encode("utf-8-sig"),
                    file_name=f"ClubPark_Cobros_{ahora_mx().strftime('%Y%m%d_%H%M%S')}.csv",
                    mime="text/csv",
                    key="descargar_respaldo_cobros_csv",
                )
            else:
                st.info("Todavía no hay cobros en el libro de respaldo. Los siguientes cobros se registrarán aquí.")

            st.divider()
            st.markdown("#### Conteo y consecutivo de folios por sucursal")
            st.caption("Cada sucursal tiene un prefijo exclusivo. El consecutivo continúa al día siguiente y no se reinicia al cerrar caja.")
            config_folios = cargar_config()
            resumen_folios = []
            for sucursal_folio_resumen in SUCURSALES:
                maximo_folio_resumen = siguiente_numero_usado(sucursal_folio_resumen, st.session_state.vehiculos)
                siguiente_configurado_resumen = int(config_folios["siguiente_folio"].get(sucursal_folio_resumen, 1))
                resumen_folios.append({
                    "Sucursal": sucursal_folio_resumen,
                    "Prefijo exclusivo": f"CP-{codigo_sucursal(sucursal_folio_resumen)}-",
                    "Último consecutivo utilizado": maximo_folio_resumen,
                    "Siguiente consecutivo": max(siguiente_configurado_resumen, maximo_folio_resumen + 1),
                    "Vehículos registrados": sum(1 for reg in st.session_state.vehiculos if reg.get("sucursal") == sucursal_folio_resumen),
                })
            st.dataframe(pd.DataFrame(resumen_folios), use_container_width=True, hide_index=True)
            st.markdown("#### Configuración de folios por sucursal")
            st.caption("Define el siguiente consecutivo. El sistema nunca bajará del máximo ya utilizado en esa sucursal para evitar reutilizar folios.")
            suc_folio_admin = st.selectbox("Sucursal para configurar folio", SUCURSALES, key="admin_sucursal_folio")
            valor_actual = int(config_folios["siguiente_folio"].get(suc_folio_admin, 1))
            st.write(f"**Siguiente folio configurado:** {valor_actual}")
            nuevo_inicio = st.number_input(
                "Número del siguiente folio",
                min_value=1,
                value=max(1, valor_actual),
                step=1,
                key="admin_nuevo_inicio_folio",
            )
            if st.button("💾 Guardar siguiente folio", type="primary", key="guardar_folio_sucursal"):
                maximo_usado = siguiente_numero_usado(suc_folio_admin, st.session_state.vehiculos)
                if int(nuevo_inicio) <= maximo_usado:
                    st.error(f"No se puede iniciar en {int(nuevo_inicio)}: ya existe un folio de esta sucursal hasta {maximo_usado}. Usa al menos {maximo_usado + 1}.")
                else:
                    anterior = int(config_folios["siguiente_folio"].get(suc_folio_admin, 1))
                    config_folios["siguiente_folio"][suc_folio_admin] = int(nuevo_inicio)
                    config_folios.setdefault("historial_folios", []).insert(0, {
                        "fecha_hora": ahora_mx().strftime("%Y-%m-%d %H:%M:%S"),
                        "usuario": st.session_state.usuario_autenticado,
                        "sucursal": suc_folio_admin,
                        "anterior": anterior,
                        "nuevo": int(nuevo_inicio),
                    })
                    try:
                        guardar_config(config_folios)
                        guardar_evento_auditoria("CAMBIO DE FOLIO", f"{suc_folio_admin}: {anterior} -> {int(nuevo_inicio)}")
                        st.success("Siguiente folio actualizado y cambio registrado en auditoría.")
                        st.rerun()
                    except OSError as exc:
                        st.error(f"No se pudo guardar la configuración: {exc}")

            st.markdown("#### Historial de cambios de folios")
            historial_folios = cargar_config().get("historial_folios", [])
            if historial_folios:
                st.dataframe(pd.DataFrame(historial_folios), use_container_width=True, hide_index=True)
            else:
                st.info("Todavía no hay cambios manuales de folio.")

            st.markdown("#### Reinicio administrativo de datos operativos")
            st.warning(
                "Esta herramienta es exclusiva del Contador Auditor. Antes de reiniciar, descarga el respaldo. "
                "El historial de auditoría se conserva y los consecutivos no se reinician salvo que marques la opción correspondiente."
            )
            cierres_file_admin = DATA_DIR / "cierres_diarios_sucursal.json"
            gastos_file_admin = DATA_DIR / "gastos_operativos.json"
            cortes_file_admin = DATA_DIR / "cortes_caja.json"
            config_admin_backup = cargar_config()
            # Blindaje institucional: los respaldos siempre quedan definidos antes de los
            # intentos de lectura, de modo que ninguna ruta de ejecución llegue al
            # diccionario de respaldo con nombres sin definir (NameError).
            cierres_admin_backup = []
            gastos_admin_backup = []
            cortes_admin_backup = []
            try:
                cierres_admin_backup = json.loads(cierres_file_admin.read_text(encoding="utf-8")) if cierres_file_admin.exists() else []
            except (OSError, json.JSONDecodeError):
                cierres_admin_backup = []
            try:
                gastos_admin_backup = json.loads(gastos_file_admin.read_text(encoding="utf-8")) if gastos_file_admin.exists() else []
            except (OSError, json.JSONDecodeError):
                gastos_admin_backup = []
            try:
                cortes_admin_backup = json.loads(cortes_file_admin.read_text(encoding="utf-8")) if cortes_file_admin.exists() else []
            except (OSError, json.JSONDecodeError):
                cortes_admin_backup = []
            reportes_admin_backup = _leer_lista_json(FISICOS_FILE)
            pagos_admin_backup = cargar_pagos_respaldo()
            eventos_admin_backup = _leer_lista_json(EVENTOS_FILE)
            respaldo_admin = {
                "generado_en": ahora_mx().strftime("%Y-%m-%d %H:%M:%S"),
                "reportes_fisicos": reportes_admin_backup,
                "libro_de_pagos": pagos_admin_backup,
                "eventos_extraordinarios": eventos_admin_backup,
                "vehiculos": st.session_state.vehiculos,
                "cierres_diarios": cierres_admin_backup,
                "gastos_operativos": gastos_admin_backup,
                "cortes_caja": cortes_admin_backup,
                "configuracion_folios": config_admin_backup,
            }
            st.download_button(
                "📦 Descargar respaldo antes de reiniciar",
                data=json.dumps(respaldo_admin, ensure_ascii=False, indent=2).encode("utf-8"),
                file_name=f"respaldo_clubpark_{ahora_mx():%Y%m%d_%H%M%S}.json",
                mime="application/json",
                key="descargar_respaldo_antes_reinicio",
            )
            alcance_reinicio = st.radio(
                "Alcance del reinicio",
                ["Solo una sucursal", "Toda la aplicación"],
                horizontal=True,
                key="alcance_reinicio_admin",
            )
            sucursal_reinicio = st.selectbox(
                "Sucursal que se reiniciará",
                SUCURSALES,
                key="sucursal_reinicio_admin",
                disabled=alcance_reinicio == "Toda la aplicación",
            )
            reiniciar_folios = st.checkbox(
                "También reiniciar el consecutivo de folios a 1 (puede reutilizar números históricos)",
                value=False,
                key="reiniciar_folios_admin",
            )
            texto_confirmacion = "REINICIAR TODO" if alcance_reinicio == "Toda la aplicación" else f"REINICIAR {sucursal_reinicio}"
        
            def normalizar_confirmacion(txt):
                import unicodedata
                t = unicodedata.normalize('NFKD', str(txt or '')).encode('ASCII', 'ignore').decode('ASCII')
                return re.sub(r'[^A-Z0-9]', '', t.upper())

            col_conf1, col_conf2 = st.columns([2, 1])
            with col_conf1:
                confirmacion_reinicio = st.text_input(
                    f'Para confirmar, escribe: {texto_confirmacion}',
                    key="texto_confirmacion_reinicio_admin",
                    placeholder=texto_confirmacion
                )
            with col_conf2:
                st.write("")
                st.write("")
                check_autorizar_directo = st.checkbox(
                    "✅ Autorizar reinicio",
                    key="check_autorizar_reinicio_directo",
                    help="Activa esta casilla para habilitar el botón de reinicio sin preocuparte por acentos o mayúsculas."
                )

            coincide_texto = normalizar_confirmacion(confirmacion_reinicio) == normalizar_confirmacion(texto_confirmacion)
            puede_reiniciar = coincide_texto or check_autorizar_directo

            if not puede_reiniciar:
                st.info(f"💡 Escribe exactamente `{texto_confirmacion}` o marca la casilla 'Autorizar reinicio' para activar el botón.")

            if st.button(
                "🧹 Reiniciar datos operativos seleccionados",
                type="primary" if puede_reiniciar else "secondary",
                disabled=not puede_reiniciar,
                key="ejecutar_reinicio_admin",
            ):
                reiniciar_todo = alcance_reinicio == "Toda la aplicación"
                sucursales_objetivo = set(SUCURSALES if reiniciar_todo else [sucursal_reinicio])
                registros_borrados = [r for r in st.session_state.vehiculos if r.get("sucursal") in sucursales_objetivo]
                registros_conservar = [r for r in st.session_state.vehiculos if r.get("sucursal") not in sucursales_objetivo]
                fotos_borradas = 0
                for registro_borrado in registros_borrados:
                    rutas = []
                    for vista in ("frente", "izquierda", "derecha", "trasera"):
                        ruta = ruta_foto(registro_borrado, vista)
                        if ruta:
                            rutas.append(ruta)
                    for ruta_extra in registro_borrado.get("fotos_custodia", []) or []:
                        try:
                            ruta = Path(ruta_extra)
                            if not ruta.is_absolute():
                                ruta = DATA_DIR / ruta
                            rutas.append(ruta)
                        except (TypeError, ValueError):
                            pass
                    for ruta in rutas:
                        try:
                            if ruta.exists() and ruta.is_file():
                                ruta.unlink()
                                fotos_borradas += 1
                        except OSError:
                            pass
                if guardar_bd(registros_conservar):
                    st.session_state.vehiculos = registros_conservar
                    for archivo_datos, lista_datos in (
                        (cierres_file_admin, cierres_admin_backup),
                        (gastos_file_admin, gastos_admin_backup),
                        (cortes_file_admin, cortes_admin_backup),
                        (PAYMENTS_FILE, pagos_admin_backup),
                        (EVENTOS_FILE, eventos_admin_backup),
                    ):
                        if reiniciar_todo:
                            filtrados = []
                        else:
                            filtrados = [d for d in lista_datos if d.get("sucursal") not in sucursales_objetivo]
                        try:
                            archivo_datos.write_text(json.dumps(filtrados, ensure_ascii=False, indent=2), encoding="utf-8")
                        except OSError as exc:
                            st.error(f"No se pudo actualizar {archivo_datos.name}: {exc}")
                    # Reportes físicos (ingresos, gastos y cierres subidos) y sus fotos
                    reportes_quitar = [r for r in reportes_admin_backup if reiniciar_todo or r.get("Sitio / Unidad") in sucursales_objetivo]
                    for r_q in reportes_quitar:
                        fotos_borradas += _borrar_archivos(r_q.get("fotos"))
                    for c_q in cierres_admin_backup:
                        if reiniciar_todo or c_q.get("sucursal") in sucursales_objetivo:
                            fotos_borradas += _borrar_archivos(c_q.get("fotos_evidencia"))
                    for r_v in registros_borrados:
                        fotos_borradas += _borrar_archivos([v for k, v in (r_v.get("fotos") or {}).items()
                                                            if k not in ("frente", "izquierda", "derecha", "trasera")])
                    try:
                        _guardar_lista_json(FISICOS_FILE, [] if reiniciar_todo else
                                            [r for r in reportes_admin_backup if r.get("Sitio / Unidad") not in sucursales_objetivo])
                    except OSError as exc:
                        st.error(f"No se pudieron reiniciar los reportes físicos: {exc}")
                    if reiniciar_folios:
                        config_nueva = cargar_config()
                        for suc in sucursales_objetivo:
                            anterior = int(config_nueva["siguiente_folio"].get(suc, 1))
                            config_nueva["siguiente_folio"][suc] = 1
                            config_nueva.setdefault("historial_folios", []).insert(0, {
                                "fecha_hora": ahora_mx().strftime("%Y-%m-%d %H:%M:%S"),
                                "usuario": st.session_state.usuario_autenticado,
                                "sucursal": suc,
                                "anterior": anterior,
                                "nuevo": 1,
                                "motivo": "Reinicio administrativo de datos",
                            })
                        guardar_config(config_nueva)
                    guardar_evento_auditoria(
                        "REINICIO ADMINISTRATIVO",
                        f"Alcance: {'todas las sucursales' if reiniciar_todo else sucursal_reinicio}; "
                        f"registros eliminados: {len(registros_borrados)}; fotos eliminadas: {fotos_borradas}; "
                        f"reinicio de folios: {'sí' if reiniciar_folios else 'no'}",
                    )
                    ultimo_comprobante = st.session_state.get("ultimo_comprobante_cobro", {})
                    sucursal_comprobante = (ultimo_comprobante.get("registro") or {}).get("sucursal")
                    if reiniciar_todo or sucursal_comprobante in sucursales_objetivo:
                        st.session_state.pop("ultimo_comprobante_cobro", None)
                    ultimo_folio_sesion = st.session_state.get("ultimo_folio")
                    if any(r.get("id") == ultimo_folio_sesion for r in registros_borrados):
                        st.session_state.pop("ultimo_folio", None)
                    st.success(
                        f"Reinicio terminado. Registros eliminados: {len(registros_borrados)}; "
                        f"fotografías eliminadas: {fotos_borradas}. El historial de auditoría se conservó."
                    )
                    st.rerun()

            st.markdown("#### Administración de registros")
            st.warning("Eliminar un registro es una acción permanente. Descarga un reporte antes de continuar.")
            ids = [r.get("id", "") for r in st.session_state.vehiculos]
            if ids:
                id_borrar = st.selectbox("Folio para eliminar", ids, key="aud_borrar_id")
                confirmar = st.checkbox("Confirmo que deseo eliminar este registro y sus fotografías.")
                if st.button("🗑️ Eliminar registro seleccionado", type="secondary", disabled=not confirmar):
                    registro = next((r for r in st.session_state.vehiculos if r.get("id") == id_borrar), None)
                    if registro:
                        for vista in ("frente", "izquierda", "derecha", "trasera"):
                            p = ruta_foto(registro, vista)
                            if p:
                                try:
                                    p.unlink()
                                except OSError:
                                    pass
                        st.session_state.vehiculos = [
                            r for r in st.session_state.vehiculos if r.get("id") != id_borrar
                        ]
                        if guardar_bd(st.session_state.vehiculos):
                            st.success(f"Registro {id_borrar} eliminado.")
                            st.rerun()

if _SECCION == "aud" and "aud" in _TABS:
    _seccion_aud()

st.sidebar.caption(
    "Datos guardados en: " + str(DATA_DIR) +
    "\n\nEn Streamlit Cloud, configura almacenamiento persistente externo para evitar pérdidas."
)


# ---------------------------------------------------------------------
# REGISTRO DE AUDITORÍA (eventos recientes)
# ---------------------------------------------------------------------
if es_contador and _SECCION == "aud":
    st.divider()
    st.subheader("Historial de auditoría del sistema")
    st.caption(
        "Bitácora institucional: registra altas, cobros, salidas, cierres, cambios de folio e impresión, "
        "con usuario, rol, fecha y detalle del cambio."
    )
    eventos = cargar_auditoria()
    if eventos:
        col_f1, col_f2, col_f3 = st.columns(3)
        tipos = ["Todos"] + sorted({str(e.get("tipo", "OTRO")) for e in eventos})
        usuarios = ["Todos"] + sorted({str(e.get("usuario", "Sistema")) for e in eventos})
        filtro_tipo = col_f1.selectbox("Filtrar por evento", tipos, key="aud_filtro_tipo")
        filtro_usuario = col_f2.selectbox("Filtrar por usuario", usuarios, key="aud_filtro_usuario")
        busqueda = col_f3.text_input("Buscar en el detalle", key="aud_busqueda").strip().lower()
        eventos_filtrados = eventos
        if filtro_tipo != "Todos":
            eventos_filtrados = [e for e in eventos_filtrados if str(e.get("tipo", "OTRO")) == filtro_tipo]
        if filtro_usuario != "Todos":
            eventos_filtrados = [e for e in eventos_filtrados if str(e.get("usuario", "Sistema")) == filtro_usuario]
        if busqueda:
            eventos_filtrados = [e for e in eventos_filtrados if busqueda in json.dumps(e, ensure_ascii=False).lower()]
        st.metric("Eventos en la bitácora", len(eventos_filtrados))
        st.dataframe(pd.DataFrame(eventos_filtrados), use_container_width=True, hide_index=True)
        st.download_button(
            "📥 Descargar bitácora de auditoría (CSV)",
            data=pd.DataFrame(eventos_filtrados).to_csv(index=False).encode("utf-8-sig"),
            file_name=f"Bitacora_Auditoria_ClubPark_{ahora_mx():%Y%m%d_%H%M}.csv",
            mime="text/csv",
            key="descargar_bitacora_auditoria",
        )
    else:
        st.info("La bitácora registrará aquí altas, cobros, salidas, cierres y cambios de folios.")


# ---------------------------------------------------------------------
# TAB 10: CREDENCIALES Y CÓDIGOS QR INSTITUCIONALES
# ---------------------------------------------------------------------
@_fragmento
def _seccion_qrs():
    global col, data, idx
    globals()["_CONTADOR_LLAVES"] = 1000  # llaves estables entre recargas completas y de sección
    with tab_qrs:
        if es_contador:
            st.markdown(f"""
        <div class="clubpark-header-card">
          <div class="clubpark-badge" style="margin-top:4px">CATÁLOGO OFICIAL DE ACCESOS Y CREDENCIALES</div>
          <div class="cpv-titulo-blanco" style="margin:6px 0 2px 0">CÓDIGOS QR OFICIALES DE SUCURSALES Y SUPERVISIÓN</div>
          <div style="font-size:13px;color:#EBF3EE">
            CLUB PARK VALET Y ESTACIONAMIENTOS S.A. DE C.V. &nbsp;|&nbsp; SERVIR CON HONESTIDAD, ACTUAR CON INTEGRIDAD Y GENERAR CONFIANZA
          </div>
        </div>
        """, unsafe_allow_html=True)

            st.write(
                "A continuación se presentan todas las credenciales QR para el Contador Auditor, "
                "los Supervisores Globales y los Líderes de las sucursales operativas. "
                "Puedes descargar individualmente cada código QR o descargarlos todos juntos en un archivo ZIP."
            )

            if ACCESOS_FUENTE == "por_defecto":
                st.warning(
                    "⚠️ **Credenciales en modo contingencia:** los tokens por defecto todavía viven en el código fuente. "
                    "Muévelos a `.streamlit/secrets.toml` (clave `CLUBPARK_TOKENS`) y gírelos de vez en cuando. "
                    "Consulta la guía de implementación adjunta."
                )
            else:
                st.success(
                    "✅ Credenciales protegidas: se cargan desde "
                    + ("Streamlit Secrets" if ACCESOS_FUENTE == "secrets" else "variables de entorno")
                    + " y no están expuestas en el código fuente."
                )

            # Paquete ZIP completo generado en memoria (sin rutas de máquina de desarrollo)
            if QR_OK and qrcode:
                buffer_zip = io.BytesIO()
                with zipfile.ZipFile(buffer_zip, "w", zipfile.ZIP_DEFLATED) as archivo_zip:
                    for token_zip in ACCESOS:
                        archivo_zip.writestr(f"QR_{token_zip}.png", qr_png_cache(contenido_qr_para(token_zip)))
                st.download_button(
                    f"📦 Descargar paquete completo de {len(ACCESOS)} Códigos QR (ZIP)",
                    data=buffer_zip.getvalue(),
                    file_name="Credenciales_QR_ClubPark_Completas.zip",
                    mime="application/zip",
                    use_container_width=True,
                    type="primary",
                    key="descargar_todos_qrs_zip",
                )
            if APP_URL:
                st.info(f"Los QR incluyen el enlace de acceso directo: {APP_URL} — al escanearlos con la cámara del teléfono la sesión se abre sola.")

            filtro_rol = st.radio(
                "Filtrar catálogo:",
                [f"Todos ({len(ACCESOS)} credenciales)", "Auditoría y Supervisión Global", "Líderes de Sucursal"],
                horizontal=True,
                key="filtro_credenciales_qr"
            )

            # Filtrar tokens
            tokens_mostrar = {}
            for tok, data in ACCESOS.items():
                es_sup = "Global" in data.get("rol", "") or "Auditor" in data.get("rol", "")
                if filtro_rol == "Auditoría y Supervisión Global" and not es_sup:
                    continue
                if filtro_rol == "Líderes de Sucursal" and es_sup:
                    continue
                tokens_mostrar[tok] = data

            buscar_cred = st.text_input("🔍 Buscar credencial", placeholder="Nombre, sucursal, rol o token (ej. barranca, josé, CPV-QR-...)",
                                        key="buscar_credencial_qr")
            if buscar_cred.strip():
                patron = _sin_acentos(buscar_cred.strip())
                tokens_mostrar = {t: d for t, d in tokens_mostrar.items()
                                  if patron in _sin_acentos(" ".join([t, str(d.get("nombre", "")), str(d.get("sucursal", "")), str(d.get("rol", ""))]))}
                st.caption(f"{len(tokens_mostrar)} credencial(es) encontrada(s).")
                if not tokens_mostrar:
                    st.info("No se encontró ninguna credencial con ese texto.")

            cols = st.columns(3)
            for idx, (tok, data) in enumerate(tokens_mostrar.items()):
                col = cols[idx % 3]
                with col:
                    st.markdown(f"""
                <div style="border:2px solid #0B3B24;border-radius:10px;padding:12px;background:#FFFFFF;margin-bottom:14px;box-shadow:0 3px 8px rgba(11,59,36,0.12)">
                  <div style="background:#0B3B24;color:#FFFFFF;padding:6px 10px;border-radius:6px;font-size:12px;font-weight:700;margin-bottom:8px;text-align:center">
                    {data.get('rol', 'Supervisor')}
                  </div>
                  <div style="font-size:14px;font-weight:800;color:#0B3B24;text-align:center">{data.get('nombre', '')}</div>
                  <div style="font-size:12px;color:#5B8F6E;font-weight:700;text-align:center;margin-bottom:6px">SUCURSAL: {data.get('sucursal', '')}</div>
                  <div style="font-family:monospace;font-size:11px;background:#EBF3EE;padding:4px;border-radius:4px;text-align:center;word-break:break-all;color:#072617;margin-bottom:8px">
                    {tok}
                  </div>
                </div>
                """, unsafe_allow_html=True)

                    if QR_OK and qrcode:
                        qr_bytes = qr_png_cache(contenido_qr_para(tok))
                        st.image(qr_bytes, width=160)
                        st.download_button(
                            label=f"⬇️ Descargar QR",
                            data=qr_bytes,
                            file_name=f"QR_{tok}.png",
                            mime="image/png",
                            key=f"dl_qr_tab_{tok}",
                            use_container_width=True
                        )
        else:
            st.error("🔒 Catálogo restringido: solo el Contador Auditor (máxima jerarquía) puede consultar y descargar las credenciales QR oficiales. Pide la rotación y descarga de credenciales al auditor.")

if _SECCION == "qrs" and "qrs" in _TABS:
    _seccion_qrs()


# ---------------------------------------------------------------------
# TAB: PRUEBAS OPERATIVAS
# ---------------------------------------------------------------------
@_fragmento
def _seccion_pruebas():
    globals()["_CONTADOR_LLAVES"] = 1000  # llaves estables entre recargas completas y de sección
    with tab_pruebas:
        encabezado_institucional("PRUEBAS OPERATIVAS DEL SISTEMA", "VERIFICACIÓN PREVIA", "Verificación formal de componentes críticos antes de abrir la jornada")
        st.caption(
            "Verifica en segundos que los componentes críticos funcionan antes de abrir la jornada: "
            "base de datos, folios, códigos QR, boletos térmicos, credenciales, bitácora y fotografías."
        )
        panel_prueba_impresion(key_sufijo="pruebas", expandido=False)
        if st.button("▶️ Ejecutar pruebas operativas", type="primary", use_container_width=True, key="btn_ejecutar_pruebas"):
            resultados_pruebas = ejecutar_pruebas_operativas()
            st.session_state["resultados_pruebas"] = resultados_pruebas
        if st.session_state.get("resultados_pruebas"):
            resultados_pruebas = st.session_state["resultados_pruebas"]
            exitos = sum(1 for r in resultados_pruebas if r["estado"] == "✅")
            col_a, col_b = st.columns(2)
            col_a.metric("Pruebas aprobadas", f"{exitos}/{len(resultados_pruebas)}")
            col_b.metric("Estado general", "OPERATIVO ✅" if exitos == len(resultados_pruebas) else "REVISAR ❌")
            st.dataframe(pd.DataFrame(resultados_pruebas), use_container_width=True, hide_index=True)
            if exitos == len(resultados_pruebas):
                st.success("Todas las pruebas operativas pasaron. El sistema está listo para la jornada.")
            else:
                st.error("Hay pruebas fallidas. Atiende los detalles antes de operar.")

if _SECCION == "pruebas" and "pruebas" in _TABS:
    _seccion_pruebas()


@_fragmento
def _seccion_intel():
    global alertas, clave, col, e, kpis, recomendaciones, resumen, titulo, valor
    globals()["_CONTADOR_LLAVES"] = 1000  # llaves estables entre recargas completas y de sección
    with tab_intel:
        encabezado_institucional("INTELIGENCIA OPERATIVA", "ANÁLISIS CONTINUO", "El sistema se observa a sí mismo: indicadores, alertas, recomendaciones y evolución")
        st.caption(
            "Panel ejecutivo que se actualiza automáticamente con cada operación registrada: "
            "indicadores, alertas tempranas, resumen del día y recomendaciones generadas a partir "
            "de la información real del sistema."
        )
        registros_intel = vehiculos_visibles(st.session_state.vehiculos)
        if not registros_intel:
            st.info("El análisis se activará automáticamente al registrar los primeros vehículos.")
        else:
            kpis, alertas, resumen, recomendaciones = analisis_inteligente(registros_intel)
            col_k1, col_k2, col_k3, col_k4 = st.columns(4)
            for col, (titulo, valor) in zip((col_k1, col_k2, col_k3, col_k4), kpis):
                col.metric(titulo, valor)
            st.markdown("**Resumen ejecutivo del día**")
            st.info(resumen)
            st.markdown("**Alertas automáticas**")
            if alertas:
                st.dataframe(pd.DataFrame(alertas), use_container_width=True, hide_index=True)
                if st.button("📤 Registrar alertas en la bitácora de auditoría", key="btn_intel_bitacora", use_container_width=True):
                    for alerta in alertas:
                        guardar_evento_auditoria("Inteligencia Operativa", f"{alerta['tipo']}: {alerta['detalle']}")
                    st.success("Alertas registradas en la bitácora de auditoría.")
            else:
                st.success("Sin alertas: no hay estancias críticas, placas duplicadas ni cobros pendientes detectados.")
            st.markdown("**Recomendaciones automáticas**")
            for recomendacion in recomendaciones:
                st.markdown(f"- {recomendacion}")

            st.markdown("**Evolución del sistema (aprendizaje continuo)**")
            notas_evolucion = evolucion_aprendizaje(registros_intel)
            if notas_evolucion:
                for nota in notas_evolucion[:5]:
                    st.markdown(f"- **{nota['fecha']}** · {nota['vehiculos_totales']} vehículos: {nota['observaciones']}")
            else:
                st.caption("Cada 25 vehículos registrados, el sistema documentará aquí lo que ha aprendido de la operación.")

        st.markdown("**Salud de la impresión térmica (RawBT)**")
        eventos_imp = cargar_registro_impresion()
        fallos_imp = [e for e in eventos_imp if e.get("resultado") == "FALLO"]
        col_i1, col_i2, col_i3 = st.columns(3)
        col_i1.metric("Intentos registrados", len(eventos_imp))
        col_i2.metric("Fallos registrados", len(fallos_imp))
        col_i3.metric("Último evento", eventos_imp[0]["fecha_hora"] if eventos_imp else "—")
        if fallos_imp:
            por_etapa = {}
            for e in fallos_imp:
                clave = e.get("etapa", "OTRO")
                por_etapa[clave] = por_etapa.get(clave, 0) + 1
            detalle_etapas = " · ".join(
                f"{ETAPAS_IMPRESION.get(k, k)}: {v}" for k, v in por_etapa.items()
            )
            st.warning("Fallos de impresión detectados → " + detalle_etapas)
            st.dataframe(pd.DataFrame(fallos_imp[:10]), use_container_width=True, hide_index=True)
        elif eventos_imp:
            st.success("La impresión térmica opera sin fallos registrados.")
        else:
            st.caption("Aún no hay eventos de impresión: ejecuta la prueba de impresión en 🧪 Pruebas Operativas o en 🖨️ Re-Impresión Térmica.")

if _SECCION == "intel" and "intel" in _TABS:
    _seccion_intel()


# ---------------------------------------------------------------------
# TAB: REPORTES FÍSICOS Y EXCEL OFICIAL
# ---------------------------------------------------------------------
# Sin fragmento: las cámaras y la subida de fotos fallan dentro de fragmentos de Streamlit
def _seccion_reportes():
    global celda, exc, fila, letra
    globals()["_CONTADOR_LLAVES"] = 1000  # llaves estables entre recargas completas y de sección
    with tab_reportes:
        encabezado_institucional("REPORTES FÍSICOS Y ENTREGA", "INGRESOS · GASTOS · EXCEL", "Sube la hoja física de ingresos o de gastos y entrega el Excel oficial")
        if "rf_nonce" not in st.session_state:
            st.session_state.rf_nonce = 0
        nonce = st.session_state.rf_nonce
        sucursales_rf = SUCURSALES if es_global else [st.session_state.sucursal_asignada]
        # Menú de apartados: solo se ejecuta el apartado abierto (con pestañas se calculaban los 5 en cada toque)
        _opciones_sub_rf = {"subir": "📤 Subir reporte", "lista": "🗂️ Reportes subidos", "reporte": "📈 Reporte por sucursal / global",
                            "superv": "👔 Reporte por supervisor", "excel": "📊 Excel oficial"}
        _sub_rf = st.radio("Apartado", list(_opciones_sub_rf), format_func=lambda k: _opciones_sub_rf[k],
                           horizontal=True, key="rf_apartado", label_visibility="collapsed")
        aviso_rf = st.session_state.pop("rf_aviso_guardado", None)
        if aviso_rf:
            st.success(aviso_rf + " Ya puedes subir el siguiente reporte.")

        # ---------------- SUBIR ----------------
        if _sub_rf == "subir":
            puede_subir_rf = es_contador or st.session_state.get("rol") == "Supervisor Global"
            if not puede_subir_rf:
                st.error("Solo el Contador Auditor y los supervisores pueden subir reportes físicos. Aquí puedes consultar los reportes de tu sucursal.")
            else:
                if st.button("🧹 Limpiar formulario y subir otro reporte", key=f"rf_limpiar_{nonce}"):
                    _limpiar_formulario_rf()
                    st.rerun()
                tipo_rf = st.radio("¿Qué reporte vas a subir?", ["📥 Reporte de ingresos", "💸 Reporte de gastos"], horizontal=True, key=f"rf_tipo_{nonce}")
                es_ingresos = tipo_rf.startswith("📥")
                rc1, rc2 = st.columns(2)
                with rc1:
                    suc_rf = st.selectbox("Sucursal", sucursales_rf, key=f"rf_suc_{nonce}")
                with rc2:
                    fecha_rf = st.date_input("Fecha de operación", value=ahora_mx().date(), key=f"rf_fecha_{nonce}")
                fecha_rf_str = fecha_rf.strftime("%d/%m/%Y")

                modo_rf = st.radio("Captura de la hoja física", ["📁 Galería", "📸 Cámara"], horizontal=True, key=f"rf_modo_{nonce}")
                fotos_rf = []
                if modo_rf.startswith("📸"):
                    foto_rf = st.camera_input("Foto de la hoja física", key=f"rf_cam_{nonce}")
                    if foto_rf:
                        fotos_rf = [foto_rf]
                else:
                    fotos_rf = list(st.file_uploader("Fotos de la hoja (puedes subir varias)", type=["png", "jpg", "jpeg", "webp"],
                                                     accept_multiple_files=True, key=f"rf_gal_{nonce}") or [])
                if fotos_rf:
                    cols_rf = st.columns(3)
                    for i_rf, f_rf in enumerate(fotos_rf):
                        cols_rf[i_rf % 3].image(f_rf, use_container_width=True)

                if fotos_rf:
                    if not ia_disponible():
                        st.caption("🤖 La lectura automática con IA no está activada (falta ANTHROPIC_API_KEY).")
                    elif st.button("🤖 Leer la hoja con IA y prellenar", key=f"rf_ia_btn_{nonce}", type="primary"):
                        try:
                            with st.spinner("Leyendo la hoja…"):
                                lectura_rf = ia_leer_reporte(fotos_rf, es_ingresos)
                            if es_ingresos:
                                st.session_state[f"rf_fini_{nonce}"] = _entero_seguro(lectura_rf.get("folio_inicial"))
                                st.session_state[f"rf_ffin_{nonce}"] = _entero_seguro(lectura_rf.get("folio_final"))
                                st.session_state[f"rf_sin_{nonce}"] = _entero_seguro(lectura_rf.get("boletos_sin_cobro"))
                                st.session_state[f"rf_canc_{nonce}"] = _entero_seguro(lectura_rf.get("boletos_cancelados"))
                                st.session_state[f"rf_perd_{nonce}"] = _entero_seguro(lectura_rf.get("boletos_perdidos"))
                                st.session_state[f"rf_directo_{nonce}"] = _flotante_seguro(lectura_rf.get("ingreso_bruto"))
                                st.session_state[f"rf_nomina_{nonce}"] = _flotante_seguro(lectura_rf.get("nomina"))
                                st.session_state[f"rf_otros_{nonce}"] = _flotante_seguro(lectura_rf.get("otros_gastos"))
                            else:
                                filas_ia = []
                                for g_ia in lectura_rf.get("gastos", []) or []:
                                    cat_ia = str(g_ia.get("categoria", "Otros")).strip().capitalize()
                                    cat_ia = next((c for c in CATEGORIAS_GASTO_RF if c.lower() == cat_ia.lower()), "Otros")
                                    filas_ia.append({"Categoría": cat_ia, "Concepto": str(g_ia.get("concepto", "")).strip(), "Importe": _flotante_seguro(g_ia.get("importe"))})
                                st.session_state[f"rf_gastos_ia_{nonce}"] = filas_ia
                                st.session_state[f"rf_ia_ver_{nonce}"] = st.session_state.get(f"rf_ia_ver_{nonce}", 0) + 1
                            st.session_state[f"rf_ia_aviso_{nonce}"] = lectura_rf
                            guardar_evento_auditoria("LECTURA IA REPORTE FISICO", f"{'ingresos' if es_ingresos else 'gastos'} · {suc_rf} · {fecha_rf_str}")
                        except RuntimeError as exc:
                            st.error(str(exc))
                    lectura_previa = st.session_state.get(f"rf_ia_aviso_{nonce}")
                    if lectura_previa:
                        st.success("Datos prellenados por IA. **Revísalos contra la hoja antes de guardar**: la IA puede equivocarse con letra a mano.")
                        if str(lectura_previa.get("dudas", "")).strip():
                            st.warning("La IA marcó dudas: " + str(lectura_previa["dudas"]))
                        sitio_ia = str(lectura_previa.get("sitio", "")).strip().upper()
                        if sitio_ia and sitio_ia not in str(suc_rf).upper() and str(suc_rf).upper() not in sitio_ia:
                            st.warning(f"La hoja parece ser de «{sitio_ia}» pero elegiste {suc_rf}.")
                        fecha_ia = str(lectura_previa.get("fecha", "")).strip()
                        if fecha_ia and fecha_ia != fecha_rf_str:
                            st.warning(f"La hoja parece ser del {fecha_ia} pero elegiste {fecha_rf_str}.")

                reportes_actuales = _leer_lista_json(FISICOS_FILE)

                if es_ingresos:
                    directo = es_sucursal_directa(suc_rf)
                    cfg_rf = tarifa_de(suc_rf)
                    # Salones de eventos: si ese día hay 2 o más eventos, se puede subir un reporte global o uno por evento
                    evento_num_rf, eventos_dia_rf = 0, 1
                    if suc_rf in SALONES_EVENTOS:
                        eventos_dia_rf = sum(max(1, _entero_seguro(e.get("num_eventos"))) for e in cargar_eventos()
                                             if e.get("sucursal") == suc_rf and fecha_flexible(e.get("fecha")) == fecha_rf
                                             and e.get("estado") != "Cancelado")
                        if eventos_dia_rf >= 2:
                            st.info(f"🎉 {suc_rf} tiene **{eventos_dia_rf} eventos** el {fecha_rf_str}. Puedes subir un reporte global del día o uno por cada evento.")
                            modo_ev_rf = st.radio("Tipo de reporte", ["📋 Un reporte global del día", f"🎉 Un reporte por evento ({eventos_dia_rf} reportes)"],
                                                  horizontal=True, key=f"rf_modo_ev_{nonce}_{suc_rf}_{fecha_rf_str}")
                            if modo_ev_rf.startswith("🎉"):
                                subidos_ev = {_entero_seguro(r.get("evento_num")) for r in reportes_actuales
                                              if r.get("tipo_reporte") == "ingresos" and r.get("Sitio / Unidad") == suc_rf
                                              and r.get("Fecha") == fecha_rf_str and r.get("origen") != "Cierre del sistema"}
                                evento_num_rf = st.selectbox(
                                    "¿De qué evento es este reporte?", list(range(1, eventos_dia_rf + 1)),
                                    format_func=lambda k: f"Evento {k} de {eventos_dia_rf}" + (" · ✅ ya subido" if k in subidos_ev else ""),
                                    key=f"rf_ev_num_{nonce}_{suc_rf}_{fecha_rf_str}")
                        elif eventos_dia_rf == 0:
                            st.caption("No hay eventos registrados para este salón en esa fecha: se sube un reporte global del día.")
                            eventos_dia_rf = 1
                    dia_rf = [r for r in reportes_actuales if r.get("tipo_reporte") == "ingresos"
                              and r.get("Sitio / Unidad") == suc_rf and r.get("Fecha") == fecha_rf_str]
                    auto_rf = [r for r in dia_rf if r.get("origen") == "Cierre del sistema"]
                    fisicos_dia_rf = [r for r in dia_rf if r.get("origen") != "Cierre del sistema"]
                    globales_rf = [r for r in fisicos_dia_rf if not _entero_seguro(r.get("evento_num"))]
                    por_evento_rf = [r for r in fisicos_dia_rf if _entero_seguro(r.get("evento_num"))]
                    if evento_num_rf == 0:
                        choque_rf = fisicos_dia_rf  # el global no puede convivir con otros reportes del día
                    else:
                        choque_rf = globales_rf + [r for r in por_evento_rf if _entero_seguro(r.get("evento_num")) == evento_num_rf]
                    reemplazar_ids_rf = {r.get("ID") for r in auto_rf} | ({r.get("ID") for r in choque_rf} if es_contador else set())
                    existente_rf = choque_rf[0] if choque_rf else None
                    existente_auto = bool(auto_rf) and not choque_rf
                    etiqueta_rep = f"del evento {evento_num_rf} de {eventos_dia_rf}" if evento_num_rf else "global del día"
                    if existente_auto:
                        st.info("Existe un reporte automático del cierre del sistema para esta sucursal y fecha. Al guardar, el reporte físico lo reemplaza.")
                    elif choque_rf:
                        detalle_choque = ", ".join((r.get("Tipo") if _entero_seguro(r.get("evento_num")) else "reporte global") + f" ({r.get('ID')})"
                                                   for r in choque_rf)
                        if es_contador:
                            st.warning(f"Ya existe: {detalle_choque}. Al guardar el reporte {etiqueta_rep}, se reemplazará.")
                        elif evento_num_rf and globales_rf:
                            st.error(f"Ya se subió un reporte global de {suc_rf} para {fecha_rf_str}. No se pueden subir reportes por evento; solo el Contador Auditor puede borrarlo o reemplazarlo.")
                        elif not evento_num_rf and por_evento_rf:
                            st.error(f"Ya se subieron reportes por evento de {suc_rf} para {fecha_rf_str}. Sube los eventos que falten o pide al Contador Auditor que los cambie.")
                        else:
                            st.error(f"Ya existe el reporte {etiqueta_rep} de {suc_rf} para {fecha_rf_str}. Solo el Contador Auditor puede reemplazarlo o borrarlo.")
                    if directo:
                        st.info("📌 Esta sucursal reporta ingreso bruto directo (sin control por folio).")
                        ingreso_directo = st.number_input("Ingreso bruto del evento ($)" if evento_num_rf else "Ingreso bruto del día ($)",
                                                          min_value=0.0, step=50.0, key=f"rf_directo_{nonce}")
                        costo_rf, f_ini_rf, f_fin_rf, sin_rf, canc_rf, perd_rf = 0.0, 0, 0, 0, 0, 0
                        total_bol_rf, cobrados_rf, bruto_rf = 0, 0, float(ingreso_directo)
                    else:
                        costo_rf = st.number_input("Tarifa por boleto ($)", min_value=0.0, value=float(cfg_rf["precio_fijo"]), step=5.0, key=f"rf_costo_{nonce}_{suc_rf}")
                        n1, n2, n3 = st.columns(3)
                        f_ini_rf = n1.number_input("Folio inicial", min_value=0, step=1, key=f"rf_fini_{nonce}")
                        f_fin_rf = n1.number_input("Folio final", min_value=0, step=1, key=f"rf_ffin_{nonce}")
                        sin_rf = n2.number_input("Boletos sin cobro (cortesía)", min_value=0, step=1, key=f"rf_sin_{nonce}")
                        canc_rf = n2.number_input("Boletos cancelados", min_value=0, step=1, key=f"rf_canc_{nonce}")
                        perd_rf = n3.number_input("Boletos perdidos", min_value=0, step=1, key=f"rf_perd_{nonce}")
                        total_bol_rf = calcular_boletos_block100(int(f_ini_rf), int(f_fin_rf))
                        cobrados_rf = max(0, total_bol_rf - int(sin_rf) - int(canc_rf))
                        bruto_rf = cobrados_rf * float(costo_rf)
                        if int(f_ini_rf) > 0 or int(f_fin_rf) > 0:
                            st.info(f"Folio {int(f_ini_rf)} → {int(f_fin_rf)}: {total_bol_rf} boletos (ciclo de 100); menos cortesías y cancelados = {cobrados_rf} cobrados.")
                    exceso_rf = (not directo) and (int(sin_rf) + int(canc_rf) + int(perd_rf) > total_bol_rf)
                    if exceso_rf:
                        st.error("Cortesías, cancelados y perdidos superan el total de boletos. Corrige antes de guardar.")
                    st.metric("Ingreso bruto del reporte", f"${bruto_rf:,.2f} MXN")
                    g1, g2 = st.columns(2)
                    nomina_rf = g1.number_input("Gasto de nómina del día ($)", min_value=0.0, step=50.0, key=f"rf_nomina_{nonce}")
                    otros_rf = g2.number_input("Otros gastos del día ($)", min_value=0.0, step=10.0, key=f"rf_otros_{nonce}")
                    st.metric("Utilidad del día (ingreso − gastos)", f"${bruto_rf - nomina_rf - otros_rf:,.2f} MXN")
                    e1, e2 = st.columns(2)
                    # Elaboró = Líder de la sucursal · Recibe = supervisor (el de la sucursal por defecto)
                    lider_rf = next((d.get("nombre") for d in ACCESOS.values()
                                     if d.get("sucursal") == suc_rf and normalizar_rol(d.get("rol")) == ROL_LIDER), "")
                    elaboro_rf = e1.text_input("Elaboró (Líder de sucursal)", value=lider_rf or f"Líder {suc_rf.title()}",
                                               key=f"rf_elab_{nonce}_{suc_rf}")
                    sup_rf = next((i.get("supervisor") for i in CATALOGO_OFICIAL_23 if i["sitio"] == suc_rf), "")
                    nombres_sup_rf = list(SUPERVISORES_NOMBRE.values())
                    recibe_rf = e2.selectbox("Recibe (Supervisor)", nombres_sup_rf,
                                             index=nombres_sup_rf.index(SUPERVISORES_NOMBRE[sup_rf]) if sup_rf in SUPERVISORES_NOMBRE else 0,
                                             key=f"rf_recibe_{nonce}_{suc_rf}")
                    cero_ok = True
                    if bruto_rf <= 0:
                        cero_ok = st.checkbox("Confirmo que el ingreso del día fue $0", key=f"rf_cero_{nonce}")
                    bloquear_rf = bool(choque_rf and not es_contador) or exceso_rf or not cero_ok
                    if st.button("💾 Guardar reporte de ingresos", type="primary", use_container_width=True, disabled=bloquear_rf, key=f"rf_btn_ing_{nonce}"):
                        if not fotos_rf and not es_contador:
                            st.error("Adjunta la foto de la hoja física para guardar el reporte.")
                        else:
                            rid = f"REP_{uuid.uuid4().hex[:8]}"
                            try:
                                rutas_rf = guardar_fotos_reporte(fotos_rf, rid)
                                registro_rf = _fila_flujo(suc_rf, fecha_rf, bruto_rf, cobrados_rf, costo_rf, f_ini_rf, f_fin_rf, total_bol_rf,
                                                          elaboro_rf or st.session_state.usuario_autenticado, st.session_state.usuario_autenticado, rid)
                                if eventos_dia_rf >= 2:
                                    registro_rf["Tipo"] = f"Evento {evento_num_rf} de {eventos_dia_rf}" if evento_num_rf else f"Global del día ({eventos_dia_rf} eventos)"
                                registro_rf.update({"evento_num": int(evento_num_rf), "eventos_dia": int(eventos_dia_rf)})
                                registro_rf.update({"Sin Cobro": int(sin_rf), "Cancelados": int(canc_rf), "Perdidos": int(perd_rf),
                                                    "Recibe": recibe_rf, "tipo_reporte": "ingresos", "fotos": rutas_rf,
                                                    "Nómina": float(nomina_rf), "Otros gastos": float(otros_rf),
                                                    "Timestamp": ahora_mx().strftime("%Y-%m-%d %H:%M:%S")})
                                reportes_actuales = [r for r in reportes_actuales if r.get("ID") not in reemplazar_ids_rf]
                                reportes_actuales.insert(0, registro_rf)
                                _guardar_lista_json(FISICOS_FILE, reportes_actuales)
                                # los gastos del reporte pasan al libro de gastos (reemplazo = se borran los del reporte anterior)
                                gastos_archivo = _leer_lista_json(GASTOS_FILE_RF)
                                if reemplazar_ids_rf:
                                    gastos_archivo = [g for g in gastos_archivo if g.get("reporte_fisico_id") not in reemplazar_ids_rf]
                                for cat_g, monto_g in (("Nómina", nomina_rf), ("Otros", otros_rf)):
                                    if monto_g > 0:
                                        gastos_archivo.insert(0, {
                                            "id": uuid.uuid4().hex[:12].upper(), "fecha": str(fecha_rf), "sucursal": suc_rf,
                                            "categoria": cat_g, "concepto": f"{cat_g} (reporte de ingresos)", "importe": float(monto_g),
                                            "registrado_por": st.session_state.usuario_autenticado,
                                            "registrado_en": ahora_mx().strftime("%Y-%m-%d %H:%M:%S"),
                                            "origen": "Reporte físico de ingresos", "reporte_fisico_id": rid,
                                        })
                                _guardar_lista_json(GASTOS_FILE_RF, gastos_archivo)
                                guardar_evento_auditoria("REPORTE FISICO DE INGRESOS", f"{suc_rf} · {fecha_rf_str} · {etiqueta_rep} · ${bruto_rf:,.2f} · {rid}" + (" · REEMPLAZO" if (choque_rf and es_contador) else ""))
                                _limpiar_formulario_rf()
                                st.session_state.rf_aviso_guardado = f"✅ Reporte de ingresos guardado ({rid})."
                                st.rerun()
                            except Exception as exc:
                                st.error(f"No se pudo guardar el reporte: {exc}")
                else:
                    existente_gas = next((r for r in reportes_actuales if r.get("tipo_reporte") == "gastos"
                                          and r.get("Sitio / Unidad") == suc_rf and r.get("Fecha") == fecha_rf_str), None)
                    if existente_gas:
                        if es_contador:
                            st.warning(f"Ya existe un reporte de gastos de {suc_rf} para {fecha_rf_str} ({existente_gas.get('ID')}). Al guardar, se reemplazará.")
                        else:
                            st.error(f"Ya existe un reporte de gastos de {suc_rf} para {fecha_rf_str}. Solo el Contador Auditor puede reemplazarlo o borrarlo.")
                    st.markdown("**Captura los gastos de la hoja** (agrega una fila por gasto):")
                    tabla_gastos = st.data_editor(
                        pd.DataFrame(st.session_state.get(f"rf_gastos_ia_{nonce}") or [{"Categoría": "Nómina", "Concepto": "", "Importe": 0.0}]),
                        num_rows="dynamic", use_container_width=True, hide_index=True,
                        key=f"rf_gastos_{nonce}_{st.session_state.get(f'rf_ia_ver_{nonce}', 0)}",
                        column_config={
                            "Categoría": st.column_config.SelectboxColumn("Categoría", options=CATEGORIAS_GASTO_RF, required=True),
                            "Concepto": st.column_config.TextColumn("Concepto"),
                            "Importe": st.column_config.NumberColumn("Importe ($)", min_value=0.0, step=10.0, format="$%.2f"),
                        },
                    )
                    filas_gasto = []
                    for fila in tabla_gastos.to_dict("records"):
                        concepto_g = str(fila.get("Concepto") or "").strip()
                        try:
                            importe_g = float(fila.get("Importe") or 0)
                        except (TypeError, ValueError):
                            importe_g = 0.0
                        if concepto_g and importe_g > 0:
                            filas_gasto.append({"categoria": fila.get("Categoría") or "Otros", "concepto": concepto_g, "importe": importe_g})
                    total_gastos_rf = sum(f["importe"] for f in filas_gasto)
                    st.metric("Total de gastos del reporte", f"${total_gastos_rf:,.2f} MXN")
                    if st.button("💾 Guardar reporte de gastos", type="primary", use_container_width=True, key=f"rf_btn_gas_{nonce}",
                                 disabled=bool(existente_gas and not es_contador)):
                        if not filas_gasto:
                            st.error("Captura al menos un gasto con concepto e importe mayor a cero.")
                        elif not fotos_rf and not es_contador:
                            st.error("Adjunta la foto de la hoja o comprobante para guardar el reporte.")
                        else:
                            rid = f"GAS_{uuid.uuid4().hex[:8]}"
                            try:
                                if existente_gas:  # reemplazo del Contador: se quita el reporte anterior y sus gastos
                                    borrar_reporte_fisico(existente_gas.get("ID"))
                                    reportes_actuales = _leer_lista_json(FISICOS_FILE)
                                rutas_rf = guardar_fotos_reporte(fotos_rf, rid)
                                gastos_archivo = _leer_lista_json(GASTOS_FILE_RF)
                                for fila in filas_gasto:
                                    gastos_archivo.insert(0, {
                                        "id": uuid.uuid4().hex[:12].upper(), "fecha": str(fecha_rf), "sucursal": suc_rf,
                                        "categoria": fila["categoria"], "concepto": fila["concepto"], "importe": fila["importe"],
                                        "registrado_por": st.session_state.usuario_autenticado,
                                        "registrado_en": ahora_mx().strftime("%Y-%m-%d %H:%M:%S"),
                                        "origen": "Reporte físico", "reporte_fisico_id": rid, "fotos": rutas_rf,
                                    })
                                _guardar_lista_json(GASTOS_FILE_RF, gastos_archivo)
                                reportes_actuales.insert(0, {
                                    "ID": rid, "tipo_reporte": "gastos", "Fecha": fecha_rf_str, "Sitio / Unidad": suc_rf,
                                    "Total Gastos": total_gastos_rf, "detalle_gastos": filas_gasto, "fotos": rutas_rf,
                                    "Usuario": st.session_state.usuario_autenticado,
                                    "Timestamp": ahora_mx().strftime("%Y-%m-%d %H:%M:%S"),
                                })
                                _guardar_lista_json(FISICOS_FILE, reportes_actuales)
                                guardar_evento_auditoria("REPORTE FISICO DE GASTOS", f"{suc_rf} · {fecha_rf_str} · ${total_gastos_rf:,.2f} · {len(filas_gasto)} conceptos · {rid}" + (" · REEMPLAZO" if existente_gas else ""))
                                _limpiar_formulario_rf()
                                st.session_state.rf_aviso_guardado = f"✅ Reporte de gastos guardado ({rid}). Ya cuenta en balances y cierres."
                                st.rerun()
                            except Exception as exc:
                                st.error(f"No se pudo guardar el reporte: {exc}")


        # ---------------- LISTA ----------------
        if _sub_rf == "lista":
            todos_rf = [r for r in _leer_lista_json(FISICOS_FILE) if es_global or r.get("Sitio / Unidad") == st.session_state.sucursal_asignada]
            if not todos_rf:
                st.info("Aún no hay reportes físicos subidos.")
            else:
                filtro_suc = st.selectbox("Sucursal", ["Todas"] + sorted({r.get("Sitio / Unidad", "") for r in todos_rf}), key="rf_filtro_suc")
                visibles_rf = [r for r in todos_rf if filtro_suc == "Todas" or r.get("Sitio / Unidad") == filtro_suc]
                st.dataframe(pd.DataFrame([{
                    "Fecha": r.get("Fecha"), "Sucursal": r.get("Sitio / Unidad"),
                    "Tipo": "Ingresos" if r.get("tipo_reporte") == "ingresos" else "Gastos",
                    "Evento": r.get("Tipo") if int(r.get("eventos_dia") or 1) >= 2 else "",
                    "Origen": r.get("origen") or "Reporte físico",
                    "Importe ($)": r.get("Ingreso Bruto") if r.get("tipo_reporte") == "ingresos" else r.get("Total Gastos"),
                    "Nómina ($)": r.get("Nómina") if r.get("tipo_reporte") == "ingresos" else None,
                    "Subió": r.get("Usuario"), "Fecha de subida": r.get("Timestamp"), "ID": r.get("ID"),
                } for r in visibles_rf]), use_container_width=True, hide_index=True)
                etiquetas_rf = {f"{r.get('Fecha')} · {r.get('Sitio / Unidad')} · {'Ingresos' if r.get('tipo_reporte') == 'ingresos' else 'Gastos'}"
                                + (f" · {r.get('Tipo')}" if int(r.get("eventos_dia") or 1) >= 2 else "") + f" · {r.get('ID')}": r for r in visibles_rf}
                elegido = etiquetas_rf[st.selectbox("Ver fotos y detalle", list(etiquetas_rf.keys()), key="rf_ver_detalle")]
                if elegido.get("tipo_reporte") == "gastos":
                    st.dataframe(pd.DataFrame(elegido.get("detalle_gastos", [])), use_container_width=True, hide_index=True)
                fotos_rf_zoom = [(f"Foto {i_fx}", _ruta_absoluta(r_fx)) for i_fx, r_fx in enumerate(elegido.get("fotos", []), start=1)]
                visor_zoom_fotos([f for f in fotos_rf_zoom if f[1]], key=f"rf_{elegido.get('ID')}", alto=480)
                if es_contador:
                    with st.expander("🗑️ Borrar este reporte (solo Contador Auditor)"):
                        st.caption("Se borra el reporte, los gastos que generó y sus fotos. Después se puede volver a subir el de esa sucursal y fecha.")
                        ok_borrar_rf = st.checkbox("Confirmo que quiero borrar este reporte", key=f"rf_borrar_ok_{elegido.get('ID')}")
                        if st.button("🗑️ Borrar reporte", disabled=not ok_borrar_rf, key=f"rf_borrar_btn_{elegido.get('ID')}"):
                            borrado = borrar_reporte_fisico(elegido.get("ID"))
                            if borrado:
                                guardar_evento_auditoria("REPORTE FISICO BORRADO", f"{borrado.get('Sitio / Unidad')} · {borrado.get('Fecha')} · {borrado.get('tipo_reporte')} · {borrado.get('ID')}")
                                st.success("Reporte borrado.")
                                st.rerun()

        # ---------------- REPORTE POR SUCURSAL / GLOBAL ----------------
        if _sub_rf == "reporte":
            st.caption("Ventas, gastos y utilidad (ventas menos gastos) por sucursal o de todas juntas, con la venta de cada día de lunes a domingo. Descárgalo por semana, quincena, mes o un periodo a tu medida.")
            hoy_r = ahora_mx().date()
            pr1, pr2 = st.columns(2)
            tipo_periodo = pr1.radio("Periodo", ["Semana", "Quincena", "Mes", "Personalizado"], horizontal=True, key="rp_tipo")
            referencia_r = pr2.date_input("Cualquier día del periodo", value=hoy_r, key="rp_ref")
            desde_pers = hasta_pers = None
            if tipo_periodo == "Personalizado":
                pp1, pp2 = st.columns(2)
                desde_pers = pp1.date_input("Desde", value=hoy_r - timedelta(days=6), key="rp_desde_pers")
                hasta_pers = pp2.date_input("Hasta", value=hoy_r, key="rp_hasta_pers")
            desde_r, hasta_r = rango_periodo(tipo_periodo, referencia_r, desde_pers, hasta_pers)
            if hasta_r < desde_r:
                st.error("La fecha final no puede ser anterior a la inicial.")
            else:
                periodo_txt = f"Periodo: {desde_r:%d/%m/%Y} al {hasta_r:%d/%m/%Y} ({(hasta_r - desde_r).days + 1} días)"
                st.info(periodo_txt)
                if es_global:
                    alcance = st.radio("Alcance", ["🌐 Global (todas las sucursales)", "🏢 Una sucursal"], horizontal=True, key="rp_alcance")
                else:
                    alcance = "🏢 Una sucursal"
                if alcance.startswith("🌐"):
                    sucursales_r = SUCURSALES
                    titulo_r = "CLUB PARK · REPORTE GLOBAL DE TODAS LAS SUCURSALES"
                else:
                    suc_unica = st.selectbox("Sucursal", sucursales_rf, key="rp_sucursal")
                    sucursales_r = [suc_unica]
                    titulo_r = f"CLUB PARK · REPORTE DE {suc_unica}"
                rep_r = construir_reporte(desde_r, hasta_r, sucursales_r)
                u_tot = rep_r["ventas_total"] - rep_r["gastos_total"]

                comparar = st.checkbox("📊 Comparar contra el periodo anterior", value=True, key="rp_comparar")
                comp = None
                desde_p = hasta_p = None
                hasta_c, hasta_pc = hasta_r, hasta_r
                rep_c, rep_p = rep_r, None
                if comparar:
                    desde_p, hasta_p = rango_anterior(tipo_periodo, desde_r, hasta_r)
                    hasta_c, hasta_pc = hasta_r, hasta_p
                    if hasta_r > hoy_r >= desde_r:
                        st.warning("El periodo actual aún no termina. Comparar contra el periodo anterior completo haría ver una caída por los días pendientes.")
                        if st.checkbox("Comparar solo los días transcurridos (mismos días en ambos periodos)", value=True, key="rp_mismos_dias"):
                            hasta_c = hoy_r
                            hasta_pc = desde_p + timedelta(days=(hoy_r - desde_r).days)
                    rep_c = rep_r if hasta_c == hasta_r else construir_reporte(desde_r, hasta_c, sucursales_r)
                    rep_p = construir_reporte(desde_p, hasta_pc, sucursales_r)
                    comp = construir_comparativo(rep_c, rep_p)
                    st.info(f"Comparando {desde_r:%d/%m/%Y} al {hasta_c:%d/%m/%Y} contra {desde_p:%d/%m/%Y} al {hasta_pc:%d/%m/%Y}")

                def _delta_txt(actual, anterior):
                    pct = _variacion_pct(actual, anterior)
                    return None if pct is None or rep_p is None else f"{pct:+.1f}% vs periodo anterior"

                k1, k2, k3, k4 = st.columns(4)
                if rep_p is not None:
                    v_c, g_c = rep_c["ventas_total"], rep_c["gastos_total"]
                    v_p, g_p = rep_p["ventas_total"], rep_p["gastos_total"]
                    k1.metric("Vendió", f"${rep_r['ventas_total']:,.2f}", delta=_delta_txt(v_c, v_p))
                    k2.metric("Gastos", f"${rep_r['gastos_total']:,.2f}", delta=_delta_txt(g_c, g_p), delta_color="inverse")
                    k3.metric("Utilidad", f"${u_tot:,.2f}", delta=_delta_txt(v_c - g_c, v_p - g_p))
                else:
                    k1.metric("Vendió", f"${rep_r['ventas_total']:,.2f}")
                    k2.metric("Gastos", f"${rep_r['gastos_total']:,.2f}")
                    k3.metric("Utilidad", f"${u_tot:,.2f}")
                k4.metric("Margen", f"{(u_tot / rep_r['ventas_total'] * 100) if rep_r['ventas_total'] else 0:,.1f}%")
                st.caption("Utilidad = ventas − gastos. Ventas: reporte físico de ingresos de la sucursal y fecha; si no hay, los cobros digitales. Gastos: todos los capturados en la app y los reportes de gastos.")

                st.markdown("#### 🏢 Resultado por sucursal")
                st.dataframe(rep_r["resumen_total"], use_container_width=True, hide_index=True, column_config=formato_columnas_df(rep_r["resumen_total"]))
                if len(sucursales_r) > 1 and rep_r["ventas_total"] > 0:
                    st.bar_chart(rep_r["resumen"].set_index("Sucursal")[["Ventas", "Gastos", "Utilidad"]])

                if comp is not None:
                    st.markdown("#### 📊 Comparativo contra el periodo anterior")
                    df_comp = comp["resumen"]
                    st.dataframe(df_comp, use_container_width=True, hide_index=True, column_config=formato_columnas_df(df_comp))
                    datos_comp = df_comp[df_comp["Sucursal"] != "TOTAL"]
                    datos_comp = datos_comp[(datos_comp["Ventas actual"] > 0) | (datos_comp["Ventas anterior"] > 0)]
                    if len(datos_comp) > 1:
                        mejor = datos_comp.loc[datos_comp["Var. ventas $"].idxmax()]
                        peor = datos_comp.loc[datos_comp["Var. ventas $"].idxmin()]
                        st.success(f"Mayor alza en ventas: {mejor['Sucursal']} ({mejor['Var. ventas $']:+,.2f} pesos)")
                        st.error(f"Mayor caída en ventas: {peor['Sucursal']} ({peor['Var. ventas $']:+,.2f} pesos)")
                    if comp["diario"][["Ventas actual", "Ventas anterior"]].to_numpy().sum() > 0:
                        st.caption("Venta día a día: periodo actual contra periodo anterior")
                        st.line_chart(comp["diario"].set_index("N° de día")[["Ventas actual", "Ventas anterior"]])
                    st.dataframe(comp["diario"], use_container_width=True, hide_index=True, column_config=formato_columnas_df(comp["diario"]))

                st.markdown("#### 📅 Venta por día (lunes a domingo)")
                df_diario = rep_r["diario_sucursal"]
                st.dataframe(df_diario, use_container_width=True, hide_index=True, column_config=formato_columnas_df(df_diario))

                st.markdown("#### 💵 Ventas, gastos y utilidad por día")
                df_util = rep_r["utilidad_diaria"]
                st.dataframe(df_util, use_container_width=True, hide_index=True, column_config=formato_columnas_df(df_util))
                if rep_r["ventas_total"] > 0:
                    st.bar_chart(df_util.set_index("Fecha")[["Ventas", "Gastos"]])

                st.markdown("#### 📆 Venta por día de la semana")
                df_sem = rep_r["por_dia_semana"]
                st.dataframe(df_sem, use_container_width=True, hide_index=True, column_config=formato_columnas_df(df_sem))

                st.markdown("#### 🧾 Gastos por categoría")
                if rep_r["gastos_categoria"].empty:
                    st.caption("No hay gastos registrados en este periodo.")
                else:
                    st.dataframe(rep_r["gastos_categoria"], use_container_width=True, hide_index=True, column_config=formato_columnas_df(rep_r["gastos_categoria"]))

                st.markdown("#### ⬇️ Descargas")
                nombre_base = "GLOBAL" if len(sucursales_r) > 1 else re.sub(r"[^A-Za-z0-9]+", "_", sucursales_r[0])
                sufijo_archivo = f"{nombre_base}_{desde_r:%Y%m%d}_{hasta_r:%Y%m%d}"
                if OPENPYXL_OK:
                    periodo_xl = periodo_txt + (f" · Comparado con {desde_p:%d/%m/%Y} al {hasta_p:%d/%m/%Y}" if comp is not None else "")
                    st.download_button("📊 Descargar Excel del reporte", data=generar_excel_reporte(rep_r, titulo_r, periodo_xl, comp),
                                       file_name=f"REPORTE_{sufijo_archivo}.xlsx",
                                       mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", key="rp_dl_excel")
                else:
                    st.warning("Falta instalar openpyxl para descargar el Excel. Agrégalo a requirements.txt.")
                st.download_button("📄 Descargar resumen en CSV", data=rep_r["resumen_total"].to_csv(index=False).encode("utf-8-sig"),
                                   file_name=f"RESUMEN_{sufijo_archivo}.csv", mime="text/csv", key="rp_dl_csv")
                if OPENPYXL_OK and len(sucursales_r) > 1:
                    if st.button("🗜️ Preparar ZIP con un Excel por sucursal", key="rp_zip_btn"):
                        memoria_zip = io.BytesIO()
                        with zipfile.ZipFile(memoria_zip, "w", zipfile.ZIP_DEFLATED) as zf:
                            for suc_zip in sucursales_r:
                                rep_suc = construir_reporte(desde_r, hasta_r, [suc_zip])
                                comp_suc = None
                                if comp is not None:
                                    comp_suc = construir_comparativo(
                                        rep_suc if hasta_c == hasta_r else construir_reporte(desde_r, hasta_c, [suc_zip]),
                                        construir_reporte(desde_p, hasta_pc, [suc_zip]))
                                zf.writestr(f"{re.sub(r'[^A-Za-z0-9]+', '_', suc_zip)}_{desde_r:%Y%m%d}_{hasta_r:%Y%m%d}.xlsx",
                                            generar_excel_reporte(rep_suc, f"CLUB PARK · REPORTE DE {suc_zip}", periodo_xl, comp_suc))
                        st.session_state["rp_zip_bytes"] = memoria_zip.getvalue()
                        st.session_state["rp_zip_nombre"] = f"REPORTES_POR_SUCURSAL_{desde_r:%Y%m%d}_{hasta_r:%Y%m%d}.zip"
                    if st.session_state.get("rp_zip_bytes"):
                        st.download_button("⬇️ Descargar ZIP por sucursal", data=st.session_state["rp_zip_bytes"],
                                           file_name=st.session_state.get("rp_zip_nombre", "REPORTES.zip"), mime="application/zip", key="rp_dl_zip")

        # ---------------- REPORTE POR SUPERVISOR ----------------
        if _sub_rf == "superv":
            clave_propia = supervisor_del_usuario()
            if not (es_contador or clave_propia):
                st.info("Este reporte es para el Contador Auditor y los supervisores (José de Jesús y Víctor).")
            else:
                st.caption("Ingreso bruto total contra gastos totales y utilidad de las sucursales a cargo de cada supervisor, con el detalle por sucursal.")
                hoy_s = ahora_mx().date()
                ps1, ps2 = st.columns(2)
                tipo_s = ps1.radio("Periodo", ["Semana", "Quincena", "Mes", "Personalizado"], horizontal=True, key="sup_tipo")
                ref_s = ps2.date_input("Cualquier día del periodo", value=hoy_s, key="sup_ref")
                dp_s = hp_s = None
                if tipo_s == "Personalizado":
                    pp_s1, pp_s2 = st.columns(2)
                    dp_s = pp_s1.date_input("Desde", value=hoy_s - timedelta(days=6), key="sup_desde")
                    hp_s = pp_s2.date_input("Hasta", value=hoy_s, key="sup_hasta")
                desde_s, hasta_s = rango_periodo(tipo_s, ref_s, dp_s, hp_s)
                if hasta_s < desde_s:
                    st.error("La fecha final no puede ser anterior a la inicial.")
                else:
                    periodo_s = f"Periodo: {desde_s:%d/%m/%Y} al {hasta_s:%d/%m/%Y} ({(hasta_s - desde_s).days + 1} días)"
                    st.info(periodo_s)
                    claves_s = list(SUPERVISORES_NOMBRE) if es_contador else [clave_propia]
                    reportes_s = {}
                    filas_s = []
                    for clave_s in claves_s:
                        sucs_s = sucursales_de_supervisor(clave_s)
                        rep_s = construir_reporte(desde_s, hasta_s, sucs_s) if sucs_s else None
                        reportes_s[clave_s] = rep_s
                        v_s = rep_s["ventas_total"] if rep_s else 0.0
                        g_s = rep_s["gastos_total"] if rep_s else 0.0
                        filas_s.append({"Supervisor": SUPERVISORES_NOMBRE[clave_s], "Sucursales": len(sucs_s),
                                        "Ingreso bruto": v_s, "Gastos": g_s, "Utilidad": v_s - g_s,
                                        "Margen %": round((v_s - g_s) / v_s * 100, 1) if v_s else 0.0})
                    df_sup = pd.DataFrame(filas_s)
                    if len(filas_s) > 1:
                        v_t, g_t = float(df_sup["Ingreso bruto"].sum()), float(df_sup["Gastos"].sum())
                        df_sup = pd.concat([df_sup, pd.DataFrame([{
                            "Supervisor": "TOTAL", "Sucursales": int(df_sup["Sucursales"].sum()), "Ingreso bruto": v_t,
                            "Gastos": g_t, "Utilidad": v_t - g_t, "Margen %": round((v_t - g_t) / v_t * 100, 1) if v_t else 0.0,
                        }])], ignore_index=True)
                    st.markdown("#### 👔 Resultado por supervisor")
                    cfg_cols_sup = formato_columnas_df(df_sup)
                    cfg_cols_sup.pop("Sucursales", None)
                    st.dataframe(df_sup, use_container_width=True, hide_index=True, column_config=cfg_cols_sup)
                    if len(filas_s) > 1 and df_sup["Ingreso bruto"].sum() > 0:
                        st.bar_chart(pd.DataFrame(filas_s).set_index("Supervisor")[["Ingreso bruto", "Gastos", "Utilidad"]])

                    for clave_s in claves_s:
                        rep_s = reportes_s[clave_s]
                        st.markdown(f"#### 🏢 Sucursales de {SUPERVISORES_NOMBRE[clave_s]}")
                        if not rep_s:
                            st.caption("Este supervisor no tiene sucursales asignadas. Asígnalas en la pestaña 💲 Tarifas.")
                            continue
                        m_s1, m_s2, m_s3 = st.columns(3)
                        m_s1.metric("Ingreso bruto", f"${rep_s['ventas_total']:,.2f}")
                        m_s2.metric("Gastos", f"${rep_s['gastos_total']:,.2f}")
                        m_s3.metric("Utilidad", f"${rep_s['ventas_total'] - rep_s['gastos_total']:,.2f}")
                        st.dataframe(rep_s["resumen_total"], use_container_width=True, hide_index=True,
                                     column_config=formato_columnas_df(rep_s["resumen_total"]))

                    st.markdown("#### ⬇️ Descargas")
                    if OPENPYXL_OK:
                        salida_sup = io.BytesIO()
                        verde_s = PatternFill(start_color="0B3B24", end_color="0B3B24", fill_type="solid")
                        hojas_s = [("Por supervisor", df_sup)]
                        for clave_s in claves_s:
                            if reportes_s[clave_s]:
                                hojas_s.append((f"Sucursales {SUPERVISORES_NOMBRE[clave_s]}"[:31], reportes_s[clave_s]["resumen_total"]))
                                hojas_s.append((f"Diario {SUPERVISORES_NOMBRE[clave_s]}"[:31], reportes_s[clave_s]["utilidad_diaria"]))
                        with pd.ExcelWriter(salida_sup, engine="openpyxl") as writer_s:
                            for nombre_h, df_h in hojas_s:
                                df_h.to_excel(writer_s, sheet_name=nombre_h, index=False, startrow=3)
                                ws_s = writer_s.sheets[nombre_h]
                                ws_s["A1"] = "CLUB PARK · REPORTE POR SUPERVISOR"
                                ws_s["A1"].font = Font(size=14, bold=True, color="0B3B24")
                                ws_s["A2"] = periodo_s
                                ws_s["A2"].font = Font(size=10, italic=True)
                                for celda in ws_s[4]:
                                    celda.fill = verde_s
                                    celda.font = Font(bold=True, color="FFFFFF")
                                    celda.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                                for i_c, col_c in enumerate(df_h.columns, start=1):
                                    letra = get_column_letter(i_c)
                                    ws_s.column_dimensions[letra].width = min(max(len(str(col_c)) + 4, 14), 34)
                                    for celda in list(ws_s[letra])[4:]:
                                        if isinstance(celda.value, (int, float)) and str(col_c).endswith("%"):
                                            celda.number_format = '0.0"%"'
                                        elif isinstance(celda.value, (int, float)) and col_c not in COLUMNAS_NO_MONEDA and col_c != "Sucursales":
                                            celda.number_format = '"$"#,##0.00'
                                for fila_x in ws_s.iter_rows(min_row=5):
                                    if str(fila_x[0].value) == "TOTAL":
                                        for celda in fila_x:
                                            celda.font = Font(bold=True)
                                ws_s.freeze_panes = "A5"
                        nombre_sup = "SUPERVISORES" if len(claves_s) > 1 else re.sub(r"[^A-Za-z0-9]+", "_", _sin_acentos(SUPERVISORES_NOMBRE[claves_s[0]]))
                        st.download_button("📊 Descargar Excel por supervisor", data=salida_sup.getvalue(),
                                           file_name=f"REPORTE_{nombre_sup}_{desde_s:%Y%m%d}_{hasta_s:%Y%m%d}.xlsx",
                                           mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", key="sup_dl_excel")
                    st.download_button("📄 Descargar resumen en CSV", data=df_sup.to_csv(index=False).encode("utf-8-sig"),
                                       file_name=f"SUPERVISORES_{desde_s:%Y%m%d}_{hasta_s:%Y%m%d}.csv", mime="text/csv", key="sup_dl_csv")

        # ---------------- EXCEL ----------------
        if _sub_rf == "excel":
            st.caption("Une los reportes físicos de ingresos, los cobros digitales y todos los gastos en el Excel oficial (DETALLE, CONCENTRADO X DÍA y Rentabilidad). Si hay reporte físico de ingresos para una sucursal y fecha, ese manda; si no, se usan los cobros digitales.")
            hoy_x = ahora_mx().date()
            lunes_x = hoy_x - timedelta(days=hoy_x.weekday())
            x1, x2 = st.columns(2)
            desde_x = x1.date_input("Desde", value=lunes_x, key="rf_x_desde")
            hasta_x = x2.date_input("Hasta", value=lunes_x + timedelta(days=6), key="rf_x_hasta")
            registros_x = registros_flujo_periodo(desde_x, hasta_x, sucursales_rf)
            bruto_x = sum(r["Ingreso Bruto"] for r in registros_x)
            gastos_x = sum(r["Total Gastos"] for r in registros_x)
            m1, m2, m3 = st.columns(3)
            m1.metric("Ingresos", f"${bruto_x:,.2f}")
            m2.metric("Gastos", f"${gastos_x:,.2f}")
            m3.metric("Neto", f"${bruto_x - gastos_x:,.2f}")
            etiqueta_x = f"{desde_x:%d/%m/%Y} al {hasta_x:%d/%m/%Y}"
            if not registros_x:
                st.info("No hay información en el periodo seleccionado.")
            else:
                st.dataframe(pd.DataFrame(registros_x).reindex(columns=["Fecha", "Sitio / Unidad", "Folio Inicial", "Folio Final", "Boletos Cobrados",
                                                                          "Ingreso Bruto", "Total Gastos", "Ingreso Neto"]), use_container_width=True, hide_index=True)
                if OPENPYXL_OK:
                    if st.button("📊 Generar Excel oficial", type="primary", key="rf_x_generar"):
                        try:
                            st.session_state["rf_excel_bytes"] = generar_excel_formato_oficial(registros_x, etiqueta_x)
                            st.session_state["rf_excel_nombre"] = f"CONTROL_FLUJO_{desde_x:%Y%m%d}_{hasta_x:%Y%m%d}.xlsx"
                            guardar_evento_auditoria("EXCEL OFICIAL DE FLUJO", f"{etiqueta_x} · {len(registros_x)} registros")
                        except Exception as exc:
                            st.error(f"No se pudo generar el Excel: {exc}")
                    if st.session_state.get("rf_excel_bytes"):
                        st.download_button("⬇️ Descargar Excel oficial", data=st.session_state["rf_excel_bytes"],
                                           file_name=st.session_state.get("rf_excel_nombre", "CONTROL_FLUJO.xlsx"),
                                           mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", key="rf_x_descarga")
                else:
                    st.warning("Falta instalar openpyxl para generar el Excel oficial. Agrégalo a requirements.txt.")
                st.download_button("⬇️ Descargar CSV", data=generar_csv_oficial(registros_x, etiqueta_x),
                                   file_name=f"CONTROL_FLUJO_{desde_x:%Y%m%d}_{hasta_x:%Y%m%d}.csv", mime="text/csv", key="rf_x_csv")

if _SECCION == "reportes" and "reportes" in _TABS:
    _seccion_reportes()


# ---------------------------------------------------------------------
# TAB: TARIFAS (solo Contador Auditor)
# ---------------------------------------------------------------------
@_fragmento
def _seccion_tarifas():
    global exc
    globals()["_CONTADOR_LLAVES"] = 1000  # llaves estables entre recargas completas y de sección
    with tab_tarifas:
        encabezado_institucional("TARIFAS POR SUCURSAL", "SOLO CONTADOR", "Tarifa fija, cobro por tiempo y promociones")
        if not es_contador:
            st.info("Solo el Contador Auditor puede modificar las tarifas.")
        else:
            tarifas_cfg = cargar_tarifas()
            suc_t = st.selectbox("Sucursal", SUCURSALES, key="tarifas_sucursal")
            cfg_t = tarifas_cfg[suc_t]
            tipo_t = st.radio(
                "Tipo de cobro", ["fija", "tiempo"], index=0 if cfg_t["tipo"] == "fija" else 1, horizontal=True,
                format_func=lambda t: "Tarifa fija (ej. $50)" if t == "fija" else "Por tiempo (1ª hora + fracciones)",
                key=f"tarifas_tipo_{suc_t}",
            )
            nueva = dict(cfg_t, tipo=tipo_t)
            if tipo_t == "fija":
                nueva["precio_fijo"] = st.number_input("Precio fijo ($)", min_value=0.0, value=float(cfg_t["precio_fijo"]), step=5.0, key=f"tarifas_fijo_{suc_t}")
            else:
                t1, t2, t3 = st.columns(3)
                nueva["primera_hora"] = t1.number_input("Primera hora ($)", min_value=0.0, value=float(cfg_t["primera_hora"]), step=1.0, key=f"tarifas_1h_{suc_t}")
                nueva["precio_bloque"] = t2.number_input("Precio por fracción ($)", min_value=0.0, value=float(cfg_t["precio_bloque"]), step=1.0, key=f"tarifas_bloque_{suc_t}")
                nueva["bloque_min"] = t3.number_input("Minutos por fracción", min_value=1, value=int(cfg_t["bloque_min"]), step=5, key=f"tarifas_min_{suc_t}")
            nueva["promo_activa"] = st.checkbox("Ofrecer promoción en esta sucursal", value=bool(cfg_t["promo_activa"]), key=f"tarifas_promo_on_{suc_t}")
            if nueva["promo_activa"]:
                if tipo_t == "fija":
                    st.warning("La promoción solo se ofrece en sucursales con cobro por tiempo. Cambia el tipo de cobro a 'Por tiempo'.")
                p1, p2 = st.columns(2)
                nueva["promo_precio"] = p1.number_input("Precio de la promoción ($)", min_value=0.0, value=float(cfg_t["promo_precio"]), step=5.0, key=f"tarifas_promo_precio_{suc_t}")
                nueva["promo_horas"] = p2.number_input("Horas incluidas", min_value=1, value=int(cfg_t["promo_horas"]), step=1, key=f"tarifas_promo_horas_{suc_t}")
                st.markdown("**Al terminar la promoción se cobra:**")
                p3, p4 = st.columns(2)
                nueva["promo_extra_precio"] = p3.number_input(
                    "Costo después de la promoción ($)", min_value=0.0,
                    value=float(cfg_t.get("promo_extra_precio", cfg_t["precio_bloque"])), step=1.0,
                    key=f"tarifas_promo_extra_precio_{suc_t}")
                nueva["promo_extra_min"] = p4.number_input(
                    "Cada cuántos minutos", min_value=1,
                    value=int(cfg_t.get("promo_extra_min", cfg_t["bloque_min"])), step=5,
                    key=f"tarifas_promo_extra_min_{suc_t}")
                ejemplo_min = int(nueva["promo_horas"]) * 60 + 30
                st.caption(f"Ejemplo: un auto que se queda {texto_duracion(ejemplo_min)} con promoción paga "
                           f"${calcular_cobro(nueva, 'promocion', ejemplo_min):,.2f}.")
            st.caption("Vista previa: " + " · ".join(etiqueta_modalidad(nueva, m) for m in modalidades_disponibles(nueva)))
            st.caption("Los cambios aplican a ingresos nuevos. Los vehículos ya registrados conservan su modalidad; al dar salida el Contador puede cambiarla o ajustar el importe.")
            if st.button("💾 Guardar tarifa", type="primary", key=f"tarifas_guardar_{suc_t}"):
                tarifas_cfg[suc_t] = nueva
                try:
                    guardar_tarifas(tarifas_cfg)
                    guardar_evento_auditoria("CAMBIO DE TARIFA", f"{suc_t}: {cfg_t} -> {nueva}")
                    st.success(f"Tarifa de {suc_t} actualizada.")
                except OSError as exc:
                    st.error(f"No se pudo guardar la tarifa: {exc}")

            st.divider()
            st.markdown("### 🏢 Agregar o quitar sucursales")
            cfg_suc = cargar_sucursales_config()
            a1, a2, a3 = st.columns([3, 2, 2])
            nueva_suc = a1.text_input("Nombre de la nueva sucursal", key="suc_nueva_nombre").strip().upper()
            nueva_tarifa = a2.number_input("Tarifa por boleto ($)", min_value=0.0, value=50.0, step=5.0, key="suc_nueva_tarifa")
            a3.write("")
            if a3.button("➕ Agregar sucursal", type="primary", key="suc_nueva_btn"):
                if not nueva_suc:
                    st.error("Escribe el nombre de la sucursal.")
                elif nueva_suc in PRECIOS_SUCURSALES:
                    st.error("Esa sucursal ya existe.")
                else:
                    cfg_suc["agregadas"][nueva_suc] = float(nueva_tarifa)
                    if nueva_suc in cfg_suc["quitadas"]:
                        cfg_suc["quitadas"].remove(nueva_suc)
                    try:
                        guardar_sucursales_config(cfg_suc)
                        guardar_evento_auditoria("SUCURSAL AGREGADA", f"{nueva_suc} · tarifa ${nueva_tarifa:,.2f}")
                        st.success(f"{nueva_suc} agregada. Reinicia la vista para verla en todos los menús.")
                        st.rerun()
                    except OSError as exc:
                        st.error(f"No se pudo guardar: {exc}")
            q1, q2 = st.columns([5, 2])
            quitar = q1.selectbox("Sucursal a quitar", ["—"] + SUCURSALES, key="suc_quitar_sel")
            confirma_quitar = q1.checkbox("Confirmo que quiero quitarla (el historial ya registrado se conserva)", key="suc_quitar_ok")
            q2.write("")
            if q2.button("🗑️ Quitar sucursal", key="suc_quitar_btn", disabled=(quitar == "—" or not confirma_quitar)):
                cfg_suc["agregadas"].pop(quitar, None)
                if quitar not in cfg_suc["quitadas"]:
                    cfg_suc["quitadas"].append(quitar)
                try:
                    guardar_sucursales_config(cfg_suc)
                    guardar_evento_auditoria("SUCURSAL QUITADA", quitar)
                    st.success(f"{quitar} quitada.")
                    st.rerun()
                except OSError as exc:
                    st.error(f"No se pudo guardar: {exc}")
            st.divider()
            st.markdown("### 👔 Supervisor a cargo de cada sucursal")
            st.caption("Define qué sucursales entran en el reporte de cada supervisor.")
            s1, s2, s3 = st.columns([3, 2, 2])
            suc_sup = s1.selectbox("Sucursal", SUCURSALES, key="sup_asig_suc")
            actual_sup = next((i["supervisor"] for i in CATALOGO_OFICIAL_23 if i["sitio"] == suc_sup), "-")
            opciones_sup = list(SUPERVISORES_NOMBRE)
            nuevo_sup = s2.selectbox(
                "Supervisor", opciones_sup,
                index=opciones_sup.index(actual_sup) if actual_sup in opciones_sup else 0,
                format_func=lambda c: SUPERVISORES_NOMBRE[c], key=f"sup_asig_sel_{suc_sup}")
            s3.write("")
            if s3.button("💾 Asignar", key="sup_asig_btn"):
                cfg_suc.setdefault("supervisores", {})[suc_sup] = nuevo_sup
                try:
                    guardar_sucursales_config(cfg_suc)
                    guardar_evento_auditoria("SUPERVISOR ASIGNADO", f"{suc_sup}: {actual_sup} -> {nuevo_sup}")
                    st.success(f"{suc_sup} queda a cargo de {SUPERVISORES_NOMBRE[nuevo_sup]}.")
                    st.rerun()
                except OSError as exc:
                    st.error(f"No se pudo guardar: {exc}")
            st.caption("Las sucursales nuevas aparecen en ingresos, reportes, Excel oficial y tarifas. Para que un encargado entre solo a su sucursal agrega su token en CLUBPARK_TOKENS_JSON; los supervisores globales la ven de inmediato.")

if _SECCION == "tarifas" and "tarifas" in _TABS:
    _seccion_tarifas()


# ---------------------------------------------------------------------
# TAB: EVENTOS EXTRAORDINARIOS (Contador captura; supervisores consultan)
# ---------------------------------------------------------------------
@_fragmento
def _seccion_eventos():
    global celda, e, exc, fila, ing, letra
    globals()["_CONTADOR_LLAVES"] = 1000  # llaves estables entre recargas completas y de sección
    with _TABS["eventos"]:
        encabezado_institucional("EVENTOS EXTRAORDINARIOS EN SALONES", "AGENDA · AUDITORÍA · UTILIDAD",
                                 "Programación semanal, supervisor a cargo, costo del evento, boletaje y utilidad por salón")
        clave_sup_ev = supervisor_del_usuario()
        if not (es_contador or clave_sup_ev):
            st.info("Esta sección es para el Contador Auditor y los supervisores.")
        else:
            todos_ev = cargar_eventos()
            # El supervisor solo ve los eventos que tiene a cargo
            mis_ev = todos_ev if es_contador else [e for e in todos_ev if e.get("supervisor") == clave_sup_ev]
            nombre_sup = lambda c: SUPERVISORES_NOMBRE.get(c, c or "—")
            sub_agenda, sub_alta, sub_aud_ev, sub_suc = st.tabs(["📅 Semana", "➕ Agregar evento", "📊 Auditoría y utilidad", "🏢 Sucursales y salones"])

            # ---------------- SEMANA ----------------
            with sub_agenda:
                ref_ev = st.date_input("Cualquier día de la semana", value=ahora_mx().date(), key="ev_semana_ref")
                lunes_ev = ref_ev - timedelta(days=ref_ev.weekday())
                domingo_ev = lunes_ev + timedelta(days=6)
                st.info(f"Semana del {lunes_ev:%d/%m/%Y} al {domingo_ev:%d/%m/%Y}")
                salon_f = st.selectbox("Salón", ["Todos"] + SALONES_EVENTOS, key="ev_semana_salon")
                semana_ev = [e for e in mis_ev if (lambda d: d and lunes_ev <= d <= domingo_ev)(fecha_flexible(e.get("fecha")))
                             and (salon_f == "Todos" or e.get("sucursal") == salon_f)]
                activos_sem = [e for e in semana_ev if e.get("estado") != "Cancelado"]
                k_e1, k_e2, k_e3, k_e4 = st.columns(4)
                k_e1.metric("Eventos", sum(_entero_seguro(e.get("num_eventos")) for e in activos_sem))
                k_e2.metric("Vehículos", sum(e["vehiculos"] for e in activos_sem))
                k_e3.metric("Ingreso", f"${sum(e['ingreso_total'] for e in activos_sem):,.2f}")
                k_e4.metric("Utilidad", f"${sum(e['utilidad'] for e in activos_sem):,.2f}")
                if not semana_ev:
                    st.caption("No hay eventos en esta semana.")
                elif es_contador:
                    st.caption("Edita directamente en la tabla y pulsa «Guardar cambios». Para borrar un evento, selecciona su fila y elimínala.")
                    df_ev = pd.DataFrame([{
                        "ID": e.get("id"), "Fecha": fecha_flexible(e.get("fecha")),
                        "Día": NOMBRES_DIA[fecha_flexible(e.get("fecha")).weekday()],
                        "Inicio": datetime.strptime(e.get("hora_inicio") or "00:00", "%H:%M").time(),
                        "Fin": datetime.strptime(e.get("hora_fin") or "00:00", "%H:%M").time(),
                        "Salón": e.get("sucursal"), "Eventos": _entero_seguro(e.get("num_eventos")) or 1,
                        "Descripción": e.get("descripcion", ""), "Supervisor": e.get("supervisor"),
                        "Costo evento": e["costo_evento"], "Vehículos": e["vehiculos"], "Costo boleto": e["costo_boleto"],
                        "Nómina": _flotante_seguro(e.get("nomina")), "Otros gastos": _flotante_seguro(e.get("otros_gastos")),
                        "Ingreso": e["ingreso_total"], "Utilidad": e["utilidad"], "Estado": e.get("estado", "Programado"),
                    } for e in semana_ev])
                    dinero_col = lambda t: st.column_config.NumberColumn(t, min_value=0.0, step=10.0, format="$%.2f")
                    editado_ev = st.data_editor(
                        df_ev, num_rows="dynamic", use_container_width=True, hide_index=True, key=f"ev_editor_{lunes_ev:%Y%m%d}_{salon_f}",
                        column_config={
                            "ID": st.column_config.TextColumn("ID", disabled=True),
                            "Fecha": st.column_config.DateColumn("Fecha", format="DD/MM/YYYY", required=True),
                            "Día": st.column_config.TextColumn("Día", disabled=True),
                            "Inicio": st.column_config.TimeColumn("Inicio", format="HH:mm"),
                            "Fin": st.column_config.TimeColumn("Fin", format="HH:mm"),
                            "Salón": st.column_config.SelectboxColumn("Salón", options=SALONES_EVENTOS, required=True),
                            "Eventos": st.column_config.NumberColumn("Eventos", min_value=1, step=1),
                            "Supervisor": st.column_config.SelectboxColumn("Supervisor", options=list(SUPERVISORES_NOMBRE), required=True),
                            "Costo evento": dinero_col("Costo evento"), "Costo boleto": dinero_col("Costo boleto"),
                            "Vehículos": st.column_config.NumberColumn("Vehículos", min_value=0, step=1),
                            "Nómina": dinero_col("Nómina"), "Otros gastos": dinero_col("Otros gastos"),
                            "Ingreso": st.column_config.NumberColumn("Ingreso", disabled=True, format="$%.2f"),
                            "Utilidad": st.column_config.NumberColumn("Utilidad", disabled=True, format="$%.2f"),
                            "Estado": st.column_config.SelectboxColumn("Estado", options=ESTADOS_EVENTO, required=True),
                        },
                    )
                    if st.button("💾 Guardar cambios de la semana", type="primary", key="ev_guardar_semana"):
                        ids_semana = {e.get("id") for e in semana_ev}
                        anteriores = {e.get("id"): e for e in semana_ev}
                        nuevos_ev, errores_ev = [], []
                        for fila in editado_ev.to_dict("records"):
                            fecha_f = fila.get("Fecha")
                            if fecha_f is None or pd.isna(fecha_f) or not fila.get("Salón"):
                                errores_ev.append("Hay filas sin fecha o sin salón.")
                                continue
                            if hasattr(fecha_f, "date") and not isinstance(fecha_f, date):
                                fecha_f = fecha_f.date()
                            elif isinstance(fecha_f, datetime):
                                fecha_f = fecha_f.date()
                            hora = lambda v: v.strftime("%H:%M") if hasattr(v, "strftime") else ""
                            id_f = fila.get("ID") if isinstance(fila.get("ID"), str) and fila.get("ID") else f"EV_{uuid.uuid4().hex[:8].upper()}"
                            base_f = dict(anteriores.get(id_f, {}))
                            base_f.update({
                                "id": id_f, "fecha": f"{fecha_f:%Y-%m-%d}", "hora_inicio": hora(fila.get("Inicio")),
                                "hora_fin": hora(fila.get("Fin")), "sucursal": fila.get("Salón"),
                                "num_eventos": _entero_seguro(fila.get("Eventos")) or 1,
                                "descripcion": str(fila.get("Descripción") or "").strip(),
                                "supervisor": fila.get("Supervisor") or "", "costo_evento": fila.get("Costo evento"),
                                "vehiculos": fila.get("Vehículos"), "costo_boleto": fila.get("Costo boleto"),
                                "nomina": _flotante_seguro(fila.get("Nómina")), "otros_gastos": _flotante_seguro(fila.get("Otros gastos")),
                                "estado": fila.get("Estado") or "Programado",
                                "modificado_por": st.session_state.usuario_autenticado,
                                "modificado_en": ahora_mx().strftime("%Y-%m-%d %H:%M:%S"),
                            })
                            base_f.setdefault("creado_por", st.session_state.usuario_autenticado)
                            nuevos_ev.append(calcular_evento(base_f))
                        if errores_ev:
                            st.error(errores_ev[0] + " Corrige antes de guardar.")
                        else:
                            borrados_ev = ids_semana - {e["id"] for e in nuevos_ev}
                            resto_ev = [e for e in todos_ev if e.get("id") not in ids_semana]
                            try:
                                guardar_eventos(resto_ev + nuevos_ev)
                                guardar_evento_auditoria("EVENTOS EXTRAORDINARIOS EDITADOS",
                                                         f"Semana {lunes_ev:%d/%m/%Y}: {len(nuevos_ev)} eventos guardados"
                                                         + (f", borrados: {', '.join(sorted(borrados_ev))}" if borrados_ev else ""))
                                st.success("Cambios guardados.")
                                st.rerun()
                            except OSError as exc:
                                st.error(f"No se pudieron guardar los eventos: {exc}")
                else:
                    st.dataframe(pd.DataFrame([{
                        "Fecha": fecha_flexible(e.get("fecha")).strftime("%d/%m/%Y"), "Día": NOMBRES_DIA[fecha_flexible(e.get("fecha")).weekday()],
                        "Horario": f"{e.get('hora_inicio', '')} a {e.get('hora_fin', '')}", "Salón": e.get("sucursal"),
                        "Eventos": e.get("num_eventos"), "Descripción": e.get("descripcion", ""), "Vehículos": e["vehiculos"],
                        "Ingreso": e["ingreso_total"], "Utilidad": e["utilidad"], "Estado": e.get("estado"),
                    } for e in semana_ev]), use_container_width=True, hide_index=True,
                        column_config={c: st.column_config.NumberColumn(c, format="$%.2f") for c in ("Ingreso", "Utilidad")})
                    st.caption("Solo el Contador Auditor puede agregar, cambiar o borrar eventos.")

            # ---------------- AGREGAR ----------------
            with sub_alta:
                if not es_contador:
                    st.info("Solo el Contador Auditor (máxima jerarquía) puede agregar eventos extraordinarios.")
                elif not SALONES_EVENTOS:
                    st.warning("No hay salones de eventos dados de alta en el catálogo de sucursales.")
                else:
                    a_salon = st.selectbox("Salón", SALONES_EVENTOS, key="ev_alta_salon")
                    a1e, a2e, a3e = st.columns(3)
                    a_fecha = a1e.date_input("Fecha", value=ahora_mx().date(), key="ev_alta_fecha")
                    a_ini = a2e.time_input("Hora de inicio", value=datetime.strptime("14:00", "%H:%M").time(), key="ev_alta_ini")
                    a_fin = a3e.time_input("Hora de término", value=datetime.strptime("20:00", "%H:%M").time(), key="ev_alta_fin")
                    b1e, b2e = st.columns(2)
                    a_num = b1e.number_input("Número de eventos ese día en el salón", min_value=1, value=1, step=1, key="ev_alta_num")
                    sup_catalogo = next((i["supervisor"] for i in CATALOGO_OFICIAL_23 if i["sitio"] == a_salon), "")
                    claves_sup_ev = list(SUPERVISORES_NOMBRE)
                    a_sup = b2e.selectbox("Supervisor a cargo", claves_sup_ev, format_func=nombre_sup,
                                          index=claves_sup_ev.index(sup_catalogo) if sup_catalogo in claves_sup_ev else 0,
                                          key=f"ev_alta_sup_{a_salon}")
                    a_desc = st.text_input("Descripción (cliente, tipo de evento)", key="ev_alta_desc")
                    st.markdown("**Costos y boletaje**")
                    c1e, c2e, c3e = st.columns(3)
                    a_costo = c1e.number_input("Costo del evento ($)", min_value=0.0, value=0.0, step=100.0, key="ev_alta_costo")
                    a_veh = c2e.number_input("Número de vehículos", min_value=0, value=0, step=1, key="ev_alta_veh")
                    a_bol = c3e.number_input("Costo por boleto ($)", min_value=0.0, value=float(tarifa_de(a_salon).get("precio_fijo", 0)),
                                             step=5.0, key=f"ev_alta_bol_{a_salon}")
                    d1e, d2e = st.columns(2)
                    a_nom = d1e.number_input("Nómina del evento ($)", min_value=0.0, value=0.0, step=50.0, key="ev_alta_nom")
                    a_otros = d2e.number_input("Otros gastos ($)", min_value=0.0, value=0.0, step=50.0, key="ev_alta_otros")
                    a_estado = st.radio("Estado", ESTADOS_EVENTO, horizontal=True, key="ev_alta_estado")
                    vista_ev = calcular_evento({"costo_evento": a_costo, "vehiculos": a_veh, "costo_boleto": a_bol,
                                                "nomina": a_nom, "otros_gastos": a_otros})
                    v1e, v2e, v3e = st.columns(3)
                    v1e.metric("Ingreso total", f"${vista_ev['ingreso_total']:,.2f}",
                               help=f"Evento ${vista_ev['costo_evento']:,.2f} + boletaje ${vista_ev['ingreso_boletaje']:,.2f}")
                    v2e.metric("Gastos", f"${vista_ev['gastos_total']:,.2f}")
                    v3e.metric("Utilidad", f"${vista_ev['utilidad']:,.2f}")
                    if st.button("➕ Agregar evento", type="primary", use_container_width=True, key="ev_alta_btn"):
                        if a_fin <= a_ini:
                            st.warning("La hora de término es menor o igual a la de inicio: se guardará como evento que termina al día siguiente.")
                        nuevo_ev = calcular_evento({
                            "id": f"EV_{uuid.uuid4().hex[:8].upper()}", "fecha": f"{a_fecha:%Y-%m-%d}",
                            "hora_inicio": a_ini.strftime("%H:%M"), "hora_fin": a_fin.strftime("%H:%M"),
                            "sucursal": a_salon, "num_eventos": int(a_num), "descripcion": a_desc.strip(),
                            "supervisor": a_sup, "costo_evento": a_costo, "vehiculos": int(a_veh), "costo_boleto": a_bol,
                            "nomina": a_nom, "otros_gastos": a_otros, "estado": a_estado,
                            "creado_por": st.session_state.usuario_autenticado,
                            "creado_en": ahora_mx().strftime("%Y-%m-%d %H:%M:%S"),
                        })
                        try:
                            guardar_eventos(cargar_eventos() + [nuevo_ev])
                            guardar_evento_auditoria("EVENTO EXTRAORDINARIO AGREGADO",
                                                     f"{a_salon} · {a_fecha:%d/%m/%Y} {a_ini:%H:%M}-{a_fin:%H:%M} · {int(a_num)} evento(s) · "
                                                     f"{nombre_sup(a_sup)} · ingreso ${nuevo_ev['ingreso_total']:,.2f} · {nuevo_ev['id']}")
                            st.success(f"Evento agregado ({nuevo_ev['id']}) a la semana del {a_fecha - timedelta(days=a_fecha.weekday()):%d/%m/%Y}.")
                        except OSError as exc:
                            st.error(f"No se pudo guardar el evento: {exc}")

            # ---------------- AUDITORÍA Y UTILIDAD ----------------
            with sub_aud_ev:
                hoy_ev = ahora_mx().date()
                pe1, pe2 = st.columns(2)
                tipo_ev = pe1.radio("Periodo", ["Semana", "Quincena", "Mes", "Personalizado"], horizontal=True, key="ev_aud_tipo")
                ref_aud_ev = pe2.date_input("Cualquier día del periodo", value=hoy_ev, key="ev_aud_ref")
                dpe = hpe = None
                if tipo_ev == "Personalizado":
                    pq1, pq2 = st.columns(2)
                    dpe = pq1.date_input("Desde", value=hoy_ev - timedelta(days=6), key="ev_aud_desde")
                    hpe = pq2.date_input("Hasta", value=hoy_ev, key="ev_aud_hasta")
                desde_ev, hasta_ev = rango_periodo(tipo_ev, ref_aud_ev, dpe, hpe)
                if hasta_ev < desde_ev:
                    st.error("La fecha final no puede ser anterior a la inicial.")
                else:
                    periodo_ev = f"Periodo: {desde_ev:%d/%m/%Y} al {hasta_ev:%d/%m/%Y}"
                    st.info(periodo_ev + " · los eventos cancelados no suman")
                    del_periodo = [e for e in complementar_eventos_con_reportes(mis_ev)
                                   if (lambda d: d and desde_ev <= d <= hasta_ev)(fecha_flexible(e.get("fecha")))
                                   and e.get("estado") != "Cancelado"]
                    con_rep = [e for e in del_periodo if e["fuente"].startswith("Reporte")]
                    sin_rep_pasados = [e for e in del_periodo if not e["fuente"].startswith("Reporte")
                                       and fecha_flexible(e.get("fecha")) and fecha_flexible(e.get("fecha")) < ahora_mx().date()]
                    st.caption(f"📄 {len(con_rep)} de {len(del_periodo)} evento(s) ya tienen reporte físico: para ellos se usan "
                               "los vehículos, boletaje y gastos REALES del reporte. Los demás muestran lo programado.")
                    if sin_rep_pasados:
                        st.warning("⚠️ Eventos ya pasados sin reporte físico: " + ", ".join(
                            f"{e.get('sucursal')} {fecha_flexible(e.get('fecha')):%d/%m}" for e in sin_rep_pasados[:12])
                            + ("…" if len(sin_rep_pasados) > 12 else ""))
                    filas_salon = []
                    for salon in SALONES_EVENTOS:
                        evs = [e for e in del_periodo if e.get("sucursal") == salon]
                        ing = sum(e["ingreso_total"] for e in evs)
                        gas = sum(e["gastos_total"] for e in evs)
                        filas_salon.append({
                            "Salón": salon, "Eventos": sum(_entero_seguro(e.get("num_eventos")) for e in evs),
                            "Vehículos": sum(e["vehiculos"] for e in evs),
                            "Costo de eventos": sum(e["costo_evento"] for e in evs),
                            "Boletaje": sum(e["ingreso_boletaje"] for e in evs), "Ingreso total": ing,
                            "Gastos": gas, "Utilidad": ing - gas, "Margen %": round((ing - gas) / ing * 100, 1) if ing else 0.0,
                        })
                    df_salon = pd.DataFrame(filas_salon)
                    tot_ing, tot_gas = float(df_salon["Ingreso total"].sum()), float(df_salon["Gastos"].sum())
                    df_salon_tot = pd.concat([df_salon, pd.DataFrame([{
                        "Salón": "TOTAL", "Eventos": int(df_salon["Eventos"].sum()), "Vehículos": int(df_salon["Vehículos"].sum()),
                        "Costo de eventos": float(df_salon["Costo de eventos"].sum()), "Boletaje": float(df_salon["Boletaje"].sum()),
                        "Ingreso total": tot_ing, "Gastos": tot_gas, "Utilidad": tot_ing - tot_gas,
                        "Margen %": round((tot_ing - tot_gas) / tot_ing * 100, 1) if tot_ing else 0.0,
                    }])], ignore_index=True)
                    m_e1, m_e2, m_e3, m_e4 = st.columns(4)
                    m_e1.metric("Eventos", int(df_salon["Eventos"].sum()))
                    m_e2.metric("Ingreso total", f"${tot_ing:,.2f}")
                    m_e3.metric("Gastos", f"${tot_gas:,.2f}")
                    m_e4.metric("Utilidad", f"${tot_ing - tot_gas:,.2f}")
                    cfg_money_ev = {c: st.column_config.NumberColumn(c, format="$%.2f")
                                    for c in ("Costo de eventos", "Boletaje", "Ingreso total", "Gastos", "Utilidad")}
                    cfg_money_ev["Margen %"] = st.column_config.NumberColumn("Margen %", format="%.1f%%")
                    st.markdown("#### 🏛️ Utilidad por salón")
                    st.dataframe(df_salon_tot, use_container_width=True, hide_index=True, column_config=cfg_money_ev)
                    if tot_ing > 0:
                        st.bar_chart(df_salon.set_index("Salón")[["Ingreso total", "Gastos", "Utilidad"]])

                    filas_sup_ev = []
                    for clave_x in (SUPERVISORES_NOMBRE if es_contador else [clave_sup_ev]):
                        evs = [e for e in del_periodo if e.get("supervisor") == clave_x]
                        ing = sum(e["ingreso_total"] for e in evs)
                        gas = sum(e["gastos_total"] for e in evs)
                        filas_sup_ev.append({"Supervisor": nombre_sup(clave_x), "Eventos": sum(_entero_seguro(e.get("num_eventos")) for e in evs),
                                             "Vehículos": sum(e["vehiculos"] for e in evs), "Ingreso total": ing, "Gastos": gas,
                                             "Utilidad": ing - gas, "Margen %": round((ing - gas) / ing * 100, 1) if ing else 0.0})
                    df_sup_ev = pd.DataFrame(filas_sup_ev)
                    st.markdown("#### 👔 Por supervisor a cargo")
                    st.dataframe(df_sup_ev, use_container_width=True, hide_index=True, column_config=cfg_money_ev)

                    df_det_ev = pd.DataFrame([{
                        "ID": e.get("id"), "Fecha": fecha_flexible(e.get("fecha")).strftime("%d/%m/%Y"),
                        "Día": NOMBRES_DIA[fecha_flexible(e.get("fecha")).weekday()],
                        "Horario": f"{e.get('hora_inicio', '')} a {e.get('hora_fin', '')}", "Salón": e.get("sucursal"),
                        "Eventos": e.get("num_eventos"), "Descripción": e.get("descripcion", ""), "Supervisor": nombre_sup(e.get("supervisor")),
                        "Costo de eventos": e["costo_evento"], "Vehículos prog.": e["vehiculos_est"], "Vehículos": e["vehiculos"],
                        "Costo boleto": e["costo_boleto"], "Boletaje prog.": e["boletaje_est"],
                        "Boletaje": e["ingreso_boletaje"], "Diferencia boletaje": round(e["ingreso_boletaje"] - e["boletaje_est"], 2),
                        "Ingreso total": e["ingreso_total"], "Gastos": e["gastos_total"],
                        "Utilidad": e["utilidad"], "Fuente": e["fuente"], "Reporte": e["reportes_ids"],
                        "Estado": e.get("estado"), "Capturó": e.get("modificado_por") or e.get("creado_por"),
                    } for e in sorted(del_periodo, key=lambda x: (x.get("fecha", ""), x.get("hora_inicio", "")))])
                    st.markdown("#### 🧾 Detalle de eventos")
                    if df_det_ev.empty:
                        st.caption("No hay eventos en el periodo.")
                    else:
                        cfg_det = dict(cfg_money_ev)
                        cfg_det["Costo boleto"] = st.column_config.NumberColumn("Costo boleto", format="$%.2f")
                        for c_d in ("Boletaje prog.", "Diferencia boletaje"):
                            cfg_det[c_d] = st.column_config.NumberColumn(c_d, format="$%.2f")
                        st.dataframe(df_det_ev, use_container_width=True, hide_index=True, column_config=cfg_det)

                    # ---------- Modificar o borrar un evento (solo Contador) ----------
                    if es_contador:
                        st.markdown("#### ✏️ Modificar o borrar un evento registrado")
                        eventos_edit = sorted([e for e in todos_ev if (lambda d: d and desde_ev <= d <= hasta_ev)(fecha_flexible(e.get("fecha")))],
                                              key=lambda x: (x.get("fecha", ""), x.get("hora_inicio", "")))
                        if not eventos_edit:
                            st.caption("No hay eventos en el periodo para modificar.")
                        else:
                            etiquetas_edit = {
                                f"{fecha_flexible(e.get('fecha')):%d/%m/%Y} · {e.get('hora_inicio', '')} · {e.get('sucursal')} · "
                                f"{e.get('num_eventos', 1)} evento(s) · {e.get('estado', 'Programado')} · {e.get('id')}": e
                                for e in eventos_edit}
                            ev_sel = etiquetas_edit[st.selectbox("Evento", list(etiquetas_edit.keys()), key="ev_mod_sel")]
                            k_m = ev_sel.get("id")
                            md1, md2, md3 = st.columns(3)
                            m_salon = md1.selectbox("Salón", SALONES_EVENTOS + ([ev_sel["sucursal"]] if ev_sel.get("sucursal") not in SALONES_EVENTOS else []),
                                                    index=(SALONES_EVENTOS + [ev_sel.get("sucursal")]).index(ev_sel.get("sucursal")), key=f"ev_mod_salon_{k_m}")
                            m_fecha = md2.date_input("Fecha", value=fecha_flexible(ev_sel.get("fecha")), key=f"ev_mod_fecha_{k_m}")
                            m_num = md3.number_input("Eventos ese día", min_value=1, value=max(1, _entero_seguro(ev_sel.get("num_eventos"))),
                                                     step=1, key=f"ev_mod_num_{k_m}")
                            me1, me2, me3 = st.columns(3)
                            m_ini = me1.time_input("Hora de inicio", value=datetime.strptime(ev_sel.get("hora_inicio") or "00:00", "%H:%M").time(),
                                                   key=f"ev_mod_ini_{k_m}")
                            m_fin = me2.time_input("Hora de término", value=datetime.strptime(ev_sel.get("hora_fin") or "00:00", "%H:%M").time(),
                                                   key=f"ev_mod_fin_{k_m}")
                            claves_m = list(SUPERVISORES_NOMBRE)
                            m_sup = me3.selectbox("Supervisor a cargo", claves_m, format_func=nombre_sup,
                                                  index=claves_m.index(ev_sel.get("supervisor")) if ev_sel.get("supervisor") in claves_m else 0,
                                                  key=f"ev_mod_sup_{k_m}")
                            m_desc = st.text_input("Descripción", value=ev_sel.get("descripcion", ""), key=f"ev_mod_desc_{k_m}")
                            mf1, mf2, mf3 = st.columns(3)
                            m_costo = mf1.number_input("Costo del evento ($)", min_value=0.0, value=float(ev_sel["costo_evento"]), step=100.0, key=f"ev_mod_costo_{k_m}")
                            m_veh = mf2.number_input("Vehículos", min_value=0, value=int(ev_sel["vehiculos"]), step=1, key=f"ev_mod_veh_{k_m}")
                            m_bol = mf3.number_input("Tarifa por boleto ($)", min_value=0.0, value=float(ev_sel["costo_boleto"]), step=5.0, key=f"ev_mod_bol_{k_m}")
                            mg1, mg2, mg3 = st.columns(3)
                            m_nom = mg1.number_input("Nómina ($)", min_value=0.0, value=_flotante_seguro(ev_sel.get("nomina")), step=50.0, key=f"ev_mod_nom_{k_m}")
                            m_otros = mg2.number_input("Otros gastos ($)", min_value=0.0, value=_flotante_seguro(ev_sel.get("otros_gastos")), step=50.0, key=f"ev_mod_otros_{k_m}")
                            m_estado = mg3.selectbox("Estado", ESTADOS_EVENTO,
                                                     index=ESTADOS_EVENTO.index(ev_sel.get("estado")) if ev_sel.get("estado") in ESTADOS_EVENTO else 0,
                                                     key=f"ev_mod_estado_{k_m}")
                            vista_m = calcular_evento({"costo_evento": m_costo, "vehiculos": m_veh, "costo_boleto": m_bol,
                                                       "nomina": m_nom, "otros_gastos": m_otros})
                            vm1, vm2, vm3 = st.columns(3)
                            vm1.metric("Ingreso total", f"${vista_m['ingreso_total']:,.2f}",
                                       delta=f"{vista_m['ingreso_total'] - ev_sel['ingreso_total']:+,.2f}" if vista_m['ingreso_total'] != ev_sel['ingreso_total'] else None)
                            vm2.metric("Gastos", f"${vista_m['gastos_total']:,.2f}")
                            vm3.metric("Utilidad", f"${vista_m['utilidad']:,.2f}",
                                       delta=f"{vista_m['utilidad'] - ev_sel['utilidad']:+,.2f}" if vista_m['utilidad'] != ev_sel['utilidad'] else None)
                            bm1, bm2 = st.columns(2)
                            if bm1.button("💾 Guardar cambios del evento", type="primary", use_container_width=True, key=f"ev_mod_guardar_{k_m}"):
                                nuevo_m = dict(ev_sel)
                                nuevo_m.update({
                                    "sucursal": m_salon, "fecha": f"{m_fecha:%Y-%m-%d}", "num_eventos": int(m_num),
                                    "hora_inicio": m_ini.strftime("%H:%M"), "hora_fin": m_fin.strftime("%H:%M"),
                                    "supervisor": m_sup, "descripcion": m_desc.strip(), "costo_evento": m_costo,
                                    "vehiculos": int(m_veh), "costo_boleto": m_bol, "nomina": m_nom, "otros_gastos": m_otros,
                                    "estado": m_estado, "modificado_por": st.session_state.usuario_autenticado,
                                    "modificado_en": ahora_mx().strftime("%Y-%m-%d %H:%M:%S"),
                                })
                                nuevo_m = calcular_evento(nuevo_m)
                                campos_log = {"sucursal": "salón", "fecha": "fecha", "num_eventos": "eventos", "hora_inicio": "inicio",
                                              "hora_fin": "fin", "supervisor": "supervisor", "costo_evento": "costo",
                                              "vehiculos": "vehículos", "costo_boleto": "tarifa", "nomina": "nómina",
                                              "otros_gastos": "otros gastos", "estado": "estado", "descripcion": "descripción"}
                                cambios_m = [f"{nombre}: {ev_sel.get(campo)} → {nuevo_m.get(campo)}" for campo, nombre in campos_log.items()
                                             if str(ev_sel.get(campo)) != str(nuevo_m.get(campo))]
                                if not cambios_m:
                                    st.info("No hay cambios que guardar.")
                                else:
                                    try:
                                        guardar_eventos([nuevo_m if e.get("id") == k_m else e for e in cargar_eventos()])
                                        guardar_evento_auditoria("EVENTO EXTRAORDINARIO MODIFICADO", f"{k_m} · " + " · ".join(cambios_m))
                                        st.success("Evento actualizado.")
                                        st.rerun()
                                    except OSError as exc:
                                        st.error(f"No se pudo guardar: {exc}")
                            ok_borrar_m = bm2.checkbox("Confirmo borrar este evento", key=f"ev_mod_borrar_ok_{k_m}")
                            if bm2.button("🗑️ Borrar evento", use_container_width=True, disabled=not ok_borrar_m, key=f"ev_mod_borrar_{k_m}"):
                                try:
                                    guardar_eventos([e for e in cargar_eventos() if e.get("id") != k_m])
                                    guardar_evento_auditoria("EVENTO EXTRAORDINARIO BORRADO",
                                                             f"{k_m} · {ev_sel.get('sucursal')} · {ev_sel.get('fecha')} · "
                                                             f"{ev_sel.get('num_eventos')} evento(s) · ingreso ${ev_sel['ingreso_total']:,.2f}")
                                    st.success("Evento borrado.")
                                    st.rerun()
                                except OSError as exc:
                                    st.error(f"No se pudo borrar: {exc}")

                    st.markdown("#### 💬 Enviar a cada supervisor sus eventos por WhatsApp")
                    cfg_tel = cargar_sucursales_config()
                    tels_sup = cfg_tel.get("telefonos_sup", {})
                    for clave_w in (SUPERVISORES_NOMBRE if es_contador else [clave_sup_ev]):
                        evs_w = sorted([e for e in del_periodo if e.get("supervisor") == clave_w],
                                       key=lambda x: (x.get("fecha", ""), x.get("hora_inicio", "")))
                        lineas_w = [f"*CLUB PARK · EVENTOS EXTRAORDINARIOS*",
                                    f"Supervisor a cargo: *{nombre_sup(clave_w)}*",
                                    f"{periodo_ev}", ""]
                        if not evs_w:
                            lineas_w.append("No tienes eventos asignados en este periodo.")
                        for n_w, e in enumerate(evs_w, start=1):
                            f_w = fecha_flexible(e.get("fecha"))
                            lineas_w.append(f"*{n_w}. {e.get('sucursal')}*")
                            lineas_w.append(f"📅 {NOMBRES_DIA[f_w.weekday()]} {f_w:%d/%m/%Y} · 🕒 {e.get('hora_inicio', '')} a {e.get('hora_fin', '')}")
                            detalle_w = f"Eventos: {e.get('num_eventos', 1)} · Vehículos estimados: {e.get('vehiculos_est', e['vehiculos'])}"
                            if e.get("descripcion"):
                                detalle_w += f" · {e['descripcion']}"
                            lineas_w.append(detalle_w)
                            if e.get("estado") and e.get("estado") != "Programado":
                                lineas_w.append(f"Estado: {e['estado']}")
                            lineas_w.append("")
                        if evs_w:
                            lineas_w.append(f"Total: {sum(_entero_seguro(e.get('num_eventos')) for e in evs_w)} evento(s) · "
                                            f"{sum(e.get('vehiculos_est', e['vehiculos']) for e in evs_w)} vehículos")
                            lineas_w.append("Favor de confirmar personal y boletaje. Gracias.")
                        texto_w = "\n".join(lineas_w)
                        with st.expander(f"👔 {nombre_sup(clave_w)} · {len(evs_w)} evento(s)", expanded=bool(evs_w)):
                            st.text_area("Mensaje", texto_w, height=220, key=f"ev_wa_txt_{clave_w}_{desde_ev:%Y%m%d}")
                            tel_w = re.sub(r"\D", "", str(tels_sup.get(clave_w, "")))
                            if es_contador:
                                tw1, tw2 = st.columns([3, 1])
                                nuevo_tel = tw1.text_input("WhatsApp del supervisor (10 dígitos)", value=tel_w[-10:] if tel_w else "",
                                                           key=f"ev_wa_tel_{clave_w}")
                                tw2.write("")
                                if tw2.button("💾 Guardar", key=f"ev_wa_tel_btn_{clave_w}"):
                                    digitos = re.sub(r"\D", "", nuevo_tel)
                                    if digitos and len(digitos) != 10:
                                        st.error("El número debe tener 10 dígitos.")
                                    else:
                                        cfg_tel.setdefault("telefonos_sup", {})[clave_w] = digitos
                                        guardar_sucursales_config(cfg_tel)
                                        st.success("Número guardado.")
                                        st.rerun()
                            destino_w = f"52{tel_w[-10:]}" if len(tel_w) >= 10 else ""
                            st.link_button(f"💬 Enviar por WhatsApp a {nombre_sup(clave_w)}",
                                           f"https://wa.me/{destino_w}?text=" + urllib.parse.quote(texto_w),
                                           use_container_width=True, disabled=not evs_w)
                            if not destino_w:
                                st.caption("Sin número guardado: WhatsApp te pedirá elegir el contacto.")

                    if OPENPYXL_OK:
                        salida_ev = io.BytesIO()
                        verde_ev = PatternFill(start_color="0B3B24", end_color="0B3B24", fill_type="solid")
                        no_moneda_ev = {"Eventos", "Vehículos", "Margen %"}
                        with pd.ExcelWriter(salida_ev, engine="openpyxl") as wr_ev:
                            for nombre_h, df_h in (("Por salón", df_salon_tot), ("Por supervisor", df_sup_ev), ("Detalle", df_det_ev)):
                                if df_h.empty:
                                    df_h = pd.DataFrame({"Sin información": ["No hay eventos en el periodo"]})
                                df_h.to_excel(wr_ev, sheet_name=nombre_h, index=False, startrow=3)
                                ws_ev = wr_ev.sheets[nombre_h]
                                ws_ev["A1"] = "CLUB PARK · EVENTOS EXTRAORDINARIOS EN SALONES"
                                ws_ev["A1"].font = Font(size=14, bold=True, color="0B3B24")
                                ws_ev["A2"] = periodo_ev
                                ws_ev["A2"].font = Font(size=10, italic=True)
                                for celda in ws_ev[4]:
                                    celda.fill = verde_ev
                                    celda.font = Font(bold=True, color="FFFFFF")
                                    celda.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                                for i_c, col_c in enumerate(df_h.columns, start=1):
                                    letra = get_column_letter(i_c)
                                    ws_ev.column_dimensions[letra].width = min(max(len(str(col_c)) + 4, 13), 34)
                                    for celda in list(ws_ev[letra])[4:]:
                                        if isinstance(celda.value, (int, float)) and col_c == "Margen %":
                                            celda.number_format = '0.0"%"'
                                        elif isinstance(celda.value, (int, float)) and col_c not in no_moneda_ev:
                                            celda.number_format = '"$"#,##0.00'
                                for fila_x in ws_ev.iter_rows(min_row=5):
                                    if str(fila_x[0].value) == "TOTAL":
                                        for celda in fila_x:
                                            celda.font = Font(bold=True)
                                ws_ev.freeze_panes = "A5"
                        st.download_button("📊 Descargar Excel de eventos", data=salida_ev.getvalue(),
                                           file_name=f"EVENTOS_{desde_ev:%Y%m%d}_{hasta_ev:%Y%m%d}.xlsx",
                                           mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", key="ev_dl_excel")
                    st.download_button("📄 Descargar detalle en CSV", data=df_det_ev.to_csv(index=False).encode("utf-8-sig"),
                                       file_name=f"EVENTOS_{desde_ev:%Y%m%d}_{hasta_ev:%Y%m%d}.csv", mime="text/csv", key="ev_dl_csv")

            # ---------------- SUCURSALES Y SALONES (solo Contador) ----------------
            with sub_suc:
                if not es_contador:
                    st.info("Solo el Contador Auditor (máxima jerarquía) puede agregar o quitar sucursales y salones.")
                else:
                    cfg_s = cargar_sucursales_config()
                    salones_act = list(SALONES_EVENTOS)

                    st.markdown("#### ➕ Agregar sucursal o salón")
                    tipo_ns = st.radio("Tipo", ["🚗 Sucursal de servicio (valet / estacionamiento)", "🏛️ Salón de eventos"],
                                       horizontal=True, key="ns_tipo")
                    es_salon_ns = tipo_ns.startswith("🏛️")
                    ns1, ns2 = st.columns([3, 2])
                    nombre_ns = re.sub(r"\s+", " ", ns1.text_input("Nombre", key="ns_nombre")).strip().upper()
                    tarifa_ns = ns2.number_input("Tarifa por boleto ($)", min_value=0.0, value=100.0 if es_salon_ns else 50.0,
                                                 step=5.0, key=f"ns_tarifa_{es_salon_ns}")
                    ns3, ns4, ns5 = st.columns(3)
                    sup_ns = ns3.selectbox("Supervisor a cargo", list(SUPERVISORES_NOMBRE), format_func=lambda c: SUPERVISORES_NOMBRE[c], key="ns_sup")
                    dias_ns = ns4.text_input("Días que opera", value="SÁBADO Y DOMINGO" if es_salon_ns else "LUNES A DOMINGO",
                                             key=f"ns_dias_{es_salon_ns}").strip().upper()
                    horario_ns = ns5.text_input("Horario", value="-", key="ns_horario").strip().upper()
                    st.caption("Al agregarla se crea su credencial QR de Líder de Sucursal (la ves en 🏷️ Credenciales QR). "
                               "Si cobra por tiempo o con promoción, ajústalo después en 💲 Tarifas.")
                    if st.button("➕ Agregar", type="primary", key="ns_agregar"):
                        if not nombre_ns:
                            st.error("Escribe el nombre.")
                        elif nombre_ns in PRECIOS_SUCURSALES:
                            if es_salon_ns and nombre_ns not in salones_act:
                                cfg_s["salones"] = salones_act + [nombre_ns]
                                guardar_sucursales_config(cfg_s)
                                guardar_evento_auditoria("SALÓN AGREGADO", f"{nombre_ns} (sucursal existente marcada como salón)")
                                st.success(f"{nombre_ns} ya existía como sucursal; ahora también es salón de eventos.")
                                st.rerun()
                            else:
                                st.error(f"{nombre_ns} ya existe.")
                        else:
                            cfg_s["agregadas"][nombre_ns] = float(tarifa_ns)
                            if nombre_ns in cfg_s["quitadas"]:
                                cfg_s["quitadas"].remove(nombre_ns)
                            cfg_s.setdefault("supervisores", {})[nombre_ns] = sup_ns
                            cfg_s.setdefault("detalles", {})[nombre_ns] = {"dias": dias_ns or "LUNES A DOMINGO", "horarios": horario_ns or "-"}
                            if es_salon_ns:
                                cfg_s["salones"] = salones_act + [nombre_ns]
                            token_ns = f"CPV-QR-{codigo_sucursal(nombre_ns)}-{uuid.uuid4().hex[:4].upper()}"
                            cfg_s.setdefault("tokens_extra", {})[token_ns] = {
                                "nombre": f"Líder {nombre_ns.title()}", "sucursal": nombre_ns, "rol": "Líder de Sucursal"}
                            try:
                                guardar_sucursales_config(cfg_s)
                                guardar_evento_auditoria("SALÓN AGREGADO" if es_salon_ns else "SUCURSAL AGREGADA",
                                                         f"{nombre_ns} · tarifa ${tarifa_ns:,.2f} · {SUPERVISORES_NOMBRE[sup_ns]} · credencial {token_ns}")
                                st.success(f"{nombre_ns} agregada. Credencial del líder: {token_ns}")
                                st.rerun()
                            except OSError as exc:
                                st.error(f"No se pudo guardar: {exc}")

                    st.divider()
                    st.markdown("#### 🗑️ Quitar sucursal o salón")
                    quitar_ns = st.selectbox("Sucursal o salón", ["—"] + SUCURSALES,
                                             format_func=lambda s: s if s == "—" else (f"{s} (salón)" if s in salones_act else s),
                                             key="ns_quitar_sel")
                    modo_q = "Quitarla de todo el sistema"
                    if quitar_ns in salones_act:
                        modo_q = st.radio("¿Qué quieres hacer?", ["Quitarla de todo el sistema", "Solo dejar de ser salón de eventos (sigue como sucursal)"],
                                          key="ns_quitar_modo")
                    st.caption("El historial ya registrado (vehículos, cobros, reportes, eventos) se conserva. "
                               "Al quitarla de todo el sistema, su credencial de líder deja de funcionar.")
                    ok_q = st.checkbox("Confirmo que quiero quitarla", key="ns_quitar_ok")
                    if st.button("🗑️ Quitar", disabled=(quitar_ns == "—" or not ok_q), key="ns_quitar_btn"):
                        cfg_s["salones"] = [s for s in salones_act if s != quitar_ns]
                        if modo_q.startswith("Quitarla"):
                            cfg_s["agregadas"].pop(quitar_ns, None)
                            if quitar_ns not in cfg_s["quitadas"]:
                                cfg_s["quitadas"].append(quitar_ns)
                            cfg_s.get("supervisores", {}).pop(quitar_ns, None)
                            cfg_s.get("detalles", {}).pop(quitar_ns, None)
                            cfg_s["tokens_extra"] = {t: d for t, d in cfg_s.get("tokens_extra", {}).items() if d.get("sucursal") != quitar_ns}
                        try:
                            guardar_sucursales_config(cfg_s)
                            guardar_evento_auditoria("SUCURSAL QUITADA" if modo_q.startswith("Quitarla") else "SALÓN QUITADO",
                                                     f"{quitar_ns} · {modo_q}")
                            st.success(f"{quitar_ns}: {modo_q.lower()}.")
                            st.rerun()
                        except OSError as exc:
                            st.error(f"No se pudo guardar: {exc}")

                    st.divider()
                    st.markdown("#### 📋 Catálogo actual")
                    tarifas_cat = cargar_tarifas()
                    st.dataframe(pd.DataFrame([{
                        "No.": i["num"], "Nombre": i["sitio"], "Tipo": "Salón de eventos" if i["sitio"] in salones_act else "Servicio",
                        "Tarifa": etiqueta_modalidad(tarifas_cat.get(i["sitio"], tarifa_de(i["sitio"])),
                                                     modalidades_disponibles(tarifas_cat.get(i["sitio"], tarifa_de(i["sitio"])))[0]),
                        "Supervisor": SUPERVISORES_NOMBRE.get(i.get("supervisor"), i.get("supervisor", "-")),
                        "Días": i.get("dias", "-"), "Horario": i.get("horarios", "-"),
                    } for i in CATALOGO_OFICIAL_23 if i["sitio"] in PRECIOS_SUCURSALES]), use_container_width=True, hide_index=True)

if _SECCION == "eventos" and "eventos" in _TABS:
    _seccion_eventos()
